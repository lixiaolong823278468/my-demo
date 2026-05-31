from __future__ import annotations

import argparse
import gc
import json
import math
import mimetypes
import os
import shutil
import stat
import threading
import time
import traceback
import urllib.parse
import webbrowser
from dataclasses import dataclass, field
from datetime import datetime
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any, Callable

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from data_quality import latest_quality_report, list_quality_reports, load_quality_report
from dayahead_core import (
    DEFAULT_FORECAST_FILE,
    DEFAULT_HISTORY_DIR,
    DEFAULT_KNN_SIMILARITY_CONFIG,
    DEFAULT_MODEL_ROOT,
    DEFAULT_REALTIME_MODEL_ROOT,
    DEFAULT_MULTI_DAY_FORECAST_FILE,
    DEFAULT_MULTI_DAY_OUTPUT_FILE,
    DEFAULT_OUTPUT_FILE,
    DEFAULT_SEGMENT_CONFIG,
    DAYAHEAD_DATA_MODE,
    MODEL_TARGET_DAYAHEAD,
    MODEL_TARGET_REALTIME,
    PRICE_CAP,
    PRICE_FLOOR,
    REALTIME_DATA_MODE,
    TrainConfig,
    activate_model_version,
    build_history_source_signature,
    delete_model_versions,
    extract_forecast_thermal_capacity_file_value,
    list_model_versions,
    load_cached_history_collection,
    load_current_metadata,
    load_forecast_generic,
    load_models,
    load_training_log,
    load_training_preferences,
    model_target_context,
    normalize_model_target,
    normalize_knn_similarity_config,
    normalize_prediction_reference_preferences,
    normalize_segment_config,
    parse_multi_day_forecast_workbook,
    predict_multi_day_prices,
    predict_prices_compare,
    predict_thermal_capacity_value,
    rank_model_candidates,
    rollback_to_previous,
    save_training_preferences,
    train_and_register,
    training_history_signature,
    training_run_fingerprint,
)
from price_interval import normalize_price_intervals
from prediction_archive import (
    DEFAULT_PREDICTION_ARCHIVE_ROOT,
    archive_prediction_bundle,
    list_prediction_archives,
    load_prediction_archive_detail,
)


BASE_DIR = Path(__file__).resolve().parent
FRONTEND_APP_DIR = BASE_DIR / "frontend_app"
DEFAULT_HOST = "127.0.0.1"
DEFAULT_PORT = 8000

MODEL_ROOT = Path(DEFAULT_MODEL_ROOT)
REALTIME_MODEL_ROOT = Path(DEFAULT_REALTIME_MODEL_ROOT)
FORECAST_FILE = Path(DEFAULT_FORECAST_FILE)
MULTI_DAY_FORECAST_FILE = Path(DEFAULT_MULTI_DAY_FORECAST_FILE)
HISTORY_DIR = Path(DEFAULT_HISTORY_DIR)
OUTPUT_FILE = Path(DEFAULT_OUTPUT_FILE)
MULTI_DAY_OUTPUT_FILE = Path(DEFAULT_MULTI_DAY_OUTPUT_FILE)
OUTPUT_ROOT = BASE_DIR / "output"
PREDICTION_ARCHIVE_ROOT = DEFAULT_PREDICTION_ARCHIVE_ROOT
WINDOW_OPTIMIZATION_STATE_FILE = MODEL_ROOT / "window_optimization.json"
WINDOW_OPTIMIZATION_OUTPUT_ROOT = OUTPUT_ROOT / "window_optimization"
REALTIME_WINDOW_OPTIMIZATION_STATE_FILE = REALTIME_MODEL_ROOT / "window_optimization.json"
REALTIME_WINDOW_OPTIMIZATION_OUTPUT_ROOT = OUTPUT_ROOT / "realtime_window_optimization"
AUTO_WINDOW_CANDIDATES = [30, 45, 60, 75, 90, 120, 150, 180, 240, 365]
AUTO_WINDOW_FINE_RADIUS = 15
AUTO_WINDOW_MAX_HISTORY_DAYS = 100
AUTO_WINDOW_VALID_DAYS = 10
AUTO_WINDOW_NUM_BOOST_ROUND = 400
AUTO_WINDOW_RERANK_TOP_N = 3
AUTO_WINDOW_RERANK_MAX_N = 5
AUTO_WINDOW_RERANK_TIE_TOLERANCE = 0.5
AUTO_WINDOW_RERANK_TIE_PCT = 0.01
AUTO_WINDOW_ROLLING_BACKTEST_HORIZONS = (14,)
AUTO_WINDOW_CHECK_INTERVAL_SECONDS = 60
DEFAULT_TRAINING_WINDOW_DAYS = 60
LAST_AUTO_WINDOW_CHECK_AT: datetime | None = None


def target_model_root(target: str | None, model_root: Path = MODEL_ROOT) -> Path:
    context = model_target_context(target, model_root)
    return Path(context["model_root"])


def load_target_training_preferences(target: str | None, model_root: Path = MODEL_ROOT) -> dict[str, Any]:
    normalized = normalize_model_target(target)
    root = target_model_root(normalized, model_root)
    if normalized == MODEL_TARGET_REALTIME:
        preference_path = root / "training_preferences.json"
        if preference_path.exists():
            return load_training_preferences(root)
        day_ahead_preferences = load_training_preferences(model_root)
        save_training_preferences(root, day_ahead_preferences)
        return load_training_preferences(root)
    return load_training_preferences(root)


def save_target_training_preferences(target: str | None, preferences: dict[str, Any], model_root: Path = MODEL_ROOT) -> dict[str, Any]:
    root = target_model_root(target, model_root)
    return save_training_preferences(root, preferences)


def parse_segment_config_payload(payload: dict[str, Any]) -> list[dict[str, object]] | None:
    mode = str(payload.get("segment_mode") or "default")
    if mode != "custom":
        return None
    raw_segments = payload.get("segment_config")
    if not isinstance(raw_segments, list):
        raise ValueError("自定义时段配置不能为空")
    normalized = normalize_segment_config(raw_segments)
    return [
        {
            "name": name,
            "start_time": f"{((start - 1) * 15) // 60:02d}:{((start - 1) * 15) % 60:02d}",
            "end_time": f"{(end * 15) // 60:02d}:{(end * 15) % 60:02d}",
        }
        for name, (start, end) in normalized.items()
    ]


def minutes_to_time_text(minutes: int) -> str:
    if minutes >= 1440:
        return "24:00"
    hour = max(0, min(23, minutes // 60))
    minute = max(0, min(59, minutes % 60))
    return f"{hour:02d}:{minute:02d}"


def build_even_segment_config(count: int) -> list[dict[str, object]]:
    segment_count = max(1, min(24, int(count or 1)))
    base_periods = 96 // segment_count
    remainder = 96 % segment_count
    rows: list[dict[str, object]] = []
    start_period = 1
    for index in range(segment_count):
        period_count = base_periods + (1 if index < remainder else 0)
        end_period = start_period + period_count - 1
        rows.append(
            {
                "name": f"segment_{index + 1}",
                "start_time": minutes_to_time_text((start_period - 1) * 15),
                "end_time": minutes_to_time_text(end_period * 15),
            }
        )
        start_period = end_period + 1
    return rows


def segment_config_count(segment_config: list[dict[str, object]] | None) -> int:
    return len(segment_config or DEFAULT_SEGMENT_CONFIG)


def normalize_segment_search_configs(
    payload: dict[str, Any],
    current_segment_config: list[dict[str, object]] | None,
) -> list[dict[str, Any]]:
    current_count = segment_config_count(current_segment_config)
    if not payload.get("segment_search_enabled"):
        return [
            {
                "segment_count": current_count,
                "segment_config": current_segment_config,
                "source": "current",
            }
        ]

    selected_counts = {
        int(item)
        for item in payload.get("segment_search_counts") or []
        if str(item).strip().isdigit() and int(item) >= 1
    }
    raw_configs = payload.get("segment_search_configs") if isinstance(payload.get("segment_search_configs"), list) else []
    configs_by_count: dict[int, list[dict[str, object]]] = {}
    for item in raw_configs:
        if not isinstance(item, dict):
            continue
        try:
            count = int(item.get("segment_count") or 0)
        except (TypeError, ValueError):
            continue
        if count < 1:
            continue
        raw_segment_config = item.get("segment_config")
        if isinstance(raw_segment_config, list) and raw_segment_config:
            parsed = parse_segment_config_payload({"segment_mode": "custom", "segment_config": raw_segment_config})
            configs_by_count[count] = parsed or build_even_segment_config(count)
        else:
            configs_by_count[count] = build_even_segment_config(count)

    if not selected_counts:
        selected_counts = {6}
    result: list[dict[str, Any]] = []
    for count in sorted(selected_counts):
        segment_config = configs_by_count.get(count) or build_even_segment_config(count)
        result.append(
            {
                "segment_count": count,
                "segment_config": segment_config,
                "source": "selected",
            }
        )
    return result


def parse_rolling_backtest_horizons(payload: dict[str, Any]) -> tuple[int, ...]:
    raw_value = payload.get("rolling_backtest_horizons")
    if raw_value is None:
        raw_value = payload.get("rolling_backtest_days")
    if raw_value is None:
        return AUTO_WINDOW_ROLLING_BACKTEST_HORIZONS
    if isinstance(raw_value, str):
        raw_items: list[Any] = [item.strip() for item in raw_value.replace("，", ",").split(",")]
    elif isinstance(raw_value, (list, tuple, set)):
        raw_items = list(raw_value)
    else:
        raw_items = [raw_value]
    horizons: list[int] = []
    seen: set[int] = set()
    for item in raw_items:
        try:
            value = int(item)
        except (TypeError, ValueError):
            continue
        if value < 1 or value > 365 or value in seen:
            continue
        seen.add(value)
        horizons.append(value)
    return tuple(horizons) if horizons else AUTO_WINDOW_ROLLING_BACKTEST_HORIZONS


def window_optimization_config_defaults(state: dict[str, Any] | None = None) -> dict[str, Any]:
    state = state or {}
    try:
        valid_days = int(state.get("valid_days") or AUTO_WINDOW_VALID_DAYS)
    except (TypeError, ValueError):
        valid_days = AUTO_WINDOW_VALID_DAYS
    return {
        "valid_days": max(1, valid_days),
        "rolling_backtest_horizons": list(parse_rolling_backtest_horizons(state)),
    }


def parse_high_price_weighting_payload(payload: dict[str, Any]) -> dict[str, Any]:
    raw = payload.get("high_price_weighting") if isinstance(payload.get("high_price_weighting"), dict) else {}
    enabled = bool(raw.get("enabled", payload.get("high_price_weight_enabled", False)))
    try:
        quantile = float(raw.get("quantile", payload.get("high_price_quantile", 0.8)))
    except (TypeError, ValueError):
        quantile = 0.8
    try:
        multiplier = float(raw.get("multiplier", payload.get("high_price_weight_multiplier", 2.0)))
    except (TypeError, ValueError):
        multiplier = 2.0
    return {
        "enabled": enabled,
        "quantile": min(0.99, max(0.5, quantile)),
        "multiplier": max(1.0, multiplier),
    }


def parse_price_intervals_payload(payload: dict[str, Any]) -> list[dict[str, object]]:
    intervals = payload.get("price_intervals")
    if intervals is None:
        intervals = load_training_preferences(MODEL_ROOT).get("price_intervals")
    return normalize_price_intervals(intervals if isinstance(intervals, list) else None, allow_legacy_open_bounds=True)


def now_text() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def make_version_label(row: dict[str, Any]) -> str:
    flags: list[str] = []
    if row.get("is_default"):
        flags.append("默认模型")
    if row.get("is_previous"):
        flags.append("上一版")
    if row.get("is_pinned_default"):
        flags.append("已锁定默认")
    suffix = f" | {' / '.join(flags)}" if flags else ""
    return (
        f"{row.get('version_key', '-')}"
        f" | {row.get('created_at', '-')}"
        f" | {row.get('train_start_date', '-')} ~ {row.get('train_end_date', '-')}"
        f"{suffix}"
    )


def sanitize_scalar(value: Any) -> Any:
    if isinstance(value, (list, tuple)):
        return [sanitize_scalar(item) for item in value]
    if isinstance(value, dict):
        return {str(key): sanitize_scalar(item) for key, item in value.items()}
    if isinstance(value, np.ndarray):
        return [sanitize_scalar(item) for item in value.tolist()]
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, (pd.Timestamp, datetime)):
        if getattr(value, "hour", 0) == 0 and getattr(value, "minute", 0) == 0 and getattr(value, "second", 0) == 0:
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        if math.isnan(float(value)):
            return None
        return float(value)
    if isinstance(value, float) and math.isnan(value):
        return None
    if pd.isna(value):
        return None
    return value


def dataframe_to_records(df: pd.DataFrame) -> list[dict[str, Any]]:
    safe_df = df.copy().where(pd.notna(df), None)
    return [{str(key): sanitize_scalar(value) for key, value in row.items()} for row in safe_df.to_dict(orient="records")]


def to_jsonable(data: Any) -> Any:
    if isinstance(data, dict):
        return {str(key): to_jsonable(value) for key, value in data.items()}
    if isinstance(data, list):
        return [to_jsonable(item) for item in data]
    if isinstance(data, tuple):
        return [to_jsonable(item) for item in data]
    if isinstance(data, pd.DataFrame):
        return dataframe_to_records(data)
    return sanitize_scalar(data)


def read_json_file(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    return data if isinstance(data, dict) else {}


def write_json_file(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(to_jsonable(data), ensure_ascii=False, indent=2), encoding="utf-8")


def parse_optional_float(value: str | None) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    return float(value)


def update_segment_price_model_selection(model_root: Path, selected: dict[str, Any]) -> dict[str, Any]:
    current_dir = model_root / "current"
    metadata_path = current_dir / "metadata.json"
    metadata = read_json_file(metadata_path)
    if not metadata:
        raise FileNotFoundError("当前模型元数据不存在，请先训练模型")
    variants = metadata.get("model_variants") if isinstance(metadata.get("model_variants"), dict) else {}
    segments = metadata.get("segments") if isinstance(metadata.get("segments"), list) else metadata.get("segment_config")
    segment_names = {str(item.get("name")) for item in segments or [] if isinstance(item, dict) and item.get("name")}
    if not isinstance(selected, dict):
        raise ValueError("selected_segment_price_models 必须是对象")
    normalized: dict[str, str] = {}
    for raw_segment, raw_model_key in selected.items():
        segment_name = str(raw_segment).strip()
        model_key = str(raw_model_key).strip()
        if not segment_name or not model_key:
            continue
        if segment_names and segment_name not in segment_names:
            raise ValueError(f"未知时段：{segment_name}")
        if model_key not in variants:
            raise ValueError(f"未知模型：{model_key}")
        normalized[segment_name] = model_key
    metadata["selected_segment_price_models"] = normalized
    metadata["segment_price_model_selection_updated_at"] = now_text()
    write_json_file(metadata_path, metadata)
    run_id = str(metadata.get("run_id") or "").strip()
    history_path = model_root / "history" / run_id / "metadata.json" if run_id else None
    if history_path and history_path.exists():
        write_json_file(history_path, metadata)
    return metadata


def list_model_versions_for_target(target: str | None, model_root: Path = MODEL_ROOT) -> list[dict[str, Any]]:
    normalized = normalize_model_target(target)
    versions_df = list_model_versions(target_model_root(target, model_root))
    if versions_df.empty:
        return []
    records = dataframe_to_records(versions_df)
    for row in records:
        row["target_market"] = row.get("target_market") or normalized
        row["label"] = make_version_label(row)
    return records


def update_segment_price_model_selection_for_target(
    target: str | None,
    selected: dict[str, Any],
    model_root: Path = MODEL_ROOT,
) -> dict[str, Any]:
    return update_segment_price_model_selection(target_model_root(target, model_root), selected)


def load_current_metadata_for_target(target: str | None, model_root: Path = MODEL_ROOT) -> dict[str, Any]:
    return load_current_metadata(target_model_root(target, model_root)) or {}


def rank_model_candidates_for_target(
    target: str | None,
    metric: str = "default_score",
    price_min: float | None = None,
    price_max: float | None = None,
    model_root: Path = MODEL_ROOT,
) -> dict[str, Any]:
    return rank_model_candidates(target_model_root(target, model_root), metric=metric, price_min=price_min, price_max=price_max)


def activate_model_version_for_target(target: str | None, version_key: str, model_root: Path = MODEL_ROOT) -> Path:
    return activate_model_version(target_model_root(target, model_root), version_key)


def rollback_model_version_for_target(target: str | None, model_root: Path = MODEL_ROOT) -> Path:
    return rollback_to_previous(target_model_root(target, model_root))


def delete_model_versions_for_target(target: str | None, version_keys: list[str], model_root: Path = MODEL_ROOT) -> list[str]:
    return delete_model_versions(target_model_root(target, model_root), version_keys)


def read_first_sheet(file_path: Path) -> tuple[str, pd.DataFrame]:
    workbook = pd.read_excel(file_path, sheet_name=None)
    for sheet_name, sheet_df in workbook.items():
        if not sheet_df.empty or len(sheet_df.columns) > 0:
            result = sheet_df.copy().reset_index(drop=True)
            result.columns = [str(column).strip() for column in result.columns]
            return str(sheet_name), result
    raise ValueError(f"未从 {file_path} 读取到可用 Sheet")


def load_forecast_template_preview(file_path: Path = FORECAST_FILE) -> dict[str, Any]:
    sheet_name, df = read_first_sheet(file_path)
    preview_df = df.iloc[:96].copy().reset_index(drop=True)
    thermal_capacity_file_value = None
    for column in preview_df.columns:
        name = str(column)
        if "开机容量" in name or "运行机组容量" in name:
            values = pd.to_numeric(preview_df[column], errors="coerce").dropna()
            if not values.empty:
                thermal_capacity_file_value = sanitize_scalar(float(values.iloc[0]))
                break
    return {
        "file_path": str(file_path),
        "sheet_name": sheet_name,
        "columns": [str(column) for column in preview_df.columns],
        "rows": dataframe_to_records(preview_df),
        "row_count": int(len(preview_df)),
        "thermal_capacity_file_value": thermal_capacity_file_value,
    }


def load_multi_day_template_preview(file_path: Path = MULTI_DAY_FORECAST_FILE) -> dict[str, Any]:
    sheet_name, df = read_first_sheet(file_path)
    preview_df = df.iloc[:96].copy().reset_index(drop=True)
    detected_dates: list[str] = []
    baseline_date = None
    try:
        parsed = parse_multi_day_forecast_workbook(file_path)
        detected_dates = list(parsed.get("forecast_dates") or [])
        baseline_date = parsed.get("baseline_date")
    except Exception:
        detected_dates = []
    return {
        "file_path": str(file_path),
        "sheet_name": sheet_name,
        "columns": [str(column) for column in preview_df.columns],
        "rows": dataframe_to_records(preview_df),
        "row_count": int(len(preview_df)),
        "baseline_date": baseline_date,
        "detected_forecast_dates": detected_dates,
    }


def load_thermal_capacity_prediction_preview(file_path: Path = FORECAST_FILE, model_root: Path = MODEL_ROOT) -> dict[str, Any]:
    try:
        history_df, _, holiday_dates, _, _ = load_cached_history_collection(HISTORY_DIR, None, True, model_root, leading_days=1)
        forecast_df, _, _ = load_forecast_generic(history_df, file_path, holiday_dates)
        file_value = extract_forecast_thermal_capacity_file_value(forecast_df)
        metadata, models, _ = load_models(model_root)
        model_entry = models.get("__thermal_capacity__") if isinstance(models, dict) else None
        model_value = predict_thermal_capacity_value(forecast_df, model_entry, file_value)
        model_meta = metadata.get("thermal_capacity_model") if isinstance(metadata.get("thermal_capacity_model"), dict) else {}
        return {
            "available": model_value is not None,
            "file_value": sanitize_scalar(file_value),
            "model_value": sanitize_scalar(model_value),
            "model_enabled": bool(model_meta.get("enabled")),
            "metrics": model_meta.get("metrics") or {},
        }
    except Exception as exc:  # noqa: BLE001
        return {
            "available": False,
            "file_value": None,
            "model_value": None,
            "error": str(exc),
        }


def save_forecast_template_rows(rows: list[dict[str, Any]], file_path: Path = FORECAST_FILE) -> None:
    workbook = load_workbook(file_path)
    worksheet = workbook[workbook.sheetnames[0]]
    headers = [worksheet.cell(row=1, column=col).value for col in range(1, worksheet.max_column + 1)]
    header_map = {str(header).strip(): index for index, header in enumerate(headers, start=1) if header not in (None, "")}

    for row_index, row_data in enumerate(rows, start=2):
        for header, value in row_data.items():
            if header not in header_map:
                continue
            worksheet.cell(row=row_index, column=header_map[header]).value = None if value == "" else value
    workbook.save(file_path)


def load_latest_prediction_output(file_path: Path = OUTPUT_FILE) -> dict[str, Any] | None:
    if not file_path.exists():
        return None
    sheet_name, df = read_first_sheet(file_path)
    result_df = df.copy().reset_index(drop=True)
    forecast_date = None
    if "date" in result_df.columns and not result_df.empty:
        forecast_date = sanitize_scalar(pd.to_datetime(result_df.loc[0, "date"], errors="coerce"))
    return {
        "file_path": str(file_path),
        "sheet_name": sheet_name,
        "forecast_date": forecast_date,
        "columns": [str(column) for column in result_df.columns],
        "rows": dataframe_to_records(result_df),
        "row_count": int(len(result_df)),
    }


@dataclass
class JobState:
    job_id: str
    job_type: str
    status: str = "等待中"
    progress: int = 0
    running: bool = True
    cancel_requested: bool = False
    cancelled: bool = False
    logs: list[str] = field(default_factory=list)
    error: str | None = None
    result: dict[str, Any] | None = None
    started_at: str = field(default_factory=now_text)
    finished_at: str | None = None


class JobCancelledError(RuntimeError):
    pass


class AppState:
    def __init__(self) -> None:
        self._lock = threading.Lock()
        self._jobs: dict[str, JobState] = {}
        self._current_job_id: str | None = None
        self._last_prediction: dict[str, Any] | None = None
        self._last_multi_day_prediction: dict[str, Any] | None = None

    def has_running_job(self) -> bool:
        if not self._current_job_id:
            return False
        job = self._jobs.get(self._current_job_id)
        return bool(job and job.running)

    def create_job(self, job_type: str) -> JobState:
        with self._lock:
            if self.has_running_job():
                raise RuntimeError("当前已有任务在执行，请等待完成后再试。")
            job_id = datetime.now().strftime(f"{job_type}_%Y%m%d_%H%M%S")
            job = JobState(job_id=job_id, job_type=job_type, status=f"{job_type}任务已创建")
            self._jobs[job_id] = job
            self._current_job_id = job_id
            return job

    def append_log(self, job_id: str, message: str, progress: int | None = None) -> None:
        with self._lock:
            job = self._jobs[job_id]
            job.status = message
            if progress is not None:
                job.progress = max(0, min(100, int(progress)))
            job.logs.append(f"[{datetime.now().strftime('%H:%M:%S')}] {message}")
            job.logs = job.logs[-400:]

    def request_cancel(self, job_id: str) -> dict[str, Any]:
        with self._lock:
            if job_id not in self._jobs:
                raise KeyError(job_id)
            job = self._jobs[job_id]
            if not job.running:
                return to_jsonable(job.__dict__)
            job.cancel_requested = True
            job.status = "正在停止任务..."
            job.logs.append(f"[{datetime.now().strftime('%H:%M:%S')}] 已收到停止请求，正在安全中断")
            job.logs = job.logs[-400:]
            return to_jsonable(job.__dict__)

    def raise_if_cancelled(self, job_id: str) -> None:
        with self._lock:
            job = self._jobs.get(job_id)
            if job and job.cancel_requested:
                raise JobCancelledError("任务已手动停止")

    def finish_job(
        self,
        job_id: str,
        result: dict[str, Any] | None = None,
        error: str | None = None,
        cancelled: bool = False,
    ) -> None:
        with self._lock:
            job = self._jobs[job_id]
            job.running = False
            job.finished_at = now_text()
            if cancelled or job.cancel_requested:
                job.cancelled = True
                job.status = "任务已停止"
                job.progress = min(job.progress or 0, 99)
                job.logs.append(f"[{datetime.now().strftime('%H:%M:%S')}] 任务已停止")
            elif error:
                job.error = error
                job.status = "执行失败"
                job.progress = min(job.progress or 0, 99)
                job.logs.append(f"[{datetime.now().strftime('%H:%M:%S')}] 执行失败：{error}")
            else:
                job.result = result
                job.status = "执行完成"
                job.progress = 100
                job.logs.append(f"[{datetime.now().strftime('%H:%M:%S')}] 执行完成")
                if result and result.get("prediction"):
                    self._last_prediction = result["prediction"]
                if result and result.get("multi_day_prediction"):
                    self._last_multi_day_prediction = result["multi_day_prediction"]
            if self._current_job_id == job_id:
                self._current_job_id = None

    def snapshot(self, job_id: str | None = None) -> dict[str, Any] | None:
        with self._lock:
            target_id = job_id or self._current_job_id
            if not target_id or target_id not in self._jobs:
                return None
            return to_jsonable(self._jobs[target_id].__dict__)

    def status(self) -> dict[str, Any]:
        with self._lock:
            current_job = self._jobs.get(self._current_job_id) if self._current_job_id else None
            recent_jobs = sorted(self._jobs.values(), key=lambda item: item.started_at, reverse=True)[:10]
            return {
                "running": bool(current_job and current_job.running),
                "current_job": to_jsonable(current_job.__dict__) if current_job else None,
                "recent_jobs": [to_jsonable(item.__dict__) for item in recent_jobs],
                "last_prediction": self._last_prediction or load_latest_prediction_output(),
                "last_multi_day_prediction": self._last_multi_day_prediction,
            }


STATE = AppState()


def build_progress_callback(job_id: str) -> Callable[[str, int | None], None]:
    def callback(message: str, percent: int | None) -> None:
        STATE.raise_if_cancelled(job_id)
        STATE.append_log(job_id, message, percent)
        STATE.raise_if_cancelled(job_id)

    return callback


def start_background_job(job_type: str, worker: Callable[[str], dict[str, Any]]) -> dict[str, Any]:
    job = STATE.create_job(job_type)

    def runner() -> None:
        try:
            result = worker(job.job_id)
            STATE.finish_job(job.job_id, result=result)
        except JobCancelledError:
            if job.job_type == "optimize":
                mark_window_optimization_attempt_status("cancelled", "任务已手动停止")
            if job.job_type == "realtime_optimize":
                mark_realtime_window_optimization_attempt_status("cancelled", "任务已手动停止")
            STATE.finish_job(job.job_id, cancelled=True)
        except Exception as exc:  # noqa: BLE001
            if job.job_type == "optimize":
                mark_window_optimization_attempt_status("failed", str(exc))
            if job.job_type == "realtime_optimize":
                mark_realtime_window_optimization_attempt_status("failed", str(exc))
            STATE.finish_job(job.job_id, error=f"{exc}\n{traceback.format_exc(limit=3)}")

    threading.Thread(target=runner, daemon=True).start()
    return {"job_id": job.job_id, "job_type": job.job_type, "status": job.status}


def summarize_train_result(result: Any) -> dict[str, Any]:
    return {
        "run_id": result.run_id,
        "current_model_dir": str(result.current_model_dir),
        "history_model_dir": str(result.history_model_dir),
        "metrics": to_jsonable(result.metrics),
        "train_dates": result.train_dates,
        "valid_dates": result.valid_dates,
        "skipped_sheets": result.skipped_sheets,
        "metadata_path": str(result.metadata_path),
        "quality_report_path": str(result.quality_report_path) if result.quality_report_path else None,
    }


def summarize_metric_rows(metrics: dict[str, dict[str, Any]]) -> dict[str, Any]:
    total_rows = 0.0
    mae_sum = 0.0
    rmse_sum = 0.0
    rmse_square_sum = 0.0
    for row in (metrics or {}).values():
        valid_rows = float(row.get("valid_rows") or 0)
        mae = row.get("final_mae")
        rmse = row.get("final_rmse")
        if valid_rows <= 0 or mae is None or rmse is None:
            continue
        mae_value = float(mae)
        rmse_value = float(rmse)
        total_rows += valid_rows
        mae_sum += mae_value * valid_rows
        rmse_sum += rmse_value * valid_rows
        rmse_square_sum += rmse_value**2 * valid_rows
    return {
        "valid_rows": int(total_rows),
        "mae": round(mae_sum / total_rows, 4) if total_rows else None,
        "rmse": round(rmse_sum / total_rows, 4) if total_rows else None,
        "rmse_true_weighted": round(math.sqrt(rmse_square_sum / total_rows), 4) if total_rows else None,
    }


def window_score(summary: dict[str, Any]) -> float:
    mae = summary.get("mae")
    rmse = summary.get("rmse")
    if mae is None or rmse is None:
        return float("inf")
    return float(mae) + 0.2 * float(rmse)


def rolling_metric_score(metric: dict[str, Any] | None) -> float | None:
    if not isinstance(metric, dict):
        return None
    mae = metric.get("mae")
    rmse = metric.get("rmse")
    if mae is None or rmse is None:
        return None
    return float(mae) + 0.2 * float(rmse)


def interval_probability_score(metric: dict[str, Any] | None) -> float | None:
    if not isinstance(metric, dict):
        return None
    interval_accuracy = metric.get("interval_accuracy")
    top2_accuracy = metric.get("top2_accuracy")
    log_loss = metric.get("log_loss")
    high_price_recall = metric.get("high_price_recall")
    if interval_accuracy is None and top2_accuracy is None and log_loss is None and high_price_recall is None:
        return None
    penalty = 0.0
    weight = 0.0
    if interval_accuracy is not None:
        penalty += 100.0 * (1.0 - float(interval_accuracy))
        weight += 1.0
    if top2_accuracy is not None:
        penalty += 40.0 * (1.0 - float(top2_accuracy))
        weight += 0.4
    if log_loss is not None:
        penalty += 12.0 * float(log_loss)
        weight += 0.3
    if high_price_recall is not None:
        penalty += 60.0 * (1.0 - float(high_price_recall))
        weight += 0.6
    if weight <= 0:
        return None
    return penalty / weight


def business_window_score(
    static_metrics: dict[str, Any],
    rolling_metrics: dict[str, Any],
) -> tuple[float, dict[str, Any]]:
    parts: list[tuple[str, float, float]] = []
    static_score = window_score(static_metrics)
    if math.isfinite(static_score):
        parts.append(("static_validation", 0.2, static_score))

    rolling_items = sorted(
        [
            (int(key), value)
            for key, value in (rolling_metrics or {}).items()
            if str(key).isdigit() and isinstance(value, dict)
        ],
        key=lambda item: item[0],
    )
    for index, (horizon, row) in enumerate(rolling_items):
        rolling_score = rolling_metric_score(row.get("overall"))
        if rolling_score is not None:
            parts.append((f"rolling_{horizon}_days", 0.3 if index == 0 else 0.2, rolling_score))

    high_scores = [
        rolling_metric_score((row.get("spike_errors") or {}).get("high"))
        for _, row in rolling_items
    ]
    high_scores = [value for value in high_scores if value is not None]
    if high_scores:
        parts.append(("high_price_spike", 0.2, sum(high_scores) / len(high_scores)))

    evening_scores = [
        rolling_metric_score((row.get("segments") or {}).get("evening_peak"))
        for _, row in rolling_items
    ]
    evening_scores = [value for value in evening_scores if value is not None]
    if evening_scores:
        parts.append(("evening_peak", 0.1, sum(evening_scores) / len(evening_scores)))

    if not parts:
        return float("inf"), {"parts": [], "basis": "none"}
    total_weight = sum(weight for _, weight, _ in parts)
    score = sum(weight * value for _, weight, value in parts) / total_weight
    return score, {
        "basis": "rolling_backtest" if rolling_metrics else "static_validation",
        "parts": [
            {"name": name, "weight": weight, "score": round(value, 4)}
            for name, weight, value in parts
        ],
    }


def summarize_window_result(result: Any, model_root: Path, window_days: int, phase: str) -> dict[str, Any]:
    metadata = read_json_file(Path(result.current_model_dir) / "metadata.json")
    variant_metrics = metadata.get("variant_metrics") or {}
    selected_key = metadata.get("selected_model_key")
    direct_metrics = summarize_metric_rows(
        variant_metrics.get(selected_key)
        or variant_metrics.get("no_lag_96")
        or variant_metrics.get("direct_price")
        or metadata.get("metrics")
        or {}
    )
    static_score = window_score(direct_metrics)
    rolling_metrics = metadata.get("rolling_backtest_metrics") or {}
    interval_metrics = ((metadata.get("price_interval_model") or {}).get("metrics") or {}).get("overall") or {}
    thermal_capacity_metrics = ((metadata.get("thermal_capacity_model") or {}).get("metrics") or {}).get("overall") or {}
    score, score_detail = business_window_score(direct_metrics, rolling_metrics)
    interval_score = interval_probability_score(interval_metrics)
    thermal_capacity_score = thermal_capacity_metrics.get("score")
    return {
        "phase": phase,
        "window_days": int(window_days),
        "score": round(score, 4) if math.isfinite(score) else None,
        "static_score": round(static_score, 4) if math.isfinite(static_score) else None,
        "interval_score": round(interval_score, 4) if interval_score is not None else None,
        "thermal_capacity_score": round(float(thermal_capacity_score), 4) if thermal_capacity_score is not None else None,
        "score_detail": score_detail,
        "run_id": result.run_id,
        "model_root": str(model_root),
        "train_start": result.train_dates[0] if result.train_dates else None,
        "train_end": result.train_dates[-1] if result.train_dates else None,
        "valid_start": result.valid_dates[0] if result.valid_dates else None,
        "valid_end": result.valid_dates[-1] if result.valid_dates else None,
        "train_date_count": len(result.train_dates),
        "valid_date_count": len(result.valid_dates),
        "sample_rows": metadata.get("sample_rows"),
        "selected_model_key": selected_key,
        "selected_model_backend": metadata.get("selected_model_backend"),
        "metrics": direct_metrics,
        "price_interval_metrics": interval_metrics,
        "thermal_capacity_metrics": thermal_capacity_metrics,
        "rolling_backtest_metrics": rolling_metrics,
    }


def window_result_sort_key(item: dict[str, Any]) -> tuple[float, float]:
    return (
        float(item.get("score") if item.get("score") is not None else float("inf")),
        float((item.get("metrics") or {}).get("mae") or float("inf")),
    )


def interval_window_result_sort_key(item: dict[str, Any]) -> tuple[float, float]:
    return (
        float(item.get("interval_score") if item.get("interval_score") is not None else float("inf")),
        float(item.get("window_days") or float("inf")),
    )


def thermal_capacity_window_result_sort_key(item: dict[str, Any]) -> tuple[float, float]:
    return (
        float(item.get("thermal_capacity_score") if item.get("thermal_capacity_score") is not None else float("inf")),
        float(item.get("window_days") or float("inf")),
    )


def select_window_rerank_candidates(
    results: list[dict[str, Any]],
    *,
    top_n: int = AUTO_WINDOW_RERANK_TOP_N,
    max_n: int = AUTO_WINDOW_RERANK_MAX_N,
    tie_tolerance: float = AUTO_WINDOW_RERANK_TIE_TOLERANCE,
    tie_pct: float = AUTO_WINDOW_RERANK_TIE_PCT,
) -> list[dict[str, Any]]:
    ranked = sorted(
        [item for item in results if item.get("score") is not None],
        key=window_result_sort_key,
    )
    if not ranked:
        return []
    base_count = min(max(1, int(top_n)), len(ranked))
    max_count = min(max(base_count, int(max_n)), len(ranked))
    selected = ranked[:base_count]
    cutoff_score = float(ranked[base_count - 1]["score"])
    allowed_gap = max(float(tie_tolerance), abs(cutoff_score) * float(tie_pct))
    for item in ranked[base_count:max_count]:
        if float(item["score"]) - cutoff_score <= allowed_gap:
            selected.append(item)
    return selected


def select_interval_rerank_candidates(
    results: list[dict[str, Any]],
    *,
    top_n: int = AUTO_WINDOW_RERANK_TOP_N,
    max_n: int = AUTO_WINDOW_RERANK_MAX_N,
    tie_tolerance: float = AUTO_WINDOW_RERANK_TIE_TOLERANCE,
    tie_pct: float = AUTO_WINDOW_RERANK_TIE_PCT,
) -> list[dict[str, Any]]:
    ranked = sorted(
        [item for item in results if item.get("interval_score") is not None],
        key=interval_window_result_sort_key,
    )
    if not ranked:
        return []
    base_count = min(max(1, int(top_n)), len(ranked))
    max_count = min(max(base_count, int(max_n)), len(ranked))
    selected = ranked[:base_count]
    cutoff_score = float(ranked[base_count - 1]["interval_score"])
    allowed_gap = max(float(tie_tolerance), abs(cutoff_score) * float(tie_pct))
    for item in ranked[base_count:max_count]:
        if float(item["interval_score"]) - cutoff_score <= allowed_gap:
            selected.append(item)
    return selected


def select_thermal_capacity_rerank_candidates(
    results: list[dict[str, Any]],
    *,
    top_n: int = AUTO_WINDOW_RERANK_TOP_N,
    max_n: int = AUTO_WINDOW_RERANK_MAX_N,
    tie_tolerance: float = AUTO_WINDOW_RERANK_TIE_TOLERANCE,
    tie_pct: float = AUTO_WINDOW_RERANK_TIE_PCT,
) -> list[dict[str, Any]]:
    ranked = sorted(
        [item for item in results if item.get("thermal_capacity_score") is not None],
        key=thermal_capacity_window_result_sort_key,
    )
    if not ranked:
        return []
    base_count = min(max(1, int(top_n)), len(ranked))
    max_count = min(max(base_count, int(max_n)), len(ranked))
    selected = ranked[:base_count]
    cutoff_score = float(ranked[base_count - 1]["thermal_capacity_score"])
    allowed_gap = max(float(tie_tolerance), abs(cutoff_score) * float(tie_pct))
    for item in ranked[base_count:max_count]:
        if float(item["thermal_capacity_score"]) - cutoff_score <= allowed_gap:
            selected.append(item)
    return selected


def combine_rerank_window_days(
    price_candidates: list[dict[str, Any]],
    interval_candidates: list[dict[str, Any]],
) -> list[int]:
    return sorted(
        {
            int(item["window_days"])
            for item in [*price_candidates, *interval_candidates]
            if item.get("window_days") is not None
        }
    )


def build_window_train_config(
    *,
    window_days: int,
    model_root: str | Path,
    feature_cache_root: str | Path | None = None,
    valid_days: int,
    num_boost_round: int,
    interval_valid_days: int,
    interval_num_boost_round: int,
    rolling_backtest_horizons: tuple[int, ...],
    segment_config: list[dict[str, object]] | None,
    high_price_weighting: dict[str, Any],
    price_intervals: list[dict[str, object]],
    enable_rolling_backtest: bool,
    search_target: str = "full",
    data_mode: str = DAYAHEAD_DATA_MODE,
    include_thermal_capacity: bool = True,
) -> TrainConfig:
    if search_target not in {"price", "interval", "thermal_capacity", "full"}:
        raise ValueError(f"未知窗口训练目标: {search_target}")
    return TrainConfig(
        history_dir=HISTORY_DIR,
        model_root=Path(model_root),
        feature_cache_root=Path(feature_cache_root) if feature_cache_root is not None else None,
        valid_days=valid_days,
        training_window_days=window_days,
        num_boost_round=num_boost_round,
        interval_valid_days=interval_valid_days,
        interval_training_window_days=window_days,
        interval_num_boost_round=interval_num_boost_round,
        thermal_capacity_valid_days=valid_days,
        thermal_capacity_training_window_days=window_days,
        thermal_capacity_num_boost_round=num_boost_round,
        enable_rolling_backtest=enable_rolling_backtest,
        rolling_backtest_horizons=rolling_backtest_horizons,
        segment_config=segment_config,
        high_price_weight_enabled=bool(high_price_weighting["enabled"]),
        high_price_quantile=float(high_price_weighting["quantile"]),
        high_price_weight_multiplier=float(high_price_weighting["multiplier"]),
        price_intervals=price_intervals,
        direct_feature_variant_keys=("no_lag_96",) if search_target in {"price", "interval", "thermal_capacity"} else None,
        model_backend_keys=("xgboost",) if search_target in {"price", "interval", "thermal_capacity"} else None,
        interval_backend_keys=("xgboost",) if search_target == "interval" else None,
        thermal_capacity_backend_keys=("xgboost", "ridge") if search_target == "thermal_capacity" else None,
        train_interval_model=search_target in {"interval", "full"},
        train_thermal_capacity_model=include_thermal_capacity and search_target in {"thermal_capacity", "full"},
        data_mode=data_mode,
    )


def annotate_train_result_metadata(result: Any, metadata: dict[str, Any]) -> None:
    for metadata_path in {
        Path(result.current_model_dir) / "metadata.json",
        Path(result.history_model_dir) / "metadata.json",
        Path(result.metadata_path),
    }:
        current = read_json_file(metadata_path)
        if current:
            current.update(metadata)
            write_json_file(metadata_path, current)


def selected_prediction_model_root(payload: dict[str, Any], job_id: str) -> tuple[Path, Path | None, str | None]:
    version_key = str(payload.get("model_version_key") or "").strip()
    if not version_key:
        model_root = Path(payload.get("model_root") or MODEL_ROOT)
        return model_root, None, None
    if any(part in {"", ".", ".."} for part in Path(version_key).parts) or "/" in version_key or "\\" in version_key:
        raise ValueError("模型版本参数不合法")
    selected_dir = MODEL_ROOT / "history" / version_key
    if not selected_dir.exists():
        raise FileNotFoundError(f"未找到模型版本：{version_key}")
    temp_root = OUTPUT_ROOT / "prediction_model_selection" / job_id
    if temp_root.exists():
        remove_tree_with_retry(temp_root)
    temp_current = temp_root / "current"
    shutil.copytree(selected_dir, temp_current)
    return temp_root, temp_root, version_key


def selected_realtime_prediction_model_root(payload: dict[str, Any], job_id: str) -> tuple[Path, Path | None, str | None]:
    base_model_root = Path(payload.get("model_root") or MODEL_ROOT)
    realtime_root = target_model_root(MODEL_TARGET_REALTIME, base_model_root)
    version_key = str(payload.get("realtime_model_version_key") or "").strip()
    if not version_key:
        return realtime_root, None, None
    if any(part in {"", ".", ".."} for part in Path(version_key).parts) or "/" in version_key or "\\" in version_key:
        raise ValueError("实时模型版本参数不合法")
    selected_dir = realtime_root / "history" / version_key
    if not selected_dir.exists():
        raise FileNotFoundError(f"未找到实时模型版本：{version_key}")
    temp_root = OUTPUT_ROOT / "prediction_realtime_model_selection" / job_id
    if temp_root.exists():
        remove_tree_with_retry(temp_root)
    temp_current = temp_root / "current"
    shutil.copytree(selected_dir, temp_current)
    return temp_root, temp_root, version_key


def parse_knn_similarity_payload(payload: dict[str, Any]) -> dict[str, float | int]:
    nested = payload.get("knn_similarity")
    source = nested if isinstance(nested, dict) else payload
    return normalize_knn_similarity_config(source)


def parse_prediction_reference_payload(payload: dict[str, Any]) -> dict[str, object]:
    return normalize_prediction_reference_preferences(
        {
            "recent_reference_days": payload.get("recent_reference_days") or payload.get("reference_days") or 1,
            "same_type_reference_days": payload.get("same_type_reference_days") or payload.get("reference_days") or 1,
            "knn_similarity": payload.get("knn_similarity") if isinstance(payload.get("knn_similarity"), dict) else payload,
        }
    )


def serialize_prediction_variant(result: Any) -> dict[str, Any]:
    payload = {
        "forecast_date": result.forecast_date,
        "output_file": str(result.output_file),
        "template_updated": result.template_updated,
        "reference_strategy_key": result.reference_strategy_key,
        "reference_strategy_label": result.reference_strategy_label,
        "reference_days_requested": result.reference_days_requested,
        "reference_dates": result.reference_dates,
        "quality_report_path": str(result.quality_report_path) if result.quality_report_path else None,
        "columns": list(result.result_df.columns),
        "rows": dataframe_to_records(result.result_df),
    }
    capacity_columns = {
        "thermal_capacity_source",
        "thermal_capacity_value",
        "thermal_capacity_model_value",
        "thermal_capacity_file_value",
    }
    if capacity_columns.intersection(result.result_df.columns):
        first = result.result_df.iloc[0]
        payload["thermal_capacity"] = {
            "source": sanitize_scalar(first.get("thermal_capacity_source")),
            "value": sanitize_scalar(first.get("thermal_capacity_value")),
            "model_value": sanitize_scalar(first.get("thermal_capacity_model_value")),
            "file_value": sanitize_scalar(first.get("thermal_capacity_file_value")),
        }
    return payload


def summarize_predict_result(result: Any) -> dict[str, Any]:
    prediction = serialize_prediction_variant(result.strategy_results[result.selected_strategy_key])
    prediction["comparison_predictions"] = {
        key: serialize_prediction_variant(value)
        for key, value in result.strategy_results.items()
    }
    prediction["selected_strategy_key"] = result.selected_strategy_key
    prediction["selected_strategy_label"] = result.selected_strategy_label
    return {
        "forecast_date": result.forecast_date,
        "output_file": str(result.output_file),
        "template_updated": result.template_updated,
        "selected_strategy_key": result.selected_strategy_key,
        "selected_strategy_label": result.selected_strategy_label,
        "quality_report_path": str(result.quality_report_path) if result.quality_report_path else None,
        "prediction": prediction,
    }


def serialize_multi_day_day(day: dict[str, Any]) -> dict[str, Any]:
    result_df = day.get("result_df")
    if not isinstance(result_df, pd.DataFrame):
        result_df = pd.DataFrame(day.get("rows") or [])
    return {
        "forecast_date": day.get("forecast_date"),
        "reference_strategy_key": day.get("reference_strategy_key"),
        "reference_strategy_label": day.get("reference_strategy_label"),
        "reference_days_requested": day.get("reference_days_requested"),
        "reference_dates": day.get("reference_dates") or [],
        "columns": list(result_df.columns),
        "rows": dataframe_to_records(result_df),
    }


def summarize_multi_day_predict_result(result: Any) -> dict[str, Any]:
    prediction = {
        "baseline_date": result.baseline_date,
        "forecast_dates": result.forecast_dates,
        "output_file": str(result.output_file),
        "template_updated": result.template_updated,
        "selected_strategy_key": result.selected_strategy_key,
        "selected_strategy_label": result.selected_strategy_label,
        "quality_report_path": str(result.quality_report_path) if result.quality_report_path else None,
        "days": [serialize_multi_day_day(day) for day in result.days],
    }
    return {
        "baseline_date": result.baseline_date,
        "forecast_dates": result.forecast_dates,
        "output_file": str(result.output_file),
        "template_updated": result.template_updated,
        "selected_strategy_key": result.selected_strategy_key,
        "selected_strategy_label": result.selected_strategy_label,
        "quality_report_path": str(result.quality_report_path) if result.quality_report_path else None,
        "multi_day_prediction": prediction,
    }


def build_train_worker(payload: dict[str, Any]) -> Callable[[str], dict[str, Any]]:
    def worker(job_id: str) -> dict[str, Any]:
        STATE.append_log(job_id, "收到手动重训请求", 1)
        STATE.raise_if_cancelled(job_id)
        target = normalize_model_target(payload.get("target"))
        base_model_root = Path(payload.get("model_root") or MODEL_ROOT)
        context = model_target_context(target, base_model_root)
        model_root = Path(context["model_root"])
        data_mode = str(context["data_mode"])
        train_thermal_capacity_model = bool(context["trains_capacity"])
        price_model_config = parse_model_training_config(
            payload,
            "price_model_config",
            default_window_days=resolve_default_training_window_days(),
            default_valid_days=14,
            default_num_boost_round=400,
        )
        interval_model_config = parse_model_training_config(
            payload,
            "interval_model_config",
            default_window_days=resolve_default_training_window_days(),
            default_valid_days=14,
            default_num_boost_round=400,
        )
        segment_config = parse_segment_config_payload(payload)
        high_price_weighting = parse_high_price_weighting_payload(payload)
        price_intervals = parse_price_intervals_payload(payload)
        STATE.append_log(job_id, describe_training_options(segment_config, high_price_weighting), 3)
        save_target_training_preferences(
            target,
            {
                "segment_mode": "custom" if segment_config else "default",
                "segment_config": segment_config or DEFAULT_SEGMENT_CONFIG,
                "high_price_weighting": high_price_weighting,
                "price_intervals": price_intervals,
            },
            model_root=base_model_root,
        )
        previous_metadata = load_current_metadata(model_root) or {}
        previous_run_id = previous_metadata.get("run_id")
        train_config = TrainConfig(
            history_dir=Path(payload.get("history_dir") or HISTORY_DIR),
            model_root=model_root,
            valid_days=price_model_config["valid_days"],
            training_window_days=price_model_config["training_window_days"],
            num_boost_round=price_model_config["num_boost_round"],
            start_date=price_model_config["start_date"],
            end_date=price_model_config["end_date"],
            interval_valid_days=interval_model_config["valid_days"],
            interval_training_window_days=interval_model_config["training_window_days"],
            interval_num_boost_round=interval_model_config["num_boost_round"],
            interval_start_date=interval_model_config["start_date"],
            interval_end_date=interval_model_config["end_date"],
            similarity_reference_days=int(payload.get("similarity_reference_days") or 100),
            segment_config=segment_config,
            high_price_weight_enabled=high_price_weighting["enabled"],
            high_price_quantile=high_price_weighting["quantile"],
            high_price_weight_multiplier=high_price_weighting["multiplier"],
            price_intervals=price_intervals,
            train_thermal_capacity_model=train_thermal_capacity_model,
            data_mode=data_mode,
        )
        if not bool(payload.get("force", True)):
            fingerprint = training_run_fingerprint(train_config, training_history_signature(train_config))
            if previous_metadata.get("training_run_fingerprint") == fingerprint:
                STATE.append_log(job_id, "训练输入和配置未变化，已跳过重复训练", 100)
                return {"skipped": True, "reason": "training_input_unchanged", "metadata": previous_metadata}
        result = train_and_register(train_config, progress_callback=build_progress_callback(job_id))
        validate_training_result_or_restore(model_root, result, segment_config, high_price_weighting, previous_run_id)
        return summarize_train_result(result)

    return worker


def build_predict_worker(payload: dict[str, Any]) -> Callable[[str], dict[str, Any]]:
    def worker(job_id: str) -> dict[str, Any]:
        STATE.append_log(job_id, "收到预测请求", 1)
        STATE.raise_if_cancelled(job_id)
        model_root, temp_root, version_key = selected_prediction_model_root(payload, job_id)
        realtime_model_root_override, realtime_temp_root, realtime_version_key = selected_realtime_prediction_model_root(payload, job_id)
        try:
            if version_key:
                STATE.append_log(job_id, f"使用指定模型版本预测：{version_key}", 3)
            if realtime_version_key:
                STATE.append_log(job_id, f"使用指定实时模型版本预测：{realtime_version_key}", 3)
            prediction_reference = parse_prediction_reference_payload(payload)
            preferences = save_training_preferences(MODEL_ROOT, {"prediction_reference": prediction_reference})
            STATE.append_log(job_id, "已保存本次预测参考配置为下次默认值", 2)
            result = predict_prices_compare(
                history_dir=Path(payload.get("history_dir") or HISTORY_DIR),
                forecast_file=Path(payload.get("forecast_file") or FORECAST_FILE),
                model_root=model_root,
                output_file=Path(payload.get("output_file") or OUTPUT_FILE),
                reference_days=int(payload.get("reference_days") or 1),
                recent_reference_days=int(prediction_reference["recent_reference_days"]),
                same_type_reference_days=int(prediction_reference["same_type_reference_days"]),
                selected_strategy=str(payload.get("selected_strategy") or "recent_n_days"),
                knn_similarity_config=prediction_reference["knn_similarity"],
                thermal_capacity_config=payload.get("thermal_capacity") or {
                    "mode": payload.get("thermal_capacity_mode") or "manual",
                    "manual_value": payload.get("manual_thermal_on_capacity"),
                },
                realtime_model_root_override=realtime_model_root_override,
                realtime_model_version_key=realtime_version_key,
                progress_callback=build_progress_callback(job_id),
            )
            summary = summarize_predict_result(result)
            metadata = load_current_metadata(model_root) or {}
            archive_result = archive_prediction_bundle(
                summary["prediction"],
                metadata,
                PREDICTION_ARCHIVE_ROOT,
                Path(payload.get("forecast_file") or FORECAST_FILE),
                Path(payload.get("output_file") or OUTPUT_FILE),
            )
            summary["model_version_key"] = version_key or metadata.get("run_id")
            summary["prediction_archive"] = archive_result
            summary["prediction"]["model_version_key"] = summary["model_version_key"]
            summary["prediction"]["realtime_model_version_key"] = realtime_version_key
            summary["prediction"]["archive_result"] = archive_result
            summary["prediction"]["prediction_reference"] = preferences.get("prediction_reference")
            return summary
        finally:
            if temp_root and temp_root.exists():
                remove_tree_with_retry(temp_root)
            if realtime_temp_root and realtime_temp_root.exists():
                remove_tree_with_retry(realtime_temp_root)

    return worker


def build_multi_day_predict_worker(payload: dict[str, Any]) -> Callable[[str], dict[str, Any]]:
    def worker(job_id: str) -> dict[str, Any]:
        STATE.append_log(job_id, "收到多天预测请求", 1)
        STATE.raise_if_cancelled(job_id)
        model_root, temp_root, version_key = selected_prediction_model_root(payload, job_id)
        try:
            if version_key:
                STATE.append_log(job_id, f"使用指定模型版本预测：{version_key}", 3)
            prediction_reference = parse_prediction_reference_payload(payload)
            preferences = save_training_preferences(MODEL_ROOT, {"prediction_reference": prediction_reference})
            result = predict_multi_day_prices(
                history_dir=Path(payload.get("history_dir") or HISTORY_DIR),
                forecast_file=Path(payload.get("forecast_file") or MULTI_DAY_FORECAST_FILE),
                model_root=model_root,
                output_file=Path(payload.get("output_file") or MULTI_DAY_OUTPUT_FILE),
                reference_days=int(prediction_reference["recent_reference_days"]),
                knn_similarity_config=prediction_reference["knn_similarity"],
                progress_callback=build_progress_callback(job_id),
            )
            summary = summarize_multi_day_predict_result(result)
            metadata = load_current_metadata(model_root) or {}
            summary["model_version_key"] = version_key or metadata.get("run_id")
            summary["multi_day_prediction"]["model_version_key"] = summary["model_version_key"]
            summary["multi_day_prediction"]["prediction_reference"] = preferences.get("prediction_reference")
            return summary
        finally:
            if temp_root and temp_root.exists():
                remove_tree_with_retry(temp_root)

    return worker


def current_history_signature() -> dict[str, Any]:
    return build_history_source_signature(HISTORY_DIR, None, True)


def window_optimization_state() -> dict[str, Any]:
    return read_json_file(WINDOW_OPTIMIZATION_STATE_FILE)


def save_window_optimization_state(state: dict[str, Any]) -> None:
    write_json_file(WINDOW_OPTIMIZATION_STATE_FILE, state)


def realtime_window_optimization_state() -> dict[str, Any]:
    return read_json_file(REALTIME_WINDOW_OPTIMIZATION_STATE_FILE)


def save_realtime_window_optimization_state(state: dict[str, Any]) -> None:
    write_json_file(REALTIME_WINDOW_OPTIMIZATION_STATE_FILE, state)


def path_size(path: Path) -> int:
    try:
        if not path.exists() and not path.is_symlink():
            return 0
        if path.is_file() or path.is_symlink():
            return path.stat().st_size
        return sum(item.stat().st_size for item in path.rglob("*") if item.is_file())
    except OSError:
        return 0


def count_path_files(path: Path) -> int:
    try:
        if not path.exists() and not path.is_symlink():
            return 0
        if path.is_file() or path.is_symlink():
            return 1
        return sum(1 for item in path.rglob("*") if item.is_file())
    except OSError:
        return 0


def safe_cleanup_target(path: Path) -> Path:
    target = path.resolve()
    base = BASE_DIR.resolve()
    if base not in [target, *target.parents]:
        raise RuntimeError(f"拒绝清理项目目录外路径：{target}")
    protected = {base, MODEL_ROOT.resolve(), HISTORY_DIR.resolve(), FRONTEND_APP_DIR.resolve()}
    if target in protected:
        raise RuntimeError(f"拒绝清理受保护目录：{target}")
    return target


def retry_writable_remove(function: Callable[[str], None], path: str, _exc_info: object) -> None:
    os.chmod(path, stat.S_IWRITE | stat.S_IREAD)
    function(path)


def remove_tree_with_retry(path: Path, retries: int = 8, delay_seconds: float = 0.25) -> None:
    if not path.exists():
        return
    last_error: Exception | None = None
    for attempt in range(retries):
        try:
            shutil.rmtree(path, onerror=retry_writable_remove)
            return
        except (OSError, shutil.Error) as exc:
            last_error = exc
            gc.collect()
            if attempt < retries - 1:
                time.sleep(delay_seconds * (attempt + 1))
    raise RuntimeError(f"删除临时目录失败：{path}，最后错误：{last_error}") from last_error


def remove_cleanup_target(
    path: Path,
    label: str,
    removed: list[dict[str, Any]],
    skipped: list[dict[str, Any]] | None = None,
) -> None:
    target = safe_cleanup_target(path)
    if not target.exists() and not target.is_symlink():
        return
    size = path_size(target)
    files = count_path_files(target)
    try:
        if target.is_dir() and not target.is_symlink():
            shutil.rmtree(target, onerror=retry_writable_remove)
        else:
            target.chmod(stat.S_IWRITE | stat.S_IREAD)
            target.unlink(missing_ok=True)
    except (OSError, shutil.Error) as exc:
        if skipped is None:
            raise
        skipped.append(
            {
                "label": label,
                "path": str(target),
                "bytes": size,
                "files": files,
                "error": str(exc),
            }
        )
        return
    removed.append(
        {
            "label": label,
            "path": str(target),
            "bytes": size,
            "files": files,
        }
    )


def prune_quality_reports(
    report_dir: Path,
    keep_latest: int,
    removed: list[dict[str, Any]],
    skipped: list[dict[str, Any]],
) -> None:
    if not report_dir.exists():
        return
    reports = sorted(report_dir.glob("*.json"), key=lambda item: item.stat().st_mtime, reverse=True)
    keep = {item.resolve() for item in reports[:keep_latest]}
    for report in reports[keep_latest:]:
        remove_cleanup_target(report, "旧数据异常报告", removed, skipped)
    for extra in report_dir.iterdir():
        if extra.resolve() in keep:
            continue
        if extra.suffix.lower() != ".json":
            remove_cleanup_target(extra, "数据异常报告目录杂项", removed, skipped)
    if report_dir.exists() and not any(report_dir.iterdir()):
        try:
            report_dir.rmdir()
        except OSError as exc:
            skipped.append(
                {
                    "label": "空数据异常报告目录",
                    "path": str(report_dir.resolve()),
                    "bytes": 0,
                    "files": 0,
                    "error": str(exc),
                }
            )


def cleanup_output_junk() -> dict[str, Any]:
    status = STATE.status()
    if status.get("running"):
        raise RuntimeError("当前有训练、寻优或预测任务正在运行，请任务结束后再清理")

    removed: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    output_file = OUTPUT_FILE.resolve()
    quality_report_dir = (OUTPUT_ROOT / "data_quality_reports").resolve()
    prediction_archive_dir = PREDICTION_ARCHIVE_ROOT.resolve()

    if OUTPUT_ROOT.exists():
        for item in list(OUTPUT_ROOT.iterdir()):
            resolved = item.resolve()
            if resolved == output_file:
                continue
            if resolved == prediction_archive_dir:
                continue
            if resolved == quality_report_dir:
                continue
            if item.name == "window_optimization":
                remove_cleanup_target(item, "自动寻优中间结果", removed, skipped)
                continue

    for folder_name in [".test_tmp", ".pytest_cache", ".mypy_cache", ".ruff_cache", "__pycache__"]:
        remove_cleanup_target(BASE_DIR / folder_name, folder_name, removed, skipped)

    for log_path in [BASE_DIR / "api_server.stdout.log", BASE_DIR / "api_server.stderr.log"]:
        if log_path.exists() and log_path.stat().st_size == 0:
            remove_cleanup_target(log_path, "空日志文件", removed, skipped)

    total_bytes = sum(item["bytes"] for item in removed)
    total_files = sum(item["files"] for item in removed)
    return {
        "message": "清理完成",
        "removed": removed,
        "removed_count": len(removed),
        "skipped": skipped,
        "skipped_count": len(skipped),
        "total_files": total_files,
        "total_bytes": total_bytes,
        "kept": [
            str(output_file),
            str(quality_report_dir),
            str(prediction_archive_dir),
            str(MODEL_ROOT.resolve()),
        ],
    }


def mark_window_optimization_attempt(
    signature: dict[str, Any],
    reason: str,
    force: bool = False,
    payload: dict[str, Any] | None = None,
) -> None:
    state = window_optimization_state()
    max_history_days = resolve_window_optimization_max_history_days(payload)
    config_defaults = window_optimization_config_defaults(payload or {})
    state.update(
        {
            "enabled": True,
            "last_attempt_at": now_text(),
            "last_attempt_reason": reason,
            "last_attempt_status": "running",
            "last_attempt_force": bool(force),
            "last_attempted_history_cache_key": signature.get("cache_key"),
            "max_search_history_days": max_history_days,
            "valid_days": config_defaults["valid_days"],
            "rolling_backtest_horizons": config_defaults["rolling_backtest_horizons"],
        }
    )
    save_window_optimization_state(state)


def mark_realtime_window_optimization_attempt(
    signature: dict[str, Any],
    reason: str,
    *,
    force: bool = False,
    payload: dict[str, Any] | None = None,
) -> None:
    payload = payload or {}
    max_history_days = resolve_window_optimization_max_history_days(payload)
    config_defaults = window_optimization_config_defaults(payload or {})
    state = realtime_window_optimization_state()
    state.update(
        {
            "enabled": True,
            "last_attempt_at": now_text(),
            "last_attempt_reason": reason,
            "last_attempt_status": "running",
            "last_attempt_message": "实时价格模型自动寻优执行中",
            "last_attempted_history_cache_key": signature.get("cache_key"),
            "max_search_history_days": max_history_days,
            "valid_days": config_defaults["valid_days"],
            "num_boost_round": int(payload.get("num_boost_round") or AUTO_WINDOW_NUM_BOOST_ROUND),
            "rolling_backtest_horizons": config_defaults["rolling_backtest_horizons"],
            "force": bool(force),
            "data_mode": REALTIME_DATA_MODE,
        }
    )
    save_realtime_window_optimization_state(state)


def resolve_window_optimization_max_history_days(payload: dict[str, Any] | None = None) -> int:
    payload = payload or {}
    state = window_optimization_state()
    raw_value = payload.get("max_history_days", state.get("max_search_history_days", AUTO_WINDOW_MAX_HISTORY_DAYS))
    try:
        value = int(raw_value)
    except (TypeError, ValueError):
        value = AUTO_WINDOW_MAX_HISTORY_DAYS
    return max(1, value)


def resolve_default_training_window_days(state: dict[str, Any] | None = None) -> int:
    optimization_state = state if state is not None else window_optimization_state()
    raw_value = optimization_state.get("best_window_days", DEFAULT_TRAINING_WINDOW_DAYS)
    try:
        value = int(raw_value)
    except (TypeError, ValueError):
        value = DEFAULT_TRAINING_WINDOW_DAYS
    return value if value >= 1 else DEFAULT_TRAINING_WINDOW_DAYS


def positive_int(value: Any, default: int) -> int:
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        parsed = default
    return parsed if parsed >= 1 else default


def parse_model_training_config(
    payload: dict[str, Any],
    key: str,
    *,
    default_window_days: int,
    default_valid_days: int,
    default_num_boost_round: int,
) -> dict[str, Any]:
    raw_config = payload.get(key)
    config = raw_config if isinstance(raw_config, dict) else {}
    training_mode = str(config.get("training_mode") or payload.get("training_mode") or "rolling_window")
    training_window_days = (
        positive_int(config.get("training_window_days", payload.get("training_window_days")), default_window_days)
        if training_mode == "rolling_window"
        else None
    )
    return {
        "training_mode": training_mode,
        "training_window_days": training_window_days,
        "start_date": (config.get("start_date") or payload.get("start_date") or None)
        if training_mode != "rolling_window" and config.get("enable_start", payload.get("enable_start", True))
        else None,
        "end_date": (config.get("end_date") or payload.get("end_date") or None)
        if training_mode != "rolling_window" and config.get("enable_end", payload.get("enable_end", True))
        else None,
        "valid_days": positive_int(config.get("valid_days", payload.get("valid_days")), default_valid_days),
        "num_boost_round": positive_int(config.get("num_boost_round", payload.get("num_boost_round")), default_num_boost_round),
    }


def mark_window_optimization_attempt_status(status: str, message: str | None = None) -> None:
    state = window_optimization_state()
    state["last_attempt_status"] = status
    state["last_attempt_finished_at"] = now_text()
    if message:
        state["last_attempt_message"] = message
    save_window_optimization_state(state)


def mark_realtime_window_optimization_attempt_status(status: str, message: str | None = None) -> None:
    state = realtime_window_optimization_state()
    state["last_attempt_status"] = status
    state["last_attempt_finished_at"] = now_text()
    if message:
        state["last_attempt_message"] = message
    save_realtime_window_optimization_state(state)


def build_window_optimization_worker_payload(
    reason: str,
    payload: dict[str, Any] | None = None,
    *,
    force: bool | None = None,
) -> dict[str, Any]:
    payload = payload or {}
    preferences = load_training_preferences(MODEL_ROOT)

    has_segment_payload = "segment_mode" in payload or "segment_config" in payload
    has_high_price_payload = (
        "high_price_weighting" in payload
        or "high_price_weight_enabled" in payload
        or "high_price_quantile" in payload
        or "high_price_weight_multiplier" in payload
    )
    segment_source = payload if has_segment_payload else preferences
    high_price_source = payload if has_high_price_payload else preferences
    price_interval_source = payload if isinstance(payload.get("price_intervals"), list) else preferences

    worker_payload: dict[str, Any] = {
        "reason": reason,
        "max_history_days": resolve_window_optimization_max_history_days(payload),
        "price_model_config": parse_model_training_config(
            payload,
            "price_model_config",
            default_window_days=DEFAULT_TRAINING_WINDOW_DAYS,
            default_valid_days=AUTO_WINDOW_VALID_DAYS,
            default_num_boost_round=AUTO_WINDOW_NUM_BOOST_ROUND,
        ),
        "interval_model_config": parse_model_training_config(
            payload,
            "interval_model_config",
            default_window_days=DEFAULT_TRAINING_WINDOW_DAYS,
            default_valid_days=AUTO_WINDOW_VALID_DAYS,
            default_num_boost_round=AUTO_WINDOW_NUM_BOOST_ROUND,
        ),
        "thermal_capacity_model_config": parse_model_training_config(
            payload,
            "thermal_capacity_model_config",
            default_window_days=DEFAULT_TRAINING_WINDOW_DAYS,
            default_valid_days=AUTO_WINDOW_VALID_DAYS,
            default_num_boost_round=AUTO_WINDOW_NUM_BOOST_ROUND,
        ),
        "fine_radius": int(payload.get("fine_radius") or AUTO_WINDOW_FINE_RADIUS),
        "segment_mode": segment_source.get("segment_mode", "default"),
        "segment_config": segment_source.get("segment_config"),
        "high_price_weighting": parse_high_price_weighting_payload(high_price_source),
        "price_intervals": normalize_price_intervals(price_interval_source.get("price_intervals"), allow_legacy_open_bounds=True),
        "rolling_backtest_horizons": list(parse_rolling_backtest_horizons(payload)),
        "segment_search_enabled": bool(payload.get("segment_search_enabled")),
        "segment_search_counts": payload.get("segment_search_counts") or [],
        "segment_search_configs": payload.get("segment_search_configs") or [],
    }
    worker_payload["valid_days"] = worker_payload["price_model_config"]["valid_days"]
    worker_payload["num_boost_round"] = worker_payload["price_model_config"]["num_boost_round"]
    if force is not None:
        worker_payload["force"] = force
    if payload.get("candidate_windows"):
        worker_payload["candidate_windows"] = payload.get("candidate_windows")
    return worker_payload


def normalize_segments_for_compare(segment_config: list[dict[str, object]] | None) -> list[dict[str, object]]:
    return parse_segment_config_payload(
        {
            "segment_mode": "custom",
            "segment_config": segment_config or DEFAULT_SEGMENT_CONFIG,
        }
    ) or []


def describe_training_options(segment_config: list[dict[str, object]] | None, high_price_weighting: dict[str, Any]) -> str:
    segments = normalize_segments_for_compare(segment_config)
    segment_text = " / ".join(
        f"{item.get('name')}:{item.get('start_time')}-{item.get('end_time')}"
        for item in segments
    )
    high_text = (
        f"开启，高价分位={high_price_weighting.get('quantile')}，权重倍数={high_price_weighting.get('multiplier')}"
        if high_price_weighting.get("enabled")
        else "关闭"
    )
    return f"实际训练配置：{len(segments)} 个时段（{segment_text}）；高价样本加权：{high_text}"


def validate_training_result_options(
    result: Any,
    segment_config: list[dict[str, object]] | None,
    high_price_weighting: dict[str, Any],
) -> None:
    metadata = read_json_file(Path(result.history_model_dir) / "metadata.json")
    expected_segments = normalize_segments_for_compare(segment_config)
    actual_segments = normalize_segments_for_compare(metadata.get("segment_config") or metadata.get("segments"))
    expected_mode = "custom" if segment_config else "default"
    actual_mode = str(metadata.get("segment_mode") or "default")
    if actual_mode != expected_mode or actual_segments != expected_segments:
        raise RuntimeError(
            "训练配置校验失败：模型元数据中的时段配置与本次提交配置不一致，已阻止启用该模型。"
        )

    actual_high = metadata.get("high_price_weighting") if isinstance(metadata.get("high_price_weighting"), dict) else {}
    expected_enabled = bool(high_price_weighting.get("enabled"))
    actual_enabled = bool(actual_high.get("enabled"))
    expected_quantile = float(high_price_weighting.get("quantile", 0.8))
    actual_quantile = float(actual_high.get("quantile", 0.8))
    expected_multiplier = float(high_price_weighting.get("multiplier", 2.0))
    actual_multiplier = float(actual_high.get("multiplier", 2.0))
    if (
        actual_enabled != expected_enabled
        or abs(actual_quantile - expected_quantile) > 1e-9
        or abs(actual_multiplier - expected_multiplier) > 1e-9
        or (expected_enabled and actual_high.get("threshold") is None)
    ):
        raise RuntimeError(
            "训练配置校验失败：模型元数据中的高价样本加权配置与本次提交配置不一致，已阻止启用该模型。"
        )


def validate_training_result_or_restore(
    model_root: Path,
    result: Any,
    segment_config: list[dict[str, object]] | None,
    high_price_weighting: dict[str, Any],
    previous_run_id: str | None,
) -> None:
    try:
        validate_training_result_options(result, segment_config, high_price_weighting)
    except Exception:
        if previous_run_id:
            try:
                activate_model_version(model_root, previous_run_id, persist_default=True)
            except Exception:
                pass
        raise


def build_window_optimization_worker(
    payload: dict[str, Any] | None = None,
    *,
    target_model_root: Path = MODEL_ROOT,
    optimization_output_root: Path = WINDOW_OPTIMIZATION_OUTPUT_ROOT,
    state_loader: Callable[[], dict[str, Any]] | None = None,
    state_writer: Callable[[dict[str, Any]], None] | None = None,
    data_mode: str = DAYAHEAD_DATA_MODE,
    include_thermal_capacity: bool = True,
    task_label: str = "自动寻优",
) -> Callable[[str], dict[str, Any]]:
    payload = payload or {}
    state_loader = state_loader or window_optimization_state
    state_writer = state_writer or save_window_optimization_state

    def worker(job_id: str) -> dict[str, Any]:
        price_model_config = parse_model_training_config(
            payload,
            "price_model_config",
            default_window_days=DEFAULT_TRAINING_WINDOW_DAYS,
            default_valid_days=AUTO_WINDOW_VALID_DAYS,
            default_num_boost_round=AUTO_WINDOW_NUM_BOOST_ROUND,
        )
        interval_model_config = parse_model_training_config(
            payload,
            "interval_model_config",
            default_window_days=DEFAULT_TRAINING_WINDOW_DAYS,
            default_valid_days=AUTO_WINDOW_VALID_DAYS,
            default_num_boost_round=AUTO_WINDOW_NUM_BOOST_ROUND,
        )
        thermal_capacity_model_config = parse_model_training_config(
            payload,
            "thermal_capacity_model_config",
            default_window_days=DEFAULT_TRAINING_WINDOW_DAYS,
            default_valid_days=AUTO_WINDOW_VALID_DAYS,
            default_num_boost_round=AUTO_WINDOW_NUM_BOOST_ROUND,
        )
        valid_days = int(price_model_config["valid_days"])
        num_boost_round = int(price_model_config["num_boost_round"])
        interval_valid_days = int(interval_model_config["valid_days"])
        interval_num_boost_round = int(interval_model_config["num_boost_round"])
        thermal_capacity_valid_days = int(thermal_capacity_model_config["valid_days"])
        thermal_capacity_num_boost_round = int(thermal_capacity_model_config["num_boost_round"])
        fine_radius = int(payload.get("fine_radius") or AUTO_WINDOW_FINE_RADIUS)
        rolling_backtest_horizons = parse_rolling_backtest_horizons(payload)
        max_history_days = resolve_window_optimization_max_history_days(payload)
        base_segment_config = parse_segment_config_payload(payload)
        segment_search_configs = normalize_segment_search_configs(payload, base_segment_config)
        high_price_weighting = parse_high_price_weighting_payload(payload)
        price_intervals = normalize_price_intervals(payload.get("price_intervals"), allow_legacy_open_bounds=True)

        STATE.append_log(job_id, f"开始{task_label}：参与时段数 {', '.join(str(item['segment_count']) for item in segment_search_configs)}", 1)
        STATE.raise_if_cancelled(job_id)
        save_training_preferences(
            target_model_root,
            {
                "segment_mode": "custom" if base_segment_config else "default",
                "segment_config": base_segment_config or DEFAULT_SEGMENT_CONFIG,
                "high_price_weighting": high_price_weighting,
                "price_intervals": price_intervals,
            },
        )

        history_signature = current_history_signature()
        history_df, _, _, _, _ = load_cached_history_collection(
            HISTORY_DIR,
            None,
            True,
            target_model_root,
            leading_days=0,
            data_mode=data_mode,
        )
        unique_dates = sorted(pd.to_datetime(history_df["date"], errors="coerce").dropna().dt.strftime("%Y-%m-%d").unique().tolist())
        max_validation_days = max(valid_days, interval_valid_days, thermal_capacity_valid_days if include_thermal_capacity else 0)
        if len(unique_dates) <= max_validation_days + 1:
            raise ValueError("历史数据天数不足，无法执行自动寻优。")
        available_train_days = len(unique_dates) - max_validation_days
        max_train_days = min(available_train_days, max_history_days)
        requested_candidates = payload.get("candidate_windows")
        if requested_candidates:
            candidate_windows = sorted({max(1, min(max_train_days, int(item))) for item in requested_candidates})
        else:
            candidate_windows = sorted({item for item in AUTO_WINDOW_CANDIDATES if item <= max_train_days} | {max_train_days})
            if max_train_days < 30:
                candidate_windows = sorted({1, max_train_days})

        run_stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        search_root = optimization_output_root / run_stamp
        search_root.mkdir(parents=True, exist_ok=True)
        shared_feature_cache_root = search_root / "shared_feature_cache"
        all_results: list[dict[str, Any]] = []
        segment_results: list[dict[str, Any]] = []
        total_segments = max(1, len(segment_search_configs))

        def segment_progress(segment_index: int, local_index: int, local_total: int) -> int:
            local_fraction = min(1.0, max(0.0, (local_index - 1) / max(local_total, 1)))
            return 5 + int(((segment_index - 1 + local_fraction) / total_segments) * 78)

        def run_segment_search(segment_item: dict[str, Any], segment_index: int) -> dict[str, Any]:
            current_segment_count = int(segment_item["segment_count"])
            current_segment_config = segment_item.get("segment_config")
            segment_root = search_root / f"{current_segment_count}_segments"
            segment_root.mkdir(parents=True, exist_ok=True)
            STATE.append_log(job_id, describe_training_options(current_segment_config, high_price_weighting), segment_progress(segment_index, 1, 1))
            price_results: list[dict[str, Any]] = []
            interval_results: list[dict[str, Any]] = []
            thermal_capacity_results: list[dict[str, Any]] = []
            thermal_capacity_fine_candidates: list[int] = []
            rerank_results: list[dict[str, Any]] = []

            def run_window(
                window_days: int,
                phase: str,
                index: int,
                total: int,
                enable_rolling_backtest: bool = False,
                search_target: str = "price",
            ) -> dict[str, Any]:
                STATE.raise_if_cancelled(job_id)
                mode_text = {
                    "price": "价格窗口粗筛",
                    "interval": "区间窗口粗筛",
                    "full": "完整复核",
                }.get(search_target, "窗口训练")
                STATE.append_log(
                    job_id,
                    f"{current_segment_count} 段 {phase} {mode_text} {index}/{total}：最近 {window_days} 天",
                    segment_progress(segment_index, index, total),
                )
                window_model_root = segment_root / f"{phase}_{window_days}"
                result = train_and_register(
                    build_window_train_config(
                        window_days=window_days,
                        model_root=window_model_root,
                        feature_cache_root=shared_feature_cache_root,
                        valid_days=thermal_capacity_valid_days if search_target == "thermal_capacity" else valid_days,
                        num_boost_round=thermal_capacity_num_boost_round if search_target == "thermal_capacity" else num_boost_round,
                        interval_valid_days=interval_valid_days,
                        interval_num_boost_round=interval_num_boost_round,
                        rolling_backtest_horizons=rolling_backtest_horizons,
                        segment_config=current_segment_config,
                        high_price_weighting=high_price_weighting,
                        price_intervals=price_intervals,
                        enable_rolling_backtest=enable_rolling_backtest,
                        search_target=search_target,
                        data_mode=data_mode,
                        include_thermal_capacity=include_thermal_capacity,
                    ),
                    progress_callback=lambda _message, _percent: STATE.raise_if_cancelled(job_id),
                )
                summary = summarize_window_result(result, window_model_root, window_days, phase)
                summary["segment_count"] = current_segment_count
                summary["segment_config"] = current_segment_config or DEFAULT_SEGMENT_CONFIG
                all_results.append(summary)
                return summary

            total_estimated = len(candidate_windows) + fine_radius * 2 + 1
            coarse_summaries = [
                run_window(window, "price_coarse", idx, total_estimated, search_target="price")
                for idx, window in enumerate(candidate_windows, start=1)
            ]
            price_results.extend(coarse_summaries)
            finite_coarse = [item for item in coarse_summaries if item.get("score") is not None]
            if not finite_coarse:
                raise ValueError(f"{current_segment_count} 段所有候选训练天数均未得到有效指标。")
            best_coarse = min(finite_coarse, key=window_result_sort_key)
            best_coarse_window = int(best_coarse["window_days"])
            fine_start = max(1 if max_train_days < 30 else 30, best_coarse_window - fine_radius)
            fine_end = min(max_train_days, best_coarse_window + fine_radius)
            searched_windows = {int(row["window_days"]) for row in price_results}
            fine_candidates = [item for item in range(fine_start, fine_end + 1) if item not in searched_windows]
            for offset, window in enumerate(fine_candidates, start=1):
                price_results.append(
                    run_window(
                        window,
                        "price_fine",
                        len(candidate_windows) + offset,
                        len(candidate_windows) + len(fine_candidates),
                        search_target="price",
                    )
                )

            finite_price_results = [item for item in price_results if item.get("score") is not None]
            if not finite_price_results:
                raise ValueError(f"{current_segment_count} 段所有候选训练天数均未得到有效指标。")
            interval_total = len(candidate_windows) + fine_radius * 2 + 1
            interval_coarse_summaries = [
                run_window(window, "interval_coarse", idx, interval_total, search_target="interval")
                for idx, window in enumerate(candidate_windows, start=1)
            ]
            interval_results.extend(interval_coarse_summaries)
            finite_interval_coarse = [item for item in interval_coarse_summaries if item.get("interval_score") is not None]
            if not finite_interval_coarse:
                raise ValueError(f"{current_segment_count} 段所有候选区间训练天数均未得到有效指标。")
            best_interval_coarse = min(finite_interval_coarse, key=interval_window_result_sort_key)
            best_interval_coarse_window = int(best_interval_coarse["window_days"])
            interval_fine_start = max(1 if max_train_days < 30 else 30, best_interval_coarse_window - fine_radius)
            interval_fine_end = min(max_train_days, best_interval_coarse_window + fine_radius)
            searched_interval_windows = {int(row["window_days"]) for row in interval_results}
            interval_fine_candidates = [
                item
                for item in range(interval_fine_start, interval_fine_end + 1)
                if item not in searched_interval_windows
            ]
            for offset, window in enumerate(interval_fine_candidates, start=1):
                interval_results.append(
                    run_window(
                        window,
                        "interval_fine",
                        len(candidate_windows) + offset,
                        len(candidate_windows) + len(interval_fine_candidates),
                        search_target="interval",
                    )
                )

            finite_interval_results = [item for item in interval_results if item.get("interval_score") is not None]
            if not finite_interval_results:
                raise ValueError(f"{current_segment_count} 段所有候选区间训练天数均未得到有效指标。")

            if include_thermal_capacity:
                thermal_capacity_total = len(candidate_windows) + fine_radius * 2 + 1
                thermal_capacity_coarse_summaries = [
                    run_window(window, "thermal_capacity_coarse", idx, thermal_capacity_total, search_target="thermal_capacity")
                    for idx, window in enumerate(candidate_windows, start=1)
                ]
                thermal_capacity_results.extend(thermal_capacity_coarse_summaries)
                finite_thermal_capacity_coarse = [
                    item for item in thermal_capacity_coarse_summaries if item.get("thermal_capacity_score") is not None
                ]
                if not finite_thermal_capacity_coarse:
                    raise ValueError(f"{current_segment_count} 段所有候选开机容量训练天数均未得到有效指标。")
                best_thermal_capacity_coarse = min(finite_thermal_capacity_coarse, key=thermal_capacity_window_result_sort_key)
                best_thermal_capacity_coarse_window = int(best_thermal_capacity_coarse["window_days"])
                thermal_capacity_fine_start = max(1 if max_train_days < 30 else 30, best_thermal_capacity_coarse_window - fine_radius)
                thermal_capacity_fine_end = min(max_train_days, best_thermal_capacity_coarse_window + fine_radius)
                searched_thermal_capacity_windows = {int(row["window_days"]) for row in thermal_capacity_results}
                thermal_capacity_fine_candidates = [
                    item
                    for item in range(thermal_capacity_fine_start, thermal_capacity_fine_end + 1)
                    if item not in searched_thermal_capacity_windows
                ]
                for offset, window in enumerate(thermal_capacity_fine_candidates, start=1):
                    thermal_capacity_results.append(
                        run_window(
                            window,
                            "thermal_capacity_fine",
                            len(candidate_windows) + offset,
                            len(candidate_windows) + len(thermal_capacity_fine_candidates),
                            search_target="thermal_capacity",
                        )
                    )
                finite_thermal_capacity_results = [
                    item for item in thermal_capacity_results if item.get("thermal_capacity_score") is not None
                ]
                if not finite_thermal_capacity_results:
                    raise ValueError(f"{current_segment_count} 段所有候选开机容量训练天数均未得到有效指标。")
            else:
                finite_thermal_capacity_results = []

            price_top = select_window_rerank_candidates(finite_price_results)
            interval_top = select_interval_rerank_candidates(finite_interval_results)
            rerank_window_days = combine_rerank_window_days(price_top, interval_top)
            for offset, window in enumerate(rerank_window_days, start=1):
                STATE.append_log(job_id, f"{current_segment_count} 段完整复核 {offset}/{len(rerank_window_days)}：最近 {window} 天", 82)
                rerank_results.append(
                    run_window(
                        window,
                        "rerank",
                        offset,
                        max(len(rerank_window_days), 1),
                        enable_rolling_backtest=True,
                        search_target="full",
                    )
                )

            best_pool = [item for item in rerank_results if item.get("score") is not None] or finite_price_results
            best = min(best_pool, key=window_result_sort_key)
            best_window_days = int(best["window_days"])
            best_interval = min(finite_interval_results, key=interval_window_result_sort_key)
            best_interval_window_days = int(best_interval["window_days"])
            best_thermal_capacity = min(finite_thermal_capacity_results, key=thermal_capacity_window_result_sort_key) if finite_thermal_capacity_results else {}
            best_thermal_capacity_window_days = int(best_thermal_capacity["window_days"]) if best_thermal_capacity else best_window_days
            STATE.append_log(job_id, f"{current_segment_count} 段最优训练天数 {best_window_days} 天，正在保存候选模型版本", 86)
            previous_metadata = load_current_metadata(target_model_root) or {}
            previous_run_id = previous_metadata.get("run_id")
            final_result = train_and_register(
                TrainConfig(
                    history_dir=HISTORY_DIR,
                    model_root=target_model_root,
                    valid_days=valid_days,
                    training_window_days=best_window_days,
                    num_boost_round=num_boost_round,
                    interval_valid_days=interval_valid_days,
                    interval_training_window_days=best_interval_window_days,
                    interval_num_boost_round=interval_num_boost_round,
                    thermal_capacity_valid_days=thermal_capacity_valid_days,
                    thermal_capacity_training_window_days=best_thermal_capacity_window_days,
                    thermal_capacity_num_boost_round=thermal_capacity_num_boost_round,
                    rolling_backtest_horizons=rolling_backtest_horizons,
                    segment_config=current_segment_config,
                    high_price_weight_enabled=high_price_weighting["enabled"],
                    high_price_quantile=high_price_weighting["quantile"],
                    high_price_weight_multiplier=high_price_weighting["multiplier"],
                    price_intervals=price_intervals,
                    train_thermal_capacity_model=include_thermal_capacity,
                    data_mode=data_mode,
                ),
                progress_callback=lambda _message, _percent: STATE.raise_if_cancelled(job_id),
            )
            validate_training_result_or_restore(target_model_root, final_result, current_segment_config, high_price_weighting, previous_run_id)
            annotate_train_result_metadata(
                final_result,
                {
                    "window_optimization": {
                        "run_stamp": run_stamp,
                        "segment_search_enabled": bool(payload.get("segment_search_enabled")),
                        "segment_count": current_segment_count,
                        "best_window_days": best_window_days,
                        "best_interval_window_days": best_interval_window_days,
                        "best_thermal_capacity_window_days": best_thermal_capacity_window_days,
                        "rolling_backtest_horizons": list(rolling_backtest_horizons),
                        "score": best.get("score"),
                        "interval_score": best_interval.get("interval_score"),
                        "thermal_capacity_score": best_thermal_capacity.get("thermal_capacity_score"),
                    }
                },
            )
            final_summary = summarize_train_result(final_result)
            combined_results = price_results + interval_results + thermal_capacity_results + rerank_results
            return {
                "segment_count": current_segment_count,
                "segment_mode": "custom" if current_segment_config else "default",
                "segment_config": current_segment_config or DEFAULT_SEGMENT_CONFIG,
                "best": best,
                "best_interval": best_interval,
                "best_thermal_capacity": best_thermal_capacity,
                "best_window_days": best_window_days,
                "best_interval_window_days": best_interval_window_days,
                "best_thermal_capacity_window_days": best_thermal_capacity_window_days,
                "fine_candidates": fine_candidates,
                "interval_fine_candidates": interval_fine_candidates,
                "thermal_capacity_fine_candidates": thermal_capacity_fine_candidates,
                "price_rerank_candidates": [int(item["window_days"]) for item in price_top],
                "interval_rerank_candidates": [int(item["window_days"]) for item in interval_top],
                "thermal_capacity_rerank_candidates": [int(item["window_days"]) for item in select_thermal_capacity_rerank_candidates(finite_thermal_capacity_results)],
                "rerank_window_days": rerank_window_days,
                "results": sorted(
                    combined_results,
                    key=lambda item: (
                        str(item.get("phase") or ""),
                        int(item.get("window_days") or 0),
                    ),
                ),
                "final_model": final_summary,
                "final_model_version": final_result.run_id,
            }

        for index, segment_item in enumerate(segment_search_configs, start=1):
            segment_results.append(run_segment_search(segment_item, index))

        best_segment = min(segment_results, key=lambda item: window_result_sort_key(item["best"]))
        best = best_segment["best"]
        best_interval = best_segment["best_interval"]
        best_thermal_capacity = best_segment["best_thermal_capacity"]
        best_window_days = int(best_segment["best_window_days"])
        best_interval_window_days = int(best_segment["best_interval_window_days"])
        best_thermal_capacity_window_days = int(best_segment["best_thermal_capacity_window_days"])
        activate_model_version(target_model_root, str(best_segment["final_model_version"]), persist_default=True)
        final_summary = best_segment["final_model"]
        state = {
            "enabled": True,
            "last_run_at": now_text(),
            "last_attempt_at": now_text(),
            "last_attempt_status": "success",
            "last_attempt_finished_at": now_text(),
            "last_attempted_history_cache_key": history_signature.get("cache_key"),
            "history_cache_key": history_signature.get("cache_key"),
            "data_mode": data_mode,
            "include_thermal_capacity": include_thermal_capacity,
            "history_date_start": unique_dates[0],
            "history_date_end": unique_dates[-1],
            "history_date_count": len(unique_dates),
            "available_train_days": available_train_days,
            "max_search_history_days": max_history_days,
            "valid_days": valid_days,
            "num_boost_round": num_boost_round,
            "price_model_config": {**price_model_config, "training_window_days": best_window_days},
            "interval_model_config": {**interval_model_config, "training_window_days": best_interval_window_days},
            "thermal_capacity_model_config": {**thermal_capacity_model_config, "training_window_days": best_thermal_capacity_window_days},
            "interval_valid_days": interval_valid_days,
            "interval_num_boost_round": interval_num_boost_round,
            "thermal_capacity_valid_days": thermal_capacity_valid_days,
            "thermal_capacity_num_boost_round": thermal_capacity_num_boost_round,
            "candidate_window_start": candidate_windows[0] if candidate_windows else None,
            "candidate_window_end": candidate_windows[-1] if candidate_windows else None,
            "candidate_window_count": len(candidate_windows),
            "coarse_candidates": candidate_windows,
            "fine_radius": fine_radius,
            "rolling_backtest_horizons": list(rolling_backtest_horizons),
            "fine_candidates": best_segment.get("fine_candidates") or [],
            "rerank_top_n": AUTO_WINDOW_RERANK_TOP_N,
            "selection_method": "segment_search_rolling_backtest_rerank",
            "segment_mode": best_segment.get("segment_mode") or "custom",
            "segment_config": best_segment["segment_config"],
            "segment_search_enabled": bool(payload.get("segment_search_enabled")),
            "segment_search_counts": [item["segment_count"] for item in segment_search_configs],
            "segment_search_results": segment_results,
            "best_segment_count": best_segment["segment_count"],
            "best_segment_model_version": best_segment["final_model_version"],
            "high_price_weighting": high_price_weighting,
            "best_window_days": best_window_days,
            "best_interval_window_days": best_interval_window_days,
            "best_thermal_capacity_window_days": best_thermal_capacity_window_days,
            "best_score": best.get("score"),
            "best_static_score": best.get("static_score"),
            "best_score_detail": best.get("score_detail"),
            "best_interval_score": best_interval.get("interval_score"),
            "best_interval_metrics": best_interval.get("price_interval_metrics"),
            "best_thermal_capacity_score": best_thermal_capacity.get("thermal_capacity_score"),
            "best_thermal_capacity_metrics": best_thermal_capacity.get("thermal_capacity_metrics"),
            "price_model_recommended": {
                "training_mode": "rolling_window",
                "training_window_days": best_window_days,
                "valid_days": valid_days,
                "num_boost_round": num_boost_round,
                "score": best.get("score"),
            },
            "interval_model_recommended": {
                "training_mode": "rolling_window",
                "training_window_days": best_interval_window_days,
                "valid_days": interval_valid_days,
                "num_boost_round": interval_num_boost_round,
                "score": best_interval.get("interval_score"),
            },
            "thermal_capacity_model_recommended": {
                "training_mode": "rolling_window",
                "training_window_days": best_thermal_capacity_window_days,
                "valid_days": thermal_capacity_valid_days,
                "num_boost_round": thermal_capacity_num_boost_round,
                "score": best_thermal_capacity.get("thermal_capacity_score"),
            },
            "best_rolling_backtest_metrics": best.get("rolling_backtest_metrics"),
            "best_metrics": {"no_lag_96": best.get("metrics")},
            "search_root": str(search_root),
            "results": sorted(all_results, key=lambda item: (float(item.get("score") or float("inf")), int(item.get("window_days") or 0))),
            "final_model": final_summary,
        }
        state_writer(state)
        STATE.append_log(job_id, f"{task_label}完成：启用 {best_segment['segment_count']} 段、最近 {best_window_days} 天模型", 100)
        return {"window_optimization": state, "train_result": final_summary}

    return worker


def maybe_start_window_optimization(reason: str = "startup", payload: dict[str, Any] | None = None) -> dict[str, Any]:
    payload = payload or {}
    state = window_optimization_state()
    try:
        signature = current_history_signature()
    except Exception as exc:  # noqa: BLE001
        return {"started": False, "reason": f"history_signature_failed: {exc}", "state": state}
    if STATE.has_running_job():
        return {"started": False, "reason": "job_running", "state": state}
    if state.get("history_cache_key") == signature.get("cache_key"):
        return {"started": False, "reason": "history_unchanged", "state": state}
    if (
        state.get("last_attempted_history_cache_key") == signature.get("cache_key")
        and state.get("last_attempt_status") == "failed"
    ):
        return {"started": False, "reason": "last_attempt_failed", "state": state}
    worker_payload = build_window_optimization_worker_payload(reason, payload)
    mark_window_optimization_attempt(signature, reason, payload=worker_payload)
    job = start_background_job("optimize", build_window_optimization_worker(worker_payload))
    return {"started": True, "job": job, "state": state}


def check_window_optimization(reason: str = "status_check", force: bool = False, payload: dict[str, Any] | None = None) -> dict[str, Any]:
    payload = payload or {}
    global LAST_AUTO_WINDOW_CHECK_AT
    if not force:
        checked_at = LAST_AUTO_WINDOW_CHECK_AT
        now = datetime.now()
        if checked_at and (now - checked_at).total_seconds() < AUTO_WINDOW_CHECK_INTERVAL_SECONDS:
            return {"started": False, "reason": "throttled", "state": window_optimization_state()}
        LAST_AUTO_WINDOW_CHECK_AT = now
        return maybe_start_window_optimization(reason, payload)
    if STATE.has_running_job():
        return {"started": False, "reason": "job_running", "state": window_optimization_state()}
    LAST_AUTO_WINDOW_CHECK_AT = datetime.now()
    worker_payload = build_window_optimization_worker_payload(reason, payload, force=True)
    try:
        mark_window_optimization_attempt(current_history_signature(), reason, force=True, payload=worker_payload)
    except Exception:
        pass
    job = start_background_job("optimize", build_window_optimization_worker(worker_payload))
    return {"started": True, "job": job, "state": window_optimization_state()}


def check_realtime_window_optimization(reason: str = "manual", force: bool = True, payload: dict[str, Any] | None = None) -> dict[str, Any]:
    payload = payload or {}
    if STATE.has_running_job():
        return {"started": False, "reason": "job_running", "state": realtime_window_optimization_state()}
    worker_payload = build_window_optimization_worker_payload(reason, payload, force=force)
    try:
        mark_realtime_window_optimization_attempt(current_history_signature(), reason, force=force, payload=worker_payload)
    except Exception:
        pass
    job = start_background_job(
        "realtime_optimize",
        build_window_optimization_worker(
            worker_payload,
            target_model_root=REALTIME_MODEL_ROOT,
            optimization_output_root=REALTIME_WINDOW_OPTIMIZATION_OUTPUT_ROOT,
            state_loader=realtime_window_optimization_state,
            state_writer=save_realtime_window_optimization_state,
            data_mode=REALTIME_DATA_MODE,
            include_thermal_capacity=False,
            task_label="实时价格模型自动寻优",
        ),
    )
    return {"started": True, "job": job, "state": realtime_window_optimization_state()}


def dual_window_target_payload(payload: dict[str, Any], target: str) -> dict[str, Any]:
    normalized = normalize_model_target(target)
    target_key = "realtime_payload" if normalized == MODEL_TARGET_REALTIME else "dayahead_payload"
    target_payload = payload.get(target_key)
    if not isinstance(target_payload, dict):
        return payload
    merged = dict(target_payload)
    if "force" in payload and "force" not in merged:
        merged["force"] = payload["force"]
    return merged


def run_window_optimization_for_target(payload: dict[str, Any], target: str, job_id: str) -> dict[str, Any]:
    normalized = normalize_model_target(target)
    target_payload = dual_window_target_payload(payload, normalized)
    worker_payload = build_window_optimization_worker_payload("dual_control", target_payload, force=bool(target_payload.get("force", True)))
    try:
        signature = current_history_signature()
    except Exception:
        signature = {}
    if normalized == MODEL_TARGET_REALTIME:
        if signature:
            mark_realtime_window_optimization_attempt(signature, "dual_control", force=True, payload=worker_payload)
        worker = build_window_optimization_worker(
            worker_payload,
            target_model_root=REALTIME_MODEL_ROOT,
            optimization_output_root=REALTIME_WINDOW_OPTIMIZATION_OUTPUT_ROOT,
            state_loader=realtime_window_optimization_state,
            state_writer=save_realtime_window_optimization_state,
            data_mode=REALTIME_DATA_MODE,
            include_thermal_capacity=False,
            task_label="实时价格模型自动寻优",
        )
        return worker(job_id)
    if signature:
        mark_window_optimization_attempt(signature, "dual_control", force=True, payload=worker_payload)
    worker = build_window_optimization_worker(worker_payload)
    return worker(job_id)


def run_dual_window_optimization_sequence(
    payload: dict[str, Any],
    runner: Callable[[str], dict[str, Any]] | None = None,
) -> dict[str, Any]:
    target_runner = runner or (lambda target: run_window_optimization_for_target(payload, target, "dual_optimize"))
    results: dict[str, Any] = {}
    for target in (MODEL_TARGET_DAYAHEAD, MODEL_TARGET_REALTIME):
        try:
            outcome = target_runner(target)
            results[target] = {"status": "success", "result": outcome}
        except Exception as exc:  # noqa: BLE001
            results[target] = {"status": "failed", "error": str(exc)}
    success_count = sum(1 for item in results.values() if item.get("status") == "success")
    if success_count == 2:
        status = "success"
    elif success_count == 1:
        status = "partial_success"
    else:
        status = "failed"
    return {"status": status, "targets": results}


def build_dual_window_optimization_worker(payload: dict[str, Any]) -> Callable[[str], dict[str, Any]]:
    def worker(job_id: str) -> dict[str, Any]:
        STATE.append_log(job_id, "开始一键训练两套价格模型", 1)

        def runner(target: str) -> dict[str, Any]:
            label = "日前价格模型" if target == MODEL_TARGET_DAYAHEAD else "实时价格模型"
            progress = 5 if target == MODEL_TARGET_DAYAHEAD else 55
            STATE.append_log(job_id, f"开始{label}自动寻优", progress)
            result = run_window_optimization_for_target(payload, target, job_id)
            STATE.append_log(job_id, f"{label}自动寻优完成", 50 if target == MODEL_TARGET_DAYAHEAD else 98)
            return result

        result = run_dual_window_optimization_sequence(payload, runner=runner)
        STATE.append_log(job_id, f"一键训练两套价格模型完成：{result['status']}", 100)
        return result

    return worker


def protected_versions(target: str | None = MODEL_TARGET_DAYAHEAD, model_root: Path = MODEL_ROOT) -> set[str]:
    versions_df = list_model_versions(target_model_root(target, model_root))
    protected: set[str] = set()
    if versions_df.empty:
        return protected
    protected.update(versions_df.loc[versions_df["is_default"], "version_key"].astype(str).tolist())
    protected.update(versions_df.loc[versions_df["is_previous"], "version_key"].astype(str).tolist())
    return protected


class ApiHandler(BaseHTTPRequestHandler):
    server_version = "DayAheadPriceApi/1.0"

    def do_OPTIONS(self) -> None:  # noqa: N802
        self.send_response(HTTPStatus.NO_CONTENT)
        self.send_common_headers()
        self.end_headers()

    def do_GET(self) -> None:  # noqa: N802
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path
        try:
            if path == "/api/config":
                return self.send_json(self.get_config())
            if path == "/api/status":
                return self.send_json(self.get_status())
            if path == "/api/window-optimization/status":
                return self.send_json({"state": window_optimization_state(), "current_job": STATE.snapshot()})
            if path == "/api/realtime-window-optimization/status":
                return self.send_json({"state": realtime_window_optimization_state(), "current_job": STATE.snapshot()})
            if path == "/api/model/current":
                query = urllib.parse.parse_qs(parsed.query)
                target = str((query.get("target") or [MODEL_TARGET_DAYAHEAD])[0] or MODEL_TARGET_DAYAHEAD)
                return self.send_json({"metadata": load_current_metadata_for_target(target)})
            if path == "/api/model/candidate-rankings":
                query = urllib.parse.parse_qs(parsed.query)
                target = str((query.get("target") or [MODEL_TARGET_DAYAHEAD])[0] or MODEL_TARGET_DAYAHEAD)
                metric = str((query.get("metric") or ["default_score"])[0] or "default_score")
                price_min = parse_optional_float((query.get("price_min") or [None])[0])
                price_max = parse_optional_float((query.get("price_max") or [None])[0])
                return self.send_json(rank_model_candidates_for_target(target, metric=metric, price_min=price_min, price_max=price_max))
            if path == "/api/model/versions":
                query = urllib.parse.parse_qs(parsed.query)
                target = str((query.get("target") or [MODEL_TARGET_DAYAHEAD])[0] or MODEL_TARGET_DAYAHEAD)
                return self.send_json(self.get_versions(target))
            if path == "/api/training/logs":
                return self.send_json(self.get_training_logs())
            if path in {"/api/training/preferences", "/api/prediction/preferences"}:
                query = urllib.parse.parse_qs(parsed.query)
                target = str((query.get("target") or [MODEL_TARGET_DAYAHEAD])[0] or MODEL_TARGET_DAYAHEAD)
                return self.send_json({"preferences": load_target_training_preferences(target)})
            if path == "/api/data-quality/reports":
                return self.send_json({"reports": list_quality_reports()})
            if path == "/api/data-quality/reports/latest":
                return self.send_json({"report": latest_quality_report()})
            if path.startswith("/api/data-quality/reports/"):
                report_id = urllib.parse.unquote(path.rsplit("/", 1)[-1])
                return self.send_json({"report": load_quality_report(report_id)})
            if path == "/api/forecast/template":
                return self.send_json(load_forecast_template_preview())
            if path == "/api/forecast/thermal-capacity-preview":
                return self.send_json(load_thermal_capacity_prediction_preview())
            if path == "/api/forecast/multi-day-template":
                return self.send_json(load_multi_day_template_preview())
            if path == "/api/prediction/latest":
                return self.send_json({"prediction": load_latest_prediction_output()})
            if path == "/api/prediction-archive":
                query = urllib.parse.parse_qs(parsed.query)
                forecast_date = str((query.get("date") or [""])[0] or "").strip() or None
                return self.send_json({"records": list_prediction_archives(PREDICTION_ARCHIVE_ROOT, forecast_date)})
            if path.startswith("/api/prediction-archive/"):
                archive_id = urllib.parse.unquote(path.rsplit("/", 1)[-1])
                return self.send_json(
                    load_prediction_archive_detail(
                        PREDICTION_ARCHIVE_ROOT,
                        archive_id,
                        history_dir=HISTORY_DIR,
                        model_root=MODEL_ROOT,
                    )
                )
            if path.startswith("/api/jobs/"):
                job_id = path.rsplit("/", 1)[-1]
                snapshot = STATE.snapshot(job_id)
                if snapshot is None:
                    return self.send_error_json(HTTPStatus.NOT_FOUND, "未找到任务")
                return self.send_json(snapshot)
            if path == "/download/output":
                return self.send_file(OUTPUT_FILE, OUTPUT_FILE.name)
            if path == "/download/multi-day-output" or path == "/api/predict-multi-day/export":
                return self.send_file(MULTI_DAY_OUTPUT_FILE, MULTI_DAY_OUTPUT_FILE.name)
            return self.serve_static(path)
        except Exception as exc:  # noqa: BLE001
            return self.send_error_json(HTTPStatus.INTERNAL_SERVER_ERROR, str(exc))

    def do_POST(self) -> None:  # noqa: N802
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path.rstrip("/") or "/"
        try:
            payload = self.read_json_body()
            if path == "/api/jobs/cancel":
                job_id = str(payload.get("job_id") or "").strip()
                if not job_id:
                    return self.send_error_json(HTTPStatus.BAD_REQUEST, "缺少 job_id")
                try:
                    return self.send_json(STATE.request_cancel(job_id))
                except KeyError:
                    return self.send_error_json(HTTPStatus.NOT_FOUND, "未找到任务")
            if path.startswith("/api/jobs/") and path.endswith("/cancel"):
                job_id = urllib.parse.unquote(path.split("/")[-2])
                try:
                    return self.send_json(STATE.request_cancel(job_id))
                except KeyError:
                    return self.send_error_json(HTTPStatus.NOT_FOUND, "未找到任务")
            if path == "/api/train":
                return self.send_json(start_background_job("train", build_train_worker(payload)), HTTPStatus.ACCEPTED)
            if path == "/api/window-optimization/run":
                return self.send_json(check_window_optimization("manual", force=bool(payload.get("force", True)), payload=payload), HTTPStatus.ACCEPTED)
            if path == "/api/realtime-window-optimization/run":
                return self.send_json(check_realtime_window_optimization("manual", force=bool(payload.get("force", True)), payload=payload), HTTPStatus.ACCEPTED)
            if path == "/api/dual-window-optimization/run":
                return self.send_json(start_background_job("dual_optimize", build_dual_window_optimization_worker(payload)), HTTPStatus.ACCEPTED)
            if path == "/api/cleanup/output":
                return self.send_json(cleanup_output_junk())
            if path in {"/api/training/preferences", "/api/prediction/preferences"}:
                target = str(payload.get("target") or MODEL_TARGET_DAYAHEAD)
                preference_payload: dict[str, Any] = {}
                if path == "/api/training/preferences":
                    segment_config = parse_segment_config_payload(payload)
                    preference_payload.update(
                        {
                            "segment_mode": "custom" if segment_config else "default",
                            "segment_config": segment_config or DEFAULT_SEGMENT_CONFIG,
                            "high_price_weighting": parse_high_price_weighting_payload(payload),
                            "price_intervals": parse_price_intervals_payload(payload),
                        }
                    )
                if path == "/api/prediction/preferences":
                    preference_payload["prediction_reference"] = parse_prediction_reference_payload(payload)
                preferences = save_target_training_preferences(target, preference_payload)
                return self.send_json({"preferences": preferences})
            if path == "/api/predict":
                return self.send_json(start_background_job("predict", build_predict_worker(payload)), HTTPStatus.ACCEPTED)
            if path == "/api/predict-multi-day":
                return self.send_json(start_background_job("predict_multi_day", build_multi_day_predict_worker(payload)), HTTPStatus.ACCEPTED)
            if path == "/api/forecast/template/save":
                save_forecast_template_rows(payload.get("rows") or [])
                return self.send_json({"message": "预测文件已保存"})
            if path == "/api/forecast/multi-day-template/save":
                save_forecast_template_rows(payload.get("rows") or [], MULTI_DAY_FORECAST_FILE)
                return self.send_json({"message": "预测文件-N天已保存"})
            if path == "/api/model/activate":
                target = str(payload.get("target") or MODEL_TARGET_DAYAHEAD)
                version_key = str(payload.get("version_key") or "").strip()
                if not version_key:
                    return self.send_error_json(HTTPStatus.BAD_REQUEST, "缺少 version_key")
                activate_model_version_for_target(target, version_key)
                return self.send_json({"message": f"已切换默认模型版本：{version_key}"})
            if path == "/api/model/segment-selection":
                target = str(payload.get("target") or MODEL_TARGET_DAYAHEAD)
                metadata = update_segment_price_model_selection_for_target(target, payload.get("selected_segment_price_models") or {})
                return self.send_json({"message": "当前版本内分时段算法选择已保存", "metadata": metadata})
            if path == "/api/model/rollback":
                target = str(payload.get("target") or MODEL_TARGET_DAYAHEAD)
                current_dir = rollback_model_version_for_target(target)
                return self.send_json({"message": "已回退到上一版模型", "current_model_dir": str(current_dir)})
            if path == "/api/model/delete":
                target = str(payload.get("target") or MODEL_TARGET_DAYAHEAD)
                version_keys = [str(item).strip() for item in payload.get("version_keys") or [] if str(item).strip()]
                if not version_keys:
                    return self.send_error_json(HTTPStatus.BAD_REQUEST, "请选择要删除的模型版本")
                invalid = [key for key in version_keys if key in protected_versions(target)]
                if invalid:
                    return self.send_error_json(HTTPStatus.BAD_REQUEST, f"以下版本当前受保护，不能直接删除：{', '.join(invalid)}")
                deleted = delete_model_versions_for_target(target, version_keys)
                return self.send_json({"message": "删除完成", "deleted": deleted})
            return self.send_error_json(HTTPStatus.NOT_FOUND, "未找到接口")
        except RuntimeError as exc:
            return self.send_error_json(HTTPStatus.CONFLICT, str(exc))
        except ValueError as exc:
            return self.send_error_json(HTTPStatus.BAD_REQUEST, str(exc))
        except Exception as exc:  # noqa: BLE001
            return self.send_error_json(HTTPStatus.INTERNAL_SERVER_ERROR, str(exc))

    def log_message(self, format: str, *args: Any) -> None:
        return

    def get_config(self) -> dict[str, Any]:
        optimization_state = window_optimization_state()
        optimization_defaults = window_optimization_config_defaults(optimization_state)
        preferences = load_training_preferences(MODEL_ROOT)
        prediction_reference = normalize_prediction_reference_preferences(preferences.get("prediction_reference"))
        return {
            "app_name": "日前电价预测平台",
            "history_dir": str(HISTORY_DIR),
            "model_root": str(MODEL_ROOT),
            "realtime_model_root": str(REALTIME_MODEL_ROOT),
            "forecast_file": str(FORECAST_FILE),
            "output_file": str(OUTPUT_FILE),
            "frontend_dir": str(FRONTEND_APP_DIR),
            "price_floor": PRICE_FLOOR,
            "price_cap": PRICE_CAP,
            "default_reference_days": max(
                int(prediction_reference["recent_reference_days"]),
                int(prediction_reference["same_type_reference_days"]),
            ),
            "default_recent_reference_days": prediction_reference["recent_reference_days"],
            "default_same_type_reference_days": prediction_reference["same_type_reference_days"],
            "default_knn_similarity": prediction_reference["knn_similarity"],
            "default_training_mode": "rolling_window",
            "default_training_window_days": resolve_default_training_window_days(optimization_state),
            "window_optimization": optimization_state,
            "realtime_window_optimization": realtime_window_optimization_state(),
            "default_window_optimization_max_history_days": resolve_window_optimization_max_history_days(),
            "default_window_optimization_valid_days": optimization_defaults["valid_days"],
            "default_window_optimization_rolling_backtest_horizons": optimization_defaults["rolling_backtest_horizons"],
            "max_reference_days": 100,
            "prediction_preferences": preferences,
            "training_preferences": preferences,
            "default_segment_config": DEFAULT_SEGMENT_CONFIG,
            "reference_strategy_options": [
                {"key": "recent_n_days", "label": "最近 N 天"},
                {"key": "recent_same_type_days", "label": "最近 N 个同类型日"},
            ],
            "default_reference_strategy": "recent_n_days",
            "server_time": now_text(),
        }

    def get_status(self) -> dict[str, Any]:
        status = STATE.status()
        status["current_model"] = load_current_metadata(MODEL_ROOT) or {}
        status["current_realtime_model"] = load_current_metadata(REALTIME_MODEL_ROOT) or {}
        status["window_optimization"] = window_optimization_state()
        status["realtime_window_optimization"] = realtime_window_optimization_state()
        return status

    def get_versions(self, target: str | None = MODEL_TARGET_DAYAHEAD) -> dict[str, Any]:
        return {"versions": list_model_versions_for_target(target)}

    def get_training_logs(self) -> dict[str, Any]:
        log_df = load_training_log(MODEL_ROOT)
        if log_df.empty:
            return {"logs": []}
        log_df = log_df.sort_values("created_at", ascending=False).reset_index(drop=True)
        return {"logs": dataframe_to_records(log_df)}

    def read_json_body(self) -> dict[str, Any]:
        content_length = int(self.headers.get("Content-Length", "0") or 0)
        if content_length <= 0:
            return {}
        raw = self.rfile.read(content_length)
        if not raw:
            return {}
        return json.loads(raw.decode("utf-8"))

    def send_common_headers(self) -> None:
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Cache-Control", "no-store")

    def send_json(self, data: Any, status: int = HTTPStatus.OK) -> None:
        body = json.dumps(to_jsonable(data), ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_common_headers()
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def send_error_json(self, status: int, message: str) -> None:
        self.send_json({"error": message}, status=status)

    def send_file(self, file_path: Path, download_name: str | None = None) -> None:
        if not file_path.exists():
            return self.send_error_json(HTTPStatus.NOT_FOUND, f"文件不存在：{file_path}")
        content = file_path.read_bytes()
        content_type, _ = mimetypes.guess_type(str(file_path))
        self.send_response(HTTPStatus.OK)
        self.send_common_headers()
        self.send_header("Content-Type", content_type or "application/octet-stream")
        if download_name:
            self.send_header("Content-Disposition", f'attachment; filename="{download_name}"')
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)

    def serve_static(self, path: str) -> None:
        decoded = urllib.parse.unquote(path)
        if decoded in {"", "/"}:
            target = FRONTEND_APP_DIR / "index.html"
        else:
            target = (FRONTEND_APP_DIR / decoded.lstrip("/")).resolve()
            if FRONTEND_APP_DIR.resolve() not in [target, *target.parents]:
                return self.send_error_json(HTTPStatus.FORBIDDEN, "非法路径")
        if not target.exists() or target.is_dir():
            target = FRONTEND_APP_DIR / "index.html"
        content = target.read_bytes()
        content_type, _ = mimetypes.guess_type(str(target))
        self.send_response(HTTPStatus.OK)
        self.send_common_headers()
        self.send_header("Content-Type", f"{content_type or 'text/html'}; charset=utf-8")
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)


def build_server(host: str, port: int) -> ThreadingHTTPServer:
    return ThreadingHTTPServer((host, port), ApiHandler)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="日前电价预测平台 API 服务")
    parser.add_argument("--host", default=DEFAULT_HOST, help="监听地址")
    parser.add_argument("--port", type=int, default=DEFAULT_PORT, help="监听端口")
    parser.add_argument("--open-browser", action="store_true", help="启动后自动打开浏览器")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if not FRONTEND_APP_DIR.exists():
        raise FileNotFoundError(f"未找到前端目录：{FRONTEND_APP_DIR}")

    server = build_server(args.host, args.port)
    url = f"http://{args.host}:{args.port}"
    print(f"前端地址：{url}")
    print(f"历史数据目录：{HISTORY_DIR}")
    print(f"预测文件：{FORECAST_FILE}")
    print(f"模型目录：{MODEL_ROOT}")
    if args.open_browser:
        threading.Timer(1.0, lambda: webbrowser.open(url)).start()
    server.serve_forever()


if __name__ == "__main__":
    main()
