from __future__ import annotations

import json
import hashlib
import pickle
import gc
import re
import shutil
import time
import uuid
import warnings
from collections import OrderedDict
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable, Iterable

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from data_quality import (
    DataQualityIssue,
    DataQualityValidationError,
    has_blocking_issues,
    save_quality_report,
    validate_forecast_template,
    validate_history_sheet,
)
from price_interval import (
    DEFAULT_PRICE_INTERVALS,
    append_interval_prediction_columns,
    available_interval_backends,
    load_interval_model_bundle,
    normalize_price_intervals,
    predict_interval_probabilities,
    train_interval_models,
)


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_HISTORY_DIR = BASE_DIR / "数据"
DEFAULT_MODEL_ROOT = BASE_DIR / "models"
DEFAULT_REALTIME_MODEL_ROOT = DEFAULT_MODEL_ROOT / "realtime_price"
DEFAULT_OUTPUT_FILE = BASE_DIR / "output" / "dayahead_price_prediction.xlsx"
DEFAULT_FORECAST_FILE = BASE_DIR / "预测文件" / "预测文件.xlsx"
DEFAULT_MULTI_DAY_FORECAST_FILE = DEFAULT_FORECAST_FILE.parent / "预测文件-N天.xlsx"
DEFAULT_MULTI_DAY_OUTPUT_FILE = BASE_DIR / "output" / "multi_day_price_prediction.xlsx"
PRICE_FLOOR = 0.0
PRICE_CAP = 1500.0
HISTORY_CACHE_VERSION = 4

TARGET_COLUMN = "日前出清价格(元/MWh)"
REALTIME_TARGET_SOURCE_COLUMN = "实时出清价格(元/MWh)"
REALTIME_DATA_MODE = "realtime"
DAYAHEAD_DATA_MODE = "dayahead"
MODEL_TARGET_DAYAHEAD = "dayahead"
MODEL_TARGET_REALTIME = "realtime"
NET_LOAD_ONLY_SIMILAR_PRICE_COLUMN = "net_load_only_similar_price"
NET_LOAD_ONLY_SIMILAR_GAP_COLUMN = "net_load_only_similar_gap"
KNN_SIMILAR_PRICE_COLUMN = "knn_similar_price"
WEIGHTED_KNN_PRICE_COLUMN = "weighted_knn_regression_price"
KNN_SIMILAR_COUNT_COLUMN = "knn_similar_count"
KNN_SIMILAR_DISTANCE_COLUMN = "knn_similar_min_distance"
TOTAL_LOAD_COLUMN = "total_load"
RENEWABLE_POWER_COLUMN = "renewable_power"
DAY_TYPE_COLUMN = "day_type"
THERMAL_SPACE_WEIGHT_KEY = "thermal_space"
THERMAL_SPACE_LOAD_RATIO_COLUMN = "thermal_space_load_ratio"
THERMAL_CAPACITY_SOURCE_COLUMN = "thermal_capacity_source"
THERMAL_CAPACITY_VALUE_COLUMN = "thermal_capacity_value"
THERMAL_CAPACITY_MODEL_VALUE_COLUMN = "thermal_capacity_model_value"
THERMAL_CAPACITY_FILE_VALUE_COLUMN = "thermal_capacity_file_value"
DEFAULT_SIMILARITY_REFERENCE_DAYS = 100
DEFAULT_KNN_SIMILARITY_CONFIG = {
    "knn_k": 5,
    "knn_max_distance": 3.0,
    "weighted_knn_k": 5,
    "weighted_knn_max_distance": 3.0,
}


def normalize_model_target(target: str | None) -> str:
    normalized = str(target or MODEL_TARGET_DAYAHEAD).strip().lower()
    if normalized in {"", "dayahead", "day_ahead", "日前"}:
        return MODEL_TARGET_DAYAHEAD
    if normalized in {"realtime", "real_time", "实时"}:
        return MODEL_TARGET_REALTIME
    raise ValueError(f"未知模型目标: {target}")


def model_target_context(target: str | None, base_model_root: str | Path = DEFAULT_MODEL_ROOT) -> dict[str, object]:
    normalized = normalize_model_target(target)
    base_root = Path(base_model_root)
    if normalized == MODEL_TARGET_DAYAHEAD:
        model_root = base_root
    elif base_root.name == "realtime_price":
        model_root = base_root
    else:
        model_root = base_root / "realtime_price"
    return {
        "target": normalized,
        "data_mode": REALTIME_DATA_MODE if normalized == MODEL_TARGET_REALTIME else DAYAHEAD_DATA_MODE,
        "model_root": model_root,
        "target_column": REALTIME_TARGET_SOURCE_COLUMN if normalized == MODEL_TARGET_REALTIME else TARGET_COLUMN,
        "trains_capacity": normalized == MODEL_TARGET_DAYAHEAD,
    }


FEATURE_COLUMNS = [
    TOTAL_LOAD_COLUMN,
    "net_load",
    RENEWABLE_POWER_COLUMN,
    THERMAL_SPACE_LOAD_RATIO_COLUMN,
    "thermal_on_capacity",
    "hour",
    "period",
    "weekday",
    "month",
    "is_weekend",
    "is_holiday",
]
LAG_PRICE_COLUMNS = [
    "price_lag_1d",
    "price_lag_2d",
    "price_lag_3d",
    "price_lag_7d",
    "price_ma_3d",
    "price_ma_7d",
    "price_std_7d",
]
LAG_LOAD_COLUMNS = [
    "net_load_lag_1d",
]
LAG_FEATURE_COLUMNS = list(FEATURE_COLUMNS) + LAG_PRICE_COLUMNS + LAG_LOAD_COLUMNS
THERMAL_CAPACITY_FEATURE_COLUMNS = [
    "total_load_mean",
    "total_load_max",
    "total_load_min",
    "renewable_power_mean",
    "renewable_power_max",
    "renewable_power_min",
    "net_load_mean",
    "net_load_max",
    "net_load_min",
    "net_load_peak_valley",
    "net_load_max_ramp",
    "thermal_space_load_ratio_mean",
    "renewable_share_mean",
    "weekday",
    "month",
    "is_weekend",
    "is_holiday",
    "previous_thermal_on_capacity",
]
KNN_FEATURE_COLUMNS = [
    KNN_SIMILAR_PRICE_COLUMN,
    WEIGHTED_KNN_PRICE_COLUMN,
    KNN_SIMILAR_COUNT_COLUMN,
    KNN_SIMILAR_DISTANCE_COLUMN,
]
HIGH_PRICE_ADJUSTED_PRICE_COLUMN = "high_price_adjusted_price"
HIGH_PRICE_ADJUSTMENT_COLUMN = "high_price_adjustment"
HIGH_PRICE_ADJUSTMENT_REASON_COLUMN = "high_price_adjustment_reason"
SCENARIO_SIMILARITY_ADJUSTED_PRICE_COLUMN = "scenario_similarity_adjusted_price"
SIMILARITY_BLEND_WEIGHT_COLUMN = "similarity_blend_weight"
SIMILARITY_ADJUSTMENT_COLUMN = "similarity_adjustment"
SIMILARITY_ADJUSTMENT_REASON_COLUMN = "similarity_adjustment_reason"
REFERENCE_STRATEGIES = OrderedDict(
    [
        ("recent_n_days", "最近 N 天"),
        ("recent_same_type_days", "最近 N 个同类型日"),
    ]
)


def period_to_time(period_boundary: int) -> str:
    minutes = max(0, min(96, int(period_boundary) - 1)) * 15
    if int(period_boundary) >= 97:
        minutes = 24 * 60
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


SEGMENTS = OrderedDict(
    [
        ("segment_1", (1, 16)),
        ("segment_2", (17, 32)),
        ("segment_3", (33, 48)),
        ("segment_4", (49, 64)),
        ("segment_5", (65, 80)),
        ("segment_6", (81, 96)),
    ]
)
DEFAULT_SEGMENT_CONFIG = [
    {"name": name, "start": start, "end": end, "start_time": period_to_time(start), "end_time": period_to_time(end + 1)}
    for name, (start, end) in SEGMENTS.items()
]
DEFAULT_XGB_PARAMS = {
    "objective": "reg:squarederror",
    "eval_metric": "rmse",
    "eta": 0.05,
    "max_depth": 5,
    "min_child_weight": 3,
    "subsample": 0.85,
    "colsample_bytree": 0.85,
    "seed": 42,
}
MODEL_BACKENDS = OrderedDict(
    [
        ("xgboost", "XGBoost"),
        ("lightgbm", "LightGBM"),
        ("catboost", "CatBoost"),
    ]
)


@dataclass
class TrainConfig:
    history_dir: Path = DEFAULT_HISTORY_DIR
    model_root: Path = DEFAULT_MODEL_ROOT
    holiday_file: Path | None = None
    valid_days: int = 14
    training_window_days: int | None = None
    num_boost_round: int = 400
    start_date: str | None = None
    end_date: str | None = None
    interval_valid_days: int | None = None
    interval_training_window_days: int | None = None
    interval_num_boost_round: int | None = None
    interval_start_date: str | None = None
    interval_end_date: str | None = None
    thermal_capacity_valid_days: int | None = None
    thermal_capacity_training_window_days: int | None = None
    thermal_capacity_num_boost_round: int | None = None
    thermal_capacity_start_date: str | None = None
    thermal_capacity_end_date: str | None = None
    similarity_reference_days: int = DEFAULT_SIMILARITY_REFERENCE_DAYS
    enable_rolling_backtest: bool = True
    rolling_backtest_horizons: tuple[int, ...] = (14, 30)
    segment_config: list[dict[str, object]] | None = None
    high_price_weight_enabled: bool = False
    high_price_quantile: float = 0.8
    high_price_weight_multiplier: float = 2.0
    price_intervals: list[dict[str, object]] | None = None
    direct_feature_variant_keys: tuple[str, ...] | None = None
    model_backend_keys: tuple[str, ...] | None = None
    interval_backend_keys: tuple[str, ...] | None = None
    thermal_capacity_backend_keys: tuple[str, ...] | None = None
    train_interval_model: bool = True
    train_thermal_capacity_model: bool = True
    feature_cache_root: Path | None = None
    data_mode: str = DAYAHEAD_DATA_MODE


@dataclass
class TrainResult:
    run_id: str
    current_model_dir: Path
    history_model_dir: Path
    metrics: dict[str, dict[str, float | int | None]]
    train_dates: list[str]
    valid_dates: list[str]
    skipped_sheets: list[str]
    metadata_path: Path
    quality_report_path: Path | None = None


@dataclass
class PredictResult:
    forecast_date: str
    output_file: Path
    template_updated: bool
    reference_strategy_key: str
    reference_strategy_label: str
    reference_days_requested: int
    reference_dates: list[str]
    result_df: pd.DataFrame
    quality_report_path: Path | None = None


@dataclass
class PredictCompareResult:
    forecast_date: str
    output_file: Path
    template_updated: bool
    selected_strategy_key: str
    selected_strategy_label: str
    strategy_results: dict[str, PredictResult]
    quality_report_path: Path | None = None


@dataclass
class MultiDayPredictResult:
    baseline_date: str
    forecast_dates: list[str]
    output_file: Path
    template_updated: bool
    selected_strategy_key: str
    selected_strategy_label: str
    days: list[dict[str, object]]
    quality_report_path: Path | None = None


ProgressCallback = Callable[[str, int | None], None]


def ensure_xgboost_available() -> None:
    try:
        import xgboost  # noqa: F401
    except ModuleNotFoundError as exc:
        raise ModuleNotFoundError("缺少 xgboost，请先执行: pip install xgboost") from exc


def emit_progress(callback: ProgressCallback | None, message: str, percent: int | None = None) -> None:
    if callback is not None:
        callback(message, percent)


def normalize_reference_days(reference_days: int | None) -> int:
    value = int(reference_days or 1)
    if value <= 0:
        raise ValueError("reference_days 必须大于等于 1")
    return value



def normalize_reference_days_by_strategy(
    reference_days: int | None = None,
    recent_reference_days: int | None = None,
    same_type_reference_days: int | None = None,
) -> dict[str, int]:
    fallback = normalize_reference_days(reference_days)
    return {
        "recent_n_days": normalize_reference_days(recent_reference_days or fallback),
        "recent_same_type_days": normalize_reference_days(same_type_reference_days or fallback),
    }


def normalize_knn_similarity_config(config: dict[str, object] | None = None) -> dict[str, float | int]:
    config = config or {}

    def positive_int(key: str) -> int:
        default = int(DEFAULT_KNN_SIMILARITY_CONFIG[key])
        try:
            value = int(round(float(config.get(key, default))))
        except (TypeError, ValueError):
            return default
        return max(1, min(value, 20))

    def positive_float(key: str) -> float:
        default = float(DEFAULT_KNN_SIMILARITY_CONFIG[key])
        try:
            value = float(config.get(key, default))
        except (TypeError, ValueError):
            return default
        if not np.isfinite(value) or value <= 0:
            return default
        return max(0.1, min(value, 20.0))

    return {
        "knn_k": positive_int("knn_k"),
        "knn_max_distance": positive_float("knn_max_distance"),
        "weighted_knn_k": positive_int("weighted_knn_k"),
        "weighted_knn_max_distance": positive_float("weighted_knn_max_distance"),
    }


def normalize_reference_strategy(reference_strategy: str | None) -> str:
    strategy_key = str(reference_strategy or "recent_n_days").strip() or "recent_n_days"
    if strategy_key not in REFERENCE_STRATEGIES:
        raise ValueError(f"不支持的参考日策略: {strategy_key}")
    return strategy_key


def time_text_to_minutes(value: object) -> int:
    text = str(value or "").strip()
    match = re.fullmatch(r"(\d{1,2}):([0-5]\d)", text)
    if not match:
        raise ValueError(f"时段时间格式不正确：{text}")
    hour = int(match.group(1))
    minute = int(match.group(2))
    if hour == 24 and minute == 0:
        return 24 * 60
    if hour < 0 or hour > 23 or minute % 15 != 0:
        raise ValueError(f"时段时间必须在 00:00~24:00 且按 15 分钟粒度：{text}")
    return hour * 60 + minute


def minutes_to_time_text(minutes: int) -> str:
    minutes = max(0, min(24 * 60, int(minutes)))
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


def normalize_segment_config(segment_config: list[dict[str, object]] | None = None) -> OrderedDict[str, tuple[int, int]]:
    if not segment_config:
        return OrderedDict((name, (start, end)) for name, (start, end) in SEGMENTS.items())
    normalized: OrderedDict[str, tuple[int, int]] = OrderedDict()
    expected_start = 0
    seen_names: set[str] = set()
    if len(segment_config) > 8:
        raise ValueError("自定义时段最多支持 8 段，避免每段样本过少")
    for index, item in enumerate(segment_config, start=1):
        name = str(item.get("name") or f"segment_{index}").strip()
        safe_name = re.sub(r"[^0-9A-Za-z_\u4e00-\u9fff-]+", "_", name).strip("_") or f"segment_{index}"
        if safe_name in seen_names:
            raise ValueError(f"时段名称不能重复：{name}")
        seen_names.add(safe_name)
        if item.get("start_time") is not None:
            start_minutes = time_text_to_minutes(item.get("start_time"))
        elif item.get("start") is not None and str(item.get("start")).strip().isdigit():
            start_minutes = (int(item.get("start")) - 1) * 15
        else:
            start_minutes = 0
        if item.get("end_time") is not None:
            end_minutes = time_text_to_minutes(item.get("end_time"))
        elif item.get("end") is not None and str(item.get("end")).strip().isdigit():
            end_minutes = int(item.get("end")) * 15
        else:
            end_minutes = 24 * 60
        if start_minutes != expected_start:
            raise ValueError("自定义时段必须连续覆盖 00:00~24:00，不能空缺或重叠")
        if end_minutes <= start_minutes:
            raise ValueError(f"时段 {name} 的结束时间必须晚于开始时间")
        if end_minutes - start_minutes < 60:
            raise ValueError(f"时段 {name} 至少需要 1 小时，避免训练样本过少")
        normalized[safe_name] = (start_minutes // 15 + 1, end_minutes // 15)
        expected_start = end_minutes
    if expected_start != 24 * 60:
        raise ValueError("自定义时段必须覆盖到 24:00")
    return normalized


def segment_metadata(segment_definitions: OrderedDict[str, tuple[int, int]]) -> list[dict[str, object]]:
    return [
        {
            "name": name,
            "start": int(start),
            "end": int(end),
            "start_time": minutes_to_time_text((int(start) - 1) * 15),
            "end_time": minutes_to_time_text(int(end) * 15),
        }
        for name, (start, end) in segment_definitions.items()
    ]


def training_feature_columns(config: TrainConfig | None = None) -> list[str]:
    return list(FEATURE_COLUMNS)


def direct_model_variants() -> OrderedDict[str, list[str]]:
    return OrderedDict(
        [
            ("no_lag_96", training_feature_columns()),
            ("with_lag_96", list(LAG_FEATURE_COLUMNS)),
        ]
    )


def metadata_has_direct_variants(metadata: dict) -> bool:
    return isinstance(metadata.get("model_variants"), dict) and bool(metadata.get("model_variants"))


def select_prediction_model_variant(metadata: dict, forecast_df: pd.DataFrame) -> str:
    variants = metadata.get("model_variants") or {}
    selected_model_key = metadata.get("selected_model_key")
    if selected_model_key in variants:
        return str(selected_model_key)
    if "no_lag_96" in variants:
        return "no_lag_96"
    if "direct_price" in variants:
        return "direct_price"
    allowed_features = set(LAG_FEATURE_COLUMNS)
    for key, value in variants.items():
        feature_columns = value.get("feature_columns") or []
        if set(feature_columns).issubset(allowed_features):
            return key
    raise ValueError("当前模型不是多模型直接预测价格版本，请重新训练模型")


def select_segment_prediction_model_variant(metadata: dict, segment_name: str, forecast_df: pd.DataFrame) -> str:
    segment_models = metadata.get("selected_segment_price_models")
    variants = metadata.get("model_variants") or {}
    if isinstance(segment_models, dict):
        segment_key = str(segment_models.get(segment_name) or "").strip()
        if segment_key in variants:
            return segment_key
    return select_prediction_model_variant(metadata, forecast_df)


def normalize_text(text: object) -> str:
    return re.sub(r"\s+", "", str(text)).strip().lower()


def list_excel_files(source: str | Path) -> list[Path]:
    path = Path(source)
    if path.is_file():
        return [path]
    files = sorted(
        [
            file
            for file in [*path.rglob("*.xlsx"), *path.rglob("*.xls")]
            if not file.name.startswith("~$")
        ]
    )
    if not files:
        raise FileNotFoundError(f"未在 {path} 下找到 Excel 文件")
    return files


def parse_optional_date(value: str | Path | pd.Timestamp | None) -> pd.Timestamp | None:
    if value is None:
        return None
    timestamp = pd.to_datetime(value, errors="coerce")
    if pd.isna(timestamp):
        return None
    return pd.Timestamp(timestamp).normalize()


def excel_file_month(file_path: str | Path) -> pd.Timestamp | None:
    path = Path(file_path)
    candidates = [path.stem, path.name, *[part for part in path.parts[-3:]]]
    for text in candidates:
        standard_match = re.search(r"(20\d{2})[-_/\.](1[0-2]|0?[1-9])", str(text))
        if standard_match:
            return pd.Timestamp(year=int(standard_match.group(1)), month=int(standard_match.group(2)), day=1)
        match = re.search(r"(20\d{2})\s*年?\D{0,8}(1[0-2]|0?[1-9])\s*月?", str(text))
        if match:
            return pd.Timestamp(year=int(match.group(1)), month=int(match.group(2)), day=1)
    year_value = None
    month_value = None
    for part in reversed(path.parts):
        year_match = re.search(r"(20\d{2})\s*年", part)
        month_match = re.search(r"(1[0-2]|0?[1-9])\s*月", part)
        if year_match and year_value is None:
            year_value = int(year_match.group(1))
        if month_match and month_value is None:
            month_value = int(month_match.group(1))
        if year_value is not None and month_value is not None:
            return pd.Timestamp(year=year_value, month=month_value, day=1)
    return None


def list_excel_files_for_date_window(
    source: str | Path,
    start_date: str | pd.Timestamp | None,
    end_date: str | pd.Timestamp | None,
    leading_days: int = 0,
) -> list[Path]:
    path = Path(source)
    files = list_excel_files(path)
    if path.is_file():
        return files
    window_start = parse_optional_date(start_date)
    window_end = parse_optional_date(end_date)
    if window_start is None and window_end is None:
        return files
    if window_start is not None and leading_days > 0:
        window_start = window_start - pd.Timedelta(days=leading_days)
    selected: list[Path] = []
    unknown: list[Path] = []
    for file_path in files:
        month_start = excel_file_month(file_path)
        if month_start is None:
            unknown.append(file_path)
            continue
        month_end = month_start + pd.offsets.MonthEnd(0)
        if window_start is not None and month_end < window_start:
            continue
        if window_end is not None and month_start > window_end:
            continue
        selected.append(file_path)
    scoped_files = [*selected, *unknown]
    if not scoped_files:
        raise FileNotFoundError(f"未找到覆盖日期范围的历史 Excel 文件: {source}")
    return sorted(scoped_files)


def excel_file_signature(
    history_dir: str | Path,
    start_date: str | pd.Timestamp | None = None,
    end_date: str | pd.Timestamp | None = None,
    leading_days: int = 0,
    excel_files: list[Path] | None = None,
) -> list[dict[str, object]]:
    files = excel_files if excel_files is not None else list_excel_files_for_date_window(history_dir, start_date, end_date, leading_days)
    return [
        {
            "path": str(path.resolve()),
            "size": path.stat().st_size,
            "mtime_ns": path.stat().st_mtime_ns,
        }
        for path in files
    ]


def build_history_source_signature(
    source: str | Path,
    holiday_file: str | Path | None,
    require_target: bool,
    excel_files: list[Path] | None = None,
    scope: dict | None = None,
) -> dict:
    source_path = Path(source).resolve()
    signature_files = excel_files if excel_files is not None else list_excel_files(source_path)
    files = excel_file_signature(source_path, excel_files=signature_files)
    holiday_signature = None
    if holiday_file:
        holiday_path = Path(holiday_file).resolve()
        holiday_stat = holiday_path.stat()
        holiday_signature = {
            "path": str(holiday_path),
            "size": holiday_stat.st_size,
            "mtime_ns": holiday_stat.st_mtime_ns,
        }
    payload = {
        "version": HISTORY_CACHE_VERSION,
        "source": str(source_path),
        "require_target": bool(require_target),
        "files": files,
        "holiday_file": holiday_signature,
        "scope": scope or {},
    }
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True)
    payload["cache_key"] = hashlib.sha256(raw.encode("utf-8")).hexdigest()
    return payload


def history_cache_signature(
    history_dir: str | Path,
    holiday_file: str | Path | None = None,
    require_target: bool = True,
    start_date: str | pd.Timestamp | None = None,
    end_date: str | pd.Timestamp | None = None,
    leading_days: int = 0,
    data_mode: str = DAYAHEAD_DATA_MODE,
    excel_files: list[Path] | None = None,
) -> dict:
    scoped_files = excel_files if excel_files is not None else list_excel_files_for_date_window(history_dir, start_date, end_date, leading_days)
    scope = {
        "start_date": str(parse_optional_date(start_date).date()) if parse_optional_date(start_date) is not None else None,
        "end_date": str(parse_optional_date(end_date).date()) if parse_optional_date(end_date) is not None else None,
        "leading_days": int(leading_days or 0),
    }
    signature = build_history_source_signature(history_dir, holiday_file, require_target, excel_files=scoped_files, scope=scope)
    normalized_data_mode = str(data_mode or DAYAHEAD_DATA_MODE)
    signature["excel_files"] = signature["files"]
    if normalized_data_mode != DAYAHEAD_DATA_MODE:
        signature["cache_key"] = derived_cache_key(
            signature["cache_key"],
            json.dumps({"history_data_mode": normalized_data_mode}, ensure_ascii=False, sort_keys=True),
        )
    signature["history_data_mode"] = normalized_data_mode
    return signature


def derived_cache_key(source_key: str, namespace: str) -> str:
    raw = json.dumps(
        {
            "version": HISTORY_CACHE_VERSION,
            "source_key": source_key,
            "namespace": namespace,
        },
        sort_keys=True,
    )
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def training_run_fingerprint(config: TrainConfig, history_signature: dict[str, object]) -> str:
    payload = {
        "history": history_signature,
        "data_mode": config.data_mode,
        "training_window_days": config.training_window_days,
        "valid_days": config.valid_days,
        "num_boost_round": config.num_boost_round,
        "start_date": config.start_date,
        "end_date": config.end_date,
        "interval_training_window_days": config.interval_training_window_days,
        "interval_valid_days": config.interval_valid_days,
        "interval_num_boost_round": config.interval_num_boost_round,
        "interval_start_date": config.interval_start_date,
        "interval_end_date": config.interval_end_date,
        "thermal_capacity_training_window_days": config.thermal_capacity_training_window_days,
        "thermal_capacity_valid_days": config.thermal_capacity_valid_days,
        "thermal_capacity_num_boost_round": config.thermal_capacity_num_boost_round,
        "thermal_capacity_start_date": config.thermal_capacity_start_date,
        "thermal_capacity_end_date": config.thermal_capacity_end_date,
        "similarity_reference_days": config.similarity_reference_days,
        "rolling_backtest_horizons": config.rolling_backtest_horizons,
        "segment_config": config.segment_config,
        "high_price_weight_enabled": config.high_price_weight_enabled,
        "high_price_quantile": config.high_price_quantile,
        "high_price_weight_multiplier": config.high_price_weight_multiplier,
        "price_intervals": config.price_intervals,
        "direct_feature_variant_keys": config.direct_feature_variant_keys,
        "model_backend_keys": config.model_backend_keys,
        "interval_backend_keys": config.interval_backend_keys,
        "thermal_capacity_backend_keys": config.thermal_capacity_backend_keys,
        "train_interval_model": config.train_interval_model,
        "train_thermal_capacity_model": config.train_thermal_capacity_model,
    }
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def dataframe_cache_path(cache_dir: str | Path, cache_name: str, cache_key: str) -> Path:
    return Path(cache_dir) / f"{cache_name}_{cache_key}.pkl"


def missing_dataframe_columns(frame: pd.DataFrame, required_columns: list[str] | tuple[str, ...] | set[str]) -> list[str]:
    return [column for column in required_columns if column not in frame.columns]


def load_dataframe_cache(cache_dir: str | Path, cache_name: str, cache_key: str) -> tuple[pd.DataFrame, list[str]] | None:
    cache_path = dataframe_cache_path(cache_dir, cache_name, cache_key)
    if not cache_path.exists():
        return None
    try:
        with cache_path.open("rb") as handle:
            payload = pickle.load(handle)
    except (OSError, pickle.PickleError, EOFError, AttributeError, ValueError):
        return None
    if not isinstance(payload, dict):
        return None
    if payload.get("version") != HISTORY_CACHE_VERSION or payload.get("cache_key") != cache_key:
        return None
    frame = payload.get("frame")
    if not isinstance(frame, pd.DataFrame):
        return None
    skipped_sheets = payload.get("skipped_sheets") or []
    if not isinstance(skipped_sheets, list):
        skipped_sheets = []
    return frame.copy(), [str(item) for item in skipped_sheets]


def save_dataframe_cache(cache_dir: str | Path, cache_name: str, cache_key: str, frame: pd.DataFrame, skipped_sheets: list[str]) -> None:
    cache_root = Path(cache_dir)
    cache_root.mkdir(parents=True, exist_ok=True)
    cache_path = dataframe_cache_path(cache_root, cache_name, cache_key)
    temp_path = cache_path.with_suffix(".tmp")
    payload = {
        "version": HISTORY_CACHE_VERSION,
        "cache_key": cache_key,
        "frame": frame,
        "skipped_sheets": skipped_sheets,
    }
    with temp_path.open("wb") as handle:
        pickle.dump(payload, handle, protocol=pickle.HIGHEST_PROTOCOL)
    temp_path.replace(cache_path)


def required_raw_history_columns(require_target: bool) -> list[str]:
    columns = [
        "date",
        "period",
        TOTAL_LOAD_COLUMN,
        "net_load",
        RENEWABLE_POWER_COLUMN,
        THERMAL_SPACE_LOAD_RATIO_COLUMN,
        "thermal_on_capacity",
        "segment",
    ]
    if require_target:
        columns.append(TARGET_COLUMN)
    return columns


def required_training_feature_cache_columns() -> list[str]:
    return sorted(set(LAG_FEATURE_COLUMNS + [TARGET_COLUMN, "date", "period", "segment"]))


def quality_report_status(issues: list[DataQualityIssue], blocked: bool = False) -> str:
    if blocked:
        return "blocked"
    if issues:
        return "completed_with_issues"
    return "clean"


def load_cached_history_collection(
    history_dir: str | Path,
    holiday_file: str | Path | None,
    require_target: bool,
    model_root: str | Path,
    progress_callback: ProgressCallback | None = None,
    progress_start: int | None = None,
    progress_end: int | None = None,
    progress_label: str = "正在加载历史数据",
    start_date: str | pd.Timestamp | None = None,
    end_date: str | pd.Timestamp | None = None,
    leading_days: int = 0,
    cache_dir: str | Path | None = None,
    data_mode: str = DAYAHEAD_DATA_MODE,
) -> tuple[pd.DataFrame, list[str], set[pd.Timestamp], dict, list[DataQualityIssue]]:
    holiday_dates = load_holiday_calendar(holiday_file)
    scoped_files = list_excel_files_for_date_window(history_dir, start_date, end_date, leading_days=leading_days)
    normalized_data_mode = str(data_mode or DAYAHEAD_DATA_MODE)
    signature = history_cache_signature(
        history_dir,
        holiday_file=holiday_file,
        require_target=require_target,
        start_date=start_date,
        end_date=end_date,
        leading_days=leading_days,
        data_mode=normalized_data_mode,
        excel_files=scoped_files,
    )
    history_cache_dir = Path(cache_dir) if cache_dir is not None else Path(model_root) / "cache"
    cache_name = "history_raw" if normalized_data_mode == DAYAHEAD_DATA_MODE else f"history_raw_{normalized_data_mode}"
    cached = load_dataframe_cache(history_cache_dir, cache_name, signature["cache_key"])
    if cached is not None:
        cached_frame, skipped_sheets = cached
        missing_columns = missing_dataframe_columns(cached_frame, required_raw_history_columns(require_target))
        if not missing_columns:
            emit_progress(progress_callback, "历史数据未变化，已复用缓存", progress_end)
            return cached_frame, skipped_sheets, holiday_dates, signature, []
        emit_progress(progress_callback, f"历史缓存缺少字段 {missing_columns}，正在重新读取历史数据", progress_start)

    builder = DayAheadDataBuilder(holiday_dates)
    try:
        history_df = builder.load_excel_collection(
            history_dir,
            require_target=require_target,
            progress_callback=progress_callback,
            progress_start=progress_start,
            progress_end=progress_end,
            progress_label=progress_label,
            excel_files=scoped_files,
            data_mode=normalized_data_mode,
        )
    except ValueError as exc:
        if builder.quality_issues:
            raise DataQualityValidationError(str(exc), builder.quality_issues) from exc
        raise
    save_dataframe_cache(history_cache_dir, cache_name, signature["cache_key"], history_df, builder.skipped_sheets)
    return history_df, builder.skipped_sheets, holiday_dates, signature, builder.quality_issues


def to_numeric(series: pd.Series) -> pd.Series:
    cleaned = (
        series.astype(str)
        .str.replace(",", "", regex=False)
        .str.replace("，", "", regex=False)
        .str.strip()
        .replace({"": np.nan, "nan": np.nan, "None": np.nan, "--": np.nan, "-": np.nan})
    )
    return pd.to_numeric(cleaned, errors="coerce")


def calculate_thermal_space_load_ratio(thermal_space: pd.Series, total_load: pd.Series) -> pd.Series:
    ratio = pd.to_numeric(thermal_space, errors="coerce") / pd.to_numeric(total_load, errors="coerce").replace(0, np.nan)
    return ratio.replace([np.inf, -np.inf], np.nan)


def load_holiday_calendar(holiday_file: str | Path | None) -> dict[pd.Timestamp, str]:
    if not holiday_file:
        return {}
    holiday_path = Path(holiday_file)
    if not holiday_path.exists():
        raise FileNotFoundError(f"节假日文件不存在: {holiday_path}")
    if holiday_path.suffix.lower() == ".csv":
        holiday_df = pd.read_csv(holiday_path)
    else:
        holiday_df = pd.read_excel(holiday_path)
    if holiday_df.empty:
        return {}
    date_series = holiday_df["date"] if "date" in holiday_df.columns else holiday_df.iloc[:, 0]
    parsed_dates = pd.to_datetime(date_series, errors="coerce").dt.normalize()
    if "type" in holiday_df.columns:
        type_series = holiday_df["type"].astype(str).str.strip().str.lower()
    else:
        type_series = pd.Series(["holiday"] * len(holiday_df), index=holiday_df.index)
    calendar: dict[pd.Timestamp, str] = {}
    for date_value, type_value in zip(parsed_dates, type_series):
        if pd.isna(date_value):
            continue
        normalized_type = str(type_value).strip().lower()
        if normalized_type in {"holiday", "节假日", "法定节假日", "休息日"}:
            calendar[pd.Timestamp(date_value)] = "holiday"
        elif normalized_type in {"workday", "working_day", "调休上班", "工作日"}:
            calendar[pd.Timestamp(date_value)] = "workday"
    return calendar


def load_holiday_dates(holiday_file: str | Path | None) -> set[pd.Timestamp]:
    return {date_value for date_value, day_type in load_holiday_calendar(holiday_file).items() if day_type == "holiday"}


def day_type_of(date_value: pd.Timestamp, holiday_dates: set[pd.Timestamp] | dict[pd.Timestamp, str]) -> str:
    normalized = pd.Timestamp(date_value).normalize()
    if isinstance(holiday_dates, dict):
        override = holiday_dates.get(normalized)
        if override in {"holiday", "workday"}:
            return override
    elif normalized in holiday_dates:
        return "holiday"
    if normalized.weekday() >= 5:
        return "weekend"
    return "workday"


def resolve_column(columns: Iterable[object], candidates: list[str], required: bool) -> str | None:
    normalized = {normalize_text(column): str(column) for column in columns}
    for candidate in candidates:
        key = normalize_text(candidate)
        if key in normalized:
            return normalized[key]
    for candidate in candidates:
        key = normalize_text(candidate)
        for normalized_name, original_name in normalized.items():
            if key in normalized_name or normalized_name in key:
                return original_name
    if required:
        raise KeyError(f"未找到字段: {candidates}")
    return None


def parse_sheet_date(sheet_name: str) -> pd.Timestamp | None:
    match = re.search(r"(\d{4})\D+(\d{1,2})\D+(\d{1,2})", sheet_name)
    if match:
        year, month, day = map(int, match.groups())
        return pd.Timestamp(year=year, month=month, day=day).normalize()
    parsed = pd.to_datetime(sheet_name, errors="coerce")
    if pd.isna(parsed):
        return None
    return parsed.normalize()


def infer_date_from_header(columns: Iterable[object], default_year: int) -> pd.Timestamp | None:
    for column in columns:
        text = str(column).strip()
        if not text or normalize_text(text).startswith("unnamed:"):
            continue
        match = re.search(r"(?<!\d)(\d{1,2})[./月-](\d{1,2})日?(?!\d)", text)
        if not match:
            continue
        month, day = map(int, match.groups())
        return pd.Timestamp(year=default_year, month=month, day=day).normalize()
    return None


def extract_month_day(text: object) -> tuple[int, int] | None:
    match = re.search(r"(?<!\d)(\d{1,2})[./月-](\d{1,2})日?(?!\d)", str(text).strip())
    if not match:
        return None
    return int(match.group(1)), int(match.group(2))


def convert_capacity_to_mw(value: float, unit: str | None) -> float:
    normalized_unit = (unit or "MW").lower()
    if normalized_unit == "gw":
        return value * 1000
    if normalized_unit in {"kw", "千瓦"}:
        return value / 1000
    if normalized_unit in {"万千瓦", "万kw", "万kw", "万kW".lower()}:
        return value * 10
    return value


def extract_thermal_on_capacity(text: object) -> float:
    if pd.isna(text):
        return np.nan
    content = str(text)
    patterns = [
        r"火电开机容量[^0-9]*([0-9]+(?:\.[0-9]+)?)\s*(万千瓦|万kw|万kW|MW|mw|GW|gw|kW|kw|KW)?",
        r"运行机组容量[^0-9]*([0-9]+(?:\.[0-9]+)?)\s*(万千瓦|万kw|万kW|MW|mw|GW|gw|kW|kw|KW)?",
    ]
    for pattern in patterns:
        match = re.search(pattern, content)
        if not match:
            continue
        return convert_capacity_to_mw(float(match.group(1)), match.group(2))
    return np.nan


def assign_segments(periods: pd.Series, segment_definitions: OrderedDict[str, tuple[int, int]] | None = None) -> pd.Series:
    segment_definitions = segment_definitions or normalize_segment_config()
    segment = pd.Series(index=periods.index, dtype="object")
    for name, (start, end) in segment_definitions.items():
        segment.loc[periods.between(start, end)] = name
    if segment.isna().any():
        invalid = periods[segment.isna()].tolist()
        raise ValueError(f"存在无效 period: {invalid[:5]}")
    return segment


class DayAheadDataBuilder:
    def __init__(self, holiday_dates: set[pd.Timestamp] | None = None) -> None:
        self.holiday_dates = holiday_dates or set()
        self.skipped_sheets: list[str] = []
        self.quality_issues: list[DataQualityIssue] = []

    def load_excel_collection(
        self,
        source: str | Path,
        require_target: bool,
        progress_callback: ProgressCallback | None = None,
        progress_start: int | None = None,
        progress_end: int | None = None,
        progress_label: str = "正在读取数据",
        excel_files: list[Path] | None = None,
        data_mode: str = DAYAHEAD_DATA_MODE,
    ) -> pd.DataFrame:
        all_days: list[pd.DataFrame] = []
        self.skipped_sheets = []
        self.quality_issues = []
        excel_files = excel_files if excel_files is not None else list_excel_files(source)
        total_files = len(excel_files)
        for file_index, file_path in enumerate(excel_files, start=1):
            if progress_callback is not None and progress_start is not None and progress_end is not None and total_files > 0:
                percent = progress_start + int((file_index - 1) / total_files * max(progress_end - progress_start, 1))
                emit_progress(progress_callback, f"{progress_label} {file_index}/{total_files}: {file_path.name}", percent)
            workbook = pd.read_excel(file_path, sheet_name=None)
            for sheet_name, raw_df in workbook.items():
                trade_date = parse_sheet_date(str(sheet_name))
                if trade_date is None:
                    self.quality_issues.append(
                        DataQualityIssue(
                            task_type="train",
                            severity="error",
                            action="skipped",
                            issue_type="unrecognized_sheet_date",
                            file_path=str(file_path),
                            sheet_name=str(sheet_name),
                            message="Sheet 名称无法识别为日期",
                        )
                    )
                    if require_target:
                        self.skipped_sheets.append(f"{file_path.name} | {sheet_name} | Sheet 名称无法识别为日期")
                    continue
                validation_issues = validate_history_sheet(
                    raw_df,
                    file_path=file_path,
                    sheet_name=str(sheet_name),
                    trade_date=trade_date,
                    require_target=require_target,
                )
                self.quality_issues.extend(validation_issues)
                if has_blocking_issues(validation_issues):
                    if require_target:
                        first_issue = validation_issues[0]
                        self.skipped_sheets.append(f"{file_path.name} | {sheet_name} | {first_issue.message}")
                        continue
                    raise DataQualityValidationError("历史数据存在关键字段异常", validation_issues)
                try:
                    day_df = self.prepare_single_sheet(
                        raw_df,
                        file_path,
                        str(sheet_name),
                        trade_date,
                        require_target,
                        data_mode=data_mode,
                    )
                    all_days.append(day_df)
                except (KeyError, ValueError) as exc:
                    self.quality_issues.append(
                        DataQualityIssue(
                            task_type="train",
                            severity="error",
                            action="skipped" if require_target else "blocked",
                            issue_type="parse_error",
                            file_path=str(file_path),
                            sheet_name=str(sheet_name),
                            date=trade_date.strftime("%Y-%m-%d"),
                            message=str(exc),
                        )
                    )
                    if require_target:
                        self.skipped_sheets.append(f"{file_path.name} | {sheet_name} | {exc}")
                        continue
                    raise
        if not all_days:
            raise ValueError(f"未解析到可用的日级 Sheet: {source}")
        if self.skipped_sheets:
            preview = "\n".join(f"  - {item}" for item in self.skipped_sheets[:10])
            warnings.warn(
                f"训练数据中跳过了 {len(self.skipped_sheets)} 个异常 Sheet，这些 Sheet 缺少关键字段或数据不完整：\n{preview}",
                stacklevel=2,
            )
        return pd.concat(all_days, ignore_index=True).sort_values(["date", "period"]).reset_index(drop=True)

    def prepare_single_sheet(
        self,
        raw_df: pd.DataFrame,
        file_path: Path,
        sheet_name: str,
        trade_date: pd.Timestamp,
        require_target: bool,
        data_mode: str = DAYAHEAD_DATA_MODE,
    ) -> pd.DataFrame:
        day_df = raw_df.copy().iloc[:96].reset_index(drop=True)
        if len(day_df) < 96:
            raise ValueError(f"{file_path.name} - {sheet_name} 不是 96 行时段数据")
        day_df.columns = [str(column).strip() for column in day_df.columns]

        normalized_data_mode = str(data_mode or DAYAHEAD_DATA_MODE)
        if normalized_data_mode == REALTIME_DATA_MODE:
            total_power_col = resolve_column(day_df.columns, ["用电负荷(MW)", "用电负荷"], required=True)
            power_col = resolve_column(day_df.columns, ["新能源总计划(MW)", "新能源总计划"], required=True)
            thermal_power_col = resolve_column(day_df.columns, ["火电总计划(MW)", "火电总计划"], required=True)
            target_col = resolve_column(
                day_df.columns,
                [REALTIME_TARGET_SOURCE_COLUMN, "实时-出清价格(元/MWh)", "实时出清价格"],
                required=require_target,
            )
        else:
            total_power_col = resolve_column(day_df.columns, ["总加电力值(MW)", "总加电力值"], required=True)
            power_col = resolve_column(day_df.columns, ["电力值(MW)", "电力值"], required=True)
            thermal_power_col = None
            target_col = resolve_column(
                day_df.columns,
                ["日前出清价格(元/MWh)", "日前-出清价格(元/MWh)", "日前出清价格", "出清价格"],
                required=require_target,
            )
        overview_col = resolve_column(day_df.columns, ["日前-出清概况", "日前出清概况", "出清概况"], required=False)

        overview_series = pd.Series([np.nan] * len(day_df))
        thermal_series = pd.Series([np.nan] * len(day_df))
        if overview_col:
            overview_series = day_df[overview_col].ffill().bfill()
            thermal_series = overview_series.map(extract_thermal_on_capacity)

        periods = np.arange(1, 97)
        date_value = trade_date.normalize()
        day_type = day_type_of(date_value, self.holiday_dates)
        total_load = to_numeric(day_df[total_power_col])
        renewable_power = to_numeric(day_df[power_col])
        thermal_space = to_numeric(day_df[thermal_power_col]) if thermal_power_col else total_load - renewable_power
        result = pd.DataFrame(
            {
                "date": date_value,
                "period": periods,
                "hour": ((periods - 1) // 4).astype(int),
                "weekday": int(date_value.weekday()),
                "month": int(date_value.month),
                "is_weekend": int(day_type == "weekend"),
                "is_holiday": int(day_type == "holiday"),
                DAY_TYPE_COLUMN: day_type,
                TOTAL_LOAD_COLUMN: total_load,
                "net_load": thermal_space,
                RENEWABLE_POWER_COLUMN: renewable_power,
                THERMAL_SPACE_LOAD_RATIO_COLUMN: calculate_thermal_space_load_ratio(thermal_space, total_load),
                "thermal_on_capacity": to_numeric(thermal_series),
                "overview_text": overview_series,
                "sheet_name": sheet_name,
                "source_file": str(file_path),
            }
        )
        result["segment"] = assign_segments(result["period"])
        if target_col:
            result[TARGET_COLUMN] = to_numeric(day_df[target_col])
        return result


def split_train_valid(data: pd.DataFrame, valid_days: int) -> tuple[pd.DataFrame, pd.DataFrame]:
    if valid_days <= 0:
        return data, data.iloc[0:0].copy()
    unique_dates = sorted(pd.to_datetime(data["date"]).drop_duplicates().tolist())
    if len(unique_dates) <= valid_days:
        return data, data.iloc[0:0].copy()
    valid_dates = set(unique_dates[-valid_days:])
    valid_df = data[data["date"].isin(valid_dates)].copy()
    train_df = data[~data["date"].isin(valid_dates)].copy()
    if train_df.empty or valid_df.empty:
        return data, data.iloc[0:0].copy()
    return train_df, valid_df


def attach_lag_features(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    for column in LAG_PRICE_COLUMNS + LAG_LOAD_COLUMNS:
        if column not in result.columns:
            result[column] = np.nan
    if result.empty or "date" not in result.columns or "period" not in result.columns:
        return result

    result["date"] = pd.to_datetime(result["date"], errors="coerce").dt.normalize()
    result["period"] = pd.to_numeric(result["period"], errors="coerce")
    target_values = pd.to_numeric(result[TARGET_COLUMN], errors="coerce") if TARGET_COLUMN in result.columns else pd.Series(np.nan, index=result.index)
    net_load_values = pd.to_numeric(result["net_load"], errors="coerce") if "net_load" in result.columns else pd.Series(np.nan, index=result.index)
    source = pd.DataFrame(
        {
            "date": result["date"],
            "period": result["period"],
            TARGET_COLUMN: target_values,
            "net_load": net_load_values,
        }
    )
    lookup_source = (
        source.groupby(["date", "period"], as_index=False)
        .agg({TARGET_COLUMN: "mean", "net_load": "mean"})
        .sort_values(["period", "date"])
        .reset_index(drop=True)
    )

    base_index = pd.Series(np.arange(len(result)), name="__lag_row")
    working = pd.concat([base_index, result[["date", "period"]].reset_index(drop=True)], axis=1)
    for days in (1, 2, 3, 7):
        lag_source = lookup_source[["date", "period", TARGET_COLUMN]].copy()
        lag_source["date"] = lag_source["date"] + pd.Timedelta(days=days)
        lag_source = lag_source.rename(columns={TARGET_COLUMN: f"price_lag_{days}d"})
        working = working.merge(lag_source, on=["date", "period"], how="left")

    load_source = lookup_source[["date", "period", "net_load"]].copy()
    load_source["date"] = load_source["date"] + pd.Timedelta(days=1)
    load_source = load_source.rename(columns={"net_load": "net_load_lag_1d"})
    working = working.merge(load_source, on=["date", "period"], how="left")

    ordered = lookup_source.copy()
    shifted_prices = ordered.groupby("period", sort=False)[TARGET_COLUMN].shift(1)
    ordered["price_ma_3d"] = shifted_prices.groupby(ordered["period"], sort=False).transform(
        lambda series: series.rolling(3, min_periods=1).mean()
    )
    ordered["price_ma_7d"] = shifted_prices.groupby(ordered["period"], sort=False).transform(
        lambda series: series.rolling(7, min_periods=1).mean()
    )
    ordered["price_std_7d"] = shifted_prices.groupby(ordered["period"], sort=False).transform(
        lambda series: series.rolling(7, min_periods=1).std(ddof=0)
    )
    rolling_source = ordered[["date", "period", "price_ma_3d", "price_ma_7d", "price_std_7d"]]
    working = working.merge(rolling_source, on=["date", "period"], how="left")

    working = working.sort_values("__lag_row").reset_index(drop=True)
    for column in LAG_PRICE_COLUMNS + LAG_LOAD_COLUMNS:
        result[column] = pd.to_numeric(working[column], errors="coerce").to_numpy()
    return result


def attach_forecast_lag_features(forecast_df: pd.DataFrame, history_df: pd.DataFrame) -> pd.DataFrame:
    forecast = forecast_df.copy()
    forecast_columns = set(forecast.columns)
    forecast["__forecast_lag_row"] = np.arange(len(forecast))
    forecast["__is_forecast_row"] = True
    history = history_df.copy()
    history["__forecast_lag_row"] = np.nan
    history["__is_forecast_row"] = False
    if TARGET_COLUMN not in forecast.columns:
        forecast[TARGET_COLUMN] = np.nan
    combined = pd.concat([history, forecast], ignore_index=True, sort=False)
    combined = attach_lag_features(combined)
    result = (
        combined[combined["__is_forecast_row"]]
        .sort_values("__forecast_lag_row")
        .drop(columns=["__forecast_lag_row", "__is_forecast_row"], errors="ignore")
        .reset_index(drop=True)
    )
    if TARGET_COLUMN not in forecast_columns and TARGET_COLUMN in result.columns:
        result = result.drop(columns=[TARGET_COLUMN])
    return result


def apply_template_reference_lag_features(forecast_df: pd.DataFrame, template_reference_df: pd.DataFrame | None) -> pd.DataFrame:
    if template_reference_df is None or template_reference_df.empty:
        return forecast_df
    result = forecast_df.copy()
    reference = template_reference_df.copy()
    reference["period"] = pd.to_numeric(reference["period"], errors="coerce")
    reference = reference.dropna(subset=["period"]).drop_duplicates("period").set_index("period")
    period_index = pd.to_numeric(result["period"], errors="coerce")
    baseline_prices = period_index.map(pd.to_numeric(reference.get(TARGET_COLUMN), errors="coerce"))
    baseline_net_load = period_index.map(pd.to_numeric(reference.get("net_load"), errors="coerce"))
    for column in LAG_PRICE_COLUMNS:
        result[column] = baseline_prices.to_numpy()
    result["price_std_7d"] = 0.0
    result["net_load_lag_1d"] = baseline_net_load.to_numpy()
    return result


def calculate_metrics(actual: np.ndarray, predicted: np.ndarray) -> dict[str, float]:
    errors = actual - predicted
    mae = float(np.mean(np.abs(errors)))
    rmse = float(np.sqrt(np.mean(np.square(errors))))
    return {"mae": round(mae, 4), "rmse": round(rmse, 4)}


def _first_numeric_value(series: pd.Series) -> float | None:
    values = pd.to_numeric(series, errors="coerce").dropna()
    return float(values.iloc[0]) if not values.empty else None


def extract_forecast_thermal_capacity_file_value(forecast_df: pd.DataFrame) -> float | None:
    if "thermal_on_capacity" not in forecast_df.columns:
        return None
    return _first_numeric_value(forecast_df["thermal_on_capacity"])


def build_daily_thermal_capacity_frame(frame: pd.DataFrame) -> pd.DataFrame:
    if frame.empty or "date" not in frame.columns:
        return pd.DataFrame()
    working = frame.copy()
    working["date"] = pd.to_datetime(working["date"], errors="coerce").dt.normalize()
    working = working.dropna(subset=["date"]).sort_values(["date", "period"] if "period" in working.columns else ["date"])
    if working.empty:
        return pd.DataFrame()
    for column in [TOTAL_LOAD_COLUMN, RENEWABLE_POWER_COLUMN, "net_load", THERMAL_SPACE_LOAD_RATIO_COLUMN, "thermal_on_capacity"]:
        if column not in working.columns:
            working[column] = np.nan
        working[column] = pd.to_numeric(working[column], errors="coerce")
    daily_rows: list[dict[str, object]] = []
    for date_value, day_df in working.groupby("date", sort=True):
        total_load = day_df[TOTAL_LOAD_COLUMN]
        renewable = day_df[RENEWABLE_POWER_COLUMN]
        net_load = day_df["net_load"]
        total_load_mean = float(total_load.mean()) if total_load.notna().any() else np.nan
        renewable_mean = float(renewable.mean()) if renewable.notna().any() else np.nan
        net_diff = net_load.diff().abs()
        capacity_value = _first_numeric_value(day_df["thermal_on_capacity"])
        date_row = day_df.iloc[0]
        daily_rows.append(
            {
                "date": pd.Timestamp(date_value).normalize(),
                "total_load_mean": total_load_mean,
                "total_load_max": float(total_load.max()) if total_load.notna().any() else np.nan,
                "total_load_min": float(total_load.min()) if total_load.notna().any() else np.nan,
                "renewable_power_mean": renewable_mean,
                "renewable_power_max": float(renewable.max()) if renewable.notna().any() else np.nan,
                "renewable_power_min": float(renewable.min()) if renewable.notna().any() else np.nan,
                "net_load_mean": float(net_load.mean()) if net_load.notna().any() else np.nan,
                "net_load_max": float(net_load.max()) if net_load.notna().any() else np.nan,
                "net_load_min": float(net_load.min()) if net_load.notna().any() else np.nan,
                "net_load_peak_valley": float(net_load.max() - net_load.min()) if net_load.notna().any() else np.nan,
                "net_load_max_ramp": float(net_diff.max()) if net_diff.notna().any() else 0.0,
                "thermal_space_load_ratio_mean": float(day_df[THERMAL_SPACE_LOAD_RATIO_COLUMN].mean()) if day_df[THERMAL_SPACE_LOAD_RATIO_COLUMN].notna().any() else np.nan,
                "renewable_share_mean": float(renewable_mean / total_load_mean) if total_load_mean and np.isfinite(total_load_mean) else np.nan,
                "weekday": int(date_row.get("weekday", pd.Timestamp(date_value).weekday())),
                "month": int(date_row.get("month", pd.Timestamp(date_value).month)),
                "is_weekend": int(date_row.get("is_weekend", int(pd.Timestamp(date_value).weekday() >= 5))),
                "is_holiday": int(date_row.get("is_holiday", 0)),
                "thermal_on_capacity": capacity_value,
            }
        )
    daily = pd.DataFrame(daily_rows).sort_values("date").reset_index(drop=True)
    previous = pd.to_numeric(daily["thermal_on_capacity"], errors="coerce").shift(1)
    daily["previous_thermal_on_capacity"] = previous.combine_first(pd.to_numeric(daily["thermal_on_capacity"], errors="coerce"))
    return daily


def split_daily_train_valid(daily_df: pd.DataFrame, valid_days: int) -> tuple[pd.DataFrame, pd.DataFrame]:
    clean = daily_df.dropna(subset=["thermal_on_capacity"]).copy()
    clean = clean.dropna(subset=THERMAL_CAPACITY_FEATURE_COLUMNS, how="any")
    if clean.empty:
        return clean, clean.copy()
    if valid_days > 0 and len(clean) > valid_days:
        return clean.iloc[:-valid_days].copy(), clean.iloc[-valid_days:].copy()
    return clean.copy(), clean.iloc[0:0].copy()


def thermal_capacity_metrics(actual: pd.Series | np.ndarray, predicted: np.ndarray) -> dict[str, float | int | None]:
    actual_array = np.asarray(actual, dtype=float)
    predicted_array = np.asarray(predicted, dtype=float)
    valid = np.isfinite(actual_array) & np.isfinite(predicted_array)
    if not np.any(valid):
        return {"valid_rows": 0, "mae": None, "rmse": None, "max_error": None, "bias": None, "score": None}
    errors = predicted_array[valid] - actual_array[valid]
    mae = float(np.mean(np.abs(errors)))
    rmse = float(np.sqrt(np.mean(np.square(errors))))
    return {
        "valid_rows": int(np.sum(valid)),
        "mae": mae,
        "rmse": rmse,
        "max_error": float(np.max(np.abs(errors))),
        "bias": float(np.mean(errors)),
        "score": float(mae + 0.2 * rmse),
    }


class RidgeThermalCapacityModel:
    def __init__(
        self,
        feature_columns: list[str],
        coefficients: list[float],
        intercept: float,
        means: list[float],
        scales: list[float],
    ) -> None:
        self.feature_columns = feature_columns
        self.coefficients = np.asarray(coefficients, dtype=float)
        self.intercept = float(intercept)
        self.means = np.asarray(means, dtype=float)
        self.scales = np.asarray(scales, dtype=float)

    def predict(self, frame: pd.DataFrame | np.ndarray) -> np.ndarray:
        values = frame.to_numpy(dtype=float) if isinstance(frame, pd.DataFrame) else np.asarray(frame, dtype=float)
        scaled = (values - self.means) / self.scales
        return scaled @ self.coefficients + self.intercept

    def save(self, output_path: Path) -> None:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(
            json.dumps(
                {
                    "feature_columns": self.feature_columns,
                    "coefficients": self.coefficients.tolist(),
                    "intercept": self.intercept,
                    "means": self.means.tolist(),
                    "scales": self.scales.tolist(),
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )

    @classmethod
    def load(cls, model_path: Path) -> "RidgeThermalCapacityModel":
        payload = json.loads(model_path.read_text(encoding="utf-8"))
        return cls(
            list(payload.get("feature_columns") or []),
            list(payload.get("coefficients") or []),
            float(payload.get("intercept") or 0.0),
            list(payload.get("means") or []),
            list(payload.get("scales") or []),
        )


def available_thermal_capacity_backends() -> OrderedDict[str, str]:
    backends: OrderedDict[str, str] = OrderedDict()
    try:
        import xgboost  # noqa: F401

        backends["xgboost"] = "XGBoost"
    except Exception:
        pass
    try:
        import lightgbm  # noqa: F401

        backends["lightgbm"] = "LightGBM"
    except Exception:
        pass
    try:
        import catboost  # noqa: F401

        backends["catboost"] = "CatBoost"
    except Exception:
        pass
    backends["ridge"] = "Ridge"
    return backends


def thermal_capacity_model_file_name(backend: str) -> str:
    if backend == "lightgbm":
        return "thermal_capacity_model.txt"
    if backend == "catboost":
        return "thermal_capacity_model.cbm"
    if backend == "ridge":
        return "thermal_capacity_model_ridge.json"
    return "thermal_capacity_model.json"


def train_ridge_thermal_capacity_model(
    train_x: pd.DataFrame,
    train_y: pd.Series,
    output_path: Path,
    alpha: float = 1.0,
) -> RidgeThermalCapacityModel:
    values = train_x.to_numpy(dtype=float)
    target = train_y.to_numpy(dtype=float)
    means = np.nanmean(values, axis=0)
    scales = np.nanstd(values, axis=0)
    scales = np.where(np.isfinite(scales) & (scales > 1e-9), scales, 1.0)
    scaled = (values - means) / scales
    design = np.column_stack([np.ones(len(scaled)), scaled])
    penalty = np.eye(design.shape[1]) * float(alpha)
    penalty[0, 0] = 0.0
    solution = np.linalg.pinv(design.T @ design + penalty) @ design.T @ target
    model = RidgeThermalCapacityModel(
        list(train_x.columns),
        coefficients=solution[1:].astype(float).tolist(),
        intercept=float(solution[0]),
        means=means.astype(float).tolist(),
        scales=scales.astype(float).tolist(),
    )
    model.save(output_path)
    return model


def train_thermal_capacity_backend(
    backend: str,
    train_df: pd.DataFrame,
    valid_df: pd.DataFrame,
    feature_columns: list[str],
    num_boost_round: int,
    output_path: Path,
) -> tuple[object, np.ndarray, int | None]:
    train_x = train_df[feature_columns].astype(float)
    train_y = train_df["thermal_on_capacity"].astype(float)
    valid_x = valid_df[feature_columns].astype(float) if not valid_df.empty else None

    if backend == "ridge":
        model = train_ridge_thermal_capacity_model(train_x, train_y, output_path)
        evaluation_x = valid_x if valid_x is not None else train_x
        return model, model.predict(evaluation_x), None

    if backend == "lightgbm":
        import lightgbm as lgb

        train_dataset = lgb.Dataset(train_x, label=train_y, feature_name=feature_columns)
        valid_sets = [train_dataset]
        callbacks = []
        if valid_x is not None:
            valid_sets.append(lgb.Dataset(valid_x, label=valid_df["thermal_on_capacity"].astype(float), reference=train_dataset, feature_name=feature_columns))
            callbacks.append(lgb.early_stopping(30, verbose=False))
        model = lgb.train(
            {
                "objective": "regression",
                "metric": "rmse",
                "learning_rate": 0.05,
                "max_depth": 5,
                "seed": 42,
                "verbosity": -1,
            },
            train_dataset,
            num_boost_round=max(1, int(num_boost_round)),
            valid_sets=valid_sets,
            callbacks=callbacks,
        )
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(model.model_to_string(), encoding="utf-8")
        best_iteration = int(model.best_iteration) if getattr(model, "best_iteration", 0) else None
        evaluation_x = valid_x if valid_x is not None else train_x
        prediction = model.predict(evaluation_x, num_iteration=best_iteration)
        return model, np.asarray(prediction, dtype=float), best_iteration

    if backend == "catboost":
        from catboost import CatBoostRegressor

        model = CatBoostRegressor(
            iterations=max(1, int(num_boost_round)),
            learning_rate=0.05,
            depth=5,
            loss_function="RMSE",
            random_seed=42,
            verbose=False,
            allow_writing_files=False,
        )
        fit_kwargs: dict[str, object] = {}
        if valid_x is not None:
            fit_kwargs["eval_set"] = (valid_x, valid_df["thermal_on_capacity"].astype(float))
            fit_kwargs["early_stopping_rounds"] = 30
            fit_kwargs["use_best_model"] = True
        model.fit(train_x, train_y, **fit_kwargs)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        model.save_model(str(output_path))
        best_iteration = int(model.get_best_iteration()) if model.get_best_iteration() is not None else None
        evaluation_x = valid_x if valid_x is not None else train_x
        return model, np.asarray(model.predict(evaluation_x), dtype=float), best_iteration

    import xgboost as xgb

    dtrain = xgb.DMatrix(train_x, label=train_y, feature_names=feature_columns)
    evals = [(dtrain, "train")]
    train_kwargs: dict[str, object] = {
        "params": {**DEFAULT_XGB_PARAMS, "objective": "reg:squarederror"},
        "dtrain": dtrain,
        "num_boost_round": max(1, int(num_boost_round)),
        "evals": evals,
        "verbose_eval": False,
    }
    dvalid = None
    if valid_x is not None:
        dvalid = xgb.DMatrix(valid_x, label=valid_df["thermal_on_capacity"].astype(float), feature_names=feature_columns)
        evals.append((dvalid, "valid"))
        train_kwargs["early_stopping_rounds"] = 30
    model = xgb.train(**train_kwargs)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    model.save_model(str(output_path))
    best_iteration = int(model.attr("best_iteration")) if model.attr("best_iteration") else None
    evaluation_matrix = dvalid if dvalid is not None else dtrain
    iteration_range = (0, best_iteration + 1) if best_iteration is not None else None
    prediction = model.predict(evaluation_matrix, iteration_range=iteration_range) if iteration_range else model.predict(evaluation_matrix)
    return model, np.asarray(prediction, dtype=float), best_iteration


def select_best_thermal_capacity_variant(variant_metrics: dict[str, dict[str, object]]) -> str | None:
    candidates: list[tuple[float, str]] = []
    for key, metrics in variant_metrics.items():
        overall = metrics.get("overall") if isinstance(metrics.get("overall"), dict) else metrics
        score = overall.get("score") if isinstance(overall, dict) else None
        if score is not None and np.isfinite(float(score)):
            candidates.append((float(score), key))
    if not candidates:
        return None
    return min(candidates)[1]


def train_thermal_capacity_model(
    history_df: pd.DataFrame,
    output_dir: Path,
    valid_days: int,
    num_boost_round: int,
    backends: dict[str, str] | None = None,
) -> dict[str, object]:
    daily_df = build_daily_thermal_capacity_frame(history_df)
    train_df, valid_df = split_daily_train_valid(daily_df, valid_days)
    backends = backends or available_thermal_capacity_backends()
    metadata: dict[str, object] = {
        "enabled": False,
        "feature_columns": list(THERMAL_CAPACITY_FEATURE_COLUMNS),
        "target_column": "thermal_on_capacity",
        "model_path": "thermal_capacity_model/thermal_capacity_model.json",
        "train_rows": int(len(train_df)),
        "valid_rows": int(len(valid_df)),
        "metrics": {"overall": {"valid_rows": 0, "mae": None, "rmse": None, "max_error": None, "bias": None, "score": None}},
        "model_variants": {},
        "variant_metrics": {},
    }
    if train_df.empty or len(train_df) < 2:
        metadata["reason"] = "not_enough_training_rows"
        return metadata
    output_dir.mkdir(parents=True, exist_ok=True)
    variant_metrics: dict[str, dict[str, object]] = {}
    model_variants: dict[str, dict[str, object]] = {}
    evaluation_df = valid_df if not valid_df.empty else train_df
    evaluation_actual = evaluation_df["thermal_on_capacity"]
    for backend_key, backend_label in backends.items():
        variant_key = f"thermal_capacity_{backend_key}"
        variant_dir = output_dir / variant_key
        model_file = thermal_capacity_model_file_name(backend_key)
        model_path = variant_dir / model_file
        try:
            _model, prediction, best_iteration = train_thermal_capacity_backend(
                backend_key,
                train_df,
                valid_df,
                list(THERMAL_CAPACITY_FEATURE_COLUMNS),
                num_boost_round,
                model_path,
            )
            metrics = thermal_capacity_metrics(evaluation_actual, prediction)
            metrics["train_rows"] = int(len(train_df))
            metrics["valid_rows"] = int(len(valid_df))
            metrics["best_iteration"] = best_iteration
            variant_metrics[variant_key] = {"overall": metrics}
            model_variants[variant_key] = {
                "model_dir": f"thermal_capacity_model/{variant_key}",
                "model_path": f"thermal_capacity_model/{variant_key}/{model_file}",
                "model_backend": backend_key,
                "model_backend_label": backend_label,
                "feature_columns": list(THERMAL_CAPACITY_FEATURE_COLUMNS),
            }
        except Exception as exc:
            variant_metrics[variant_key] = {
                "overall": {
                    "valid_rows": 0,
                    "mae": None,
                    "rmse": None,
                    "max_error": None,
                    "bias": None,
                    "score": None,
                    "error": str(exc),
                }
            }

    selected_model_key = select_best_thermal_capacity_variant(variant_metrics)
    if not selected_model_key or selected_model_key not in model_variants:
        metadata.update(
            {
                "reason": "no_valid_thermal_capacity_backend",
                "model_variants": model_variants,
                "variant_metrics": variant_metrics,
            }
        )
        return metadata

    selected_variant = model_variants[selected_model_key]
    metadata.update(
        {
            "enabled": True,
            "selected_model_key": selected_model_key,
            "selected_model_backend": selected_variant["model_backend"],
            "selected_model_backend_label": selected_variant["model_backend_label"],
            "model_path": selected_variant["model_path"],
            "model_variants": model_variants,
            "variant_metrics": variant_metrics,
            "train_rows": int(len(train_df)),
            "valid_rows": int(len(valid_df)),
            "train_dates": pd.to_datetime(train_df["date"]).dt.strftime("%Y-%m-%d").tolist(),
            "valid_dates": pd.to_datetime(valid_df["date"]).dt.strftime("%Y-%m-%d").tolist(),
            "metrics": variant_metrics[selected_model_key],
        }
    )
    return metadata


def predict_thermal_capacity_value(
    forecast_df: pd.DataFrame,
    model_entry: dict[str, object] | None,
    file_value: float | None = None,
) -> float | None:
    if not model_entry or not model_entry.get("model"):
        return None
    feature_columns = list((model_entry.get("metadata") or {}).get("feature_columns") or THERMAL_CAPACITY_FEATURE_COLUMNS)
    daily_df = build_daily_thermal_capacity_frame(forecast_df)
    if daily_df.empty:
        return None
    if file_value is not None and "previous_thermal_on_capacity" in daily_df.columns:
        daily_df.loc[:, "previous_thermal_on_capacity"] = float(file_value)
    feature_df = daily_df[feature_columns].astype(float)
    if feature_df.isna().any().any():
        feature_df = feature_df.ffill().bfill().fillna(0.0)
    backend = str(model_entry.get("backend") or (model_entry.get("metadata") or {}).get("selected_model_backend") or "xgboost")
    prediction = predict_thermal_capacity_backend(feature_df, model_entry["model"], backend, feature_columns)
    if len(prediction) == 0 or not np.isfinite(float(prediction[0])):
        return None
    return float(prediction[0])


def load_thermal_capacity_backend_model(model_path: Path, backend: str) -> object:
    if backend == "ridge":
        return RidgeThermalCapacityModel.load(model_path)
    if backend == "lightgbm":
        import lightgbm as lgb

        return lgb.Booster(model_str=model_path.read_text(encoding="utf-8"))
    if backend == "catboost":
        from catboost import CatBoostRegressor

        model = CatBoostRegressor()
        model.load_model(str(model_path))
        return model
    import xgboost as xgb

    model = xgb.Booster()
    model.load_model(str(model_path))
    return model


def predict_thermal_capacity_backend(
    feature_df: pd.DataFrame,
    model: object,
    backend: str,
    feature_columns: list[str],
) -> np.ndarray:
    if backend in {"ridge", "lightgbm", "catboost"}:
        return np.asarray(model.predict(feature_df), dtype=float)
    import xgboost as xgb

    dmatrix = xgb.DMatrix(feature_df, feature_names=feature_columns)
    return np.asarray(model.predict(dmatrix), dtype=float)


def build_thermal_capacity_choice(
    forecast_df: pd.DataFrame,
    config: dict[str, object] | None,
    model_predicted_value: float | None,
) -> dict[str, object]:
    config = config or {}
    mode = str(config.get("mode") or config.get("thermal_capacity_mode") or "manual").strip().lower()
    file_value = extract_forecast_thermal_capacity_file_value(forecast_df)
    manual_value_raw = config.get("manual_value", config.get("manual_thermal_on_capacity"))
    manual_value = float(manual_value_raw) if manual_value_raw not in (None, "") else file_value
    if mode == "model":
        if model_predicted_value is None:
            raise ValueError("已选择开机容量模型预测值，但当前模型版本没有可用的开机容量模型预测结果")
        source = "model"
        value = float(model_predicted_value)
    else:
        source = "manual"
        value = float(manual_value) if manual_value is not None else np.nan
    return {
        "source": source,
        "value": value,
        "model_value": float(model_predicted_value) if model_predicted_value is not None else None,
        "file_value": float(file_value) if file_value is not None else None,
        "manual_value": float(manual_value) if manual_value is not None else None,
    }


def apply_thermal_capacity_choice(forecast_df: pd.DataFrame, choice: dict[str, object]) -> pd.DataFrame:
    result = forecast_df.copy()
    value = float(choice["value"])
    result["thermal_on_capacity"] = value
    result[THERMAL_CAPACITY_SOURCE_COLUMN] = str(choice.get("source") or "")
    result[THERMAL_CAPACITY_VALUE_COLUMN] = value
    result[THERMAL_CAPACITY_MODEL_VALUE_COLUMN] = choice.get("model_value")
    result[THERMAL_CAPACITY_FILE_VALUE_COLUMN] = choice.get("file_value")
    return result


def nearest_similarity_price(target_loads: np.ndarray, ref_loads: np.ndarray, ref_prices: np.ndarray, k: int = 3) -> tuple[np.ndarray, np.ndarray]:
    if len(ref_loads) == 0:
        return np.full(len(target_loads), np.nan), np.full(len(target_loads), np.nan)
    ref_loads = ref_loads.astype(float)
    ref_prices = ref_prices.astype(float)
    target_loads = target_loads.astype(float)
    diff = np.abs(target_loads[:, None] - ref_loads[None, :])
    top_k = min(k, len(ref_loads))
    indices = np.argpartition(diff, top_k - 1, axis=1)[:, :top_k]
    top_diff = np.take_along_axis(diff, indices, axis=1)
    top_price = ref_prices[indices]
    weights = 1.0 / (top_diff + 1e-6)
    weighted_price = np.sum(weights * top_price, axis=1) / np.sum(weights, axis=1)
    min_gap = top_diff.min(axis=1)
    return weighted_price, min_gap


def multi_factor_knn_similarity_price(
    target_df: pd.DataFrame,
    reference_df: pd.DataFrame,
    k: int = 5,
    max_distance: float = 3.0,
    weighted_k: int | None = None,
    weighted_max_distance: float | None = None,
) -> pd.DataFrame:
    numeric_features: list[tuple[str, float]] = [
        ("net_load", 2.0),
        (TOTAL_LOAD_COLUMN, 1.0),
        (RENEWABLE_POWER_COLUMN, 1.0),
        ("thermal_on_capacity", 0.8),
        (THERMAL_SPACE_LOAD_RATIO_COLUMN, 1.2),
    ]
    ref_prices = pd.to_numeric(reference_df.get(TARGET_COLUMN), errors="coerce") if TARGET_COLUMN in reference_df.columns else pd.Series(dtype=float)
    base_valid = ref_prices.notna().to_numpy()
    rows: list[dict[str, float | int | None]] = []
    plain_top_k = max(1, int(k or 1))
    weighted_top_k = max(1, int(weighted_k or plain_top_k))
    plain_max_distance = max(0.1, float(max_distance or DEFAULT_KNN_SIMILARITY_CONFIG["knn_max_distance"]))
    weighted_limit = weighted_max_distance if weighted_max_distance is not None else plain_max_distance
    weighted_max_distance_value = max(0.1, float(weighted_limit or DEFAULT_KNN_SIMILARITY_CONFIG["weighted_knn_max_distance"]))

    def select_candidate_prices(
        valid_indices: np.ndarray,
        valid_distances: np.ndarray,
        order: np.ndarray,
        candidate_count: int,
        distance_limit: float,
    ) -> tuple[np.ndarray, np.ndarray, float]:
        selected_order = order[: min(candidate_count, len(order))]
        candidate_indices = valid_indices[selected_order]
        candidate_distances = valid_distances[selected_order]
        nearest_distance = float(candidate_distances[0])
        adaptive_threshold = min(float(distance_limit), max(0.35, nearest_distance * 3.0))
        keep_mask = candidate_distances <= adaptive_threshold
        if not keep_mask.any():
            keep_mask[0] = True
        kept_indices = candidate_indices[keep_mask]
        kept_distances = candidate_distances[keep_mask]
        kept_prices = ref_prices.iloc[kept_indices].to_numpy(dtype=float)
        return kept_prices, kept_distances, nearest_distance

    for _, target_row in target_df.iterrows():
        valid_mask = base_valid.copy()
        if len(valid_mask) == 0:
            rows.append({"plain_price": np.nan, "weighted_price": np.nan, "count": 0, "min_distance": np.nan})
            continue
        distance = np.zeros(len(reference_df), dtype=float)
        used_feature_count = 0
        for feature, weight in numeric_features:
            if feature not in target_df.columns or feature not in reference_df.columns:
                continue
            target_value = pd.to_numeric(pd.Series([target_row.get(feature)]), errors="coerce").iloc[0]
            if not pd.notna(target_value):
                continue
            reference_values = pd.to_numeric(reference_df[feature], errors="coerce")
            feature_valid = reference_values.notna().to_numpy() & pd.notna(target_value)
            valid_mask &= feature_valid
            if feature == THERMAL_SPACE_LOAD_RATIO_COLUMN:
                scale = max(float(reference_values.abs().median(skipna=True) or 0.0), 0.05)
            else:
                scale = max(float(reference_values.std(skipna=True) or 0.0), float(reference_values.abs().median(skipna=True) or 0.0) * 0.05, 1.0)
            normalized_gap = (reference_values.to_numpy(dtype=float) - float(target_value)) / scale
            distance += weight * np.square(normalized_gap)
            used_feature_count += 1
        if "period" in target_df.columns and "period" in reference_df.columns:
            target_period = pd.to_numeric(pd.Series([target_row.get("period")]), errors="coerce").iloc[0]
            if pd.notna(target_period):
                reference_periods = pd.to_numeric(reference_df["period"], errors="coerce")
                period_valid = reference_periods.notna().to_numpy()
                valid_mask &= period_valid
                period_gap = np.abs(reference_periods.to_numpy(dtype=float) - float(target_period))
                period_gap = np.minimum(period_gap, 96 - period_gap) / 96.0
                distance += 0.5 * np.square(period_gap)
                used_feature_count += 1
        if DAY_TYPE_COLUMN in target_df.columns and DAY_TYPE_COLUMN in reference_df.columns:
            target_type = str(target_row.get(DAY_TYPE_COLUMN) or "").strip().lower()
            reference_types = reference_df[DAY_TYPE_COLUMN].astype(str).str.strip().str.lower().to_numpy()
            if target_type:
                distance += 0.6 * (reference_types != target_type).astype(float)
                used_feature_count += 1
        if used_feature_count == 0 or not valid_mask.any():
            rows.append({"plain_price": np.nan, "weighted_price": np.nan, "count": 0, "min_distance": np.nan})
            continue
        valid_indices = np.where(valid_mask)[0]
        valid_distances = np.sqrt(distance[valid_indices])
        order = np.argsort(valid_distances)
        plain_prices, plain_distances, nearest_distance = select_candidate_prices(
            valid_indices,
            valid_distances,
            order,
            plain_top_k,
            plain_max_distance,
        )
        weighted_prices, weighted_distances, _ = select_candidate_prices(
            valid_indices,
            valid_distances,
            order,
            weighted_top_k,
            weighted_max_distance_value,
        )
        weights = 1.0 / (weighted_distances + 1e-6)
        rows.append(
            {
                "plain_price": float(np.mean(plain_prices)),
                "weighted_price": float(np.sum(weights * weighted_prices) / np.sum(weights)),
                "count": int(len(plain_prices)),
                "min_distance": nearest_distance,
            }
        )
    return pd.DataFrame(rows)


def normalize_similarity_reference_days(reference_days: int | None) -> int:
    value = int(reference_days or DEFAULT_SIMILARITY_REFERENCE_DAYS)
    if value <= 0:
        raise ValueError("similarity_reference_days 必须大于等于 1")
    return value


def day_type_penalty(target_type: object, reference_types: pd.Series) -> pd.Series:
    target_text = str(target_type or "").strip().lower()
    reference_text = reference_types.astype(str).str.strip().str.lower()
    return (reference_text != target_text).astype(float)


def normalized_numeric_gap(target_value: object, reference_values: pd.Series) -> pd.Series:
    numeric_reference = pd.to_numeric(reference_values, errors="coerce")
    numeric_target = pd.to_numeric(pd.Series([target_value]), errors="coerce").iloc[0]
    if pd.isna(numeric_target):
        return pd.Series(np.nan, index=reference_values.index, dtype=float)
    normalizer = max(float(abs(numeric_target)), float(numeric_reference.abs().median(skipna=True) or 0.0), 1.0)
    return (numeric_reference - float(numeric_target)).abs() / normalizer


def ensure_similarity_columns(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    if TOTAL_LOAD_COLUMN not in result.columns:
        result[TOTAL_LOAD_COLUMN] = np.nan
    if RENEWABLE_POWER_COLUMN not in result.columns:
        result[RENEWABLE_POWER_COLUMN] = np.nan
    if "thermal_on_capacity" not in result.columns:
        result["thermal_on_capacity"] = np.nan
    if THERMAL_SPACE_LOAD_RATIO_COLUMN not in result.columns:
        thermal_space = result["net_load"] if "net_load" in result.columns else pd.Series([np.nan] * len(result), index=result.index)
        result[THERMAL_SPACE_LOAD_RATIO_COLUMN] = calculate_thermal_space_load_ratio(thermal_space, result[TOTAL_LOAD_COLUMN])
    if DAY_TYPE_COLUMN not in result.columns:
        if "date" in result.columns:
            result[DAY_TYPE_COLUMN] = pd.to_datetime(result["date"], errors="coerce").map(lambda value: day_type_of(value, set()) if pd.notna(value) else "")
        else:
            result[DAY_TYPE_COLUMN] = ""
    return result


def attach_similarity_features(
    history_df: pd.DataFrame,
    reference_days: int | None = None,
) -> pd.DataFrame:
    result = ensure_similarity_columns(history_df)
    result[NET_LOAD_ONLY_SIMILAR_PRICE_COLUMN] = np.nan
    result[NET_LOAD_ONLY_SIMILAR_GAP_COLUMN] = np.nan
    unique_dates = sorted(pd.to_datetime(result["date"]).drop_duplicates().tolist())
    reference_days = normalize_similarity_reference_days(reference_days)
    for index in range(1, len(unique_dates)):
        current_date = unique_dates[index]
        ref_dates = unique_dates[max(0, index - reference_days):index]
        ref_day = result[result["date"].isin(ref_dates)]
        current_mask = result["date"] == current_date
        current_day = result.loc[current_mask]
        similar_price, similar_gap = nearest_similarity_price(
            current_day["net_load"].to_numpy(dtype=float),
            ref_day["net_load"].to_numpy(dtype=float),
            ref_day[TARGET_COLUMN].to_numpy(dtype=float),
        )
        result.loc[current_mask, NET_LOAD_ONLY_SIMILAR_PRICE_COLUMN] = similar_price
        result.loc[current_mask, NET_LOAD_ONLY_SIMILAR_GAP_COLUMN] = similar_gap
    return result


def filter_date_range(data: pd.DataFrame, start_date: str | None, end_date: str | None) -> pd.DataFrame:
    filtered = data.copy()
    if start_date:
        filtered = filtered[filtered["date"] >= pd.Timestamp(start_date)]
    if end_date:
        filtered = filtered[filtered["date"] <= pd.Timestamp(end_date)]
    return filtered.reset_index(drop=True)


def resolve_training_date_range(history_df: pd.DataFrame, config: TrainConfig) -> tuple[str | None, str | None]:
    if not config.training_window_days:
        return config.start_date, config.end_date
    window_days = int(config.training_window_days)
    if window_days < 1:
        raise ValueError("training_window_days 必须大于等于 1")
    valid_days = max(int(config.valid_days or 0), 0)
    dates = sorted(pd.to_datetime(history_df["date"], errors="coerce").dropna().dt.normalize().unique())
    if not dates:
        return config.start_date, config.end_date
    end_limit = parse_optional_date(config.end_date)
    eligible_dates = [date for date in dates if end_limit is None or date <= end_limit]
    if not eligible_dates:
        raise ValueError("训练使用天数内没有可用历史日期，请检查训练结束日期。")
    required_days = window_days + valid_days
    selected_dates = eligible_dates[-required_days:] if len(eligible_dates) > required_days else eligible_dates
    return pd.Timestamp(selected_dates[0]).strftime("%Y-%m-%d"), pd.Timestamp(selected_dates[-1]).strftime("%Y-%m-%d")


def training_history_signature(config: TrainConfig) -> dict:
    feature_reference_days = max(7, normalize_similarity_reference_days(config.similarity_reference_days))
    return history_cache_signature(
        config.history_dir,
        holiday_file=config.holiday_file,
        require_target=True,
        start_date=config.start_date,
        end_date=config.end_date,
        leading_days=feature_reference_days,
        data_mode=config.data_mode,
    )


def build_training_frame(
    config: TrainConfig,
    progress_callback: ProgressCallback | None = None,
) -> tuple[pd.DataFrame, list[str], Path]:
    feature_reference_days = max(7, normalize_similarity_reference_days(config.similarity_reference_days))
    feature_cache_dir = Path(config.feature_cache_root) if config.feature_cache_root is not None else config.model_root / "cache"
    try:
        history_df, skipped_sheets, _, history_signature, quality_issues = load_cached_history_collection(
            config.history_dir,
            config.holiday_file,
            True,
            config.model_root,
            progress_callback=progress_callback,
            progress_start=6,
            progress_end=18,
            progress_label="正在加载历史数据文件",
            start_date=config.start_date,
            end_date=config.end_date,
            leading_days=feature_reference_days,
            cache_dir=feature_cache_dir,
            data_mode=config.data_mode,
        )
    except DataQualityValidationError as exc:
        report_path = save_quality_report(
            task_type="train",
            status="blocked",
            issues=exc.issues,
            metadata={
                "history_dir": str(config.history_dir),
                "start_date": config.start_date,
                "end_date": config.end_date,
                "training_window_days": config.training_window_days,
            },
        )
        exc.report_path = report_path
        raise
    quality_report_path = save_quality_report(
        task_type="train",
        status=quality_report_status(quality_issues),
        issues=quality_issues,
        metadata={
            "history_dir": str(config.history_dir),
            "skipped_sheets": skipped_sheets,
            "start_date": config.start_date,
            "end_date": config.end_date,
            "training_window_days": config.training_window_days,
        },
    )
    raw_history_df = history_df
    effective_start_date, effective_end_date = resolve_training_date_range(raw_history_df, config)
    feature_context_start_date = effective_start_date
    if feature_context_start_date:
        feature_context_start_date = (
            pd.Timestamp(feature_context_start_date) - pd.Timedelta(days=feature_reference_days)
        ).strftime("%Y-%m-%d")
    feature_cache_payload = {
        "namespace": "direct_price_training_features",
        "similarity_reference_days": feature_reference_days,
        "feature_context_start_date": feature_context_start_date,
        "effective_start_date": effective_start_date,
        "effective_end_date": effective_end_date,
    }
    if str(config.data_mode or DAYAHEAD_DATA_MODE) != DAYAHEAD_DATA_MODE:
        feature_cache_payload["data_mode"] = str(config.data_mode or DAYAHEAD_DATA_MODE)
    feature_cache_key = derived_cache_key(
        history_signature["cache_key"],
        json.dumps(feature_cache_payload, ensure_ascii=False, sort_keys=True),
    )
    cached_features = load_dataframe_cache(feature_cache_dir, "training_features", feature_cache_key)
    if cached_features is not None:
        history_df, _ = cached_features
        missing_columns = missing_dataframe_columns(history_df, required_training_feature_cache_columns())
        if missing_columns:
            cached_features = None
            history_df = raw_history_df.copy()
            emit_progress(progress_callback, f"训练特征缓存缺少字段 {missing_columns}，正在重新构建", 19)
        else:
            emit_progress(progress_callback, "训练特征未变化，已复用缓存", 21)
    else:
        missing_columns = []
    if cached_features is None:
        emit_progress(progress_callback, "正在准备训练特征", 19)
        history_df = filter_date_range(raw_history_df, feature_context_start_date, effective_end_date)
        history_df = attach_lag_features(history_df)
        save_dataframe_cache(feature_cache_dir, "training_features", feature_cache_key, history_df, skipped_sheets)
        emit_progress(progress_callback, "训练特征构建完成", 21)
    full_feature_history_df = history_df.copy()
    if config.training_window_days:
        emit_progress(progress_callback, f"按最近 {config.training_window_days} 天训练使用天数过滤数据", 22)
    history_df = filter_date_range(history_df, effective_start_date, effective_end_date)
    emit_progress(progress_callback, "正在过滤日期范围", 22)
    required_columns = [TARGET_COLUMN, "net_load"]
    history_df = history_df.dropna(subset=required_columns).reset_index(drop=True)
    if history_df.empty:
        raise ValueError("数据加载后经清洗为空，请检查历史数据文件是否包含有效数据。")
    history_df.attrs["full_feature_history_df"] = full_feature_history_df
    history_df.attrs["history_signature"] = history_signature
    return history_df, skipped_sheets, quality_report_path


def model_training_frame_from_base(
    base_history_df: pd.DataFrame,
    *,
    training_window_days: int | None,
    valid_days: int,
    start_date: str | None,
    end_date: str | None,
    similarity_reference_days: int | None = None,
) -> pd.DataFrame:
    scoped_config = TrainConfig(
        valid_days=valid_days,
        training_window_days=training_window_days,
        start_date=start_date,
        end_date=end_date,
    )
    if missing_dataframe_columns(base_history_df, LAG_PRICE_COLUMNS + LAG_LOAD_COLUMNS):
        base_history_df = attach_lag_features(base_history_df)
    effective_start_date, effective_end_date = resolve_training_date_range(base_history_df, scoped_config)
    scoped_df = filter_date_range(base_history_df, effective_start_date, effective_end_date)
    scoped_df = scoped_df.dropna(subset=[TARGET_COLUMN, "net_load"]).reset_index(drop=True)
    if scoped_df.empty:
        raise ValueError("模型训练数据为空，请检查训练窗口或起止日期。")
    return scoped_df

def make_dmatrix(frame: pd.DataFrame, feature_columns: list[str], label_column: str | None = None):
    import xgboost as xgb

    feature_frame = frame[feature_columns].astype(float)
    if label_column:
        return xgb.DMatrix(feature_frame, label=frame[label_column].astype(float), feature_names=feature_columns)
    return xgb.DMatrix(feature_frame, feature_names=feature_columns)


def best_iteration_range(booster) -> tuple[int, int] | None:
    best_iteration = booster.attr("best_iteration")
    if best_iteration is None:
        return None
    return 0, int(best_iteration) + 1



def available_model_backends() -> OrderedDict[str, str]:
    available = OrderedDict()
    for backend, label in MODEL_BACKENDS.items():
        if backend == "xgboost":
            ensure_xgboost_available()
            available[backend] = label
            continue
        try:
            __import__(backend)
        except Exception:
            continue
        available[backend] = label
    return available


def model_file_name(segment_name: str, backend: str) -> str:
    if backend == "lightgbm":
        return f"{segment_name}.txt"
    if backend == "catboost":
        return f"{segment_name}.cbm"
    return f"{segment_name}.json"


def save_lightgbm_model(model, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(model.model_to_string(), encoding="utf-8")


def load_lightgbm_model(model_path: Path):
    import lightgbm as lgb

    return lgb.Booster(model_str=model_path.read_text(encoding="utf-8"))


def train_backend_model(
    backend: str,
    train_df: pd.DataFrame,
    valid_df: pd.DataFrame,
    feature_columns: list[str],
    label_column: str,
    num_boost_round: int,
    output_path: Path | None,
    sample_weight: np.ndarray | pd.Series | None = None,
) -> tuple[object, int | None, np.ndarray | None]:
    train_x = train_df[feature_columns].astype(float)
    train_y = train_df[label_column].astype(float)
    valid_x = valid_df[feature_columns].astype(float) if not valid_df.empty else None
    valid_y = valid_df[label_column].astype(float) if not valid_df.empty else None
    train_weight = np.asarray(sample_weight, dtype=float) if sample_weight is not None else None

    if backend == "lightgbm":
        import lightgbm as lgb

        train_dataset = lgb.Dataset(train_x, label=train_y, weight=train_weight, feature_name=feature_columns)
        valid_sets = [train_dataset]
        valid_names = ["train"]
        callbacks = []
        if valid_x is not None:
            valid_dataset = lgb.Dataset(valid_x, label=valid_y, feature_name=feature_columns, reference=train_dataset)
            valid_sets.append(valid_dataset)
            valid_names.append("valid")
            callbacks.append(lgb.early_stopping(50, verbose=False))
        model = lgb.train(
            {
                "objective": "regression",
                "metric": "rmse",
                "learning_rate": DEFAULT_XGB_PARAMS["eta"],
                "max_depth": DEFAULT_XGB_PARAMS["max_depth"],
                "min_data_in_leaf": DEFAULT_XGB_PARAMS["min_child_weight"],
                "feature_fraction": DEFAULT_XGB_PARAMS["colsample_bytree"],
                "bagging_fraction": DEFAULT_XGB_PARAMS["subsample"],
                "bagging_freq": 1,
                "seed": DEFAULT_XGB_PARAMS["seed"],
                "verbosity": -1,
            },
            train_dataset,
            num_boost_round=num_boost_round,
            valid_sets=valid_sets,
            valid_names=valid_names,
            callbacks=callbacks,
        )
        best_iteration = int(model.best_iteration) if getattr(model, "best_iteration", 0) else None
        prediction = model.predict(valid_x, num_iteration=best_iteration) if valid_x is not None else None
        if output_path is not None:
            save_lightgbm_model(model, output_path)
        return model, best_iteration, prediction

    if backend == "catboost":
        from catboost import CatBoostRegressor

        model = CatBoostRegressor(
            iterations=num_boost_round,
            learning_rate=DEFAULT_XGB_PARAMS["eta"],
            depth=DEFAULT_XGB_PARAMS["max_depth"],
            loss_function="RMSE",
            random_seed=DEFAULT_XGB_PARAMS["seed"],
            verbose=False,
            allow_writing_files=False,
        )
        fit_kwargs = {}
        if valid_x is not None:
            fit_kwargs["eval_set"] = (valid_x, valid_y)
            fit_kwargs["early_stopping_rounds"] = 50
            fit_kwargs["use_best_model"] = True
        if train_weight is not None:
            fit_kwargs["sample_weight"] = train_weight
        model.fit(train_x, train_y, **fit_kwargs)
        best_iteration = int(model.get_best_iteration()) if model.get_best_iteration() is not None else None
        prediction = model.predict(valid_x) if valid_x is not None else None
        if output_path is not None:
            output_path.parent.mkdir(parents=True, exist_ok=True)
            model.save_model(str(output_path))
        return model, best_iteration, prediction

    import xgboost as xgb

    dtrain = make_dmatrix(train_df, feature_columns, label_column)
    if train_weight is not None:
        dtrain.set_weight(train_weight)
    evals = [(dtrain, "train")]
    train_kwargs = {
        "params": DEFAULT_XGB_PARAMS,
        "dtrain": dtrain,
        "num_boost_round": num_boost_round,
        "evals": evals,
        "verbose_eval": False,
    }
    dvalid = None
    if valid_x is not None:
        dvalid = make_dmatrix(valid_df, feature_columns, label_column)
        evals.append((dvalid, "valid"))
        train_kwargs["early_stopping_rounds"] = 50
    booster = xgb.train(**train_kwargs)
    if output_path is not None:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        booster.save_model(str(output_path))
    best_iteration = int(booster.attr("best_iteration")) if booster.attr("best_iteration") else None
    if dvalid is None:
        prediction = None
    else:
        iteration_range = best_iteration_range(booster)
        prediction = booster.predict(dvalid, iteration_range=iteration_range) if iteration_range else booster.predict(dvalid)
    return booster, best_iteration, prediction


def high_price_sample_weights(
    frame: pd.DataFrame,
    label_column: str,
    enabled: bool,
    high_price_threshold: float | None,
    multiplier: float,
) -> np.ndarray | None:
    if not enabled or high_price_threshold is None or multiplier <= 1:
        return None
    weights = np.ones(len(frame), dtype=float)
    labels = pd.to_numeric(frame[label_column], errors="coerce").to_numpy(dtype=float)
    weights[labels >= float(high_price_threshold)] = float(multiplier)
    return weights


def train_segment_models(
    history_df: pd.DataFrame,
    output_dir: Path,
    valid_days: int,
    num_boost_round: int,
    feature_columns: list[str],
    progress_callback: ProgressCallback | None = None,
    progress_start: int = 25,
    progress_span: int = 45,
    progress_label: str = "正在训练分时段模型",
    label_column: str = TARGET_COLUMN,
    model_backend: str = "xgboost",
    segment_definitions: OrderedDict[str, tuple[int, int]] | None = None,
    high_price_weight_enabled: bool = False,
    high_price_threshold: float | None = None,
    high_price_weight_multiplier: float = 2.0,
    validation_records: list[dict] | None = None,
    model_key: str | None = None,
    model_backend_label: str | None = None,
) -> tuple[dict[str, dict[str, float | int | None]], list[str], list[str]]:
    output_dir.mkdir(parents=True, exist_ok=True)
    segment_definitions = segment_definitions or normalize_segment_config()
    metrics_summary: dict[str, dict[str, float | int | None]] = {}
    unique_dates = sorted(pd.to_datetime(history_df["date"]).dt.strftime("%Y-%m-%d").unique().tolist())
    valid_dates = unique_dates[-valid_days:] if valid_days > 0 and len(unique_dates) > valid_days else []
    train_dates = unique_dates[:-valid_days] if valid_dates else unique_dates

    total_segments = len(segment_definitions)
    for idx, segment_name in enumerate(segment_definitions, start=1):
        emit_progress(progress_callback, f"{progress_label}：{segment_name}", progress_start + int(idx / total_segments * progress_span))
        segment_df = history_df[history_df["segment"] == segment_name].copy()
        segment_df = segment_df.dropna(subset=[label_column])
        if segment_df.empty:
            raise ValueError(f"时段 {segment_name} 无训练数据")
        train_df, valid_df = split_train_valid(segment_df, valid_days)
        sample_weight = high_price_sample_weights(
            train_df,
            label_column,
            high_price_weight_enabled,
            high_price_threshold,
            high_price_weight_multiplier,
        )
        _, best_iteration, final_pred = train_backend_model(
            model_backend,
            train_df,
            valid_df,
            feature_columns,
            label_column,
            num_boost_round,
            output_dir / model_file_name(segment_name, model_backend),
            sample_weight=sample_weight,
        )

        segment_metrics: dict[str, float | int | None] = {
            "train_rows": int(len(train_df)),
            "valid_rows": int(len(valid_df)),
            "best_iteration": best_iteration,
            "high_price_train_rows": int(np.sum(sample_weight > 1)) if sample_weight is not None else 0,
        }
        if not valid_df.empty and final_pred is not None:
            actual = valid_df[TARGET_COLUMN].to_numpy(dtype=float)
            computed_metrics = calculate_metrics(actual, final_pred)
            if validation_records is not None and model_key:
                for row_index, predicted_value in enumerate(final_pred):
                    row = valid_df.iloc[row_index]
                    validation_records.append(
                        {
                            "model_key": model_key,
                            "model_backend": model_backend,
                            "model_backend_label": model_backend_label or model_backend,
                            "segment": segment_name,
                            "date": pd.Timestamp(row["date"]).strftime("%Y-%m-%d"),
                            "period": int(row["period"]),
                            "actual": float(row[TARGET_COLUMN]),
                            "predicted": float(predicted_value),
                        }
                    )
            validation_error_summary = summarize_prediction_errors(
                pd.DataFrame(
                    {
                        "date": pd.to_datetime(valid_df["date"]).dt.strftime("%Y-%m-%d").to_numpy(),
                        "period": valid_df["period"].to_numpy(),
                        "actual": actual,
                        "predicted": final_pred,
                    }
                )
            )
            segment_metrics.update(
                {
                    "final_mae": computed_metrics["mae"],
                    "final_rmse": computed_metrics["rmse"],
                    "direct_mae": computed_metrics["mae"],
                    "direct_rmse": computed_metrics["rmse"],
                    "bias": validation_error_summary["bias"],
                    "max_abs_error": validation_error_summary["max_abs_error"],
                    "p90_abs_error": validation_error_summary["p90_abs_error"],
                    "direction_accuracy": validation_error_summary["direction_accuracy"],
                    "over_rate": validation_error_summary["over_rate"],
                    "under_rate": validation_error_summary["under_rate"],
                }
            )
        metrics_summary[segment_name] = segment_metrics
    emit_progress(progress_callback, "分时段模型训练完成", 75)
    return metrics_summary, train_dates, valid_dates


def direction_accuracy(frame: pd.DataFrame, predicted_column: str = "predicted") -> float | None:
    if frame.empty or "date" not in frame.columns or "period" not in frame.columns:
        return None
    ordered = frame.sort_values(["date", "period"]).copy()
    actual_diff = ordered.groupby("date")["actual"].diff()
    predicted_diff = ordered.groupby("date")[predicted_column].diff()
    valid = actual_diff.notna() & predicted_diff.notna() & (actual_diff != 0)
    if not bool(valid.any()):
        return None
    actual_direction = np.sign(actual_diff[valid].to_numpy(dtype=float))
    predicted_direction = np.sign(predicted_diff[valid].to_numpy(dtype=float))
    return round(float(np.mean(actual_direction == predicted_direction) * 100), 2)


def summarize_prediction_errors(error_df: pd.DataFrame, predicted_column: str = "predicted") -> dict:
    if error_df.empty or predicted_column not in error_df.columns:
        return {
            "rows": 0,
            "mae": None,
            "rmse": None,
            "bias": None,
            "max_abs_error": None,
            "p90_abs_error": None,
            "direction_accuracy": None,
            "over_rate": None,
            "under_rate": None,
        }
    clean_df = error_df.dropna(subset=["actual", predicted_column]).copy()
    if clean_df.empty:
        return {
            "rows": 0,
            "mae": None,
            "rmse": None,
            "bias": None,
            "max_abs_error": None,
            "p90_abs_error": None,
            "direction_accuracy": None,
            "over_rate": None,
            "under_rate": None,
        }
    actual = clean_df["actual"].to_numpy(dtype=float)
    predicted = clean_df[predicted_column].to_numpy(dtype=float)
    errors = predicted - actual
    abs_errors = np.abs(errors)
    metrics = calculate_metrics(actual, predicted)
    return {
        "rows": int(len(clean_df)),
        "mae": metrics["mae"],
        "rmse": metrics["rmse"],
        "bias": round(float(np.mean(errors)), 4),
        "max_abs_error": round(float(np.max(abs_errors)), 4),
        "p90_abs_error": round(float(np.quantile(abs_errors, 0.9)), 4),
        "direction_accuracy": direction_accuracy(clean_df, predicted_column),
        "over_rate": round(float(np.mean(errors > 0) * 100), 2),
        "under_rate": round(float(np.mean(errors < 0) * 100), 2),
    }


def apply_high_price_probability_adjustment(
    frame: pd.DataFrame,
    *,
    price_column: str = "predicted_price",
    high_probability_column: str = "high_price_probability",
    extreme_probability_column: str = "extreme_price_probability",
    high_price_floor: float = 600.0,
    high_probability_threshold: float = 0.30,
    extreme_probability_threshold: float = 0.10,
    max_adjustment: float = 120.0,
    extreme_max_adjustment: float = 180.0,
    min_adjustable_price: float = 250.0,
) -> pd.DataFrame:
    result = frame.copy()
    if price_column not in result.columns:
        result[HIGH_PRICE_ADJUSTED_PRICE_COLUMN] = np.nan
        result[HIGH_PRICE_ADJUSTMENT_COLUMN] = np.nan
        result[HIGH_PRICE_ADJUSTMENT_REASON_COLUMN] = ""
        return result
    prices = pd.to_numeric(result[price_column], errors="coerce")
    high_probability = (
        pd.to_numeric(result[high_probability_column], errors="coerce").fillna(0.0)
        if high_probability_column in result.columns
        else pd.Series(0.0, index=result.index)
    )
    extreme_probability = (
        pd.to_numeric(result[extreme_probability_column], errors="coerce").fillna(0.0)
        if extreme_probability_column in result.columns
        else pd.Series(0.0, index=result.index)
    )

    adjusted = prices.copy()
    eligible = (
        prices.notna()
        & (prices >= float(min_adjustable_price))
        & (prices < float(high_price_floor))
        & (high_probability >= float(high_probability_threshold))
    )
    risk_target = float(high_price_floor) * high_probability + prices * (1.0 - high_probability)
    adjustment_limit = pd.Series(float(max_adjustment), index=result.index)
    adjustment_limit.loc[extreme_probability >= float(extreme_probability_threshold)] = float(extreme_max_adjustment)
    capped_target = pd.concat([risk_target, prices + adjustment_limit], axis=1).min(axis=1)
    adjusted.loc[eligible] = capped_target.loc[eligible]
    adjusted = pd.Series(clip_price_series(adjusted.to_numpy(dtype=float)), index=result.index)
    adjustment = (adjusted - prices).fillna(0.0)
    result[HIGH_PRICE_ADJUSTED_PRICE_COLUMN] = adjusted
    result[HIGH_PRICE_ADJUSTMENT_COLUMN] = adjustment
    result[HIGH_PRICE_ADJUSTMENT_REASON_COLUMN] = np.where(
        eligible & (adjustment > 0),
        np.where(
            extreme_probability >= float(extreme_probability_threshold),
            "extreme_price_probability",
            "high_price_probability",
        ),
        "",
    )
    return result


def apply_scenario_similarity_reference_adjustment(
    frame: pd.DataFrame,
    *,
    price_column: str = "predicted_price",
    similarity_column: str = NET_LOAD_ONLY_SIMILAR_PRICE_COLUMN,
    min_downward_gap: float = 50.0,
    renewable_ratio_floor: float = 0.18,
    renewable_quantile: float = 0.60,
    net_load_quantile: float = 0.40,
    high_probability_column: str = "high_price_probability",
    extreme_probability_column: str = "extreme_price_probability",
    high_probability_block_threshold: float = 0.30,
    extreme_probability_block_threshold: float = 0.10,
    blocked_segments: tuple[str, ...] = ("segment_5",),
) -> pd.DataFrame:
    result = frame.copy()
    required_columns = [price_column, similarity_column, TOTAL_LOAD_COLUMN, RENEWABLE_POWER_COLUMN, "net_load"]
    if any(column not in result.columns for column in required_columns):
        prices = pd.to_numeric(result[price_column], errors="coerce") if price_column in result.columns else pd.Series(np.nan, index=result.index)
        result[SCENARIO_SIMILARITY_ADJUSTED_PRICE_COLUMN] = prices
        result[SIMILARITY_BLEND_WEIGHT_COLUMN] = 0.0
        result[SIMILARITY_ADJUSTMENT_COLUMN] = 0.0
        result[SIMILARITY_ADJUSTMENT_REASON_COLUMN] = ""
        return result

    prices = pd.to_numeric(result[price_column], errors="coerce")
    similar = pd.to_numeric(result[similarity_column], errors="coerce")
    total_load = pd.to_numeric(result[TOTAL_LOAD_COLUMN], errors="coerce")
    renewable_power = pd.to_numeric(result[RENEWABLE_POWER_COLUMN], errors="coerce")
    net_load = pd.to_numeric(result["net_load"], errors="coerce")
    renewable_ratio = (renewable_power / total_load.replace(0, np.nan)).replace([np.inf, -np.inf], np.nan)

    renewable_trigger = pd.Series(False, index=result.index)
    net_load_trigger = pd.Series(False, index=result.index)
    group_source = result["date"] if "date" in result.columns else pd.Series("all", index=result.index)
    for _, index in result.groupby(group_source, dropna=False).groups.items():
        group_index = pd.Index(index)
        group_ratio = renewable_ratio.loc[group_index]
        group_net_load = net_load.loc[group_index]
        if group_ratio.notna().any():
            threshold = max(float(renewable_ratio_floor), float(group_ratio.quantile(float(renewable_quantile))))
            renewable_trigger.loc[group_index] = group_ratio >= threshold
        if group_net_load.notna().any():
            net_load_trigger.loc[group_index] = group_net_load <= float(group_net_load.quantile(float(net_load_quantile)))

    high_probability = (
        pd.to_numeric(result[high_probability_column], errors="coerce").fillna(0.0)
        if high_probability_column in result.columns
        else pd.Series(0.0, index=result.index)
    )
    extreme_probability = (
        pd.to_numeric(result[extreme_probability_column], errors="coerce").fillna(0.0)
        if extreme_probability_column in result.columns
        else pd.Series(0.0, index=result.index)
    )
    blocked_segment_mask = result.get("segment", pd.Series("", index=result.index)).astype(str).isin(blocked_segments)
    model_minus_similarity = prices - similar
    downward_mask = (
        prices.notna()
        & similar.notna()
        & (model_minus_similarity >= float(min_downward_gap))
        & renewable_trigger
        & (net_load_trigger | (renewable_ratio >= 0.25))
        & (high_probability < float(high_probability_block_threshold))
        & (extreme_probability < float(extreme_probability_block_threshold))
        & ~blocked_segment_mask
    )

    gap_weight = ((model_minus_similarity - float(min_downward_gap)) / 200.0).clip(lower=0.0, upper=0.25)
    ratio_weight = (renewable_ratio.fillna(0.0) - float(renewable_ratio_floor)).clip(lower=0.0) * 0.75
    blend_weight = (0.45 + ratio_weight + gap_weight).clip(lower=0.0, upper=0.85).where(downward_mask, 0.0)
    adjusted = prices * (1.0 - blend_weight) + similar * blend_weight
    adjusted = pd.Series(clip_price_series(adjusted.to_numpy(dtype=float)), index=result.index)
    adjustment = (adjusted - prices).fillna(0.0)

    result[SCENARIO_SIMILARITY_ADJUSTED_PRICE_COLUMN] = adjusted
    result[SIMILARITY_BLEND_WEIGHT_COLUMN] = blend_weight.fillna(0.0)
    result[SIMILARITY_ADJUSTMENT_COLUMN] = adjustment
    result[SIMILARITY_ADJUSTMENT_REASON_COLUMN] = np.where(
        downward_mask,
        "renewable_high_net_load_low_model_above_similarity",
        "",
    )
    return result


def summarize_high_price_adjustment_effect(
    frame: pd.DataFrame,
    *,
    original_column: str = "predicted",
    adjusted_column: str = HIGH_PRICE_ADJUSTED_PRICE_COLUMN,
    high_threshold: float = 600.0,
) -> dict:
    overall_original = summarize_prediction_errors(frame, original_column)
    overall_adjusted = summarize_prediction_errors(frame, adjusted_column)
    high_frame = filter_price_range(frame, price_min=high_threshold)
    high_original = summarize_prediction_errors(high_frame, original_column)
    high_adjusted = summarize_prediction_errors(high_frame, adjusted_column)

    def delta(before: dict, after: dict) -> dict:
        before_mae = before.get("mae")
        after_mae = after.get("mae")
        if before_mae is None or after_mae is None:
            return {"mae_delta": None, "mae_delta_pct": None}
        mae_delta = round(float(after_mae) - float(before_mae), 4)
        mae_delta_pct = None if float(before_mae) == 0 else round(mae_delta / float(before_mae) * 100, 2)
        return {"mae_delta": mae_delta, "mae_delta_pct": mae_delta_pct}

    return {
        "overall": {
            "original": overall_original,
            "adjusted": overall_adjusted,
            **delta(overall_original, overall_adjusted),
        },
        "high_price": {
            "threshold": float(high_threshold),
            "original": high_original,
            "adjusted": high_adjusted,
            **delta(high_original, high_adjusted),
        },
    }


def build_selected_validation_adjustment_frame(
    validation_records: list[dict] | pd.DataFrame,
    selected_segment_price_models: dict[str, str],
    interval_feature_frame: pd.DataFrame,
) -> pd.DataFrame:
    records = pd.DataFrame(validation_records)
    if records.empty or not selected_segment_price_models:
        return pd.DataFrame()
    records = records.copy()
    records["date"] = pd.to_datetime(records["date"], errors="coerce").dt.normalize()
    selected_mask = records.apply(
        lambda row: str(row.get("model_key")) == str(selected_segment_price_models.get(str(row.get("segment")))),
        axis=1,
    )
    selected = records.loc[selected_mask].copy()
    if selected.empty:
        return selected

    interval_columns = [
        "date",
        "period",
        "segment",
        "predicted_interval",
        "predicted_interval_probability",
        "high_price_probability",
        "extreme_price_probability",
        "interval_backtest_accuracy",
        "interval_probabilities",
    ]
    interval_features = interval_feature_frame[[column for column in interval_columns if column in interval_feature_frame.columns]].copy()
    if interval_features.empty:
        return apply_high_price_probability_adjustment(selected, price_column="predicted")
    interval_features["date"] = pd.to_datetime(interval_features["date"], errors="coerce").dt.normalize()
    merged = selected.merge(
        interval_features,
        on=["date", "period", "segment"],
        how="left",
    )
    return apply_high_price_probability_adjustment(merged, price_column="predicted")


def filter_price_range(frame: pd.DataFrame, price_min: float | None = None, price_max: float | None = None) -> pd.DataFrame:
    if frame.empty or "actual" not in frame.columns:
        return frame.iloc[0:0].copy()
    result = frame.copy()
    actual = pd.to_numeric(result["actual"], errors="coerce")
    mask = actual.notna()
    if price_min is not None:
        mask &= actual >= float(price_min)
    if price_max is not None:
        mask &= actual <= float(price_max)
    return result.loc[mask].copy()


def candidate_metric_value(summary: dict, metric: str) -> float | None:
    if metric == "default_score":
        mae = summary.get("mae")
        rmse = summary.get("rmse")
        if mae is None or rmse is None:
            return None
        return float(mae) + 0.2 * float(rmse)
    value = summary.get(metric)
    return None if value is None else float(value)


def metric_sort_reverse(metric: str) -> bool:
    return metric in {"direction_accuracy"}


def segment_model_score(metrics: dict[str, float | int | None]) -> float | None:
    mae = metrics.get("final_mae") if metrics.get("final_mae") is not None else metrics.get("mae")
    rmse = metrics.get("final_rmse") if metrics.get("final_rmse") is not None else metrics.get("rmse")
    if mae is None or rmse is None:
        return None
    return float(mae) + 0.2 * float(rmse)


def select_best_segment_price_models(
    variant_metrics: dict[str, dict[str, dict[str, float | int | None]]],
    segment_names: list[str],
    fallback_model_key: str,
) -> dict[str, str]:
    selected: dict[str, str] = {}
    for segment_name in segment_names:
        candidates: list[tuple[float, str]] = []
        for model_key, metrics_by_segment in variant_metrics.items():
            score = segment_model_score(metrics_by_segment.get(segment_name, {}))
            if score is not None:
                candidates.append((score, str(model_key)))
        selected[segment_name] = min(candidates, key=lambda item: (item[0], item[1]))[1] if candidates else fallback_model_key
    return selected


def load_validation_prediction_records(model_dir: Path, metadata: dict) -> pd.DataFrame:
    relative_path = metadata.get("validation_predictions_path") or "validation_predictions.jsonl"
    path = model_dir / str(relative_path)
    if not path.exists():
        return pd.DataFrame()
    rows = []
    with path.open("r", encoding="utf-8") as file:
        for line in file:
            line = line.strip()
            if not line:
                continue
            try:
                rows.append(json.loads(line))
            except json.JSONDecodeError:
                continue
    return pd.DataFrame(rows)


def rank_model_candidates(
    model_root: Path = DEFAULT_MODEL_ROOT,
    metric: str = "default_score",
    price_min: float | None = None,
    price_max: float | None = None,
) -> dict:
    active_dir = resolve_active_model_dir(model_root)
    metadata = read_model_metadata(active_dir) or {}
    records = load_validation_prediction_records(active_dir, metadata)
    allowed_metrics = {"default_score", "mae", "rmse", "p90_abs_error", "max_abs_error", "direction_accuracy", "mae_range"}
    metric_key = metric if metric in allowed_metrics else "default_score"
    metric_for_summary = "mae" if metric_key == "mae_range" else metric_key
    scoped_records = filter_price_range(records, price_min, price_max) if metric_key == "mae_range" else records
    selected_segment_models = metadata.get("selected_segment_price_models") if isinstance(metadata.get("selected_segment_price_models"), dict) else {}
    segments: list[dict] = []
    if scoped_records.empty:
        return {
            "metric": metric_key,
            "price_min": price_min,
            "price_max": price_max,
            "segments": [],
            "selected_segment_price_models": selected_segment_models,
            "available_metrics": sorted(allowed_metrics),
            "message": "当前模型缺少候选模型验证明细，请重新训练模型",
        }
    for segment_name, segment_df in scoped_records.groupby("segment"):
        rankings = []
        for model_key, model_df in segment_df.groupby("model_key"):
            summary = summarize_prediction_errors(model_df)
            value = candidate_metric_value(summary, metric_for_summary)
            rankings.append(
                {
                    "model_key": str(model_key),
                    "model_backend": str(model_df["model_backend"].iloc[0]) if "model_backend" in model_df.columns else "",
                    "model_backend_label": str(model_df["model_backend_label"].iloc[0]) if "model_backend_label" in model_df.columns else str(model_key),
                    "metric_value": value,
                    "rows": summary.get("rows"),
                    "mae": summary.get("mae"),
                    "rmse": summary.get("rmse"),
                    "p90_abs_error": summary.get("p90_abs_error"),
                    "max_abs_error": summary.get("max_abs_error"),
                    "direction_accuracy": summary.get("direction_accuracy"),
                }
            )
        reverse = metric_sort_reverse(metric_for_summary)
        rankings.sort(key=lambda item: (item["metric_value"] is None, -(item["metric_value"] or 0) if reverse else (item["metric_value"] or float("inf")), item["model_key"]))
        segments.append(
            {
                "segment": str(segment_name),
                "current_model_key": selected_segment_models.get(str(segment_name)) or metadata.get("selected_model_key"),
                "rankings": rankings,
            }
        )
    segment_order = {str(item.get("name")): index for index, item in enumerate(metadata.get("segments") or [])}
    segments.sort(key=lambda item: segment_order.get(item["segment"], 999))
    return {
        "metric": metric_key,
        "price_min": price_min,
        "price_max": price_max,
        "segments": segments,
        "selected_segment_price_models": selected_segment_models,
        "available_metrics": sorted(allowed_metrics),
    }


def summarize_model_vs_baseline(error_df: pd.DataFrame, baseline_column: str = "net_load_only_similar_predicted") -> dict:
    if error_df.empty or baseline_column not in error_df.columns:
        return {
            "model": summarize_prediction_errors(error_df),
            "baseline": summarize_prediction_errors(pd.DataFrame()),
            "mae_improvement": None,
            "mae_improvement_pct": None,
        }
    model_summary = summarize_prediction_errors(error_df)
    baseline_summary = summarize_prediction_errors(error_df, baseline_column)
    model_mae = model_summary.get("mae")
    baseline_mae = baseline_summary.get("mae")
    if model_mae is None or baseline_mae is None or float(baseline_mae) == 0:
        improvement = None
        improvement_pct = None
    else:
        improvement = round(float(baseline_mae) - float(model_mae), 4)
        improvement_pct = round(improvement / float(baseline_mae) * 100, 2)
    return {
        "model": model_summary,
        "baseline": baseline_summary,
        "mae_improvement": improvement,
        "mae_improvement_pct": improvement_pct,
    }


def top_daily_error_days(error_df: pd.DataFrame, limit: int = 5) -> dict[str, list[dict]]:
    if error_df.empty or "date" not in error_df.columns:
        return {"best": [], "worst": []}
    rows: list[dict] = []
    for date_value, date_df in error_df.groupby("date"):
        summary = summarize_prediction_errors(date_df)
        if summary.get("mae") is None:
            continue
        rows.append(
            {
                "date": str(date_value),
                "mae": summary.get("mae"),
                "rmse": summary.get("rmse"),
                "max_abs_error": summary.get("max_abs_error"),
                "direction_accuracy": summary.get("direction_accuracy"),
            }
        )
    return {
        "best": sorted(rows, key=lambda item: float(item["mae"]))[:limit],
        "worst": sorted(rows, key=lambda item: float(item["mae"]), reverse=True)[:limit],
    }


def rolling_backtest_cache_key(target: str, run_id: str, horizon: int, config_payload: dict[str, object]) -> str:
    payload = {
        "target": normalize_model_target(target),
        "run_id": str(run_id),
        "horizon": int(horizon),
        "config": config_payload,
    }
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def rolling_backtest_cache_path(model_root: str | Path, cache_key: str) -> Path:
    return Path(model_root) / "cache" / "rolling_backtest" / f"{cache_key}.json"


def rolling_backtest_selected_model(
    history_df: pd.DataFrame,
    feature_columns: list[str],
    model_backend: str,
    training_window_days: int | None,
    num_boost_round: int,
    horizons: tuple[int, ...] = (14, 30),
    similarity_reference_days: int = DEFAULT_SIMILARITY_REFERENCE_DAYS,
    segment_definitions: OrderedDict[str, tuple[int, int]] | None = None,
    high_price_weight_enabled: bool = False,
    high_price_threshold: float | None = None,
    high_price_weight_multiplier: float = 2.0,
    cache_root: str | Path | None = None,
    cache_target: str = DAYAHEAD_DATA_MODE,
    cache_run_id: str | None = None,
    cache_config_payload: dict[str, object] | None = None,
) -> dict:
    segment_definitions = segment_definitions or normalize_segment_config()
    unique_dates = sorted(pd.to_datetime(history_df["date"], errors="coerce").dropna().dt.strftime("%Y-%m-%d").unique().tolist())
    results: dict[str, dict] = {}
    if len(unique_dates) < 3:
        return results
    baseline_history_df: pd.DataFrame | None = None

    for horizon in horizons:
        cache_path: Path | None = None
        if cache_root is not None and cache_run_id:
            cache_key = rolling_backtest_cache_key(cache_target, cache_run_id, horizon, cache_config_payload or {})
            cache_path = rolling_backtest_cache_path(cache_root, cache_key)
            if cache_path.exists():
                try:
                    cached_result = json.loads(cache_path.read_text(encoding="utf-8"))
                except (OSError, json.JSONDecodeError):
                    cached_result = None
                if isinstance(cached_result, dict):
                    results[str(horizon)] = cached_result
                    continue
        if baseline_history_df is None:
            feature_history_df = attach_lag_features(history_df)
            baseline_history_df = attach_similarity_features(
                feature_history_df,
                reference_days=similarity_reference_days,
            )
        test_dates = unique_dates[-min(horizon, len(unique_dates) - 1):]
        predictions: list[pd.DataFrame] = []
        for test_date in test_dates:
            previous_dates = [item for item in unique_dates if item < test_date]
            if training_window_days:
                previous_dates = previous_dates[-int(training_window_days):]
            if not previous_dates:
                continue
            train_pool = baseline_history_df[baseline_history_df["date"].dt.strftime("%Y-%m-%d").isin(previous_dates)].copy()
            test_pool = baseline_history_df[baseline_history_df["date"].dt.strftime("%Y-%m-%d") == test_date].copy()
            for segment_name in segment_definitions:
                train_df = train_pool[train_pool["segment"] == segment_name].dropna(subset=[TARGET_COLUMN]).copy()
                test_df = test_pool[test_pool["segment"] == segment_name].dropna(subset=[TARGET_COLUMN]).copy()
                if train_df.empty or test_df.empty:
                    continue
                sample_weight = high_price_sample_weights(
                    train_df,
                    TARGET_COLUMN,
                    high_price_weight_enabled,
                    high_price_threshold,
                    high_price_weight_multiplier,
                )
                model, _, _ = train_backend_model(
                    model_backend,
                    train_df,
                    pd.DataFrame(columns=train_df.columns),
                    feature_columns,
                    TARGET_COLUMN,
                    num_boost_round,
                    None,
                    sample_weight=sample_weight,
                )
                pred = predict_backend_model(model, model_backend, test_df, feature_columns)
                predictions.append(
                    pd.DataFrame(
                        {
                            "date": test_date,
                            "segment": segment_name,
                            "period": test_df["period"].to_numpy(),
                            "actual": test_df[TARGET_COLUMN].to_numpy(dtype=float),
                            "predicted": pred,
                            "net_load_only_similar_predicted": test_df[NET_LOAD_ONLY_SIMILAR_PRICE_COLUMN].to_numpy(dtype=float)
                            if NET_LOAD_ONLY_SIMILAR_PRICE_COLUMN in test_df.columns
                            else np.full(len(test_df), np.nan),
                        }
                    )
                )
        if not predictions:
            results[str(horizon)] = {"rows": 0, "overall": summarize_prediction_errors(pd.DataFrame())}
            if cache_path is not None:
                write_json(cache_path, results[str(horizon)])
            continue
        error_df = pd.concat(predictions, ignore_index=True)
        segment_metrics = {
            segment_name: summarize_prediction_errors(segment_df)
            for segment_name, segment_df in error_df.groupby("segment")
        }
        segment_net_load_only_baseline_metrics = {
            segment_name: summarize_prediction_errors(segment_df, "net_load_only_similar_predicted")
            for segment_name, segment_df in error_df.groupby("segment")
        }
        high_threshold = float(error_df["actual"].quantile(0.9))
        low_threshold = float(error_df["actual"].quantile(0.1))
        high_spike_df = error_df[error_df["actual"] >= high_threshold]
        low_spike_df = error_df[error_df["actual"] <= low_threshold]
        model_vs_net_load_only_baseline = summarize_model_vs_baseline(error_df, "net_load_only_similar_predicted")
        results[str(horizon)] = {
            "rows": int(len(error_df)),
            "test_date_start": min(test_dates) if test_dates else None,
            "test_date_end": max(test_dates) if test_dates else None,
            "overall": summarize_prediction_errors(error_df),
            "net_load_only_baseline": summarize_prediction_errors(error_df, "net_load_only_similar_predicted"),
            "net_load_only_baseline_label": "净负荷相似法预测价格",
            "model_vs_net_load_only_similarity": model_vs_net_load_only_baseline,
            "segments": segment_metrics,
            "segment_net_load_only_baseline": segment_net_load_only_baseline_metrics,
            "daily_error_rank": top_daily_error_days(error_df),
            "spike_errors": {
                "high_threshold": round(high_threshold, 4),
                "high": summarize_prediction_errors(high_spike_df),
                "high_net_load_only_baseline": summarize_prediction_errors(high_spike_df, "net_load_only_similar_predicted"),
                "high_model_vs_net_load_only_similarity": summarize_model_vs_baseline(high_spike_df, "net_load_only_similar_predicted"),
                "low_threshold": round(low_threshold, 4),
                "low": summarize_prediction_errors(low_spike_df),
                "low_net_load_only_baseline": summarize_prediction_errors(low_spike_df, "net_load_only_similar_predicted"),
            },
        }
        if cache_path is not None:
            write_json(cache_path, results[str(horizon)])
    return results



def ensure_model_dirs(model_root: Path) -> tuple[Path, Path, Path]:
    current_dir = model_root / "current"
    previous_dir = model_root / "previous"
    history_dir = model_root / "history"
    model_root.mkdir(parents=True, exist_ok=True)
    history_dir.mkdir(parents=True, exist_ok=True)
    return current_dir, previous_dir, history_dir


def retry_filesystem_operation(description: str, operation: Callable[[], object], retries: int = 10, delay_seconds: float = 0.3) -> object:
    last_error: Exception | None = None
    for attempt in range(retries):
        try:
            return operation()
        except (OSError, shutil.Error) as exc:
            last_error = exc
            gc.collect()
            if attempt < retries - 1:
                time.sleep(delay_seconds * (attempt + 1))
    raise RuntimeError(f"{description}失败：{last_error}") from last_error


def remove_tree_with_retry(path: Path) -> None:
    if not path.exists():
        return
    retry_filesystem_operation(f"删除目录 {path}", lambda: shutil.rmtree(path))


def copy_tree_with_retry(source: Path, target: Path) -> None:
    def operation() -> None:
        if target.exists():
            shutil.rmtree(target)
        shutil.copytree(source, target)

    retry_filesystem_operation(f"复制目录 {source} 到 {target}", operation)


def write_json(path: Path, data: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def append_jsonl(path: Path, data: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as file:
        file.write(json.dumps(data, ensure_ascii=False) + "\n")


def training_preferences_path(model_root: str | Path = DEFAULT_MODEL_ROOT) -> Path:
    return Path(model_root) / "training_preferences.json"


def default_training_preferences() -> dict:
    return {
        "segment_mode": "default",
        "segment_config": DEFAULT_SEGMENT_CONFIG,
        "high_price_weighting": {"enabled": False, "quantile": 0.8, "multiplier": 2.0},
        "price_intervals": normalize_price_intervals(DEFAULT_PRICE_INTERVALS),
        "prediction_reference": {
            "recent_reference_days": 1,
            "same_type_reference_days": 1,
            "knn_similarity": dict(DEFAULT_KNN_SIMILARITY_CONFIG),
        },
    }


def normalize_prediction_reference_preferences(preferences: dict[str, object] | None = None) -> dict[str, object]:
    preferences = preferences or {}
    return {
        "recent_reference_days": normalize_reference_days(preferences.get("recent_reference_days") or 1),
        "same_type_reference_days": normalize_reference_days(preferences.get("same_type_reference_days") or 1),
        "knn_similarity": normalize_knn_similarity_config(
            preferences.get("knn_similarity") if isinstance(preferences.get("knn_similarity"), dict) else preferences
        ),
    }


def load_training_preferences(model_root: str | Path = DEFAULT_MODEL_ROOT) -> dict:
    path = training_preferences_path(model_root)
    if not path.exists():
        return default_training_preferences()
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return default_training_preferences()
    preferences = default_training_preferences()
    if isinstance(data, dict) and data.get("segment_mode") == "custom" and isinstance(data.get("segment_config"), list):
        preferences["segment_mode"] = "custom"
        preferences["segment_config"] = segment_metadata(normalize_segment_config(data["segment_config"]))
    if isinstance(data, dict) and isinstance(data.get("high_price_weighting"), dict):
        high_config = data["high_price_weighting"]
        preferences["high_price_weighting"] = {
            "enabled": bool(high_config.get("enabled", False)),
            "quantile": min(0.99, max(0.5, float(high_config.get("quantile", 0.8)))),
            "multiplier": max(1.0, float(high_config.get("multiplier", 2.0))),
        }
    if isinstance(data, dict) and isinstance(data.get("price_intervals"), list):
        preferences["price_intervals"] = normalize_price_intervals(data["price_intervals"], allow_legacy_open_bounds=True)
    if isinstance(data, dict) and isinstance(data.get("prediction_reference"), dict):
        preferences["prediction_reference"] = normalize_prediction_reference_preferences(data["prediction_reference"])
    return preferences


def save_training_preferences(model_root: str | Path, preferences: dict) -> dict:
    current = load_training_preferences(model_root)
    if preferences.get("segment_mode") in {"default", "custom"}:
        current["segment_mode"] = preferences.get("segment_mode")
    if isinstance(preferences.get("segment_config"), list):
        current["segment_config"] = segment_metadata(normalize_segment_config(preferences["segment_config"]))
    if isinstance(preferences.get("high_price_weighting"), dict):
        high_config = preferences["high_price_weighting"]
        current["high_price_weighting"] = {
            "enabled": bool(high_config.get("enabled", False)),
            "quantile": min(0.99, max(0.5, float(high_config.get("quantile", 0.8)))),
            "multiplier": max(1.0, float(high_config.get("multiplier", 2.0))),
        }
    if isinstance(preferences.get("price_intervals"), list):
        current["price_intervals"] = normalize_price_intervals(preferences["price_intervals"])
    if isinstance(preferences.get("prediction_reference"), dict):
        current["prediction_reference"] = normalize_prediction_reference_preferences(preferences["prediction_reference"])
    path = training_preferences_path(model_root)
    write_json(path, current)
    return current


def default_model_pointer_path(model_root: Path) -> Path:
    return model_root / "default_model.json"


def read_pinned_default_version(model_root: Path) -> str | None:
    pointer_path = default_model_pointer_path(model_root)
    if not pointer_path.exists():
        return None
    try:
        data = json.loads(pointer_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return None
    version_key = str(data.get("version_key", "")).strip()
    if not version_key:
        return None
    if not (model_root / "history" / version_key).exists():
        try:
            pointer_path.unlink()
        except OSError:
            pass
        return None
    return version_key


def write_pinned_default_version(model_root: Path, version_key: str | None) -> None:
    pointer_path = default_model_pointer_path(model_root)
    if not version_key:
        if pointer_path.exists():
            pointer_path.unlink()
        return
    write_json(
        pointer_path,
        {
            "version_key": version_key,
            "updated_at": datetime.now().isoformat(timespec="seconds"),
        },
    )


def find_version_key_by_run_id(model_root: Path, run_id: str | None) -> str | None:
    if not run_id:
        return None
    _, _, history_dir = ensure_model_dirs(model_root)
    for run_dir in history_dir.iterdir():
        if not run_dir.is_dir() or not run_dir.name.startswith("run_"):
            continue
        metadata = read_model_metadata(run_dir)
        if metadata and metadata.get("run_id") == run_id:
            return run_dir.name
    return None


def finalize_model_version(model_root: Path, staging_dir: Path, run_id: str) -> tuple[Path, Path]:
    current_dir, previous_dir, history_dir = ensure_model_dirs(model_root)
    run_dir = history_dir / run_id
    if current_dir.exists():
        if previous_dir.exists():
            remove_tree_with_retry(previous_dir)
        copy_tree_with_retry(current_dir, previous_dir)
        remove_tree_with_retry(current_dir)
    copy_tree_with_retry(staging_dir, current_dir)
    if run_dir.exists():
        remove_tree_with_retry(run_dir)
    copy_tree_with_retry(staging_dir, run_dir)
    remove_tree_with_retry(staging_dir)
    return current_dir, run_dir


def train_and_register(config: TrainConfig, progress_callback: ProgressCallback | None = None) -> TrainResult:
    ensure_xgboost_available()
    emit_progress(progress_callback, "开始加载历史数据", 5)
    history_df, skipped_sheets, quality_report_path = build_training_frame(config, progress_callback=progress_callback)
    history_signature = history_df.attrs.get("history_signature") or training_history_signature(config)
    run_fingerprint = training_run_fingerprint(config, history_signature)
    emit_progress(progress_callback, "历史数据加载完成，开始生成训练样本", 20)
    segment_definitions = normalize_segment_config(config.segment_config)
    segment_config_metadata = segment_metadata(segment_definitions)
    history_df["segment"] = assign_segments(history_df["period"], segment_definitions)
    full_feature_history_df = history_df.attrs.get("full_feature_history_df")
    if not isinstance(full_feature_history_df, pd.DataFrame):
        full_feature_history_df = history_df.copy()
    interval_valid_days = int(config.interval_valid_days if config.interval_valid_days is not None else config.valid_days)
    interval_training_window_days = (
        int(config.interval_training_window_days)
        if config.interval_training_window_days is not None
        else config.training_window_days
    )
    interval_num_boost_round = int(
        config.interval_num_boost_round if config.interval_num_boost_round is not None else config.num_boost_round
    )
    interval_start_date = config.interval_start_date if config.interval_start_date is not None else config.start_date
    interval_end_date = config.interval_end_date if config.interval_end_date is not None else config.end_date
    thermal_capacity_valid_days = int(
        config.thermal_capacity_valid_days if config.thermal_capacity_valid_days is not None else config.valid_days
    )
    thermal_capacity_training_window_days = (
        int(config.thermal_capacity_training_window_days)
        if config.thermal_capacity_training_window_days is not None
        else config.training_window_days
    )
    thermal_capacity_num_boost_round = int(
        config.thermal_capacity_num_boost_round if config.thermal_capacity_num_boost_round is not None else config.num_boost_round
    )
    thermal_capacity_start_date = config.thermal_capacity_start_date if config.thermal_capacity_start_date is not None else config.start_date
    thermal_capacity_end_date = config.thermal_capacity_end_date if config.thermal_capacity_end_date is not None else config.end_date
    interval_history_df = pd.DataFrame()
    if config.train_interval_model:
        interval_history_df = model_training_frame_from_base(
            full_feature_history_df,
            training_window_days=interval_training_window_days,
            valid_days=interval_valid_days,
            start_date=interval_start_date,
            end_date=interval_end_date,
            similarity_reference_days=config.similarity_reference_days,
        )
        interval_history_df["segment"] = assign_segments(interval_history_df["period"], segment_definitions)
    thermal_capacity_history_df = pd.DataFrame()
    if config.train_thermal_capacity_model:
        thermal_capacity_history_df = model_training_frame_from_base(
            full_feature_history_df,
            training_window_days=thermal_capacity_training_window_days,
            valid_days=thermal_capacity_valid_days,
            start_date=thermal_capacity_start_date,
            end_date=thermal_capacity_end_date,
            similarity_reference_days=config.similarity_reference_days,
        )
    high_price_threshold = None
    if config.high_price_weight_enabled:
        quantile = min(0.99, max(0.5, float(config.high_price_quantile or 0.8)))
        high_price_threshold = float(pd.to_numeric(history_df[TARGET_COLUMN], errors="coerce").quantile(quantile))
    run_id = datetime.now().strftime("run_%Y%m%d_%H%M%S")
    staging_dir = config.model_root / "_staging" / run_id
    staging_dir.mkdir(parents=True, exist_ok=True)
    feature_variants = direct_model_variants()
    if config.direct_feature_variant_keys:
        allowed_variants = set(config.direct_feature_variant_keys)
        feature_variants = OrderedDict(
            (key, value)
            for key, value in feature_variants.items()
            if key in allowed_variants
        )
    model_backends = available_model_backends()
    if config.model_backend_keys:
        allowed_backends = set(config.model_backend_keys)
        model_backends = OrderedDict(
            (key, value)
            for key, value in model_backends.items()
            if key in allowed_backends
        )
    if not feature_variants:
        raise RuntimeError("没有可用的直接价格模型特征组合")
    if not model_backends:
        raise RuntimeError("没有可用的模型训练后端")
    model_variants: OrderedDict[str, dict[str, object]] = OrderedDict()
    variant_metrics: dict[str, dict[str, dict[str, float | int | None]]] = {}
    validation_records: list[dict] = []
    train_dates: list[str] = []
    valid_dates: list[str] = []
    total_variants = max(1, len(feature_variants) * len(model_backends))
    variant_index = 0
    for feature_variant_key, feature_columns in feature_variants.items():
        for backend_key, backend_label in model_backends.items():
            variant_index += 1
            variant_key = f"{feature_variant_key}_{backend_key}"
            variant_dir = staging_dir / variant_key
            progress_start = 25 + int((variant_index - 1) / total_variants * 50)
            progress_span = max(8, int(45 / total_variants))
            metrics_summary, variant_train_dates, variant_valid_dates = train_segment_models(
                history_df,
                variant_dir,
                config.valid_days,
                config.num_boost_round,
                feature_columns,
                progress_callback=progress_callback,
                progress_start=progress_start,
                progress_span=progress_span,
                progress_label=f"正在训练 {backend_label} 直接价格模型",
                label_column=TARGET_COLUMN,
                model_backend=backend_key,
                segment_definitions=segment_definitions,
                high_price_weight_enabled=config.high_price_weight_enabled,
                high_price_threshold=high_price_threshold,
                high_price_weight_multiplier=config.high_price_weight_multiplier,
                validation_records=validation_records,
                model_key=variant_key,
                model_backend_label=backend_label,
            )
            model_variants[variant_key] = {
                "feature_columns": feature_columns,
                "model_dir": variant_key,
                "model_backend": backend_key,
                "model_backend_label": backend_label,
                "base_variant": feature_variant_key,
            }
            variant_metrics[variant_key] = metrics_summary
            if not train_dates:
                train_dates = variant_train_dates
            if not valid_dates:
                valid_dates = variant_valid_dates
    variant_quality_metrics = {key: summarize_model_quality_metrics(metrics) for key, metrics in variant_metrics.items()}

    def variant_score(item: tuple[str, dict[str, float | None]]) -> tuple[float, str]:
        key, quality = item
        mae = quality.get("final_mae")
        rmse = quality.get("final_rmse")
        if mae is None or rmse is None:
            return float("inf"), key
        return float(mae) + 0.2 * float(rmse), key

    selected_model_key = min(variant_quality_metrics.items(), key=variant_score)[0]
    segment_names = [str(item.get("name")) for item in segment_config_metadata]
    selected_segment_price_models = select_best_segment_price_models(variant_metrics, segment_names, selected_model_key)
    metrics_summary = {
        segment_name: variant_metrics.get(selected_segment_price_models.get(segment_name, selected_model_key), {}).get(
            segment_name,
            variant_metrics[selected_model_key].get(segment_name, {}),
        )
        for segment_name in segment_names
    }
    selected_variant_meta = model_variants[selected_model_key]
    price_intervals = normalize_price_intervals(config.price_intervals or load_training_preferences(config.model_root).get("price_intervals"))
    if config.train_interval_model:
        interval_backends = available_interval_backends()
        if config.interval_backend_keys:
            allowed_interval_backends = set(config.interval_backend_keys)
            interval_backends = {
                key: value
                for key, value in interval_backends.items()
                if key in allowed_interval_backends
            }
        if not interval_backends:
            raise RuntimeError("没有可用的价格区间模型训练后端")
        emit_progress(progress_callback, "正在训练价格区间分类模型", 76)
        price_interval_model = train_interval_models(
            interval_history_df,
            target_column=TARGET_COLUMN,
            feature_columns=list(selected_variant_meta["feature_columns"]),
            segment_config=segment_config_metadata,
            output_dir=staging_dir,
            intervals=price_intervals,
            valid_days=interval_valid_days,
            num_boost_round=interval_num_boost_round,
            backends=interval_backends,
        )
    else:
        price_interval_model = {"enabled": False, "metrics": {}, "model_variants": {}}
    if config.train_thermal_capacity_model:
        emit_progress(progress_callback, "正在训练日级火电开机容量模型", 77)
        thermal_capacity_backends = available_thermal_capacity_backends()
        if config.thermal_capacity_backend_keys:
            allowed_thermal_backends = set(config.thermal_capacity_backend_keys)
            thermal_capacity_backends = OrderedDict(
                (key, value)
                for key, value in thermal_capacity_backends.items()
                if key in allowed_thermal_backends
            )
        if not thermal_capacity_backends:
            raise RuntimeError("没有可用的开机容量模型训练后端")
        thermal_capacity_model = train_thermal_capacity_model(
            thermal_capacity_history_df,
            staging_dir / "thermal_capacity_model",
            valid_days=thermal_capacity_valid_days,
            num_boost_round=thermal_capacity_num_boost_round,
            backends=dict(thermal_capacity_backends),
        )
    else:
        thermal_capacity_model = {"enabled": False, "metrics": {}, "reason": "disabled"}
    high_price_adjustment_metrics = {}
    try:
        interval_bundle = load_interval_model_bundle(staging_dir, price_interval_model)
        if interval_bundle is not None and valid_dates:
            validation_feature_df = history_df[
                pd.to_datetime(history_df["date"], errors="coerce").dt.strftime("%Y-%m-%d").isin(valid_dates)
            ].copy()
            interval_prediction = predict_interval_probabilities(validation_feature_df, interval_bundle)
            validation_intervals = normalize_price_intervals(price_interval_model.get("intervals"))
            interval_feature_df = append_interval_prediction_columns(
                validation_feature_df,
                interval_prediction,
                validation_intervals,
                price_column=TARGET_COLUMN,
            )
            selected_validation_adjustment_df = build_selected_validation_adjustment_frame(
                validation_records,
                selected_segment_price_models,
                interval_feature_df,
            )
            high_price_adjustment_metrics = summarize_high_price_adjustment_effect(
                selected_validation_adjustment_df,
                original_column="predicted",
                adjusted_column=HIGH_PRICE_ADJUSTED_PRICE_COLUMN,
                high_threshold=600.0,
            )
    except Exception as exc:
        high_price_adjustment_metrics = {"error": str(exc)}
    if config.enable_rolling_backtest:
        horizon_text = "/".join(str(item) for item in config.rolling_backtest_horizons) or "14/30"
        emit_progress(progress_callback, f"正在执行最近 {horizon_text} 天滚动回测", 78)
        rolling_backtest_metrics = rolling_backtest_selected_model(
            history_df,
            list(selected_variant_meta["feature_columns"]),
            str(selected_variant_meta["model_backend"]),
            config.training_window_days,
            config.num_boost_round,
            horizons=tuple(config.rolling_backtest_horizons or (14, 30)),
            similarity_reference_days=normalize_similarity_reference_days(config.similarity_reference_days),
            segment_definitions=segment_definitions,
            high_price_weight_enabled=config.high_price_weight_enabled,
            high_price_threshold=high_price_threshold,
            high_price_weight_multiplier=config.high_price_weight_multiplier,
            cache_root=config.feature_cache_root or config.model_root,
            cache_target=str(config.data_mode or DAYAHEAD_DATA_MODE),
            cache_run_id=run_fingerprint,
            cache_config_payload={
                "feature_columns": list(selected_variant_meta["feature_columns"]),
                "model_backend": str(selected_variant_meta["model_backend"]),
                "training_window_days": config.training_window_days,
                "num_boost_round": config.num_boost_round,
                "similarity_reference_days": normalize_similarity_reference_days(config.similarity_reference_days),
                "segment_config": segment_config_metadata,
                "high_price_weight_enabled": config.high_price_weight_enabled,
                "high_price_threshold": high_price_threshold,
                "high_price_weight_multiplier": config.high_price_weight_multiplier,
            },
        )
    else:
        rolling_backtest_metrics = {}

    metadata = {
        "run_id": run_id,
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "model_type": "direct_price_multi_model",
        "target_column": TARGET_COLUMN,
        "target_market": REALTIME_DATA_MODE if str(config.data_mode or DAYAHEAD_DATA_MODE) == REALTIME_DATA_MODE else DAYAHEAD_DATA_MODE,
        "target_source_column": REALTIME_TARGET_SOURCE_COLUMN if str(config.data_mode or DAYAHEAD_DATA_MODE) == REALTIME_DATA_MODE else TARGET_COLUMN,
        "data_mode": str(config.data_mode or DAYAHEAD_DATA_MODE),
        "training_run_fingerprint": run_fingerprint,
        "feature_columns": selected_variant_meta["feature_columns"],
        "model_variants": dict(model_variants),
        "selected_model_key": selected_model_key,
        "selected_segment_price_models": selected_segment_price_models,
        "selected_model_backend": selected_variant_meta["model_backend"],
        "selected_model_backend_label": selected_variant_meta["model_backend_label"],
        "model_backend_candidates": dict(model_backends),
        "variant_quality_metrics": variant_quality_metrics,
        "price_interval_model": price_interval_model,
        "thermal_capacity_model": thermal_capacity_model,
        "rolling_backtest_metrics": rolling_backtest_metrics,
        "rolling_backtest_horizons": list(config.rolling_backtest_horizons or (14, 30)),
        "segments": segment_config_metadata,
        "segment_config": segment_config_metadata,
        "segment_mode": "custom" if config.segment_config else "default",
        "high_price_weighting": {
            "enabled": bool(config.high_price_weight_enabled),
            "quantile": float(config.high_price_quantile),
            "multiplier": float(config.high_price_weight_multiplier),
            "threshold": round(float(high_price_threshold), 4) if high_price_threshold is not None else None,
        },
        "valid_days": config.valid_days,
        "training_window_days": config.training_window_days,
        "training_mode": "rolling_window" if config.training_window_days else "manual_date_range",
        "num_boost_round": config.num_boost_round,
        "price_model_training_config": {
            "training_window_days": config.training_window_days,
            "training_mode": "rolling_window" if config.training_window_days else "manual_date_range",
            "start_date": config.start_date,
            "end_date": config.end_date,
            "valid_days": config.valid_days,
            "num_boost_round": config.num_boost_round,
        },
        "interval_model_training_config": {
            "training_window_days": interval_training_window_days,
            "training_mode": "rolling_window" if interval_training_window_days else "manual_date_range",
            "start_date": interval_start_date,
            "end_date": interval_end_date,
            "valid_days": interval_valid_days,
            "num_boost_round": interval_num_boost_round,
        },
        "thermal_capacity_model_training_config": {
            "training_window_days": thermal_capacity_training_window_days,
            "training_mode": "rolling_window" if thermal_capacity_training_window_days else "manual_date_range",
            "start_date": thermal_capacity_start_date,
            "end_date": thermal_capacity_end_date,
            "valid_days": thermal_capacity_valid_days,
            "num_boost_round": thermal_capacity_num_boost_round,
        },
        "similarity_reference_days": normalize_similarity_reference_days(config.similarity_reference_days),
        "xgboost_params": DEFAULT_XGB_PARAMS,
        "train_start_date": train_dates[0] if train_dates else None,
        "train_end_date": train_dates[-1] if train_dates else None,
        "train_dates": train_dates,
        "valid_dates": valid_dates,
        "validation_predictions_path": "validation_predictions.jsonl",
        "sample_rows": int(len(history_df)),
        "skipped_sheets": skipped_sheets,
        "quality_report_path": str(quality_report_path),
        "train_config": {
            "history_dir": str(config.history_dir),
            "holiday_file": str(config.holiday_file) if config.holiday_file else None,
            "data_mode": str(config.data_mode or DAYAHEAD_DATA_MODE),
            "start_date": config.start_date,
            "end_date": config.end_date,
            "training_window_days": config.training_window_days,
            "training_mode": "rolling_window" if config.training_window_days else "manual_date_range",
            "valid_days": config.valid_days,
            "num_boost_round": config.num_boost_round,
            "price_model_config": {
                "training_window_days": config.training_window_days,
                "training_mode": "rolling_window" if config.training_window_days else "manual_date_range",
                "start_date": config.start_date,
                "end_date": config.end_date,
                "valid_days": config.valid_days,
                "num_boost_round": config.num_boost_round,
            },
            "interval_model_config": {
                "training_window_days": interval_training_window_days,
                "training_mode": "rolling_window" if interval_training_window_days else "manual_date_range",
                "start_date": interval_start_date,
                "end_date": interval_end_date,
                "valid_days": interval_valid_days,
                "num_boost_round": interval_num_boost_round,
            },
            "thermal_capacity_model_config": {
                "training_window_days": thermal_capacity_training_window_days,
                "training_mode": "rolling_window" if thermal_capacity_training_window_days else "manual_date_range",
                "start_date": thermal_capacity_start_date,
                "end_date": thermal_capacity_end_date,
                "valid_days": thermal_capacity_valid_days,
                "num_boost_round": thermal_capacity_num_boost_round,
            },
            "similarity_reference_days": normalize_similarity_reference_days(config.similarity_reference_days),
            "segment_config": segment_config_metadata,
            "high_price_weighting": {
                "enabled": bool(config.high_price_weight_enabled),
                "quantile": float(config.high_price_quantile),
                "multiplier": float(config.high_price_weight_multiplier),
                "threshold": round(float(high_price_threshold), 4) if high_price_threshold is not None else None,
            },
            "price_intervals": price_intervals,
        },
        "prediction_method": "direct_price_multi_model",
        "similarity_usage": "prediction_reference_only",
        "blend_method": "direct_price_multi_model",
        "metrics": metrics_summary,
        "variant_metrics": variant_metrics,
        "high_price_adjustment_metrics": high_price_adjustment_metrics,
        "rolling_backtest_metrics": rolling_backtest_metrics,
    }
    metadata_path = staging_dir / "metadata.json"
    emit_progress(progress_callback, "正在写入模型元数据", 82)
    validation_path = staging_dir / "validation_predictions.jsonl"
    for row in validation_records:
        append_jsonl(validation_path, row)
    write_json(metadata_path, metadata)
    emit_progress(progress_callback, "正在注册当前模型版本", 90)
    current_dir, history_dir = finalize_model_version(config.model_root, staging_dir, run_id)
    pinned_version_key = read_pinned_default_version(config.model_root)
    if pinned_version_key and pinned_version_key != run_id:
        current_dir = activate_model_version(config.model_root, pinned_version_key, persist_default=False)

    log_entry = {
        "run_id": run_id,
        "created_at": metadata["created_at"],
        "target_market": metadata["target_market"],
        "target_source_column": metadata["target_source_column"],
        "data_mode": metadata["data_mode"],
        "train_start_date": metadata["train_start_date"],
        "train_end_date": metadata["train_end_date"],
        "valid_dates": valid_dates,
        "sample_rows": metadata["sample_rows"],
        "start_date_filter": config.start_date,
        "end_date_filter": config.end_date,
        "training_window_days": config.training_window_days,
        "training_mode": metadata["training_mode"],
        "price_model_training_config": metadata["price_model_training_config"],
        "interval_model_training_config": metadata["interval_model_training_config"],
        "thermal_capacity_model_training_config": metadata["thermal_capacity_model_training_config"],
        "similarity_reference_days": normalize_similarity_reference_days(config.similarity_reference_days),
        "metrics": metrics_summary,
        "variant_metrics": variant_metrics,
        "selected_model_key": selected_model_key,
        "selected_segment_price_models": selected_segment_price_models,
        "selected_model_backend": selected_variant_meta["model_backend"],
        "price_interval_model": price_interval_model,
        "thermal_capacity_model": thermal_capacity_model,
        "variant_quality_metrics": variant_quality_metrics,
        "high_price_adjustment_metrics": high_price_adjustment_metrics,
        "rolling_backtest_metrics": rolling_backtest_metrics,
        "segment_config": segment_config_metadata,
        "high_price_weighting": metadata["high_price_weighting"],
    }
    append_jsonl(config.model_root / "training_runs.jsonl", log_entry)
    emit_progress(progress_callback, f"训练完成，当前默认模型：{run_id}", 100)
    return TrainResult(
        run_id=run_id,
        current_model_dir=current_dir,
        history_model_dir=history_dir,
        metrics=metrics_summary,
        train_dates=train_dates,
        valid_dates=valid_dates,
        skipped_sheets=skipped_sheets,
        metadata_path=current_dir / "metadata.json",
        quality_report_path=quality_report_path,
    )


def rollback_to_previous(model_root: Path = DEFAULT_MODEL_ROOT) -> Path:
    current_dir, previous_dir, history_dir = ensure_model_dirs(model_root)
    if not previous_dir.exists():
        raise FileNotFoundError("未找到上一版模型，无法回退")
    rollback_backup = history_dir / datetime.now().strftime("rollback_backup_%Y%m%d_%H%M%S")
    if current_dir.exists():
        copy_tree_with_retry(current_dir, rollback_backup)
        remove_tree_with_retry(current_dir)
    copy_tree_with_retry(previous_dir, current_dir)
    current_metadata = read_model_metadata(current_dir)
    write_pinned_default_version(model_root, find_version_key_by_run_id(model_root, current_metadata.get("run_id") if current_metadata else None))
    return current_dir


def resolve_active_model_dir(model_root: str | Path) -> Path:
    root = Path(model_root)
    current_dir = root / "current"
    metadata_path = current_dir / "metadata.json"
    if metadata_path.exists():
        return current_dir
    legacy_metadata = root / "metadata.json"
    if legacy_metadata.exists():
        return root
    raise FileNotFoundError(f"未找到可用模型目录: {root}")



def load_backend_model(model_path: Path, backend: str):
    if backend == "lightgbm":
        return load_lightgbm_model(model_path)
    if backend == "catboost":
        from catboost import CatBoostRegressor

        model = CatBoostRegressor()
        model.load_model(str(model_path))
        return model
    import xgboost as xgb

    booster = xgb.Booster()
    booster.load_model(str(model_path))
    return booster


def predict_backend_model(model, backend: str, frame: pd.DataFrame, feature_columns: list[str]) -> np.ndarray:
    if backend == "lightgbm":
        return model.predict(frame[feature_columns].astype(float))
    if backend == "catboost":
        return model.predict(frame[feature_columns].astype(float))
    dmatrix = make_dmatrix(frame, feature_columns)
    iteration_range = best_iteration_range(model)
    return model.predict(dmatrix, iteration_range=iteration_range) if iteration_range else model.predict(dmatrix)


def load_models(model_root: str | Path) -> tuple[dict, dict[str, object], Path]:
    active_dir = resolve_active_model_dir(model_root)
    metadata = json.loads((active_dir / "metadata.json").read_text(encoding="utf-8"))
    models = {}
    if metadata_has_direct_variants(metadata):
        for variant_key, variant_meta in metadata["model_variants"].items():
            variant_dir = active_dir / variant_meta.get("model_dir", variant_key)
            backend = str(variant_meta.get("model_backend") or "xgboost")
            models[variant_key] = {}
            for segment in metadata["segments"]:
                segment_name = segment["name"]
                model_path = variant_dir / model_file_name(segment_name, backend)
                if not model_path.exists() and backend != "xgboost":
                    model_path = variant_dir / f"{segment_name}.json"
                    backend = "xgboost"
                models[variant_key][segment_name] = {"backend": backend, "model": load_backend_model(model_path, backend)}
        interval_bundle = load_interval_model_bundle(active_dir, metadata.get("price_interval_model"))
        if interval_bundle is not None:
            models["__price_interval__"] = interval_bundle
        thermal_capacity_metadata = metadata.get("thermal_capacity_model") if isinstance(metadata.get("thermal_capacity_model"), dict) else {}
        if thermal_capacity_metadata.get("enabled"):
            thermal_model_path = active_dir / str(thermal_capacity_metadata.get("model_path") or "thermal_capacity_model/thermal_capacity_model.json")
            if not thermal_model_path.exists():
                thermal_model_path = active_dir / "thermal_capacity_model" / "thermal_capacity_model.json"
            if thermal_model_path.exists():
                thermal_backend = str(thermal_capacity_metadata.get("selected_model_backend") or "xgboost")
                thermal_model = load_thermal_capacity_backend_model(thermal_model_path, thermal_backend)
                models["__thermal_capacity__"] = {"backend": thermal_backend, "model": thermal_model, "metadata": thermal_capacity_metadata}
    else:
        raise ValueError("当前模型缺少多模型直接预测配置，请重新训练模型")
    return metadata, models, active_dir


def realtime_model_root(model_root: str | Path = DEFAULT_MODEL_ROOT) -> Path:
    return Path(model_root) / "realtime_price"


def realtime_model_available(model_root: str | Path = DEFAULT_MODEL_ROOT) -> bool:
    return (realtime_model_root(model_root) / "current" / "metadata.json").exists()


def merge_realtime_prediction_columns(base_df: pd.DataFrame, realtime_df: pd.DataFrame) -> pd.DataFrame:
    result = base_df.copy()
    if result.empty or realtime_df.empty:
        return result
    columns = ["date", "period", "predicted_price", "model_variant"]
    available_columns = [column for column in columns if column in realtime_df.columns]
    if not {"date", "period", "predicted_price"}.issubset(set(available_columns)):
        return result
    right = realtime_df[available_columns].copy()
    right["date"] = pd.to_datetime(right["date"], errors="coerce").dt.normalize()
    right["period"] = pd.to_numeric(right["period"], errors="coerce")
    right = right.rename(
        columns={
            "predicted_price": "realtime_predicted_price",
            "model_variant": "realtime_model_variant",
        }
    )
    left = result.copy()
    left["date"] = pd.to_datetime(left["date"], errors="coerce").dt.normalize()
    left["period"] = pd.to_numeric(left["period"], errors="coerce")
    merged = left.merge(right, on=["date", "period"], how="left")
    if "predicted_price" in merged.columns and "realtime_predicted_price" in merged.columns:
        merged["price_spread_realtime_minus_dayahead"] = (
            pd.to_numeric(merged["realtime_predicted_price"], errors="coerce")
            - pd.to_numeric(merged["predicted_price"], errors="coerce")
        )
    return merged


def resolve_forecast_template_column(columns: Iterable[object], target_date: pd.Timestamp, candidates: list[str], required: bool) -> str | None:
    month_day_tokens = [f"{target_date.month}.{target_date.day}", f"{target_date.month}月{target_date.day}"]
    names = [str(column).strip() for column in columns]
    matched = []
    for name in names:
        if not any(token in name for token in month_day_tokens):
            continue
        normalized_name = normalize_text(name)
        for candidate in candidates:
            if normalize_text(candidate) in normalized_name:
                matched.append(name)
                break
    if matched:
        return matched[0]
    return resolve_column(names, candidates, required=required)


def parse_multi_day_header_date(header: object, default_year: int | None = None) -> pd.Timestamp | None:
    month_day = extract_month_day(header)
    if month_day is None:
        return None
    year = int(default_year or datetime.now().year)
    try:
        return pd.Timestamp(year=year, month=month_day[0], day=month_day[1]).normalize()
    except ValueError:
        return None


def _is_header_match(header: object, keywords: list[str]) -> bool:
    normalized = normalize_text(header)
    return all(normalize_text(keyword) in normalized for keyword in keywords)


def _first_valid_value(series: pd.Series, field_name: str) -> float:
    values = to_numeric(series).dropna()
    if values.empty:
        raise ValueError(f"预测文件-N天缺少有效{field_name}")
    return float(values.iloc[0])


def parse_multi_day_forecast_workbook(
    forecast_file: str | Path = DEFAULT_MULTI_DAY_FORECAST_FILE,
    default_year: int | None = None,
) -> dict[str, object]:
    workbook = pd.read_excel(forecast_file, sheet_name=None)
    non_empty_sheets = [(str(name), df) for name, df in workbook.items() if not df.empty and len(df.columns) > 0]
    if not non_empty_sheets:
        raise ValueError("预测文件-N天没有可用工作表")
    sheet_name, raw_df = non_empty_sheets[0]
    raw_df = raw_df.copy().reset_index(drop=True)
    raw_df.columns = [str(column).strip() if column is not None else "" for column in raw_df.columns]
    if len(raw_df) < 96:
        raise ValueError("预测文件-N天有效行数不足 96")
    raw_df = raw_df.iloc[:96].copy()
    columns = list(raw_df.columns)

    price_col = next((column for column in columns if _is_header_match(column, ["日前出清价格"])), None)
    if price_col is None:
        price_col = next((column for column in columns if _is_header_match(column, ["出清价格"])), None)
    if price_col is None:
        raise KeyError("预测文件-N天缺少基准日前出清价格列")
    baseline_date = parse_multi_day_header_date(price_col, default_year)
    if baseline_date is None:
        raise ValueError("预测文件-N天无法从基准日前出清价格列识别日期")

    period_col = resolve_column(columns, ["序号", "时段", "period"], required=False)
    baseline_renewable_col = resolve_forecast_template_column(raw_df.columns, baseline_date, ["风光电力值", "电力值"], required=False)
    baseline_total_col = resolve_forecast_template_column(raw_df.columns, baseline_date, ["总加电力值"], required=False)
    baseline_net_col = resolve_forecast_template_column(raw_df.columns, baseline_date, ["剩余电力值", "火电空间", "火电剩余空间"], required=False)
    thermal_col = resolve_column(columns, ["火电开机容量", "运行机组容量"], required=False)
    if thermal_col is None:
        raise KeyError("预测文件-N天缺少火电开机容量列")

    date_columns: dict[str, dict[str, str]] = {}
    for column in columns:
        column_date = parse_multi_day_header_date(column, default_year)
        if column_date is None or column_date == baseline_date:
            continue
        key = column_date.strftime("%Y-%m-%d")
        entry = date_columns.setdefault(key, {})
        if _is_header_match(column, ["风光"]) or (_is_header_match(column, ["电力值"]) and not _is_header_match(column, ["总加"])):
            entry["renewable_power_col"] = column
        if _is_header_match(column, ["总加"]):
            entry["total_load_col"] = column
    forecast_dates = [
        date_key
        for date_key, mapping in sorted(date_columns.items())
        if mapping.get("renewable_power_col") and mapping.get("total_load_col")
    ]
    if not forecast_dates:
        raise ValueError("预测文件-N天没有识别到后续预测日的风光/总加电力值列")

    periods = to_numeric(raw_df[period_col]).fillna(0).astype(int).tolist() if period_col else list(range(1, 97))
    baseline_prices = to_numeric(raw_df[price_col]).tolist()
    thermal_capacity = _first_valid_value(raw_df[thermal_col], "火电开机容量")
    if baseline_net_col:
        baseline_net = to_numeric(raw_df[baseline_net_col])
    elif baseline_total_col and baseline_renewable_col:
        baseline_net = to_numeric(raw_df[baseline_total_col]) - to_numeric(raw_df[baseline_renewable_col])
    else:
        baseline_net = pd.Series([np.nan] * len(raw_df), index=raw_df.index)
    baseline_total = to_numeric(raw_df[baseline_total_col]) if baseline_total_col else pd.Series([np.nan] * len(raw_df), index=raw_df.index)
    baseline_renewable = to_numeric(raw_df[baseline_renewable_col]) if baseline_renewable_col else pd.Series([np.nan] * len(raw_df), index=raw_df.index)

    rows: list[dict[str, object]] = []
    for index in range(96):
        forecast_days: dict[str, dict[str, float]] = {}
        for date_key in forecast_dates:
            mapping = date_columns[date_key]
            renewable_power = float(to_numeric(raw_df[mapping["renewable_power_col"]]).iloc[index])
            total_load = float(to_numeric(raw_df[mapping["total_load_col"]]).iloc[index])
            forecast_days[date_key] = {
                "renewable_power": renewable_power,
                "total_load": total_load,
                "net_load": total_load - renewable_power,
            }
        rows.append(
            {
                "period": int(periods[index]) if index < len(periods) and periods[index] else index + 1,
                "baseline_total_load": float(baseline_total.iloc[index]) if pd.notna(baseline_total.iloc[index]) else np.nan,
                "baseline_renewable_power": float(baseline_renewable.iloc[index]) if pd.notna(baseline_renewable.iloc[index]) else np.nan,
                "baseline_net_load": float(baseline_net.iloc[index]) if pd.notna(baseline_net.iloc[index]) else np.nan,
                "baseline_price": float(baseline_prices[index]),
                "thermal_on_capacity": thermal_capacity,
                "forecast_days": forecast_days,
            }
        )

    return {
        "file_path": str(forecast_file),
        "sheet_name": sheet_name,
        "baseline_date": baseline_date.strftime("%Y-%m-%d"),
        "forecast_dates": forecast_dates,
        "rows": rows,
    }


def build_multi_day_forecast_frame(
    parsed: dict[str, object],
    forecast_date: str,
    holiday_dates: set[pd.Timestamp] | dict[pd.Timestamp, str] | None = None,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    holiday_dates = holiday_dates or set()
    baseline_date = pd.Timestamp(str(parsed["baseline_date"])).normalize()
    target_date = pd.Timestamp(str(forecast_date)).normalize()
    rows = list(parsed.get("rows") or [])
    forecast_records: list[dict[str, object]] = []
    reference_records: list[dict[str, object]] = []
    target_day_type = day_type_of(target_date, holiday_dates)
    baseline_day_type = day_type_of(baseline_date, holiday_dates)
    for row in rows:
        period = int(row["period"])
        day_values = (row.get("forecast_days") or {}).get(forecast_date)
        if not isinstance(day_values, dict):
            raise KeyError(f"预测文件-N天缺少 {forecast_date} 的第 {period} 点数据")
        total_load = float(day_values["total_load"])
        renewable_power = float(day_values["renewable_power"])
        net_load = float(day_values["net_load"])
        thermal_capacity = float(row["thermal_on_capacity"])
        forecast_records.append(
            {
                "date": target_date,
                "period": period,
                "hour": int((period - 1) // 4),
                "weekday": int(target_date.weekday()),
                "month": int(target_date.month),
                "is_weekend": int(target_day_type == "weekend"),
                "is_holiday": int(target_day_type == "holiday"),
                DAY_TYPE_COLUMN: target_day_type,
                TOTAL_LOAD_COLUMN: total_load,
                "net_load": net_load,
                RENEWABLE_POWER_COLUMN: renewable_power,
                THERMAL_SPACE_LOAD_RATIO_COLUMN: net_load / total_load if total_load else np.nan,
                "thermal_on_capacity": thermal_capacity,
                "sheet_name": parsed.get("sheet_name") or "multi_day_forecast",
                "source_file": parsed.get("file_path") or "",
            }
        )
        baseline_total = float(row.get("baseline_total_load") or np.nan)
        baseline_net = float(row.get("baseline_net_load") or np.nan)
        reference_records.append(
            {
                "date": baseline_date,
                "period": period,
                "hour": int((period - 1) // 4),
                "weekday": int(baseline_date.weekday()),
                "month": int(baseline_date.month),
                "is_weekend": int(baseline_day_type == "weekend"),
                "is_holiday": int(baseline_day_type == "holiday"),
                DAY_TYPE_COLUMN: baseline_day_type,
                TOTAL_LOAD_COLUMN: baseline_total,
                "net_load": baseline_net,
                RENEWABLE_POWER_COLUMN: float(row.get("baseline_renewable_power") or np.nan),
                THERMAL_SPACE_LOAD_RATIO_COLUMN: baseline_net / baseline_total if baseline_total else np.nan,
                "thermal_on_capacity": thermal_capacity,
                TARGET_COLUMN: float(row["baseline_price"]),
                "sheet_name": parsed.get("sheet_name") or "multi_day_forecast",
                "source_file": parsed.get("file_path") or "",
            }
        )
    forecast_df = pd.DataFrame(forecast_records)
    forecast_df["segment"] = assign_segments(forecast_df["period"])
    reference_df = pd.DataFrame(reference_records)
    reference_df["segment"] = assign_segments(reference_df["period"])
    return forecast_df, reference_df


def try_load_forecast_template(
    forecast_file: str | Path,
    holiday_dates: set[pd.Timestamp],
    default_year: int,
) -> tuple[pd.DataFrame | None, pd.DataFrame | None, list[DataQualityIssue]]:
    workbook = pd.read_excel(forecast_file, sheet_name=None)
    non_empty_sheets = [(str(name), df) for name, df in workbook.items() if not df.empty and len(df.columns) > 0]
    if not non_empty_sheets:
        return None, None, []
    template_sheet_name, template_raw_df = non_empty_sheets[0]
    raw_df = template_raw_df.copy().reset_index(drop=True)
    raw_df.columns = [str(column).strip() for column in raw_df.columns]
    header_dates = []
    for column in raw_df.columns:
        month_day = extract_month_day(column)
        if month_day is None:
            continue
        header_dates.append(pd.Timestamp(year=default_year, month=month_day[0], day=month_day[1]).normalize())
    unique_dates = sorted(set(header_dates))
    if len(unique_dates) < 2:
        return None, None, []

    quality_issues = validate_forecast_template(
        raw_df,
        file_path=Path(forecast_file),
        sheet_name=template_sheet_name,
        default_year=default_year,
    )
    if has_blocking_issues(quality_issues):
        raise DataQualityValidationError("预测文件关键字段异常，已停止预测", quality_issues)

    target_date = unique_dates[-1]
    reference_date = unique_dates[-2]
    target_net_load_col = resolve_forecast_template_column(raw_df.columns, target_date, ["剩余电力值(MW)", "剩余电力值", "火电空间", "火电剩余空间"], required=False)
    target_total_col = resolve_forecast_template_column(raw_df.columns, target_date, ["总加电力值(MW)", "总加电力值"], required=False)
    target_power_col = resolve_forecast_template_column(raw_df.columns, target_date, ["风光电力值(MW)", "风光电力值", "电力值(MW)", "电力值"], required=False)
    reference_net_load_col = resolve_forecast_template_column(raw_df.columns, reference_date, ["剩余电力值(MW)", "剩余电力值", "火电空间", "火电剩余空间"], required=False)
    reference_total_col = resolve_forecast_template_column(raw_df.columns, reference_date, ["总加电力值(MW)", "总加电力值"], required=False)
    reference_power_col = resolve_forecast_template_column(raw_df.columns, reference_date, ["风光电力值(MW)", "风光电力值", "电力值(MW)", "电力值"], required=False)
    reference_price_col = resolve_forecast_template_column(raw_df.columns, reference_date, ["日前出清价格(元/MWh)", "日前-出清价格(元/MWh)", "日前出清价格", "出清价格"], required=False)
    thermal_col = resolve_column(raw_df.columns, ["火电开机容量(MW)", "火电开机容量", "运行机组容量"], required=False)
    period_col = resolve_column(raw_df.columns, ["序号", "时段", "period"], required=False)

    def build_net_load(df: pd.DataFrame, net_col: str | None, total_col: str | None, power_col: str | None) -> pd.Series:
        if net_col:
            return to_numeric(df[net_col])
        if total_col and power_col:
            return to_numeric(df[total_col]) - to_numeric(df[power_col])
        raise KeyError("预测文件缺少火电空间相关字段")

    def build_total_load(df: pd.DataFrame, total_col: str | None) -> pd.Series:
        if total_col:
            return to_numeric(df[total_col])
        return pd.Series([np.nan] * len(df), index=df.index)

    target_net_load = build_net_load(raw_df, target_net_load_col, target_total_col, target_power_col)
    reference_net_load = build_net_load(raw_df, reference_net_load_col, reference_total_col, reference_power_col)
    target_total_load = build_total_load(raw_df, target_total_col)
    reference_total_load = build_total_load(raw_df, reference_total_col)
    target_renewable_power = to_numeric(raw_df[target_power_col]) if target_power_col else pd.Series([np.nan] * len(raw_df))
    reference_renewable_power = to_numeric(raw_df[reference_power_col]) if reference_power_col else pd.Series([np.nan] * len(raw_df))
    thermal_series = pd.Series([np.nan] * len(raw_df))
    if thermal_col:
        thermal_series = to_numeric(raw_df[thermal_col]).ffill().bfill()
    periods = to_numeric(raw_df[period_col]).fillna(0).astype(int).to_numpy() if period_col else np.arange(1, len(raw_df) + 1)
    if len(periods) < 96:
        raise ValueError("预测文件有效行数不足 96")
    periods = periods[:96]
    target_date = target_date.normalize()
    reference_date = reference_date.normalize()
    target_day_type = day_type_of(target_date, holiday_dates)
    reference_day_type = day_type_of(reference_date, holiday_dates)
    forecast_df = pd.DataFrame(
        {
            "date": target_date,
            "period": periods,
            "hour": ((periods - 1) // 4).astype(int),
            "weekday": int(target_date.weekday()),
            "month": int(target_date.month),
            "is_weekend": int(target_day_type == "weekend"),
            "is_holiday": int(target_day_type == "holiday"),
            DAY_TYPE_COLUMN: target_day_type,
            TOTAL_LOAD_COLUMN: target_total_load.iloc[:96].to_numpy(),
            "net_load": target_net_load.iloc[:96].to_numpy(),
            RENEWABLE_POWER_COLUMN: target_renewable_power.iloc[:96].to_numpy(),
            THERMAL_SPACE_LOAD_RATIO_COLUMN: calculate_thermal_space_load_ratio(target_net_load, target_total_load).iloc[:96].to_numpy(),
            "thermal_on_capacity": thermal_series.iloc[:96].to_numpy(),
            "sheet_name": "forecast_template",
            "source_file": str(forecast_file),
        }
    )
    forecast_df["segment"] = assign_segments(forecast_df["period"])

    reference_df = None
    if reference_price_col:
        reference_df = pd.DataFrame(
            {
                "date": reference_date,
                "period": periods,
                "hour": ((periods - 1) // 4).astype(int),
                "weekday": int(reference_date.weekday()),
                "month": int(reference_date.month),
                "is_weekend": int(reference_day_type == "weekend"),
                "is_holiday": int(reference_day_type == "holiday"),
                DAY_TYPE_COLUMN: reference_day_type,
                TOTAL_LOAD_COLUMN: reference_total_load.iloc[:96].to_numpy(),
                "net_load": reference_net_load.iloc[:96].to_numpy(),
                RENEWABLE_POWER_COLUMN: reference_renewable_power.iloc[:96].to_numpy(),
                THERMAL_SPACE_LOAD_RATIO_COLUMN: calculate_thermal_space_load_ratio(reference_net_load, reference_total_load).iloc[:96].to_numpy(),
                "thermal_on_capacity": thermal_series.iloc[:96].to_numpy(),
                TARGET_COLUMN: to_numeric(raw_df[reference_price_col]).iloc[:96].to_numpy(),
                "sheet_name": "forecast_template",
                "source_file": str(forecast_file),
            }
        )
        reference_df["segment"] = assign_segments(reference_df["period"])

    return forecast_df, reference_df, quality_issues


def load_forecast_generic(
    history_df: pd.DataFrame,
    forecast_file: str | Path,
    holiday_dates: set[pd.Timestamp],
) -> tuple[pd.DataFrame, pd.DataFrame | None, list[DataQualityIssue]]:
    default_year = int(pd.to_datetime(history_df["date"]).max().year)
    template_df, template_reference, quality_issues = try_load_forecast_template(forecast_file, holiday_dates, default_year)
    if template_df is not None:
        return template_df, template_reference, quality_issues

    workbook = pd.read_excel(forecast_file, sheet_name=None)
    all_days = []
    quality_issues = []
    for sheet_name, raw_df in workbook.items():
        trade_date = parse_sheet_date(str(sheet_name))
        if trade_date is None:
            trade_date = infer_date_from_header(raw_df.columns, default_year)
        if trade_date is None:
            quality_issues.append(
                DataQualityIssue(
                    task_type="predict",
                    severity="warning",
                    action="recorded",
                    issue_type="unrecognized_sheet_date",
                    file_path=str(forecast_file),
                    sheet_name=str(sheet_name),
                    message="Sheet 名称和字段表头均无法识别预测日期，已忽略该 Sheet",
                )
            )
            continue
        sheet_issues = validate_history_sheet(
            raw_df,
            file_path=Path(forecast_file),
            sheet_name=str(sheet_name),
            trade_date=trade_date,
            require_target=False,
            task_type="predict",
            blocking_action="blocked",
        )
        quality_issues.extend(sheet_issues)
        if has_blocking_issues(sheet_issues):
            raise DataQualityValidationError("预测文件关键字段异常，已停止预测", sheet_issues)
        day_df = raw_df.copy().iloc[:96].reset_index(drop=True)
        if len(day_df) < 96:
            raise ValueError(f"{sheet_name} 不是 96 行时段数据")
        day_df.columns = [str(column).strip() for column in day_df.columns]
        total_power_col = resolve_column(day_df.columns, ["总加电力值(MW)", "总加电力值"], required=True)
        power_col = resolve_column(day_df.columns, ["电力值(MW)", "电力值"], required=True)
        thermal_col = resolve_column(day_df.columns, ["火电开机容量(MW)", "火电开机容量", "运行机组容量"], required=False)
        overview_col = resolve_column(day_df.columns, ["日前-出清概况", "日前出清概况", "出清概况"], required=False)
        thermal_series = pd.Series([np.nan] * len(day_df))
        if overview_col:
            thermal_series = day_df[overview_col].ffill().bfill().map(extract_thermal_on_capacity)
        if thermal_col:
            thermal_series = to_numeric(day_df[thermal_col]).combine_first(to_numeric(thermal_series))
        periods = np.arange(1, 97)
        day_type = day_type_of(trade_date.normalize(), holiday_dates)
        total_load = to_numeric(day_df[total_power_col])
        renewable_power = to_numeric(day_df[power_col])
        thermal_space = total_load - renewable_power
        result = pd.DataFrame(
            {
                "date": trade_date.normalize(),
                "period": periods,
                "hour": ((periods - 1) // 4).astype(int),
                "weekday": int(trade_date.weekday()),
                "month": int(trade_date.month),
                "is_weekend": int(day_type == "weekend"),
                "is_holiday": int(day_type == "holiday"),
                DAY_TYPE_COLUMN: day_type,
                TOTAL_LOAD_COLUMN: total_load,
                "net_load": thermal_space,
                RENEWABLE_POWER_COLUMN: renewable_power,
                THERMAL_SPACE_LOAD_RATIO_COLUMN: calculate_thermal_space_load_ratio(thermal_space, total_load),
                "thermal_on_capacity": to_numeric(thermal_series),
                "sheet_name": sheet_name,
                "source_file": str(forecast_file),
            }
        )
        result["segment"] = assign_segments(result["period"])
        all_days.append(result)
    if not all_days:
        raise ValueError(f"未解析到可用的预测 Sheet: {forecast_file}")
    forecast_df = pd.concat(all_days, ignore_index=True).sort_values(["date", "period"]).reset_index(drop=True)
    return forecast_df, None, quality_issues


def default_forecast_year_from_model(model_root: str | Path) -> int:
    metadata = load_current_metadata(Path(model_root)) or {}
    for key in ("train_end_date", "created_at"):
        value = metadata.get(key)
        if not value:
            continue
        timestamp = pd.to_datetime(value, errors="coerce")
        if not pd.isna(timestamp):
            return int(pd.Timestamp(timestamp).year)
    return datetime.now().year


def peek_forecast_target_date(forecast_file: str | Path, default_year: int) -> pd.Timestamp | None:
    workbook = pd.read_excel(forecast_file, sheet_name=None)
    header_dates: list[pd.Timestamp] = []
    sheet_dates: list[pd.Timestamp] = []
    for sheet_name, raw_df in workbook.items():
        sheet_date = parse_sheet_date(str(sheet_name))
        if sheet_date is not None:
            sheet_dates.append(sheet_date.normalize())
        if raw_df.empty:
            continue
        for column in raw_df.columns:
            month_day = extract_month_day(column)
            if month_day is not None:
                header_dates.append(pd.Timestamp(year=default_year, month=month_day[0], day=month_day[1]).normalize())
    if header_dates:
        return max(header_dates)
    if sheet_dates:
        return max(sheet_dates)
    return None


def prediction_history_window(forecast_date: pd.Timestamp | None, reference_days: int) -> tuple[str | None, str | None]:
    if forecast_date is None:
        return None, None
    normalized_days = normalize_reference_days(reference_days)
    lookback_days = max(normalized_days, normalized_days * 8)
    start_date = pd.Timestamp(forecast_date).normalize() - pd.Timedelta(days=lookback_days)
    end_date = pd.Timestamp(forecast_date).normalize() - pd.Timedelta(days=1)
    return start_date.strftime("%Y-%m-%d"), end_date.strftime("%Y-%m-%d")


def build_reference_frame(
    history_df: pd.DataFrame,
    forecast_df: pd.DataFrame,
    template_reference_df: pd.DataFrame | None,
    reference_days: int,
    reference_strategy: str,
    holiday_dates: set[pd.Timestamp],
) -> tuple[pd.DataFrame, list[str]]:
    reference_days = normalize_reference_days(reference_days)
    reference_strategy = normalize_reference_strategy(reference_strategy)
    forecast_date = pd.to_datetime(forecast_df["date"]).min().normalize()
    forecast_day_type = day_type_of(forecast_date, holiday_dates)
    history_dates = pd.to_datetime(history_df["date"], errors="coerce").dropna().dt.normalize()
    selected_dates: list[pd.Timestamp] = []
    selected_frames: list[pd.DataFrame] = []

    def is_eligible_reference(ref_date: pd.Timestamp) -> bool:
        if reference_strategy == "recent_n_days":
            return True
        return day_type_of(ref_date, holiday_dates) == forecast_day_type

    def reference_columns(frame: pd.DataFrame) -> list[str]:
        desired = [
            "date",
            "period",
            TOTAL_LOAD_COLUMN,
            "net_load",
            RENEWABLE_POWER_COLUMN,
            THERMAL_SPACE_LOAD_RATIO_COLUMN,
            "thermal_on_capacity",
            DAY_TYPE_COLUMN,
            TARGET_COLUMN,
        ]
        return [column for column in desired if column in frame.columns]

    if template_reference_df is not None and not template_reference_df.empty and len(selected_dates) < reference_days:
        template_reference = template_reference_df[reference_columns(template_reference_df)].copy()
        template_reference["date"] = pd.to_datetime(template_reference["date"], errors="coerce").dt.normalize()
        template_date = pd.to_datetime(template_reference["date"]).iloc[0].normalize()
        if pd.notna(template_date):
            selected_frames.append(template_reference)
            selected_dates.append(template_date)

    used_dates = set(selected_dates)
    recent_history_dates = [
        ref_date
        for ref_date in sorted(set(history_dates.tolist()), reverse=True)
        if ref_date < forecast_date and ref_date not in used_dates and is_eligible_reference(ref_date)
    ]
    for ref_date in recent_history_dates:
        if len(selected_dates) >= reference_days:
            break
        history_reference = history_df.loc[history_dates == ref_date, reference_columns(history_df)].copy()
        if history_reference.empty:
            continue
        selected_frames.append(history_reference)
        selected_dates.append(ref_date)

    if not selected_frames:
        raise ValueError("未找到可用于相似法预测的参考日数据")
    if len(selected_dates) < reference_days:
        warnings.warn(
            f"请求使用最近 {reference_days} 天参考数据，但仅找到 {len(selected_dates)} 天可用样本，已自动降级为现有样本。",
            stacklevel=2,
        )

    reference_df = pd.concat(selected_frames, ignore_index=True)
    reference_df["date"] = pd.to_datetime(reference_df["date"], errors="coerce").dt.normalize()
    reference_df = ensure_similarity_columns(reference_df)
    reference_df = reference_df.dropna(subset=["net_load", TARGET_COLUMN]).reset_index(drop=True)
    if reference_df.empty:
        raise ValueError("参考日样本缺少有效火电空间或价格数据，无法构造相似法特征")
    reference_dates = [pd.Timestamp(ref_date).strftime("%Y-%m-%d") for ref_date in selected_dates]
    return reference_df, reference_dates


def attach_forecast_similarity_features(
    forecast_df: pd.DataFrame,
    reference_df: pd.DataFrame,
    knn_similarity_config: dict[str, object] | None = None,
) -> pd.DataFrame:
    result = ensure_similarity_columns(forecast_df)
    similar_price, similar_gap = nearest_similarity_price(
        result["net_load"].to_numpy(dtype=float),
        reference_df["net_load"].to_numpy(dtype=float),
        reference_df[TARGET_COLUMN].to_numpy(dtype=float),
    )
    result[NET_LOAD_ONLY_SIMILAR_PRICE_COLUMN] = similar_price
    result[NET_LOAD_ONLY_SIMILAR_GAP_COLUMN] = similar_gap
    knn_config = normalize_knn_similarity_config(knn_similarity_config)
    knn_result = multi_factor_knn_similarity_price(
        result,
        ensure_similarity_columns(reference_df),
        k=int(knn_config["knn_k"]),
        max_distance=float(knn_config["knn_max_distance"]),
        weighted_k=int(knn_config["weighted_knn_k"]),
        weighted_max_distance=float(knn_config["weighted_knn_max_distance"]),
    )
    result[KNN_SIMILAR_PRICE_COLUMN] = knn_result["plain_price"].to_numpy(dtype=float)
    result[WEIGHTED_KNN_PRICE_COLUMN] = knn_result["weighted_price"].to_numpy(dtype=float)
    result[KNN_SIMILAR_COUNT_COLUMN] = knn_result["count"].to_numpy(dtype=int)
    result[KNN_SIMILAR_DISTANCE_COLUMN] = knn_result["min_distance"].to_numpy(dtype=float)
    return result


def clip_price_series(values: pd.Series | np.ndarray) -> pd.Series | np.ndarray:
    return np.clip(values, PRICE_FLOOR, PRICE_CAP)


def prepare_prediction_inputs(
    history_dir: str | Path,
    forecast_file: str | Path,
    model_root: str | Path,
    holiday_file: str | Path | None,
    reference_days: int = 1,
    progress_callback: ProgressCallback | None = None,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame | None, set[pd.Timestamp], dict, dict[str, object], Path]:
    emit_progress(progress_callback, "开始加载预测所需数据", 5)
    forecast_default_year = default_forecast_year_from_model(model_root)
    forecast_target_date = peek_forecast_target_date(forecast_file, forecast_default_year)
    history_start_date, history_end_date = prediction_history_window(forecast_target_date, reference_days)
    try:
        history_df, _, holiday_dates, _, history_quality_issues = load_cached_history_collection(
            history_dir,
            holiday_file,
            True,
            model_root,
            progress_callback=progress_callback,
            progress_start=6,
            progress_end=18,
            progress_label="正在加载历史数据",
            start_date=history_start_date,
            end_date=history_end_date,
        )
    except DataQualityValidationError as exc:
        report_path = save_quality_report(
            task_type="predict",
            status="blocked",
            issues=exc.issues,
            metadata={"forecast_file": str(forecast_file), "history_dir": str(history_dir)},
        )
        exc.report_path = report_path
        raise
    emit_progress(progress_callback, "历史数据加载完成，正在加载预测文件", 20)
    emit_progress(progress_callback, "正在加载当前默认模型", 21)
    metadata, models, _ = load_models(model_root)
    quality_issues = list(history_quality_issues)
    try:
        forecast_df, template_reference_df, forecast_quality_issues = load_forecast_generic(
            history_df,
            forecast_file,
            holiday_dates,
        )
        model_segment_definitions = normalize_segment_config(metadata.get("segment_config") or metadata.get("segments"))
        history_df["segment"] = assign_segments(history_df["period"], model_segment_definitions)
        forecast_df["segment"] = assign_segments(forecast_df["period"], model_segment_definitions)
        if template_reference_df is not None and not template_reference_df.empty:
            template_reference_df["segment"] = assign_segments(template_reference_df["period"], model_segment_definitions)
        quality_issues.extend(forecast_quality_issues)
    except DataQualityValidationError as exc:
        quality_issues.extend(exc.issues)
        report_path = save_quality_report(
            task_type="predict",
            status="blocked",
            issues=quality_issues,
            metadata={"forecast_file": str(forecast_file), "history_dir": str(history_dir)},
        )
        exc.report_path = report_path
        raise
    emit_progress(progress_callback, "预测输入准备完成", 45)
    report_path = save_quality_report(
        task_type="predict",
        status=quality_report_status(quality_issues),
        issues=quality_issues,
        metadata={"forecast_file": str(forecast_file), "history_dir": str(history_dir)},
    )
    return history_df, forecast_df, template_reference_df, holiday_dates, metadata, models, report_path


def run_prediction_with_strategy(
    history_df: pd.DataFrame,
    forecast_df: pd.DataFrame,
    template_reference_df: pd.DataFrame | None,
    holiday_dates: set[pd.Timestamp],
    metadata: dict,
    models: dict[str, object],
    reference_days: int,
    reference_strategy: str,
    knn_similarity_config: dict[str, object] | None = None,
    progress_callback: ProgressCallback | None = None,
    segment_progress_start: int = 55,
    segment_progress_span: int = 30,
) -> tuple[pd.DataFrame, list[str], str, str]:
    strategy_key = normalize_reference_strategy(reference_strategy)
    strategy_label = REFERENCE_STRATEGIES[strategy_key]
    reference_df, reference_dates = build_reference_frame(
        history_df=history_df,
        forecast_df=forecast_df,
        template_reference_df=template_reference_df,
        reference_days=reference_days,
        reference_strategy=strategy_key,
        holiday_dates=holiday_dates,
    )
    strategy_forecast_df = attach_forecast_similarity_features(forecast_df, reference_df, knn_similarity_config)
    strategy_forecast_df = attach_forecast_lag_features(strategy_forecast_df, history_df)
    if forecast_df.attrs.get("force_template_lag_reference"):
        strategy_forecast_df = apply_template_reference_lag_features(strategy_forecast_df, template_reference_df)

    prediction_frames: list[pd.DataFrame] = []
    total_segments = len(metadata["segments"])
    for idx, segment in enumerate(metadata["segments"], start=1):
        segment_name = segment["name"]
        if progress_callback is not None:
            segment_percent = segment_progress_start + int(idx / max(total_segments, 1) * segment_progress_span)
            emit_progress(progress_callback, f"正在预测 {strategy_label} / 时段{segment_name}", segment_percent)
        segment_df = strategy_forecast_df[strategy_forecast_df["segment"] == segment_name].copy()
        if segment_df.empty:
            continue
        variant_key = select_segment_prediction_model_variant(metadata, segment_name, strategy_forecast_df)
        variant_meta = metadata["model_variants"][variant_key]
        feature_columns = variant_meta["feature_columns"]
        variant_models = models[variant_key]
        model_entry = variant_models[segment_name]
        if isinstance(model_entry, dict) and "model" in model_entry:
            model_pred = predict_backend_model(model_entry["model"], str(model_entry.get("backend") or "xgboost"), segment_df, feature_columns)
        else:
            model_pred = predict_backend_model(model_entry, "xgboost", segment_df, feature_columns)
        raw_predicted_price = model_pred
        clipped_predicted_price = clip_price_series(raw_predicted_price)
        segment_df["predicted_price"] = clipped_predicted_price
        segment_df["model_variant"] = variant_key
        segment_df["dayahead_model_version"] = str(metadata.get("run_id") or "")
        segment_df["model_similarity_diff"] = segment_df["predicted_price"] - segment_df[NET_LOAD_ONLY_SIMILAR_PRICE_COLUMN]
        segment_df["residual_pred"] = segment_df["model_similarity_diff"]
        prediction_frames.append(segment_df)

    result_df = pd.concat(prediction_frames, ignore_index=True).sort_values(["date", "period"]).reset_index(drop=True)
    interval_bundle = models.get("__price_interval__") if isinstance(models, dict) else None
    if interval_bundle is not None:
        interval_metadata = metadata.get("price_interval_model") or {}
        intervals = normalize_price_intervals(interval_metadata.get("intervals"))
        interval_result = predict_interval_probabilities(result_df, interval_bundle)
        result_df = append_interval_prediction_columns(result_df, interval_result, intervals)
        result_df = apply_high_price_probability_adjustment(result_df)
    result_df = apply_scenario_similarity_reference_adjustment(result_df)
    return result_df, reference_dates, strategy_key, strategy_label


def append_realtime_prediction_if_available(
    result_df: pd.DataFrame,
    *,
    history_dir: str | Path,
    model_root: str | Path,
    holiday_file: str | Path | None,
    forecast_df: pd.DataFrame,
    template_reference_df: pd.DataFrame | None,
    holiday_dates: set[pd.Timestamp],
    reference_days: int,
    reference_strategy: str,
    knn_similarity_config: dict[str, object] | None = None,
    realtime_model_root_override: str | Path | None = None,
    realtime_model_version_key: str | None = None,
    progress_callback: ProgressCallback | None = None,
) -> pd.DataFrame:
    realtime_root = Path(realtime_model_root_override) if realtime_model_root_override is not None else realtime_model_root(model_root)
    if not (realtime_root / "current" / "metadata.json").exists():
        return result_df
    try:
        emit_progress(progress_callback, "正在加载实时价格模型", 88)
        metadata, models, _ = load_models(realtime_root)
        segment_definitions = normalize_segment_config(metadata.get("segment_config") or metadata.get("segments"))
        realtime_forecast_df = forecast_df.copy()
        realtime_forecast_df["segment"] = assign_segments(realtime_forecast_df["period"], segment_definitions)
        forecast_date = pd.to_datetime(realtime_forecast_df["date"], errors="coerce").min()
        history_start_date, history_end_date = prediction_history_window(forecast_date, reference_days)
        realtime_history_df, _, _, _, _ = load_cached_history_collection(
            history_dir,
            holiday_file,
            True,
            realtime_root,
            start_date=history_start_date,
            end_date=history_end_date,
            data_mode=REALTIME_DATA_MODE,
        )
        realtime_history_df["segment"] = assign_segments(realtime_history_df["period"], segment_definitions)
        realtime_prediction_df, _, _, _ = run_prediction_with_strategy(
            history_df=realtime_history_df,
            forecast_df=realtime_forecast_df,
            template_reference_df=None,
            holiday_dates=holiday_dates,
            metadata=metadata,
            models=models,
            reference_days=reference_days,
            reference_strategy=reference_strategy,
            knn_similarity_config=knn_similarity_config,
            progress_callback=progress_callback,
            segment_progress_start=88,
            segment_progress_span=3,
        )
        merged = merge_realtime_prediction_columns(result_df, realtime_prediction_df)
        merged["realtime_prediction_status"] = "ok"
        merged["realtime_prediction_error"] = None
        merged["realtime_model_version"] = str(metadata.get("run_id") or realtime_model_version_key or "")
        merged.attrs.update(result_df.attrs)
        merged.attrs["realtime_prediction_status"] = "ok"
        return merged
    except Exception as exc:  # noqa: BLE001
        fallback = result_df.copy()
        fallback["realtime_prediction_status"] = "failed"
        fallback["realtime_prediction_error"] = str(exc)
        fallback.attrs.update(result_df.attrs)
        fallback.attrs["realtime_prediction_status"] = "failed"
        fallback.attrs["realtime_prediction_error"] = str(exc)
        emit_progress(progress_callback, f"实时价格模型预测失败，已保留日前预测结果：{exc}", 91)
        return fallback


def predict_prices(
    history_dir: str | Path,
    forecast_file: str | Path,
    model_root: str | Path = DEFAULT_MODEL_ROOT,
    output_file: str | Path = DEFAULT_OUTPUT_FILE,
    holiday_file: str | Path | None = None,
    reference_days: int = 1,
    reference_strategy: str = "recent_n_days",
    knn_similarity_config: dict[str, object] | None = None,
    thermal_capacity_config: dict[str, object] | None = None,
    realtime_model_root_override: str | Path | None = None,
    realtime_model_version_key: str | None = None,
    progress_callback: ProgressCallback | None = None,
) -> PredictResult:
    ensure_xgboost_available()
    reference_days = normalize_reference_days(reference_days)
    history_df, forecast_df, template_reference_df, holiday_dates, metadata, models, quality_report_path = prepare_prediction_inputs(
        history_dir=history_dir,
        forecast_file=forecast_file,
        model_root=model_root,
        holiday_file=holiday_file,
        reference_days=reference_days,
        progress_callback=progress_callback,
    )
    emit_progress(progress_callback, "正在执行预测策略", 50)
    thermal_file_value = extract_forecast_thermal_capacity_file_value(forecast_df)
    thermal_model_value = predict_thermal_capacity_value(
        forecast_df,
        models.get("__thermal_capacity__") if isinstance(models, dict) else None,
        thermal_file_value,
    )
    forecast_df = apply_thermal_capacity_choice(
        forecast_df,
        build_thermal_capacity_choice(forecast_df, thermal_capacity_config, thermal_model_value),
    )
    result_df, reference_dates, strategy_key, strategy_label = run_prediction_with_strategy(
        history_df=history_df,
        forecast_df=forecast_df,
        template_reference_df=template_reference_df,
        holiday_dates=holiday_dates,
        metadata=metadata,
        models=models,
        reference_days=reference_days,
        reference_strategy=reference_strategy,
        knn_similarity_config=knn_similarity_config,
        progress_callback=progress_callback,
    )
    result_df = append_realtime_prediction_if_available(
        result_df,
        history_dir=history_dir,
        model_root=model_root,
        holiday_file=holiday_file,
        forecast_df=forecast_df,
        template_reference_df=template_reference_df,
        holiday_dates=holiday_dates,
        reference_days=reference_days,
        reference_strategy=strategy_key,
        knn_similarity_config=knn_similarity_config,
        realtime_model_root_override=realtime_model_root_override,
        realtime_model_version_key=realtime_model_version_key,
        progress_callback=progress_callback,
    )
    emit_progress(progress_callback, "正在导出预测结果", 92)
    export_prediction(result_df, output_file)
    emit_progress(progress_callback, "预测任务完成", 100)
    forecast_date = pd.to_datetime(result_df["date"]).dt.strftime("%Y-%m-%d").iloc[0]
    return PredictResult(
        forecast_date=forecast_date,
        output_file=Path(output_file),
        template_updated=False,
        reference_strategy_key=strategy_key,
        reference_strategy_label=strategy_label,
        reference_days_requested=reference_days,
        reference_dates=reference_dates,
        result_df=result_df,
        quality_report_path=quality_report_path,
    )


def predict_prices_compare(
    history_dir: str | Path,
    forecast_file: str | Path,
    model_root: str | Path = DEFAULT_MODEL_ROOT,
    output_file: str | Path = DEFAULT_OUTPUT_FILE,
    holiday_file: str | Path | None = None,
    reference_days: int = 1,
    recent_reference_days: int | None = None,
    same_type_reference_days: int | None = None,
    selected_strategy: str = "recent_n_days",
    knn_similarity_config: dict[str, object] | None = None,
    thermal_capacity_config: dict[str, object] | None = None,
    realtime_model_root_override: str | Path | None = None,
    realtime_model_version_key: str | None = None,
    progress_callback: ProgressCallback | None = None,
) -> PredictCompareResult:
    ensure_xgboost_available()
    knn_similarity_config = normalize_knn_similarity_config(knn_similarity_config)
    reference_days_by_strategy = normalize_reference_days_by_strategy(
        reference_days=reference_days,
        recent_reference_days=recent_reference_days,
        same_type_reference_days=same_type_reference_days,
    )
    max_reference_days = max(reference_days_by_strategy.values())
    selected_strategy_key = normalize_reference_strategy(selected_strategy)
    history_df, forecast_df, template_reference_df, holiday_dates, metadata, models, quality_report_path = prepare_prediction_inputs(
        history_dir=history_dir,
        forecast_file=forecast_file,
        model_root=model_root,
        holiday_file=holiday_file,
        reference_days=max_reference_days,
        progress_callback=progress_callback,
    )
    thermal_file_value = extract_forecast_thermal_capacity_file_value(forecast_df)
    thermal_model_value = predict_thermal_capacity_value(
        forecast_df,
        models.get("__thermal_capacity__") if isinstance(models, dict) else None,
        thermal_file_value,
    )
    forecast_df = apply_thermal_capacity_choice(
        forecast_df,
        build_thermal_capacity_choice(forecast_df, thermal_capacity_config, thermal_model_value),
    )

    strategy_keys = list(REFERENCE_STRATEGIES.keys())
    strategy_results: dict[str, PredictResult] = {}
    total_strategies = len(strategy_keys)
    for strategy_index, strategy_key in enumerate(strategy_keys, start=1):
        strategy_reference_days = reference_days_by_strategy[strategy_key]
        emit_progress(progress_callback, f"正在对比策略 {strategy_index}/{total_strategies}: {REFERENCE_STRATEGIES[strategy_key]}", 48 + int((strategy_index - 1) / total_strategies * 6))
        result_df, reference_dates, normalized_strategy_key, strategy_label = run_prediction_with_strategy(
            history_df=history_df,
            forecast_df=forecast_df,
            template_reference_df=template_reference_df,
            holiday_dates=holiday_dates,
            metadata=metadata,
            models=models,
            reference_days=strategy_reference_days,
            reference_strategy=strategy_key,
            knn_similarity_config=knn_similarity_config,
            progress_callback=progress_callback,
            segment_progress_start=55 + int((strategy_index - 1) * 15),
            segment_progress_span=12,
        )
        result_df = append_realtime_prediction_if_available(
            result_df,
            history_dir=history_dir,
            model_root=model_root,
            holiday_file=holiday_file,
            forecast_df=forecast_df,
            template_reference_df=template_reference_df,
            holiday_dates=holiday_dates,
            reference_days=strategy_reference_days,
            reference_strategy=strategy_key,
            knn_similarity_config=knn_similarity_config,
            realtime_model_root_override=realtime_model_root_override,
            realtime_model_version_key=realtime_model_version_key,
            progress_callback=progress_callback,
        )
        forecast_date = pd.to_datetime(result_df["date"]).dt.strftime("%Y-%m-%d").iloc[0]
        strategy_results[strategy_key] = PredictResult(
            forecast_date=forecast_date,
            output_file=Path(output_file),
            template_updated=False,
            reference_strategy_key=normalized_strategy_key,
            reference_strategy_label=strategy_label,
            reference_days_requested=strategy_reference_days,
            reference_dates=reference_dates,
            result_df=result_df,
            quality_report_path=quality_report_path,
        )

    selected_result = strategy_results[selected_strategy_key]
    emit_progress(progress_callback, "正在导出预测结果", 92)
    export_prediction(selected_result.result_df, output_file)
    emit_progress(progress_callback, f"预测完成，已选择策略 {selected_result.reference_strategy_label}", 100)

    return PredictCompareResult(
        forecast_date=selected_result.forecast_date,
        output_file=Path(output_file),
        template_updated=False,
        selected_strategy_key=selected_strategy_key,
        selected_strategy_label=REFERENCE_STRATEGIES[selected_strategy_key],
        strategy_results={
            key: PredictResult(
                forecast_date=value.forecast_date,
                output_file=value.output_file,
                template_updated=False,
                reference_strategy_key=value.reference_strategy_key,
                reference_strategy_label=value.reference_strategy_label,
                reference_days_requested=value.reference_days_requested,
                reference_dates=value.reference_dates,
                result_df=value.result_df.copy(),
                quality_report_path=value.quality_report_path,
            )
            for key, value in strategy_results.items()
        },
        quality_report_path=quality_report_path,
    )


def predict_multi_day_prices(
    history_dir: str | Path,
    forecast_file: str | Path = DEFAULT_MULTI_DAY_FORECAST_FILE,
    model_root: str | Path = DEFAULT_MODEL_ROOT,
    output_file: str | Path = DEFAULT_MULTI_DAY_OUTPUT_FILE,
    holiday_file: str | Path | None = None,
    reference_days: int = 1,
    knn_similarity_config: dict[str, object] | None = None,
    progress_callback: ProgressCallback | None = None,
) -> MultiDayPredictResult:
    ensure_xgboost_available()
    reference_days = normalize_reference_days(reference_days)
    knn_similarity_config = normalize_knn_similarity_config(knn_similarity_config)
    forecast_default_year = default_forecast_year_from_model(model_root)
    parsed = parse_multi_day_forecast_workbook(forecast_file, default_year=forecast_default_year)
    baseline_date = str(parsed.get("baseline_date") or "")
    forecast_dates = list(parsed.get("forecast_dates") or [])
    if not forecast_dates:
        raise ValueError("预测文件-N天没有可预测日期")
    emit_progress(progress_callback, f"已识别 {len(forecast_dates)} 个预测日", 5)

    latest_forecast_date = pd.Timestamp(max(forecast_dates)).normalize()
    history_start_date, history_end_date = prediction_history_window(latest_forecast_date, reference_days)
    history_df, _, holiday_dates, _, history_quality_issues = load_cached_history_collection(
        history_dir,
        holiday_file,
        True,
        model_root,
        progress_callback=progress_callback,
        progress_start=8,
        progress_end=24,
        progress_label="正在加载多天预测历史数据",
        start_date=history_start_date,
        end_date=history_end_date,
    )
    emit_progress(progress_callback, "正在加载模型", 26)
    metadata, models, _ = load_models(model_root)
    model_segment_definitions = normalize_segment_config(metadata.get("segment_config") or metadata.get("segments"))
    history_df["segment"] = assign_segments(history_df["period"], model_segment_definitions)
    quality_report_path = save_quality_report(
        task_type="predict",
        status=quality_report_status(history_quality_issues),
        issues=history_quality_issues,
        metadata={"forecast_file": str(forecast_file), "history_dir": str(history_dir), "mode": "multi_day"},
    )

    days: list[dict[str, object]] = []
    total_days = len(forecast_dates)
    for index, forecast_date in enumerate(forecast_dates, start=1):
        emit_progress(progress_callback, f"正在预测 {forecast_date} ({index}/{total_days})", 30 + int((index - 1) / total_days * 55))
        forecast_df, template_reference_df = build_multi_day_forecast_frame(parsed, forecast_date, holiday_dates)
        forecast_df["segment"] = assign_segments(forecast_df["period"], model_segment_definitions)
        forecast_df.attrs["force_template_lag_reference"] = True
        template_reference_df["segment"] = assign_segments(template_reference_df["period"], model_segment_definitions)
        result_df, reference_dates, strategy_key, strategy_label = run_prediction_with_strategy(
            history_df=history_df,
            forecast_df=forecast_df,
            template_reference_df=template_reference_df,
            holiday_dates=holiday_dates,
            metadata=metadata,
            models=models,
            reference_days=reference_days,
            reference_strategy="recent_n_days",
            knn_similarity_config=knn_similarity_config,
            progress_callback=progress_callback,
            segment_progress_start=32 + int((index - 1) / total_days * 55),
            segment_progress_span=max(1, int(45 / total_days)),
        )
        days.append(
            {
                "forecast_date": forecast_date,
                "reference_strategy_key": strategy_key,
                "reference_strategy_label": strategy_label,
                "reference_days_requested": reference_days,
                "reference_dates": reference_dates,
                "result_df": result_df,
                "rows": result_df,
            }
        )

    emit_progress(progress_callback, "正在导出多天预测结果", 94)
    export_multi_day_prediction_workbook({"days": days}, output_file)
    emit_progress(progress_callback, "多天预测完成", 100)
    return MultiDayPredictResult(
        baseline_date=baseline_date,
        forecast_dates=forecast_dates,
        output_file=Path(output_file),
        template_updated=False,
        selected_strategy_key="recent_n_days",
        selected_strategy_label=REFERENCE_STRATEGIES["recent_n_days"],
        days=days,
        quality_report_path=quality_report_path,
    )


def export_prediction(result_df: pd.DataFrame, output_file: str | Path) -> None:
    output_path = Path(output_file)
    export_columns = [
        "date",
        "period",
        "segment",
        "hour",
        TOTAL_LOAD_COLUMN,
        "net_load",
        RENEWABLE_POWER_COLUMN,
        THERMAL_SPACE_LOAD_RATIO_COLUMN,
        "thermal_on_capacity",
        THERMAL_CAPACITY_SOURCE_COLUMN,
        THERMAL_CAPACITY_VALUE_COLUMN,
        THERMAL_CAPACITY_MODEL_VALUE_COLUMN,
        THERMAL_CAPACITY_FILE_VALUE_COLUMN,
        DAY_TYPE_COLUMN,
        NET_LOAD_ONLY_SIMILAR_PRICE_COLUMN,
        KNN_SIMILAR_PRICE_COLUMN,
        WEIGHTED_KNN_PRICE_COLUMN,
        KNN_SIMILAR_COUNT_COLUMN,
        KNN_SIMILAR_DISTANCE_COLUMN,
        "model_variant",
        "model_similarity_diff",
        "predicted_price",
        "realtime_predicted_price",
        "price_spread_realtime_minus_dayahead",
        "dayahead_model_version",
        "realtime_model_version",
        "realtime_model_variant",
        "realtime_prediction_status",
        "realtime_prediction_error",
        SCENARIO_SIMILARITY_ADJUSTED_PRICE_COLUMN,
        SIMILARITY_BLEND_WEIGHT_COLUMN,
        SIMILARITY_ADJUSTMENT_COLUMN,
        SIMILARITY_ADJUSTMENT_REASON_COLUMN,
        HIGH_PRICE_ADJUSTED_PRICE_COLUMN,
        HIGH_PRICE_ADJUSTMENT_COLUMN,
        HIGH_PRICE_ADJUSTMENT_REASON_COLUMN,
        "predicted_interval",
        "predicted_interval_probability",
        "interval_probabilities",
        "high_price_probability",
        "extreme_price_probability",
        "interval_backtest_accuracy",
        "price_interval_consistency",
    ]
    export_df = result_df[[column for column in export_columns if column in result_df.columns]].copy()
    export_df["date"] = pd.to_datetime(export_df["date"]).dt.strftime("%Y-%m-%d")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.suffix.lower() == ".csv":
        export_df.to_csv(output_path, index=False, encoding="utf-8-sig")
        return
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        export_df.to_excel(writer, index=False, sheet_name="prediction")


def export_multi_day_prediction_workbook(result: dict[str, object] | MultiDayPredictResult, output_file: str | Path) -> Path:
    output_path = Path(output_file)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    days = result.days if isinstance(result, MultiDayPredictResult) else result.get("days", [])
    if not days:
        raise ValueError("\u6682\u65e0\u53ef\u5bfc\u51fa\u7684\u591a\u5929\u9884\u6d4b\u7ed3\u679c")

    temp_path = output_path.with_name(f".{output_path.stem}.{uuid.uuid4().hex}.tmp{output_path.suffix or '.xlsx'}")

    column_specs = [
        ("predicted_price", "\u6a21\u578b\u9884\u6d4b\u4ef7\u683c"),
        (SCENARIO_SIMILARITY_ADJUSTED_PRICE_COLUMN, "\u573a\u666f\u76f8\u4f3c\u6cd5\u4fee\u6b63\u53c2\u8003\u4ef7"),
        (NET_LOAD_ONLY_SIMILAR_PRICE_COLUMN, "\u51c0\u8d1f\u8377\u76f8\u4f3c\u6cd5\u9884\u6d4b\u4ef7\u683c"),
        (KNN_SIMILAR_PRICE_COLUMN, "KNN\u76f8\u4f3c\u6cd5\u9884\u6d4b\u4ef7\u683c"),
        (WEIGHTED_KNN_PRICE_COLUMN, "\u52a0\u6743KNN\u56de\u5f52\u9884\u6d4b\u4ef7\u683c"),
        ("predicted_interval", "\u4ef7\u683c\u533a\u95f4"),
        ("predicted_interval_probability", "\u4ef7\u683c\u533a\u95f4\u6982\u7387"),
        ("high_price_probability", "\u9ad8\u4ef7\u6982\u7387"),
    ]
    written_sheets = 0
    try:
        with pd.ExcelWriter(temp_path, engine="openpyxl") as writer:
            used_sheet_names: set[str] = set()
            for index, day in enumerate(days, start=1):
                forecast_date = str(day.get("forecast_date") or "")
                rows = day.get("rows")
                if rows is None:
                    rows = []
                if isinstance(rows, pd.DataFrame):
                    rows_df = rows.sort_values("period").copy() if "period" in rows.columns else rows.copy()
                else:
                    rows_df = pd.DataFrame(list(rows)) if rows else pd.DataFrame()
                    if "period" in rows_df.columns:
                        rows_df = rows_df.sort_values("period")

                export_data: dict[str, list[object]] = {}
                for key, label in column_specs:
                    if key in rows_df.columns:
                        export_data[f"{forecast_date} {label}".strip()] = rows_df[key].tolist()
                if not export_data:
                    fallback_date = forecast_date or f"\u7b2c{index}\u5929"
                    export_data[f"{fallback_date} \u6a21\u578b\u9884\u6d4b\u4ef7\u683c"] = []

                base_sheet_name = (forecast_date[:31] or f"prediction_{index}")[:31]
                sheet_name = base_sheet_name
                suffix = 2
                while sheet_name in used_sheet_names:
                    suffix_text = f"_{suffix}"
                    sheet_name = f"{base_sheet_name[:31 - len(suffix_text)]}{suffix_text}"
                    suffix += 1
                used_sheet_names.add(sheet_name)
                pd.DataFrame(export_data).to_excel(writer, index=False, sheet_name=sheet_name)
                written_sheets += 1

        if written_sheets <= 0:
            raise ValueError("\u6682\u65e0\u53ef\u5bfc\u51fa\u7684\u591a\u5929\u9884\u6d4b\u7ed3\u679c")
        workbook = load_workbook(temp_path)
        try:
            if not any(sheet.sheet_state == "visible" for sheet in workbook.worksheets):
                raise ValueError("\u591a\u5929\u9884\u6d4b\u5bfc\u51fa\u7ed3\u679c\u6ca1\u6709\u53ef\u89c1 Sheet")
        finally:
            workbook.close()
        temp_path.replace(output_path)
    except Exception:
        if temp_path.exists():
            try:
                temp_path.unlink()
            except OSError:
                pass
        raise
    return output_path

def save_workbook_with_replacement(
    workbook: object,
    target_file: str | Path,
    retries: int = 3,
    delay_seconds: float = 0.2,
) -> None:
    target_path = Path(target_file)
    target_path.parent.mkdir(parents=True, exist_ok=True)
    last_error: Exception | None = None
    for attempt in range(retries):
        temp_path = target_path.with_name(f".{target_path.stem}.{uuid.uuid4().hex}.tmp.xlsx")
        try:
            workbook.save(temp_path)  # type: ignore[attr-defined]
            validation_workbook = load_workbook(temp_path, read_only=True, data_only=True)
            try:
                if not validation_workbook.sheetnames:
                    raise ValueError(f"保存预测文件失败：临时文件无可用 Sheet，路径 {temp_path}")
            finally:
                validation_workbook.close()
            temp_path.replace(target_path)
            return
        except OSError as exc:
            last_error = exc
            gc.collect()
            if attempt < retries - 1:
                time.sleep(delay_seconds * (attempt + 1))
        except Exception:
            if temp_path.exists():
                try:
                    temp_path.unlink()
                except OSError:
                    pass
            raise
        finally:
            if temp_path.exists():
                try:
                    temp_path.unlink()
                except OSError:
                    pass
    if last_error is not None:
        raise last_error


def write_prediction_to_template(forecast_file: str | Path, result_df: pd.DataFrame) -> bool:
    workbook = load_workbook(forecast_file)
    for sheet in workbook.worksheets:
        if sheet.max_row < 2 or sheet.max_column < 1:
            continue
        headers = [sheet.cell(row=1, column=col).value for col in range(1, sheet.max_column + 1)]
        target_col = None
        for idx, header in enumerate(headers, start=1):
            if normalize_text(header) == normalize_text("预测价格"):
                target_col = idx
                break
        if target_col is None:
            continue
        predicted_values = result_df.sort_values("period")["predicted_price"].tolist()
        for row_idx, value in enumerate(predicted_values, start=2):
            sheet.cell(row=row_idx, column=target_col).value = round(float(value), 4)
        try:
            save_workbook_with_replacement(workbook, forecast_file)
        finally:
            workbook.close()
        return True
    workbook.close()
    return False


def load_training_log(model_root: Path = DEFAULT_MODEL_ROOT) -> pd.DataFrame:
    log_path = model_root / "training_runs.jsonl"
    if not log_path.exists():
        return pd.DataFrame()
    rows = [json.loads(line) for line in log_path.read_text(encoding="utf-8").splitlines() if line.strip()]
    return pd.DataFrame(rows)


def load_current_metadata(model_root: Path = DEFAULT_MODEL_ROOT) -> dict | None:
    try:
        current_dir = resolve_active_model_dir(model_root)
    except FileNotFoundError:
        return None
    return json.loads((current_dir / "metadata.json").read_text(encoding="utf-8"))


def read_model_metadata(model_dir: Path) -> dict | None:
    metadata_path = model_dir / "metadata.json"
    if not metadata_path.exists():
        return None
    return json.loads(metadata_path.read_text(encoding="utf-8"))


def summarize_model_quality_metrics(metrics: dict | None) -> dict[str, float | None]:
    if not isinstance(metrics, dict) or not metrics:
        return {
            "baseline_mae": None,
            "baseline_rmse": None,
            "final_mae": None,
            "final_rmse": None,
        }

    total_valid_rows = 0.0
    baseline_mae_weighted = 0.0
    baseline_rmse_squared_weighted = 0.0
    final_mae_weighted = 0.0
    final_rmse_squared_weighted = 0.0

    for segment_metrics in metrics.values():
        if not isinstance(segment_metrics, dict):
            continue
        valid_rows = segment_metrics.get("valid_rows")
        if valid_rows is None:
            continue
        try:
            valid_rows_value = float(valid_rows)
        except (TypeError, ValueError):
            continue
        if valid_rows_value <= 0:
            continue

        baseline_mae = segment_metrics.get("baseline_mae")
        baseline_rmse = segment_metrics.get("baseline_rmse")
        final_mae = segment_metrics.get("final_mae")
        final_rmse = segment_metrics.get("final_rmse")
        if any(value is None for value in (final_mae, final_rmse)):
            continue

        try:
            baseline_mae_value = float(baseline_mae) if baseline_mae is not None else np.nan
            baseline_rmse_value = float(baseline_rmse) if baseline_rmse is not None else np.nan
            final_mae_value = float(final_mae)
            final_rmse_value = float(final_rmse)
        except (TypeError, ValueError):
            continue

        total_valid_rows += valid_rows_value
        if not np.isnan(baseline_mae_value):
            baseline_mae_weighted += valid_rows_value * baseline_mae_value
        if not np.isnan(baseline_rmse_value):
            baseline_rmse_squared_weighted += valid_rows_value * (baseline_rmse_value**2)
        final_mae_weighted += valid_rows_value * final_mae_value
        final_rmse_squared_weighted += valid_rows_value * (final_rmse_value**2)

    if total_valid_rows <= 0:
        return {
            "baseline_mae": None,
            "baseline_rmse": None,
            "final_mae": None,
            "final_rmse": None,
        }

    return {
        "baseline_mae": round(baseline_mae_weighted / total_valid_rows, 4) if baseline_mae_weighted else None,
        "baseline_rmse": round(float(np.sqrt(baseline_rmse_squared_weighted / total_valid_rows)), 4) if baseline_rmse_squared_weighted else None,
        "final_mae": round(final_mae_weighted / total_valid_rows, 4),
        "final_rmse": round(float(np.sqrt(final_rmse_squared_weighted / total_valid_rows)), 4),
    }


def summarize_variant_display_metrics(metrics: dict | None) -> dict[str, float | int | None]:
    quality = summarize_model_quality_metrics(metrics)
    if not isinstance(metrics, dict) or not metrics:
        return {
            **quality,
            "train_rows": None,
            "valid_rows": None,
            "max_abs_error": None,
            "p90_abs_error": None,
            "direction_accuracy": None,
            "best_iteration_avg": None,
            "best_iteration_max": None,
            "best_iteration_min": None,
        }

    train_rows_total = 0
    valid_rows_total = 0
    max_abs_error = None
    p90_abs_error_weighted = 0.0
    direction_accuracy_weighted = 0.0
    p90_weight_rows = 0.0
    direction_weight_rows = 0.0
    best_iterations: list[int] = []

    for segment_metrics in metrics.values():
        if not isinstance(segment_metrics, dict):
            continue
        train_rows = segment_metrics.get("train_rows")
        valid_rows = segment_metrics.get("valid_rows")
        try:
            train_rows_value = int(float(train_rows or 0))
            valid_rows_value = int(float(valid_rows or 0))
        except (TypeError, ValueError):
            train_rows_value = 0
            valid_rows_value = 0
        train_rows_total += max(0, train_rows_value)
        valid_rows_total += max(0, valid_rows_value)

        try:
            current_max = float(segment_metrics.get("max_abs_error"))
            max_abs_error = current_max if max_abs_error is None else max(max_abs_error, current_max)
        except (TypeError, ValueError):
            pass

        try:
            p90_value = float(segment_metrics.get("p90_abs_error"))
            if valid_rows_value > 0:
                p90_abs_error_weighted += valid_rows_value * p90_value
                p90_weight_rows += valid_rows_value
        except (TypeError, ValueError):
            pass

        try:
            direction_value = float(segment_metrics.get("direction_accuracy"))
            if valid_rows_value > 0:
                direction_accuracy_weighted += valid_rows_value * direction_value
                direction_weight_rows += valid_rows_value
        except (TypeError, ValueError):
            pass

        try:
            best_iteration = segment_metrics.get("best_iteration")
            if best_iteration is not None:
                best_iterations.append(int(float(best_iteration)))
        except (TypeError, ValueError):
            pass

    return {
        **quality,
        "train_rows": train_rows_total or None,
        "valid_rows": valid_rows_total or None,
        "max_abs_error": round(max_abs_error, 4) if max_abs_error is not None else None,
        "p90_abs_error": round(p90_abs_error_weighted / p90_weight_rows, 4) if p90_weight_rows else None,
        "direction_accuracy": round(direction_accuracy_weighted / direction_weight_rows, 4) if direction_weight_rows else None,
        "best_iteration_avg": round(float(np.mean(best_iterations)), 2) if best_iterations else None,
        "best_iteration_max": max(best_iterations) if best_iterations else None,
        "best_iteration_min": min(best_iterations) if best_iterations else None,
    }


def nested_metric(data: dict | None, *keys: str) -> object | None:
    current: object = data or {}
    for key in keys:
        if not isinstance(current, dict):
            return None
        current = current.get(key)
    return current


def version_rolling_metrics(metadata: dict) -> dict[str, object | None]:
    rolling = metadata.get("rolling_backtest_metrics") if isinstance(metadata.get("rolling_backtest_metrics"), dict) else {}
    metrics: dict[str, object | None] = {}
    for days in ("14", "30"):
        row = rolling.get(days) if isinstance(rolling.get(days), dict) else {}
        prefix = f"rolling_{days}"
        metrics[f"{prefix}_mae"] = nested_metric(row, "overall", "mae")
        metrics[f"{prefix}_rmse"] = nested_metric(row, "overall", "rmse")
        metrics[f"{prefix}_direction_accuracy"] = nested_metric(row, "overall", "direction_accuracy")
        metrics[f"{prefix}_max_abs_error"] = nested_metric(row, "overall", "max_abs_error")
        metrics[f"{prefix}_p90_abs_error"] = nested_metric(row, "overall", "p90_abs_error")
        metrics[f"{prefix}_thermal_space_similarity_mae"] = nested_metric(row, "net_load_only_baseline", "mae")
        metrics[f"{prefix}_thermal_space_similarity_rmse"] = nested_metric(row, "net_load_only_baseline", "rmse")
        metrics[f"{prefix}_thermal_space_similarity_improvement"] = nested_metric(row, "model_vs_net_load_only_similarity", "mae_improvement")
        metrics[f"{prefix}_thermal_space_similarity_improvement_pct"] = nested_metric(row, "model_vs_net_load_only_similarity", "mae_improvement_pct")
        metrics[f"{prefix}_high_price_mae"] = nested_metric(row, "spike_errors", "high", "mae")
        metrics[f"{prefix}_high_price_rmse"] = nested_metric(row, "spike_errors", "high", "rmse")
        metrics[f"{prefix}_high_price_thermal_space_similarity_mae"] = nested_metric(row, "spike_errors", "high_net_load_only_baseline", "mae")
        metrics[f"{prefix}_evening_peak_mae"] = nested_metric(row, "segments", "evening_peak", "mae")
        metrics[f"{prefix}_evening_peak_rmse"] = nested_metric(row, "segments", "evening_peak", "rmse")
    return metrics


def version_algorithm_variants(metadata: dict) -> list[dict[str, object | None]]:
    variants = metadata.get("model_variants") if isinstance(metadata.get("model_variants"), dict) else {}
    variant_metrics = metadata.get("variant_metrics") if isinstance(metadata.get("variant_metrics"), dict) else {}
    quality_metrics = metadata.get("variant_quality_metrics") if isinstance(metadata.get("variant_quality_metrics"), dict) else {}
    selected_key = metadata.get("selected_model_key")
    rows: list[dict[str, object | None]] = []
    for variant_key, variant_meta in variants.items():
        if not isinstance(variant_meta, dict):
            variant_meta = {}
        quality = quality_metrics.get(variant_key)
        if not isinstance(quality, dict):
            quality = summarize_model_quality_metrics(variant_metrics.get(variant_key))
        display_metrics = summarize_variant_display_metrics(variant_metrics.get(variant_key))
        rows.append(
            {
                "variant_key": variant_key,
                "model_backend": variant_meta.get("model_backend"),
                "model_backend_label": variant_meta.get("model_backend_label") or variant_meta.get("model_backend") or variant_key,
                "is_selected": variant_key == selected_key,
                "final_mae": quality.get("final_mae") if isinstance(quality, dict) else None,
                "final_rmse": quality.get("final_rmse") if isinstance(quality, dict) else None,
                "train_rows": display_metrics.get("train_rows"),
                "valid_rows": display_metrics.get("valid_rows"),
                "max_abs_error": display_metrics.get("max_abs_error"),
                "p90_abs_error": display_metrics.get("p90_abs_error"),
                "direction_accuracy": display_metrics.get("direction_accuracy"),
                "best_iteration_avg": display_metrics.get("best_iteration_avg"),
                "best_iteration_max": display_metrics.get("best_iteration_max"),
                "best_iteration_min": display_metrics.get("best_iteration_min"),
            }
        )
    if rows:
        return rows
    display_metrics = summarize_variant_display_metrics(metadata.get("metrics"))
    return [
        {
            "variant_key": selected_key,
            "model_backend": metadata.get("selected_model_backend"),
            "model_backend_label": metadata.get("selected_model_backend_label") or metadata.get("selected_model_backend"),
            "is_selected": True,
            "final_mae": display_metrics.get("final_mae"),
            "final_rmse": display_metrics.get("final_rmse"),
            "train_rows": display_metrics.get("train_rows"),
            "valid_rows": display_metrics.get("valid_rows"),
            "max_abs_error": display_metrics.get("max_abs_error"),
            "p90_abs_error": display_metrics.get("p90_abs_error"),
            "direction_accuracy": display_metrics.get("direction_accuracy"),
            "best_iteration_avg": display_metrics.get("best_iteration_avg"),
            "best_iteration_max": display_metrics.get("best_iteration_max"),
            "best_iteration_min": display_metrics.get("best_iteration_min"),
        }
    ]


def version_price_interval_metrics(metadata: dict) -> dict[str, object | None]:
    interval_model = metadata.get("price_interval_model") if isinstance(metadata.get("price_interval_model"), dict) else {}
    metrics = interval_model.get("metrics") if isinstance(interval_model.get("metrics"), dict) else {}
    overall = metrics.get("overall") if isinstance(metrics.get("overall"), dict) else metrics
    return {
        "accuracy": nested_metric(overall, "interval_accuracy"),
        "top2_accuracy": nested_metric(overall, "top2_accuracy"),
        "logloss": nested_metric(overall, "logloss") or nested_metric(overall, "log_loss"),
        "high_price_recall": nested_metric(overall, "high_price_recall"),
        "score": nested_metric(overall, "score"),
    }


def version_interval_training_rows(interval_model: dict) -> dict[str, int | None]:
    selected_key = interval_model.get("selected_model_key") or interval_model.get("model_key")
    variant_metrics = interval_model.get("variant_metrics") if isinstance(interval_model.get("variant_metrics"), dict) else {}
    selected_metrics = variant_metrics.get(selected_key) if selected_key in variant_metrics else None
    if not isinstance(selected_metrics, dict):
        selected_metrics = interval_model.get("metrics") if isinstance(interval_model.get("metrics"), dict) else {}
    segments = selected_metrics.get("segments") if isinstance(selected_metrics.get("segments"), dict) else {}
    train_rows = 0
    valid_rows = 0
    found_train = False
    found_valid = False
    for row in segments.values():
        if not isinstance(row, dict):
            continue
        try:
            train_rows += int(float(row.get("train_rows")))
            found_train = True
        except (TypeError, ValueError):
            pass
        try:
            valid_rows += int(float(row.get("valid_rows")))
            found_valid = True
        except (TypeError, ValueError):
            pass
    return {
        "train_rows": train_rows if found_train else interval_model.get("train_rows"),
        "valid_rows": valid_rows if found_valid else interval_model.get("valid_rows"),
    }


def version_thermal_capacity_metrics(metadata: dict) -> dict[str, object | None]:
    thermal_model = metadata.get("thermal_capacity_model") if isinstance(metadata.get("thermal_capacity_model"), dict) else {}
    metrics = thermal_model.get("metrics") if isinstance(thermal_model.get("metrics"), dict) else {}
    overall = metrics.get("overall") if isinstance(metrics.get("overall"), dict) else metrics
    return {
        "mae": nested_metric(overall, "mae"),
        "rmse": nested_metric(overall, "rmse"),
        "max_error": nested_metric(overall, "max_error"),
        "bias": nested_metric(overall, "bias"),
        "score": nested_metric(overall, "score"),
    }


def version_training_config(metadata: dict, key: str, fallback_window: object | None = None, fallback_valid_days: object | None = None) -> dict[str, object | None]:
    config = metadata.get(key) if isinstance(metadata.get(key), dict) else {}
    return {
        "training_mode": config.get("training_mode") or metadata.get("training_mode"),
        "training_window_days": config.get("training_window_days", fallback_window),
        "valid_days": config.get("valid_days", fallback_valid_days),
        "num_boost_round": config.get("num_boost_round") or metadata.get("num_boost_round"),
        "start_date": config.get("start_date") or metadata.get("train_start_date"),
        "end_date": config.get("end_date") or metadata.get("train_end_date"),
    }


def build_version_summary(metadata: dict, quality_metrics: dict[str, object | None], rolling_metrics: dict[str, object | None], segment_count: int | None) -> dict[str, object]:
    interval_model = metadata.get("price_interval_model") if isinstance(metadata.get("price_interval_model"), dict) else {}
    thermal_model = metadata.get("thermal_capacity_model") if isinstance(metadata.get("thermal_capacity_model"), dict) else {}
    interval_metrics = version_price_interval_metrics(metadata)
    interval_rows = version_interval_training_rows(interval_model)
    thermal_metrics = version_thermal_capacity_metrics(metadata)
    return {
        "price_model": {
            "backend": metadata.get("selected_model_backend"),
            "backend_label": metadata.get("selected_model_backend_label") or metadata.get("selected_model_backend"),
            "selected_model_key": metadata.get("selected_model_key"),
            "mae": quality_metrics.get("final_mae"),
            "rmse": quality_metrics.get("final_rmse"),
            "rolling_30_mae": rolling_metrics.get("rolling_30_mae"),
            "rolling_30_rmse": rolling_metrics.get("rolling_30_rmse"),
            "high_price_mae": rolling_metrics.get("rolling_30_high_price_mae"),
            "training_window_days": metadata.get("training_window_days"),
            "valid_days": metadata.get("valid_days"),
            "sample_rows": metadata.get("sample_rows"),
            "segment_count": segment_count,
        },
        "price_interval_model": {
            "enabled": interval_model.get("enabled"),
            "selected_model_key": interval_model.get("selected_model_key") or interval_model.get("model_key"),
            "accuracy": interval_metrics.get("accuracy"),
            "top2_accuracy": interval_metrics.get("top2_accuracy"),
            "logloss": interval_metrics.get("logloss"),
            "high_price_recall": interval_metrics.get("high_price_recall"),
            "interval_count": len(interval_model.get("intervals") or []) if isinstance(interval_model.get("intervals"), list) else None,
            "train_rows": interval_rows.get("train_rows"),
            "valid_rows": interval_rows.get("valid_rows"),
            "feature_count": len(interval_model.get("feature_columns") or []) if isinstance(interval_model.get("feature_columns"), list) else None,
        },
        "thermal_capacity_model": {
            "enabled": thermal_model.get("enabled"),
            "selected_model_key": thermal_model.get("selected_model_key"),
            "selected_model_backend": thermal_model.get("selected_model_backend"),
            "selected_model_backend_label": thermal_model.get("selected_model_backend_label") or thermal_model.get("selected_model_backend"),
            "candidate_count": len(thermal_model.get("model_variants") or {}) if isinstance(thermal_model.get("model_variants"), dict) else None,
            "mae": thermal_metrics.get("mae"),
            "rmse": thermal_metrics.get("rmse"),
            "max_error": thermal_metrics.get("max_error"),
            "bias": thermal_metrics.get("bias"),
            "feature_count": len(thermal_model.get("feature_columns") or []) if isinstance(thermal_model.get("feature_columns"), list) else None,
        },
    }


def build_version_detail(metadata: dict, algorithm_variants: list[dict[str, object | None]], segment_config: list | None) -> dict[str, object]:
    interval_model = metadata.get("price_interval_model") if isinstance(metadata.get("price_interval_model"), dict) else {}
    thermal_model = metadata.get("thermal_capacity_model") if isinstance(metadata.get("thermal_capacity_model"), dict) else {}
    interval_metrics = version_price_interval_metrics(metadata)
    interval_rows = version_interval_training_rows(interval_model)
    thermal_metrics = version_thermal_capacity_metrics(metadata)
    intervals = interval_model.get("intervals") if isinstance(interval_model.get("intervals"), list) else []
    interval_features = interval_model.get("feature_columns") if isinstance(interval_model.get("feature_columns"), list) else []
    thermal_features = thermal_model.get("feature_columns") if isinstance(thermal_model.get("feature_columns"), list) else []
    return {
        "training": {
            "price": version_training_config(metadata, "price_model_training_config", metadata.get("training_window_days"), metadata.get("valid_days")),
            "interval": version_training_config(
                metadata,
                "interval_model_training_config",
                metadata.get("interval_training_window_days"),
                metadata.get("interval_valid_days"),
            ),
            "thermal_capacity": version_training_config(
                metadata,
                "thermal_capacity_model_training_config",
                metadata.get("thermal_capacity_training_window_days"),
                metadata.get("thermal_capacity_valid_days"),
            ),
        },
        "price_model": {
            "selected_model_key": metadata.get("selected_model_key"),
            "selected_model_backend": metadata.get("selected_model_backend"),
            "selected_model_backend_label": metadata.get("selected_model_backend_label") or metadata.get("selected_model_backend"),
            "algorithm_variants": algorithm_variants,
            "selected_segment_price_models": metadata.get("selected_segment_price_models") if isinstance(metadata.get("selected_segment_price_models"), dict) else {},
            "model_backend_candidates": metadata.get("model_backend_candidates") if isinstance(metadata.get("model_backend_candidates"), list) else [],
            "variant_quality_metrics": metadata.get("variant_quality_metrics") if isinstance(metadata.get("variant_quality_metrics"), dict) else {},
            "segment_config": segment_config or [],
            "segment_count": len(segment_config or []),
            "high_price_weighting": metadata.get("high_price_weighting") if isinstance(metadata.get("high_price_weighting"), dict) else {},
            "high_price_adjustment_metrics": metadata.get("high_price_adjustment_metrics") if isinstance(metadata.get("high_price_adjustment_metrics"), dict) else {},
            "rolling_backtest_metrics": metadata.get("rolling_backtest_metrics") if isinstance(metadata.get("rolling_backtest_metrics"), dict) else {},
        },
        "price_interval_model": {
            "enabled": interval_model.get("enabled"),
            "selected_model_key": interval_model.get("selected_model_key") or interval_model.get("model_key"),
            "selected_model_backend": interval_model.get("selected_model_backend"),
            "selected_model_backend_label": interval_model.get("selected_model_backend_label") or interval_model.get("selected_model_backend"),
            "metrics": interval_metrics,
            "intervals": intervals,
            "interval_count": len(intervals),
            "feature_columns": interval_features,
            "feature_count": len(interval_features),
            "train_rows": interval_rows.get("train_rows"),
            "valid_rows": interval_rows.get("valid_rows"),
        },
        "thermal_capacity_model": {
            "enabled": thermal_model.get("enabled"),
            "selected_model_key": thermal_model.get("selected_model_key"),
            "selected_model_backend": thermal_model.get("selected_model_backend"),
            "selected_model_backend_label": thermal_model.get("selected_model_backend_label") or thermal_model.get("selected_model_backend"),
            "model_variants": thermal_model.get("model_variants") if isinstance(thermal_model.get("model_variants"), dict) else {},
            "variant_metrics": thermal_model.get("variant_metrics") if isinstance(thermal_model.get("variant_metrics"), dict) else {},
            "metrics": thermal_metrics,
            "feature_columns": thermal_features,
            "feature_count": len(thermal_features),
            "train_rows": thermal_model.get("train_rows"),
            "valid_rows": thermal_model.get("valid_rows"),
        },
    }


def list_model_versions(model_root: Path = DEFAULT_MODEL_ROOT) -> pd.DataFrame:
    current_dir, previous_dir, history_dir = ensure_model_dirs(model_root)
    current_metadata = read_model_metadata(current_dir)
    previous_metadata = read_model_metadata(previous_dir)
    current_run_id = current_metadata.get("run_id") if current_metadata else None
    previous_run_id = previous_metadata.get("run_id") if previous_metadata else None
    pinned_version_key = read_pinned_default_version(model_root)

    rows: list[dict] = []
    for run_dir in sorted(history_dir.iterdir(), reverse=True):
        if not run_dir.is_dir():
            continue
        if not run_dir.name.startswith("run_"):
            continue
        metadata = read_model_metadata(run_dir)
        if metadata is None:
            continue
        quality_metrics = summarize_model_quality_metrics(metadata.get("metrics"))
        segment_config = metadata.get("segment_config") or metadata.get("segments") or []
        segment_count = len(segment_config) if isinstance(segment_config, list) else None
        high_price_weighting = metadata.get("high_price_weighting") if isinstance(metadata.get("high_price_weighting"), dict) else {}
        rolling_metrics = version_rolling_metrics(metadata)
        algorithm_variants = version_algorithm_variants(metadata)
        version_summary = build_version_summary(metadata, quality_metrics, rolling_metrics, segment_count)
        version_detail = build_version_detail(metadata, algorithm_variants, segment_config if isinstance(segment_config, list) else [])
        rows.append(
            {
                "version_key": run_dir.name,
                "run_id": metadata.get("run_id", run_dir.name),
                "created_at": metadata.get("created_at"),
                "train_start_date": metadata.get("train_start_date"),
                "train_end_date": metadata.get("train_end_date"),
                "training_mode": metadata.get("training_mode"),
                "training_window_days": metadata.get("training_window_days"),
                "valid_days": metadata.get("valid_days"),
                "num_boost_round": metadata.get("num_boost_round"),
                "sample_rows": metadata.get("sample_rows"),
                "is_default": metadata.get("run_id") == current_run_id,
                "is_previous": metadata.get("run_id") == previous_run_id,
                "is_pinned_default": run_dir.name == pinned_version_key,
                "model_type": metadata.get("model_type") or metadata.get("prediction_method"),
                "selected_model_key": metadata.get("selected_model_key"),
                "selected_model_backend": metadata.get("selected_model_backend"),
                "selected_model_backend_label": metadata.get("selected_model_backend_label") or metadata.get("selected_model_backend"),
                "algorithm_variants": algorithm_variants,
                "segment_mode": metadata.get("segment_mode") or ("custom" if metadata.get("segment_config") else "default"),
                "segment_count": segment_count,
                "high_price_weight_enabled": high_price_weighting.get("enabled"),
                "high_price_quantile": high_price_weighting.get("quantile"),
                "high_price_weight_multiplier": high_price_weighting.get("multiplier"),
                "high_price_threshold": high_price_weighting.get("threshold"),
                "baseline_mae": quality_metrics.get("baseline_mae"),
                "baseline_rmse": quality_metrics.get("baseline_rmse"),
                "final_mae": quality_metrics.get("final_mae"),
                "final_rmse": quality_metrics.get("final_rmse"),
                **rolling_metrics,
                "version_summary": version_summary,
                "version_detail": version_detail,
                "path": str(run_dir),
            }
        )
    df = pd.DataFrame(rows)
    if df.empty:
        return df
    df = df.sort_values("created_at", ascending=False).reset_index(drop=True)
    return df


def activate_model_version(model_root: Path, version_key: str, persist_default: bool = True) -> Path:
    current_dir, previous_dir, history_dir = ensure_model_dirs(model_root)
    selected_dir = history_dir / version_key
    if not selected_dir.exists():
        raise FileNotFoundError(f"未找到指定模型版本: {version_key}")
    if current_dir.exists():
        if previous_dir.exists():
            remove_tree_with_retry(previous_dir)
        copy_tree_with_retry(current_dir, previous_dir)
        remove_tree_with_retry(current_dir)
    copy_tree_with_retry(selected_dir, current_dir)
    if persist_default:
        write_pinned_default_version(model_root, version_key)
    return current_dir


def delete_model_versions(model_root: Path, version_keys: list[str]) -> list[str]:
    _, _, history_dir = ensure_model_dirs(model_root)
    deleted: list[str] = []
    pinned_version_key = read_pinned_default_version(model_root)
    for version_key in version_keys:
        target_dir = history_dir / version_key
        if target_dir.exists():
            remove_tree_with_retry(target_dir)
            deleted.append(version_key)
    if pinned_version_key and pinned_version_key in deleted:
        write_pinned_default_version(model_root, None)
    return deleted


def print_metrics(metrics_summary: dict[str, dict[str, float | int | None]]) -> None:
    print("训练完成，各分时段验证结果：")
    for segment_name, metric in metrics_summary.items():
        final_mae = metric.get("final_mae")
        final_rmse = metric.get("final_rmse")
        baseline_rmse = metric.get("baseline_rmse")
        if final_mae is None or final_rmse is None:
            print(f"  - {segment_name}: train_rows={metric['train_rows']}, valid_rows={metric['valid_rows']}")
        else:
            print(
                f"  - {segment_name}: baseline_RMSE={baseline_rmse}, "
                f"final_MAE={final_mae}, final_RMSE={final_rmse}"
            )
