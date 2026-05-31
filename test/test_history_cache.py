from __future__ import annotations

import json
import sys
import time
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

import pandas as pd


PROJECT_ROOT = Path(__file__).resolve().parent.parent


class HistoryCacheTests(unittest.TestCase):
    removed_previous_day_feature = "lag_" + "96"

    def test_realtime_and_dayahead_history_share_excel_file_signature_cache(self) -> None:
        from dayahead_core import history_cache_signature

        with TemporaryDirectory() as temp_dir:
            history_dir = Path(temp_dir) / "history"
            history_dir.mkdir()
            (history_dir / "2026年5月.xlsx").write_bytes(b"fake")

            day = history_cache_signature(history_dir, data_mode="dayahead")
            realtime = history_cache_signature(history_dir, data_mode="realtime")

        self.assertEqual(day["excel_files"], realtime["excel_files"])
        self.assertEqual(day["history_data_mode"], "dayahead")
        self.assertEqual(realtime["history_data_mode"], "realtime")
        self.assertNotEqual(day["cache_key"], realtime["cache_key"])

    def test_training_run_fingerprint_separates_realtime_and_dayahead(self) -> None:
        from dayahead_core import TrainConfig, training_run_fingerprint

        history_signature = {"cache_key": "same-files"}
        day = training_run_fingerprint(TrainConfig(data_mode="dayahead"), history_signature)
        realtime = training_run_fingerprint(TrainConfig(data_mode="realtime"), history_signature)

        self.assertNotEqual(day, realtime)

    def test_write_prediction_to_template_falls_back_when_direct_save_fails(self) -> None:
        from openpyxl import Workbook, load_workbook
        from openpyxl.workbook.workbook import Workbook as OpenpyxlWorkbook

        from dayahead_core import write_prediction_to_template

        with TemporaryDirectory() as temp_dir:
            template_path = Path(temp_dir) / "forecast.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.cell(row=1, column=1).value = "period"
            worksheet.cell(row=1, column=2).value = "预测价格"
            for period in range(1, 4):
                worksheet.cell(row=period + 1, column=1).value = period
            workbook.save(template_path)
            workbook.close()

            result_df = pd.DataFrame({"period": [1, 2, 3], "predicted_price": [101.12345, 202.0, 303.0]})
            original_save = OpenpyxlWorkbook.save

            def fail_direct_save(self: OpenpyxlWorkbook, filename: object) -> None:
                if Path(filename) == template_path:
                    raise OSError(22, "Invalid argument", str(template_path))
                return original_save(self, filename)

            with patch.object(OpenpyxlWorkbook, "save", fail_direct_save):
                self.assertTrue(write_prediction_to_template(template_path, result_df))

            saved = load_workbook(template_path, data_only=True)
            try:
                sheet = saved.active
                self.assertEqual(sheet.cell(row=2, column=2).value, 101.1235)
                self.assertEqual(sheet.cell(row=3, column=2).value, 202.0)
                self.assertEqual(sheet.cell(row=4, column=2).value, 303.0)
            finally:
                saved.close()

    def test_predict_prices_does_not_write_back_forecast_template(self) -> None:
        from dayahead_core import predict_prices

        result_df = pd.DataFrame({"date": ["2026-05-22"], "period": [1], "predicted_price": [321.0]})

        with (
            patch("dayahead_core.ensure_xgboost_available"),
            patch(
                "dayahead_core.prepare_prediction_inputs",
                return_value=(pd.DataFrame(), pd.DataFrame(), pd.DataFrame(), set(), {}, {}, None),
            ),
            patch(
                "dayahead_core.run_prediction_with_strategy",
                return_value=(result_df, ["2026-05-21"], "recent_n_days", "recent"),
            ),
            patch("dayahead_core.export_prediction"),
            patch("dayahead_core.write_prediction_to_template") as write_prediction_to_template,
        ):
            result = predict_prices("history", "forecast.xlsx", output_file="output.xlsx")

        write_prediction_to_template.assert_not_called()
        self.assertFalse(result.template_updated)

    def test_predict_prices_compare_does_not_write_back_forecast_template(self) -> None:
        from dayahead_core import predict_prices_compare

        result_df = pd.DataFrame({"date": ["2026-05-22"], "period": [1], "predicted_price": [321.0]})

        with (
            patch("dayahead_core.ensure_xgboost_available"),
            patch(
                "dayahead_core.prepare_prediction_inputs",
                return_value=(pd.DataFrame(), pd.DataFrame(), pd.DataFrame(), set(), {}, {}, None),
            ),
            patch(
                "dayahead_core.run_prediction_with_strategy",
                return_value=(result_df, ["2026-05-21"], "recent_n_days", "recent"),
            ),
            patch("dayahead_core.export_prediction"),
            patch("dayahead_core.write_prediction_to_template") as write_prediction_to_template,
        ):
            result = predict_prices_compare("history", "forecast.xlsx", output_file="output.xlsx")

        write_prediction_to_template.assert_not_called()
        self.assertFalse(result.template_updated)
        self.assertTrue(all(not item.template_updated for item in result.strategy_results.values()))

    def test_training_feature_columns_do_not_include_previous_day_price(self) -> None:
        from dayahead_core import TrainConfig, training_feature_columns

        self.assertNotIn(self.removed_previous_day_feature, training_feature_columns(TrainConfig()))
        self.assertNotIn("similar_price", training_feature_columns(TrainConfig()))
        self.assertNotIn("similar_gap", training_feature_columns(TrainConfig()))

    def test_direct_model_variants_define_base_and_lag_variants(self) -> None:
        from dayahead_core import LAG_FEATURE_COLUMNS, direct_model_variants

        variants = direct_model_variants()

        self.assertEqual(list(variants), ["no_lag_96", "with_lag_96"])
        self.assertNotIn(self.removed_previous_day_feature, variants["no_lag_96"])
        self.assertNotIn("similar_price", variants["no_lag_96"])
        self.assertEqual(variants["with_lag_96"], LAG_FEATURE_COLUMNS)
        self.assertIn("price_lag_1d", variants["with_lag_96"])
        self.assertNotIn(self.removed_previous_day_feature, variants["with_lag_96"])
        self.assertNotIn("similar_price", variants["with_lag_96"])

    def test_similarity_features_do_not_create_residual_training_target(self) -> None:
        from dayahead_core import TARGET_COLUMN, attach_similarity_features

        rows = []
        for date_text, price_base in [("2026-05-01", 100.0), ("2026-05-02", 120.0)]:
            for period in [1, 2]:
                rows.append(
                    {
                        "date": pd.Timestamp(date_text),
                        "period": period,
                        "net_load": 100.0 + period,
                        "renewable_power": 20.0,
                        "thermal_on_capacity": 1000.0,
                        "day_type": "workday",
                        TARGET_COLUMN: price_base + period,
                    }
                )

        frame = attach_similarity_features(pd.DataFrame(rows), reference_days=1)

        self.assertNotIn("residual_target", frame.columns)

    def test_prediction_error_summary_includes_plain_language_metrics(self) -> None:
        from dayahead_core import summarize_model_vs_baseline, summarize_prediction_errors

        frame = pd.DataFrame(
            {
                "date": ["2026-04-01"] * 4,
                "period": [1, 2, 3, 4],
                "actual": [100.0, 120.0, 90.0, 130.0],
                "predicted": [105.0, 115.0, 95.0, 125.0],
                "net_load_only_similar_predicted": [110.0, 100.0, 80.0, 120.0],
            }
        )

        summary = summarize_prediction_errors(frame)
        comparison = summarize_model_vs_baseline(frame)

        self.assertEqual(summary["rows"], 4)
        self.assertEqual(summary["mae"], 5.0)
        self.assertEqual(summary["max_abs_error"], 5.0)
        self.assertEqual(summary["direction_accuracy"], 100.0)
        self.assertGreater(comparison["mae_improvement"], 0)

    def test_custom_segment_config_maps_periods_continuously(self) -> None:
        from dayahead_core import assign_segments, normalize_segment_config

        segments = normalize_segment_config(
            [
                {"name": "low", "start_time": "00:00", "end_time": "08:00"},
                {"name": "mid", "start_time": "08:00", "end_time": "18:00"},
                {"name": "high", "start_time": "18:00", "end_time": "24:00"},
            ]
        )
        assigned = assign_segments(pd.Series([1, 32, 33, 72, 73, 96]), segments).tolist()

        self.assertEqual(assigned, ["low", "low", "mid", "mid", "high", "high"])

    def test_high_price_weighting_marks_only_expensive_samples(self) -> None:
        from dayahead_core import high_price_sample_weights

        frame = pd.DataFrame({"price": [100.0, 200.0, 500.0]})

        weights = high_price_sample_weights(frame, "price", True, 200.0, 3.0)

        self.assertEqual(weights.tolist(), [1.0, 3.0, 3.0])

    def test_window_optimization_payload_keeps_advanced_training_options(self) -> None:
        from api_server import build_window_optimization_worker_payload

        payload = {
            "valid_days": 30,
            "num_boost_round": 500,
            "fine_radius": 7,
            "segment_mode": "custom",
            "segment_config": [
                {"name": "low", "start_time": "00:00", "end_time": "08:00"},
                {"name": "mid", "start_time": "08:00", "end_time": "18:00"},
                {"name": "high", "start_time": "18:00", "end_time": "24:00"},
            ],
            "high_price_weighting": {"enabled": True, "quantile": 0.85, "multiplier": 3.0},
        }

        worker_payload = build_window_optimization_worker_payload("manual", payload, force=True)

        self.assertEqual(worker_payload["valid_days"], 30)
        self.assertEqual(worker_payload["num_boost_round"], 500)
        self.assertEqual(worker_payload["fine_radius"], 7)
        self.assertEqual(worker_payload["segment_mode"], "custom")
        self.assertEqual(len(worker_payload["segment_config"]), 3)
        self.assertTrue(worker_payload["high_price_weighting"]["enabled"])
        self.assertEqual(worker_payload["high_price_weighting"]["quantile"], 0.85)
        self.assertEqual(worker_payload["high_price_weighting"]["multiplier"], 3.0)
        self.assertTrue(worker_payload["force"])

    def test_status_refresh_does_not_start_window_optimization_check(self) -> None:
        from api_server import ApiHandler

        handler = object.__new__(ApiHandler)

        with (
            patch("api_server.check_window_optimization") as check_window_optimization,
            patch("api_server.STATE.status", return_value={}),
            patch("api_server.load_current_metadata", return_value={}),
            patch("api_server.window_optimization_state", return_value={"history_cache_key": "old"}),
        ):
            status = handler.get_status()

        check_window_optimization.assert_not_called()
        self.assertEqual(status["window_optimization"], {"history_cache_key": "old"})

    def test_window_optimization_payload_uses_saved_advanced_preferences(self) -> None:
        import api_server
        from dayahead_core import save_training_preferences

        root = Path(__file__).resolve().parent / ".test_tmp" / f"prefs_{time.time_ns()}"
        root.mkdir(parents=True, exist_ok=True)
        save_training_preferences(
            root,
            {
                "segment_mode": "custom",
                "segment_config": [
                    {"name": "low", "start_time": "00:00", "end_time": "12:00"},
                    {"name": "high", "start_time": "12:00", "end_time": "24:00"},
                ],
                "high_price_weighting": {"enabled": True, "quantile": 0.9, "multiplier": 2.5},
            },
        )

        original_model_root = api_server.MODEL_ROOT
        try:
            api_server.MODEL_ROOT = root
            worker_payload = api_server.build_window_optimization_worker_payload("startup", {})
        finally:
            api_server.MODEL_ROOT = original_model_root

        self.assertEqual(worker_payload["segment_mode"], "custom")
        self.assertEqual(len(worker_payload["segment_config"]), 2)
        self.assertTrue(worker_payload["high_price_weighting"]["enabled"])
        self.assertEqual(worker_payload["high_price_weighting"]["quantile"], 0.9)
        self.assertEqual(worker_payload["high_price_weighting"]["multiplier"], 2.5)

    def test_training_result_validation_rejects_mismatched_options(self) -> None:
        import api_server

        root = Path(__file__).resolve().parent / ".test_tmp" / f"metadata_{time.time_ns()}"
        model_dir = root / "history" / "run_test"
        model_dir.mkdir(parents=True, exist_ok=True)
        (model_dir / "metadata.json").write_text(
            "{"
            '"segment_mode":"default",'
            '"segment_config":[{"name":"night","start_time":"00:00","end_time":"24:00"}],'
            '"high_price_weighting":{"enabled":false,"quantile":0.8,"multiplier":2.0}'
            "}",
            encoding="utf-8",
        )
        result = type("Result", (), {"history_model_dir": model_dir})()

        with self.assertRaises(RuntimeError):
            api_server.validate_training_result_options(
                result,
                [{"name": "low", "start_time": "00:00", "end_time": "12:00"}, {"name": "high", "start_time": "12:00", "end_time": "24:00"}],
                {"enabled": True, "quantile": 0.8, "multiplier": 2.0},
            )

    def test_model_version_list_exposes_rolling_comparison_fields(self) -> None:
        from dayahead_core import list_model_versions

        root = Path(__file__).resolve().parent / ".test_tmp" / f"versions_{time.time_ns()}"
        run_dir = root / "history" / "run_test"
        run_dir.mkdir(parents=True, exist_ok=True)
        metadata = {
            "run_id": "run_test",
            "created_at": "2026-05-05T12:00:00",
            "train_start_date": "2026-02-01",
            "train_end_date": "2026-04-16",
            "training_window_days": 56,
            "valid_days": 14,
            "num_boost_round": 400,
            "selected_model_backend_label": "CatBoost",
            "segment_mode": "custom",
            "segment_config": [
                {"name": "low", "start_time": "00:00", "end_time": "12:00"},
                {"name": "high", "start_time": "12:00", "end_time": "24:00"},
            ],
            "high_price_weighting": {"enabled": True, "quantile": 0.8, "multiplier": 2.0, "threshold": 380.0},
            "metrics": {
                "low": {"train_rows": 100, "valid_rows": 10, "best_iteration": 12, "final_mae": 20.0, "final_rmse": 30.0},
                "high": {"train_rows": 100, "valid_rows": 10, "best_iteration": 34, "final_mae": 40.0, "final_rmse": 50.0},
            },
            "rolling_backtest_metrics": {
                "30": {
                    "overall": {"mae": 53.4, "rmse": 104.3, "direction_accuracy": 63.7, "max_abs_error": 1000.0},
                    "net_load_only_baseline": {"mae": 82.2, "rmse": 130.0},
                    "model_vs_net_load_only_similarity": {"mae_improvement": 28.8, "mae_improvement_pct": 35.0},
                    "spike_errors": {
                        "high": {"mae": 180.0, "rmse": 270.0},
                        "high_net_load_only_baseline": {"mae": 183.0},
                    },
                    "segments": {"evening_peak": {"mae": 66.0, "rmse": 104.0}},
                }
            },
        }
        (run_dir / "metadata.json").write_text(json.dumps(metadata, ensure_ascii=False), encoding="utf-8")

        versions = list_model_versions(root)

        self.assertEqual(len(versions), 1)
        row = versions.iloc[0].to_dict()
        self.assertEqual(row["selected_model_backend_label"], "CatBoost")
        self.assertEqual(row["segment_count"], 2)
        self.assertTrue(row["high_price_weight_enabled"])
        self.assertEqual(row["rolling_30_mae"], 53.4)
        self.assertEqual(row["rolling_30_thermal_space_similarity_mae"], 82.2)
        self.assertEqual(row["rolling_30_high_price_mae"], 180.0)
        self.assertEqual(row["rolling_30_evening_peak_mae"], 66.0)
        self.assertEqual(row["algorithm_variants"][0]["best_iteration_avg"], 23.0)
        self.assertEqual(row["algorithm_variants"][0]["best_iteration_max"], 34)

    def test_prediction_variant_prefers_no_lag_and_skips_legacy_lag(self) -> None:
        from dayahead_core import select_prediction_model_variant

        metadata = {
            "model_variants": {
                "no_lag_96": {"feature_columns": ["net_load"]},
                "legacy_lag": {"feature_columns": ["net_load", self.removed_previous_day_feature]},
            }
        }

        forecast_df = pd.DataFrame({"net_load": [100.0], self.removed_previous_day_feature: [pd.NA]})

        self.assertEqual(select_prediction_model_variant(metadata, forecast_df), "no_lag_96")

        legacy_metadata = {
            "model_variants": {
                "legacy_lag": {"feature_columns": ["net_load", self.removed_previous_day_feature]},
                "legacy_direct": {"feature_columns": ["net_load"]},
            }
        }
        self.assertEqual(select_prediction_model_variant(legacy_metadata, forecast_df), "legacy_direct")

    def test_segment_prediction_variant_uses_saved_segment_selection(self) -> None:
        from dayahead_core import select_segment_prediction_model_variant

        metadata = {
            "selected_model_key": "global_model",
            "selected_segment_price_models": {"peak": "peak_model"},
            "model_variants": {
                "global_model": {"feature_columns": ["net_load"]},
                "peak_model": {"feature_columns": ["net_load"]},
            },
        }

        forecast_df = pd.DataFrame({"net_load": [100.0]})

        self.assertEqual(select_segment_prediction_model_variant(metadata, "peak", forecast_df), "peak_model")
        self.assertEqual(select_segment_prediction_model_variant(metadata, "valley", forecast_df), "global_model")

    def test_training_default_segment_selection_can_choose_different_models(self) -> None:
        from dayahead_core import select_best_segment_price_models

        selected = select_best_segment_price_models(
            {
                "model_a": {
                    "night": {"final_mae": 10.0, "final_rmse": 20.0},
                    "peak": {"final_mae": 80.0, "final_rmse": 100.0},
                },
                "model_b": {
                    "night": {"final_mae": 20.0, "final_rmse": 30.0},
                    "peak": {"final_mae": 40.0, "final_rmse": 50.0},
                },
            },
            ["night", "peak"],
            "model_a",
        )

        self.assertEqual(selected, {"night": "model_a", "peak": "model_b"})

    def test_model_candidate_ranking_filters_range_by_actual_price(self) -> None:
        from dayahead_core import rank_model_candidates

        root = Path(__file__).resolve().parent / ".test_tmp" / f"ranking_{time.time_ns()}"
        current_dir = root / "current"
        current_dir.mkdir(parents=True, exist_ok=True)
        metadata = {
            "selected_model_key": "model_a",
            "selected_segment_price_models": {"peak": "model_b"},
            "segments": [{"name": "peak", "start_time": "00:00", "end_time": "24:00"}],
            "validation_predictions_path": "validation_predictions.jsonl",
        }
        (current_dir / "metadata.json").write_text(json.dumps(metadata, ensure_ascii=False), encoding="utf-8")
        rows = [
            {"model_key": "model_a", "model_backend": "xgboost", "model_backend_label": "XGBoost", "segment": "peak", "date": "2026-05-01", "period": 1, "actual": 200.0, "predicted": 1000.0},
            {"model_key": "model_b", "model_backend": "catboost", "model_backend_label": "CatBoost", "segment": "peak", "date": "2026-05-01", "period": 1, "actual": 200.0, "predicted": 200.0},
            {"model_key": "model_a", "model_backend": "xgboost", "model_backend_label": "XGBoost", "segment": "peak", "date": "2026-05-01", "period": 2, "actual": 500.0, "predicted": 510.0},
            {"model_key": "model_b", "model_backend": "catboost", "model_backend_label": "CatBoost", "segment": "peak", "date": "2026-05-01", "period": 2, "actual": 500.0, "predicted": 550.0},
        ]
        with (current_dir / "validation_predictions.jsonl").open("w", encoding="utf-8") as file:
            for row in rows:
                file.write(json.dumps(row, ensure_ascii=False) + "\n")

        ranking = rank_model_candidates(root, metric="mae_range", price_min=250, price_max=1000)

        segment = ranking["segments"][0]
        self.assertEqual(segment["current_model_key"], "model_b")
        self.assertEqual([item["model_key"] for item in segment["rankings"]], ["model_a", "model_b"])
        self.assertEqual(segment["rankings"][0]["metric_value"], 10.0)
        self.assertEqual(segment["rankings"][0]["rows"], 1)

    def test_api_segment_selection_updates_current_and_history_metadata(self) -> None:
        from api_server import update_segment_price_model_selection

        root = Path(__file__).resolve().parent / ".test_tmp" / f"segment_selection_{time.time_ns()}"
        current_dir = root / "current"
        history_dir = root / "history" / "run_test"
        current_dir.mkdir(parents=True, exist_ok=True)
        history_dir.mkdir(parents=True, exist_ok=True)
        metadata = {
            "run_id": "run_test",
            "segments": [{"name": "peak", "start_time": "00:00", "end_time": "24:00"}],
            "model_variants": {
                "model_a": {"feature_columns": ["net_load"]},
                "model_b": {"feature_columns": ["net_load"]},
            },
            "selected_model_key": "model_a",
        }
        for path in [current_dir / "metadata.json", history_dir / "metadata.json"]:
            path.write_text(json.dumps(metadata, ensure_ascii=False), encoding="utf-8")

        updated = update_segment_price_model_selection(root, {"peak": "model_b"})

        self.assertEqual(updated["selected_segment_price_models"], {"peak": "model_b"})
        history_metadata = json.loads((history_dir / "metadata.json").read_text(encoding="utf-8"))
        self.assertEqual(history_metadata["selected_segment_price_models"], {"peak": "model_b"})

    def test_prediction_variant_rejects_residual_only_legacy_metadata(self) -> None:
        from dayahead_core import select_prediction_model_variant

        with self.assertRaises(ValueError):
            select_prediction_model_variant({"feature_columns": ["net_load"]}, pd.DataFrame({"net_load": [100.0]}))

    def test_price_interval_selection_is_per_segment(self) -> None:
        from price_interval import select_interval_segment_models

        model_variants = {
            "interval_xgboost": {"model_backend": "xgboost"},
            "interval_catboost": {"model_backend": "catboost"},
        }
        variant_metrics = {
            "interval_xgboost": {
                "overall": {"score": -0.1},
                "segments": {
                    "segment_1": {"score": -0.9},
                    "segment_2": {"score": -0.2},
                },
            },
            "interval_catboost": {
                "overall": {"score": -0.2},
                "segments": {
                    "segment_1": {"score": -0.3},
                    "segment_2": {"score": -0.8},
                },
            },
        }

        selected = select_interval_segment_models(model_variants, variant_metrics, ["segment_1", "segment_2"])

        self.assertEqual(selected["segment_1"], "interval_xgboost")
        self.assertEqual(selected["segment_2"], "interval_catboost")

    def test_price_interval_loader_uses_segment_selected_variant(self) -> None:
        import price_interval

        root = Path(__file__).resolve().parent / ".test_tmp" / f"interval_loader_{time.time_ns()}"
        (root / "interval_xgboost").mkdir(parents=True, exist_ok=True)
        (root / "interval_catboost").mkdir(parents=True, exist_ok=True)
        (root / "interval_xgboost" / "segment_1.json").write_text("xgb", encoding="utf-8")
        (root / "interval_catboost" / "segment_2.cbm").write_text("cat", encoding="utf-8")
        metadata = {
            "enabled": True,
            "selected_model_key": "interval_xgboost",
            "selected_segment_models": {"segment_1": "interval_xgboost", "segment_2": "interval_catboost"},
            "model_variants": {
                "interval_xgboost": {"model_dir": "interval_xgboost", "model_backend": "xgboost"},
                "interval_catboost": {"model_dir": "interval_catboost", "model_backend": "catboost"},
            },
            "metrics": {"segments": {"segment_1": {}, "segment_2": {}}},
        }
        loaded: list[tuple[str, str]] = []
        original_loader = price_interval.load_interval_backend_model
        try:
            price_interval.load_interval_backend_model = lambda path, backend: loaded.append((path.name, backend)) or f"{backend}:{path.name}"
            bundle = price_interval.load_interval_model_bundle(root, metadata)
        finally:
            price_interval.load_interval_backend_model = original_loader

        self.assertEqual(loaded, [("segment_1.json", "xgboost"), ("segment_2.cbm", "catboost")])
        self.assertEqual(bundle.models["segment_1"]["backend"], "xgboost")
        self.assertEqual(bundle.models["segment_2"]["backend"], "catboost")

    def test_training_window_resolves_latest_train_and_validation_range(self) -> None:
        from dayahead_core import TrainConfig, resolve_training_date_range

        frame = pd.DataFrame({"date": pd.date_range("2026-01-01", periods=30, freq="D")})

        start_date, end_date = resolve_training_date_range(frame, TrainConfig(training_window_days=10, valid_days=5))

        self.assertEqual(start_date, "2026-01-16")
        self.assertEqual(end_date, "2026-01-30")

    def test_forecast_template_allows_missing_reference_price(self) -> None:
        from dayahead_core import try_load_forecast_template

        root = Path(__file__).resolve().parent / ".test_tmp"
        root.mkdir(exist_ok=True)
        forecast_file = root / f"direct_price_forecast_{time.time_ns()}.xlsx"
        frame = pd.DataFrame(
            {
                "序号": list(range(1, 97)),
                "5.1日剩余电力值": [80.0] * 96,
                "5.2日剩余电力值": [85.0] * 96,
                "火电开机容量": [36095.0] * 96,
            }
        )
        frame.to_excel(forecast_file, index=False)

        forecast_df, reference_df, issues = try_load_forecast_template(
            forecast_file,
            set(),
            2026,
        )

        self.assertIsNotNone(forecast_df)
        self.assertIsNone(reference_df)
        self.assertNotIn(self.removed_previous_day_feature, forecast_df.columns)
        self.assertFalse(any(issue.field == "参考日前日价格" for issue in issues))

    def test_apply_manual_thermal_capacity_choice_records_source_and_values(self) -> None:
        from dayahead_core import (
            THERMAL_CAPACITY_FILE_VALUE_COLUMN,
            THERMAL_CAPACITY_MODEL_VALUE_COLUMN,
            THERMAL_CAPACITY_SOURCE_COLUMN,
            THERMAL_CAPACITY_VALUE_COLUMN,
            apply_thermal_capacity_choice,
            build_thermal_capacity_choice,
        )

        forecast_df = pd.DataFrame(
            {
                "date": [pd.Timestamp("2026-05-22")] * 2,
                "period": [1, 2],
                "thermal_on_capacity": [36095.0, 36095.0],
            }
        )

        choice = build_thermal_capacity_choice(
            forecast_df,
            {"mode": "manual", "manual_value": 36500.0},
            model_predicted_value=35280.0,
        )
        result = apply_thermal_capacity_choice(forecast_df, choice)

        self.assertEqual(result["thermal_on_capacity"].tolist(), [36500.0, 36500.0])
        self.assertEqual(result[THERMAL_CAPACITY_SOURCE_COLUMN].unique().tolist(), ["manual"])
        self.assertEqual(result[THERMAL_CAPACITY_VALUE_COLUMN].unique().tolist(), [36500.0])
        self.assertEqual(result[THERMAL_CAPACITY_MODEL_VALUE_COLUMN].unique().tolist(), [35280.0])
        self.assertEqual(result[THERMAL_CAPACITY_FILE_VALUE_COLUMN].unique().tolist(), [36095.0])

    def test_apply_model_thermal_capacity_choice_uses_model_value(self) -> None:
        from dayahead_core import (
            THERMAL_CAPACITY_SOURCE_COLUMN,
            apply_thermal_capacity_choice,
            build_thermal_capacity_choice,
        )

        forecast_df = pd.DataFrame(
            {
                "date": [pd.Timestamp("2026-05-22")] * 2,
                "period": [1, 2],
                "thermal_on_capacity": [36095.0, 36095.0],
            }
        )

        choice = build_thermal_capacity_choice(
            forecast_df,
            {"mode": "model", "manual_value": 36500.0},
            model_predicted_value=35280.0,
        )
        result = apply_thermal_capacity_choice(forecast_df, choice)

        self.assertEqual(result["thermal_on_capacity"].tolist(), [35280.0, 35280.0])
        self.assertEqual(result[THERMAL_CAPACITY_SOURCE_COLUMN].unique().tolist(), ["model"])

    def test_daily_thermal_capacity_frame_uses_one_row_per_day(self) -> None:
        from dayahead_core import build_daily_thermal_capacity_frame

        rows = []
        for day_index, date_text in enumerate(["2026-05-20", "2026-05-21"]):
            for period in range(1, 5):
                rows.append(
                    {
                        "date": pd.Timestamp(date_text),
                        "period": period,
                        "total_load": 1000.0 + day_index * 100 + period,
                        "renewable_power": 200.0 + period,
                        "net_load": 800.0 + day_index * 100,
                        "thermal_space_load_ratio": 0.8,
                        "thermal_on_capacity": 30000.0 + day_index * 500,
                        "weekday": day_index,
                        "month": 5,
                        "is_weekend": 0,
                        "is_holiday": 0,
                    }
                )

        daily = build_daily_thermal_capacity_frame(pd.DataFrame(rows))

        self.assertEqual(len(daily), 2)
        self.assertEqual(daily["thermal_on_capacity"].tolist(), [30000.0, 30500.0])
        self.assertEqual(daily["previous_thermal_on_capacity"].tolist(), [30000.0, 30000.0])

    def test_template_reference_day_feeds_similarity_before_history(self) -> None:
        from dayahead_core import TARGET_COLUMN, build_reference_frame, try_load_forecast_template

        root = Path(__file__).resolve().parent / ".test_tmp"
        root.mkdir(exist_ok=True)
        forecast_file = root / f"template_reference_{time.time_ns()}.xlsx"
        periods = list(range(1, 97))
        frame = pd.DataFrame(
            {
                "序号": periods,
                "5.3日剩余电力值": [80.0] * 96,
                "5.3日日前出清价格(元/MWh)": [9000.0 + period for period in periods],
                "5.4日剩余电力值": [85.0] * 96,
                "火电开机容量": [36095.0] * 96,
            }
        )
        frame.to_excel(forecast_file, index=False)

        forecast_df, template_reference_df, _ = try_load_forecast_template(forecast_file, set(), 2026)

        self.assertIsNotNone(forecast_df)
        self.assertIsNotNone(template_reference_df)

        history_rows = []
        for date_text, price_base in [("2026-05-01", 1000.0), ("2026-05-02", 2000.0), ("2026-05-03", 3000.0)]:
            for period in periods:
                history_rows.append(
                    {
                        "date": pd.Timestamp(date_text),
                        "period": period,
                        "net_load": 70.0,
                        "thermal_on_capacity": 36095.0,
                        TARGET_COLUMN: price_base + period,
                    }
                )
        history_df = pd.DataFrame(history_rows)

        reference_df, reference_dates = build_reference_frame(
            history_df=history_df,
            forecast_df=forecast_df,
            template_reference_df=template_reference_df,
            reference_days=3,
            reference_strategy="recent_n_days",
            holiday_dates=set(),
        )

        self.assertEqual(reference_dates, ["2026-05-03", "2026-05-02", "2026-05-01"])
        template_period_one = reference_df[(reference_df["date"] == pd.Timestamp("2026-05-03")) & (reference_df["period"] == 1)]
        self.assertEqual(float(template_period_one[TARGET_COLUMN].iloc[0]), 9001.0)

    def test_local_holiday_calendar_overrides_weekend_workday(self) -> None:
        from dayahead_core import day_type_of, load_holiday_calendar

        root = Path(__file__).resolve().parent / ".test_tmp"
        root.mkdir(exist_ok=True)
        calendar_file = root / f"holiday_calendar_{time.time_ns()}.csv"
        calendar_file.write_text(
            "date,type,name\n"
            "2026-02-15,workday,春节调休上班\n"
            "2026-02-16,holiday,春节\n",
            encoding="utf-8",
        )

        calendar = load_holiday_calendar(calendar_file)

        self.assertEqual(day_type_of(pd.Timestamp("2026-02-15"), calendar), "workday")
        self.assertEqual(day_type_of(pd.Timestamp("2026-02-16"), calendar), "holiday")
        self.assertEqual(day_type_of(pd.Timestamp("2026-02-14"), calendar), "weekend")

    def test_training_preferences_ignore_removed_similarity_weights(self) -> None:
        from dayahead_core import load_training_preferences, save_training_preferences

        root = Path(__file__).resolve().parent / ".test_tmp" / f"preferences_{time.time_ns()}"

        saved = save_training_preferences(root, {"similarity_weights": {"net_load": 1.0}})
        loaded = load_training_preferences(root)

        self.assertNotIn("similarity_weights", saved)
        self.assertNotIn("similarity_weights", loaded)

    def test_cli_no_longer_accepts_prediction_similarity_weights(self) -> None:
        from dayahead_timeseg_model import parse_args

        argv = ["dayahead_timeseg_model.py", "predict", "--similarity-weights", '{"net_load":1.0}']

        with patch.object(sys, "argv", argv):
            with self.assertRaises(SystemExit):
                parse_args()

    def test_frontend_labels_distinguish_model_version_from_segment_algorithm_selection(self) -> None:
        root = PROJECT_ROOT
        index_html = (root / "frontend_app" / "index.html").read_text(encoding="utf-8")
        app_js = (root / "frontend_app" / "assets" / "app.js").read_text(encoding="utf-8")

        self.assertIn("当前默认模型版本分时段指标", index_html)
        self.assertIn("当前版本内分时段算法选择", index_html)
        self.assertIn("设为默认模型版本", index_html)
        self.assertIn("每行可选择该时段预测用的算法变体", app_js)
        self.assertIn("请先选择历史模型版本", app_js)
        self.assertIn("已切换默认模型版本", app_js)

    def test_training_page_prioritizes_model_and_interval_configuration(self) -> None:
        root = PROJECT_ROOT
        index_html = (root / "frontend_app" / "index.html").read_text(encoding="utf-8")
        app_css = (root / "frontend_app" / "assets" / "app.css").read_text(encoding="utf-8")

        model_index = index_html.index("模型时段、参数配置")
        interval_index = index_html.index("价格区间概率模型训练配置")
        optimize_index = index_html.index("自动寻优</div>")
        manual_index = index_html.index("手动训练模型")

        self.assertLess(model_index, optimize_index)
        self.assertLess(interval_index, optimize_index)
        self.assertLess(optimize_index, manual_index)
        self.assertIn('class="train-config-grid"', index_html)
        self.assertIn('class="train-action-grid"', index_html)
        self.assertIn("顶部两块配置会同时作用于自动寻优和手动训练", index_html)
        self.assertIn(".train-config-grid", app_css)
        self.assertIn(".train-action-grid", app_css)

    def test_training_page_has_independent_manual_configs_for_both_models(self) -> None:
        root = PROJECT_ROOT
        index_html = (root / "frontend_app" / "index.html").read_text(encoding="utf-8")
        app_js = (root / "frontend_app" / "assets" / "app.js").read_text(encoding="utf-8")

        self.assertIn("具体价格模型训练配置", index_html)
        self.assertIn("价格区间概率模型训练配置", index_html)
        for element_id in [
            "price-training-window-days",
            "price-valid-days",
            "price-num-boost-round",
            "interval-training-window-days",
            "interval-valid-days",
            "interval-num-boost-round",
        ]:
            self.assertIn(f'id="{element_id}"', index_html)
        self.assertIn("price_model_config", app_js)
        self.assertIn("interval_model_config", app_js)

    def test_auto_optimization_keeps_price_and_interval_scores_separate(self) -> None:
        from api_server import business_window_score, interval_probability_score

        direct_metrics = {"mae": 50.0, "rmse": 80.0}
        good_interval = {"interval_accuracy": 0.8, "top2_accuracy": 0.95, "log_loss": 0.35, "high_price_recall": 0.75}
        poor_interval = {"interval_accuracy": 0.35, "top2_accuracy": 0.6, "log_loss": 1.2, "high_price_recall": 0.2}

        price_score, price_detail = business_window_score(direct_metrics, {})
        good_interval_score = interval_probability_score(good_interval)
        poor_interval_score = interval_probability_score(poor_interval)

        self.assertEqual(price_score, business_window_score(direct_metrics, {})[0])
        self.assertNotIn("price_interval_probability", [part["name"] for part in price_detail["parts"]])
        self.assertLess(good_interval_score, poor_interval_score)

    def test_prediction_reference_days_are_independent_by_strategy(self) -> None:
        from dayahead_core import normalize_reference_days_by_strategy

        result = normalize_reference_days_by_strategy(
            reference_days=5,
            recent_reference_days=3,
            same_type_reference_days=9,
        )

        self.assertEqual(result["recent_n_days"], 3)
        self.assertEqual(result["recent_same_type_days"], 9)

    def test_prediction_page_uses_independent_reference_day_inputs(self) -> None:
        root = PROJECT_ROOT
        index_html = (root / "frontend_app" / "index.html").read_text(encoding="utf-8")
        app_js = (root / "frontend_app" / "assets" / "app.js").read_text(encoding="utf-8")

        self.assertIn('id="recent-reference-days"', index_html)
        self.assertIn('id="same-type-reference-days"', index_html)
        self.assertIn("最近天数", index_html)
        self.assertIn("同类型日个数", index_html)
        self.assertIn("currentRecentReferenceDays", app_js)
        self.assertIn("currentSameTypeReferenceDays", app_js)
        self.assertIn("recent_reference_days", app_js)
        self.assertIn("same_type_reference_days", app_js)
        self.assertNotIn('id="reference-days"', index_html)
        self.assertNotIn("相似法参考天数", index_html)
        self.assertNotIn("相似法参考天数", app_js)

    def test_forecast_similarity_keeps_net_load_only_reference_price(self) -> None:
        from dayahead_core import NET_LOAD_ONLY_SIMILAR_PRICE_COLUMN, TARGET_COLUMN, attach_forecast_similarity_features

        target_df = pd.DataFrame(
            {
                "date": pd.to_datetime(["2026-05-04"]),
                "period": [1],
                "net_load": [100.0],
                "renewable_power": [50.0],
                "thermal_on_capacity": [1000.0],
                "day_type": ["workday"],
            }
        )
        reference_df = pd.DataFrame(
            {
                "date": pd.to_datetime(["2026-05-01", "2026-05-02"]),
                "period": [1, 1],
                "net_load": [101.0, 130.0],
                "renewable_power": [500.0, 50.0],
                "thermal_on_capacity": [1000.0, 1000.0],
                "day_type": ["workday", "workday"],
                TARGET_COLUMN: [10.0, 200.0],
            }
        )

        result = attach_forecast_similarity_features(target_df, reference_df)

        self.assertLess(float(result.loc[0, NET_LOAD_ONLY_SIMILAR_PRICE_COLUMN]), 50.0)
        self.assertNotIn("similar_price", result.columns)

    def test_date_window_selects_only_overlapping_month_files(self) -> None:
        from dayahead_core import list_excel_files_for_date_window

        root = Path(__file__).resolve().parent / ".test_tmp"
        root.mkdir(exist_ok=True)
        marker = time.time_ns()
        selected_names = [
            f"{marker}_2025-12.xlsx",
            f"{marker}_2026-01.xlsx",
            f"{marker}_2026-04.xlsx",
        ]
        excluded_names = [
            f"{marker}_2025-11.xlsx",
            f"{marker}_2026-05.xlsx",
        ]
        for name in [*selected_names, *excluded_names]:
            (root / name).write_bytes(b"x")

        files = list_excel_files_for_date_window(root, "2026-01-01", "2026-04-30", leading_days=1)
        result_names = {path.name for path in files if path.name.startswith(str(marker))}

        self.assertEqual(result_names, set(selected_names))

    def test_prediction_history_window_looks_back_enough_for_same_type_days(self) -> None:
        from dayahead_core import prediction_history_window

        start_date, end_date = prediction_history_window(pd.Timestamp("2026-05-02"), 4)

        self.assertEqual(start_date, "2026-03-31")
        self.assertEqual(end_date, "2026-05-01")

    def test_dataframe_cache_round_trip_and_signature_invalidation(self) -> None:
        from dayahead_core import (
            build_history_source_signature,
            load_dataframe_cache,
            save_dataframe_cache,
        )

        root = Path(__file__).resolve().parent / ".test_tmp"
        root.mkdir(exist_ok=True)
        source_file = root / f"history_cache_source_{time.time_ns()}.xlsx"
        source_file.write_bytes(b"first")

        cache_dir = root / "cache"
        cache_name = source_file.stem
        frame = pd.DataFrame(
            {
                "date": pd.to_datetime(["2025-01-01"]),
                "period": [1],
                "net_load": [123.45],
            }
        )
        signature = build_history_source_signature(source_file, None, True)

        save_dataframe_cache(cache_dir, cache_name, signature["cache_key"], frame, ["skipped"])
        cached = load_dataframe_cache(cache_dir, cache_name, signature["cache_key"])

        self.assertIsNotNone(cached)
        cached_frame, skipped_sheets = cached
        pd.testing.assert_frame_equal(cached_frame, frame)
        self.assertEqual(skipped_sheets, ["skipped"])

        time.sleep(0.01)
        source_file.write_bytes(b"changed")
        changed_signature = build_history_source_signature(source_file, None, True)

        self.assertNotEqual(signature["cache_key"], changed_signature["cache_key"])
        self.assertIsNone(load_dataframe_cache(cache_dir, cache_name, changed_signature["cache_key"]))

    def test_build_training_frame_rebuilds_stale_feature_cache_missing_direct_columns(self) -> None:
        from dayahead_core import (
            TARGET_COLUMN,
            THERMAL_SPACE_LOAD_RATIO_COLUMN,
            TOTAL_LOAD_COLUMN,
            TrainConfig,
            build_training_frame,
            derived_cache_key,
            load_cached_history_collection,
            save_dataframe_cache,
        )

        root = Path(__file__).resolve().parent / ".test_tmp" / f"stale_feature_cache_{time.time_ns()}"
        history_dir = root / "history"
        model_root = root / "models"
        history_dir.mkdir(parents=True)
        model_root.mkdir(parents=True)
        periods = list(range(96))
        history_frame = pd.DataFrame(
            {
                "总加电力值(MW)": [1000.0 + period for period in periods],
                "电力值(MW)": [200.0 for _ in periods],
                "日前出清价格(元/MWh)": [300.0 + period * 0.1 for period in periods],
                "日前-出清概况": ["火电开机容量 500MW" for _ in periods],
            }
        )
        with pd.ExcelWriter(history_dir / "2025年1月.xlsx", engine="openpyxl") as writer:
            history_frame.to_excel(writer, sheet_name="2025-01-01", index=False)

        _, _, _, signature, _ = load_cached_history_collection(history_dir, None, True, model_root, leading_days=1)
        stale_feature_key = derived_cache_key(signature["cache_key"], "direct_price_training_features_v0")
        stale_features = pd.DataFrame(
            {
                "date": pd.to_datetime(["2025-01-01"]),
                "period": [1],
                "hour": [0],
                "weekday": [2],
                "month": [1],
                "is_weekend": [0],
                "is_holiday": [0],
                "day_type": ["workday"],
                "net_load": [800.0],
                "thermal_on_capacity": [500.0],
                "segment": ["night"],
                TARGET_COLUMN: [300.0],
            }
        )
        save_dataframe_cache(model_root / "cache", "training_features", stale_feature_key, stale_features, [])

        rebuilt_frame, _, _ = build_training_frame(TrainConfig(history_dir=history_dir, model_root=model_root))

        self.assertIn(TOTAL_LOAD_COLUMN, rebuilt_frame.columns)
        self.assertIn(THERMAL_SPACE_LOAD_RATIO_COLUMN, rebuilt_frame.columns)
        self.assertEqual(len(rebuilt_frame), 96)

    def test_build_training_frame_uses_shared_feature_cache_root(self) -> None:
        from dayahead_core import (
            TARGET_COLUMN,
            TrainConfig,
            build_training_frame,
            derived_cache_key,
            required_training_feature_cache_columns,
            save_dataframe_cache,
        )

        root = Path(__file__).resolve().parent / ".test_tmp" / f"shared_feature_cache_{time.time_ns()}"
        model_root = root / "models"
        shared_cache_root = root / "window_search_cache"
        raw_history = pd.DataFrame(
            {
                "date": pd.to_datetime(["2026-05-01"]),
                "period": [1],
                "segment": ["night"],
                TARGET_COLUMN: [320.0],
                "net_load": [800.0],
            }
        )
        cached_row = {column: 0.0 for column in required_training_feature_cache_columns()}
        cached_row.update(
            {
                "date": pd.Timestamp("2026-05-01"),
                "period": 1,
                "segment": "night",
                TARGET_COLUMN: 321.0,
                "net_load": 1234.0,
            }
        )
        cached_features = pd.DataFrame([cached_row])
        feature_cache_key = derived_cache_key(
            "raw-history-key",
            json.dumps(
                {
                    "namespace": "direct_price_training_features",
                    "similarity_reference_days": 7,
                    "feature_context_start_date": "2026-04-24",
                    "effective_start_date": "2026-05-01",
                    "effective_end_date": "2026-05-01",
                },
                ensure_ascii=False,
                sort_keys=True,
            ),
        )
        save_dataframe_cache(shared_cache_root, "training_features", feature_cache_key, cached_features, [])

        with patch(
            "dayahead_core.load_cached_history_collection",
            return_value=(raw_history, [], set(), {"cache_key": "raw-history-key"}, []),
        ), patch("dayahead_core.attach_lag_features", side_effect=AssertionError("shared cache was not used")):
            frame, _, _ = build_training_frame(
                TrainConfig(
                    history_dir=root / "history",
                    model_root=model_root,
                    feature_cache_root=shared_cache_root,
                    start_date="2026-05-01",
                    end_date="2026-05-01",
                    similarity_reference_days=7,
                )
            )

        self.assertEqual(float(frame.loc[0, "net_load"]), 1234.0)
        self.assertFalse((model_root / "cache" / f"training_features_{feature_cache_key}.pkl").exists())


if __name__ == "__main__":
    unittest.main()
