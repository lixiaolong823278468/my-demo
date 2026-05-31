from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook, Workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent


def _write_multi_day_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    headers = [
        "序号",
        "5.20日风光电力值(MW)",
        "5.20日总加电力值(MW)",
        "5.20日剩余电力值",
        "5.20日日前出清价格(元/MWh)",
        "火电开机容量 ",
        None,
        None,
        "5.21日风光电力值(MW)",
        "5.21日总加电力值(MW)",
        "5.22日风光电力值(MW)",
        "5.22日总加电力值(MW)",
    ]
    sheet.append(headers)
    for period in range(1, 97):
        sheet.append(
            [
                period,
                9000 + period,
                33000 + period,
                24000,
                280 + period / 100,
                37700 if period == 1 else None,
                None,
                None,
                8000 + period,
                32000 + period,
                8100 + period,
                32100 + period,
            ]
        )
    workbook.save(path)


def _write_dynamic_multi_day_workbook(path: Path, baseline_label: str, future_labels: tuple[str, str]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    headers = [
        "\u5e8f\u53f7",
        f"{baseline_label}\u65e5\u98ce\u5149\u7535\u529b\u503c(MW)",
        f"{baseline_label}\u65e5\u603b\u52a0\u7535\u529b\u503c(MW)",
        f"{baseline_label}\u65e5\u5269\u4f59\u7535\u529b\u503c",
        f"{baseline_label}\u65e5\u65e5\u524d\u51fa\u6e05\u4ef7\u683c(\u5143/MWh)",
        "\u706b\u7535\u5f00\u673a\u5bb9\u91cf",
        f"{future_labels[0]}\u65e5\u98ce\u5149\u7535\u529b\u503c(MW)",
        f"{future_labels[0]}\u65e5\u603b\u52a0\u7535\u529b\u503c(MW)",
        f"{future_labels[1]}\u65e5\u98ce\u5149\u7535\u529b\u503c(MW)",
        f"{future_labels[1]}\u65e5\u603b\u52a0\u7535\u529b\u503c(MW)",
    ]
    sheet.append(headers)
    for period in range(1, 97):
        sheet.append(
            [
                period,
                9000 + period,
                33000 + period,
                24000,
                280 + period / 100,
                37700 if period == 1 else None,
                8000 + period,
                32000 + period,
                8100 + period,
                32100 + period,
            ]
        )
    workbook.save(path)


class MultiDayPredictionTests(unittest.TestCase):
    def test_parse_multi_day_workbook_detects_dynamic_days_and_baseline_values(self):
        from dayahead_core import parse_multi_day_forecast_workbook

        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "预测文件-N天.xlsx"
            _write_multi_day_workbook(workbook_path)

            parsed = parse_multi_day_forecast_workbook(workbook_path, default_year=2026)

        self.assertEqual(parsed["baseline_date"], "2026-05-20")
        self.assertEqual(parsed["forecast_dates"], ["2026-05-21", "2026-05-22"])
        self.assertEqual(len(parsed["rows"]), 96)
        first = parsed["rows"][0]
        self.assertEqual(first["period"], 1)
        self.assertEqual(first["baseline_price"], 280.01)
        self.assertEqual(first["thermal_on_capacity"], 37700)
        self.assertEqual(first["forecast_days"]["2026-05-21"]["renewable_power"], 8001)
        self.assertEqual(first["forecast_days"]["2026-05-21"]["total_load"], 32001)
        self.assertEqual(first["forecast_days"]["2026-05-21"]["net_load"], 24000)
        self.assertEqual(first["forecast_days"]["2026-05-22"]["net_load"], 24000)


    def test_parse_multi_day_workbook_does_not_require_may_20_baseline(self):
        from dayahead_core import TARGET_COLUMN, build_multi_day_forecast_frame, parse_multi_day_forecast_workbook

        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "multi-day-dynamic.xlsx"
            _write_dynamic_multi_day_workbook(workbook_path, "5.24", ("5.25", "5.26"))
            parsed = parse_multi_day_forecast_workbook(workbook_path, default_year=2026)

            forecast_df, reference_df = build_multi_day_forecast_frame(parsed, "2026-05-25", holiday_dates=set())

        self.assertEqual(parsed["baseline_date"], "2026-05-24")
        self.assertEqual(parsed["forecast_dates"], ["2026-05-25", "2026-05-26"])
        self.assertEqual(forecast_df["date"].dt.strftime("%Y-%m-%d").unique().tolist(), ["2026-05-25"])
        self.assertEqual(reference_df["date"].dt.strftime("%Y-%m-%d").unique().tolist(), ["2026-05-24"])
        self.assertEqual(reference_df.loc[0, TARGET_COLUMN], 280.01)
        self.assertEqual(forecast_df.loc[0, "thermal_on_capacity"], 37700)


    def test_build_multi_day_forecast_frame_reuses_baseline_price_and_capacity(self):
        from dayahead_core import TARGET_COLUMN, build_multi_day_forecast_frame, parse_multi_day_forecast_workbook

        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "预测文件-N天.xlsx"
            _write_multi_day_workbook(workbook_path)
            parsed = parse_multi_day_forecast_workbook(workbook_path, default_year=2026)

            forecast_df, reference_df = build_multi_day_forecast_frame(parsed, "2026-05-21", holiday_dates=set())

        self.assertEqual(len(forecast_df), 96)
        self.assertEqual(len(reference_df), 96)
        self.assertEqual(forecast_df["date"].dt.strftime("%Y-%m-%d").unique().tolist(), ["2026-05-21"])
        self.assertEqual(reference_df["date"].dt.strftime("%Y-%m-%d").unique().tolist(), ["2026-05-20"])
        self.assertEqual(forecast_df.loc[0, "thermal_on_capacity"], 37700)
        self.assertEqual(reference_df.loc[0, TARGET_COLUMN], 280.01)
        self.assertEqual(reference_df.loc[10, TARGET_COLUMN], 280.11)


    def test_export_multi_day_prediction_workbook_writes_one_sheet_per_day(self):
        from dayahead_core import export_multi_day_prediction_workbook

        result = {
            "forecast_dates": ["2026-05-21", "2026-05-22"],
            "days": [
                {
                    "forecast_date": "2026-05-21",
                    "rows": [
                        {"period": period, "predicted_price": 300 + period, "net_load_only_similar_price": 290 + period}
                        for period in range(1, 97)
                    ],
                },
                {
                    "forecast_date": "2026-05-22",
                    "rows": [
                        {"period": period, "predicted_price": 400 + period, "net_load_only_similar_price": 390 + period}
                        for period in range(1, 97)
                    ],
                },
            ],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "multi-day-output.xlsx"

            export_multi_day_prediction_workbook(result, output_path)

            workbook = load_workbook(output_path, read_only=True, data_only=True)
            self.assertEqual(workbook.sheetnames, ["2026-05-21", "2026-05-22"])
            first_sheet = workbook["2026-05-21"]
            self.assertEqual(first_sheet.max_row, 97)
            self.assertEqual(first_sheet["A1"].value, "2026-05-21 模型预测价格")
            self.assertEqual(first_sheet["B1"].value, "2026-05-21 净负荷相似法预测价格")
            self.assertEqual(first_sheet["A2"].value, 301)
            self.assertEqual(first_sheet["B97"].value, 386)
            workbook.close()

    def test_export_multi_day_prediction_workbook_replaces_corrupt_file_atomically(self):
        from dayahead_core import export_multi_day_prediction_workbook

        result = {
            "days": [
                {
                    "forecast_date": "2026-05-25",
                    "rows": [{"period": period, "predicted_price": 300 + period} for period in range(1, 97)],
                }
            ],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "multi-day-output.xlsx"
            output_path.write_bytes(b"broken")

            export_multi_day_prediction_workbook(result, output_path)

            workbook = load_workbook(output_path, read_only=True, data_only=True)
            self.assertEqual(workbook.sheetnames, ["2026-05-25"])
            self.assertEqual(workbook["2026-05-25"]["A1"].value, "2026-05-25 模型预测价格")
            self.assertEqual(workbook["2026-05-25"]["A97"].value, 396)
            workbook.close()

    def test_export_multi_day_prediction_workbook_accepts_dataframe_rows(self):
        import pandas as pd
        from dayahead_core import export_multi_day_prediction_workbook

        result = {
            "days": [
                {
                    "forecast_date": "2026-05-25",
                    "rows": pd.DataFrame(
                        [{"period": period, "predicted_price": 300 + period} for period in range(1, 97)]
                    ),
                }
            ],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "multi-day-output.xlsx"

            export_multi_day_prediction_workbook(result, output_path)

            workbook = load_workbook(output_path, read_only=True, data_only=True)
            self.assertEqual(workbook.sheetnames, ["2026-05-25"])
            self.assertEqual(workbook["2026-05-25"]["A97"].value, 396)
            workbook.close()

    def test_api_multi_day_template_preview_reports_detected_dates(self):
        from api_server import load_multi_day_template_preview

        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "预测文件-N天.xlsx"
            _write_multi_day_workbook(workbook_path)

            preview = load_multi_day_template_preview(workbook_path)

        self.assertEqual(preview["row_count"], 96)
        self.assertEqual(preview["detected_forecast_dates"], ["2026-05-21", "2026-05-22"])
        self.assertIn("5.21日风光电力值(MW)", preview["columns"])

    def test_frontend_declares_independent_multi_day_prediction_page(self):
        root = PROJECT_ROOT
        index_html = (root / "frontend_app" / "index.html").read_text(encoding="utf-8")
        app_js = (root / "frontend_app" / "assets" / "app.js").read_text(encoding="utf-8")
        app_css = (root / "frontend_app" / "assets" / "app.css").read_text(encoding="utf-8")

        self.assertIn('data-tab="predict-multi-day"', index_html)
        self.assertIn('id="tab-predict-multi-day"', index_html)
        self.assertIn('id="predict-multi-day-btn"', index_html)
        self.assertIn('id="export-multi-day-btn"', index_html)
        panel = index_html.split('id="tab-predict-multi-day"', 1)[1].split('id="tab-versions"', 1)[0]
        self.assertNotIn('id="same-type-reference-days"', panel)
        self.assertIn("startMultiDayPredict", app_js)
        self.assertIn("/api/predict-multi-day", app_js)
        self.assertIn("#tab-predict-multi-day .similarity-config-card", app_css)
        self.assertIn("#tab-predict-multi-day .predict-reference-field input", app_css)


if __name__ == "__main__":
    unittest.main()
